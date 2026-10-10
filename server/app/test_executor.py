"""Execute concrete test plans; reports contain observations, never AI pass guesses."""
from __future__ import annotations

import logging
import asyncio
import ast
import base64
from copy import deepcopy
from datetime import datetime, timezone
import json
import math
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import time
from html import escape
from urllib.parse import urljoin, urlsplit
from uuid import UUID
from defusedxml import ElementTree

import httpx
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from app.config import get_settings
from app.db import connect
from app.files import resolve_object
from app.docx_export import crop_picture as _crop_report_picture
from app.docx_parser import _markdown_text

logger = logging.getLogger(__name__)


class MissingInput(ValueError):
    pass


_FOCUS_ACTIONS = {"fill", "check", "uncheck", "select_option", "assert_text", "assert_count", "assert_value", "assert_enabled"}
_PAGE_READY_TIMEOUT_MS = 10000
_PAGE_READY_JS = """() => {
    if (!document.body || document.readyState === 'loading' || document.fonts?.status === 'loading') return false;
    const visible = element => {
        if (!element) return false;
        const box = element.getBoundingClientRect();
        if (box.width <= 0 || box.height <= 0 || box.right <= 0 || box.bottom <= 0 || box.left >= innerWidth || box.top >= innerHeight) return false;
        for (let parent = element; parent; parent = parent.parentElement) {
            const style = getComputedStyle(parent);
            if (style.display === 'none' || style.visibility === 'hidden' || style.visibility === 'collapse' || Number(style.opacity) === 0) return false;
        }
        return true;
    };
    let ready = [...document.body.querySelectorAll('button,input:not([type="hidden"]),select,textarea,canvas,svg,img')]
        .some(element => visible(element) && (element.tagName !== 'IMG' || (element.complete && element.naturalWidth > 0)));
    if (!ready) {
        const text = document.createTreeWalker(document.body, NodeFilter.SHOW_TEXT);
        for (let node = text.nextNode(); node; node = text.nextNode()) {
            if (!node.nodeValue.trim() || /^(SCRIPT|STYLE|NOSCRIPT|TEMPLATE)$/.test(node.parentElement?.tagName || '') || !visible(node.parentElement)) continue;
            const range = document.createRange(); range.selectNodeContents(node);
            const box = range.getBoundingClientRect();
            if (box.width > 0 && box.height > 0 && box.right > 0 && box.bottom > 0 && box.left < innerWidth && box.top < innerHeight) { ready = true; break; }
        }
    }
    return ready;
}"""
_PAGE_PAINT_JS = """() => new Promise((resolve, reject) => {
    const timer = setTimeout(() => reject(new Error('首屏绘制等待超时。')), 1000);
    requestAnimationFrame(() => requestAnimationFrame(() => { clearTimeout(timer); resolve(true); }));
})"""


def _focus_rect(box: dict | None, width: int = 1920, height: int = 1080) -> dict | None:
    """Describe a readable region of the single full viewport capture."""
    if not isinstance(box, dict):
        return None
    values = [box.get(key) for key in ("x", "y", "width", "height")]
    if any(isinstance(value, bool) or not isinstance(value, (int, float)) or abs(value) > 1e9 or not math.isfinite(value) for value in values):
        return None
    x, y, box_width, box_height = values
    if box_width <= 0 or box_height <= 0 or x >= width or y >= height or x + box_width <= 0 or y + box_height <= 0:
        return None
    crop_width = min(width, 960, max(640, int(box_width) + 300))
    crop_height = min(height, 540, max(240, int(box_height) + 200))
    return {"x": max(0, min(width - crop_width, int(x) - 150)),
            "y": max(0, min(height - crop_height, int(y) - 100)),
            "width": crop_width, "height": crop_height}


def run_directory(run_id: str) -> Path:
    root = get_settings().file_data_dir.resolve()
    base = root / "test-runs"
    directory = base / str(UUID(run_id))
    if base.is_symlink() or base.is_junction() or directory.is_symlink() or directory.is_junction():
        raise MissingInput("测试证据目录不正确。")
    directory.mkdir(parents=True, exist_ok=True)
    return directory


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _origin(url: str) -> tuple[str, str, int | None]:
    parsed = urlsplit(url)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname or parsed.username or parsed.password:
        raise MissingInput("需要不含账号密码的 HTTP 服务地址。")
    return parsed.scheme, parsed.hostname.lower(), parsed.port or (443 if parsed.scheme == "https" else 80)


def _target(task: dict, url: str, *, frontend: bool = False) -> str:
    base = str(task.get("frontendUrl" if frontend else "backendUrl") or "")
    if not base:
        raise MissingInput("尚未填写前端地址。" if frontend else "尚未填写后端地址。")
    resolved = urljoin(base.rstrip("/") + "/", url)
    allowed = {_origin(value) for value in (task.get("frontendUrl"), task.get("backendUrl")) if value}
    if _origin(resolved) not in allowed:
        raise MissingInput("测试地址超出任务配置的前后端服务。")
    return resolved


def _test_text(value: object) -> str:
    if isinstance(value, str):
        return value
    if (not isinstance(value, dict) or set(value) != {"text", "repeat"}
            or not isinstance(value["text"], str) or not value["text"]
            or type(value["repeat"]) is not int or not 1 <= value["repeat"] <= 4000
            or len(value["text"]) * value["repeat"] > 4000):
        raise MissingInput("测试文本需为字符串，或最多 4000 字的 {text, repeat} 重复文本。")
    return value["text"] * value["repeat"]


def _value(task: dict, value: object) -> str:
    text = _test_text(value)
    for name, key in (("username", "testUsername"), ("password", "testPassword")):
        if "{{" + name + "}}" in text:
            if not task.get(key):
                raise MissingInput("该步骤需要测试账号和密码。")
            text = text.replace("{{" + name + "}}", str(task[key]))
    return text


def _redact(task: dict, value: str) -> str:
    for secret in (task.get("testPassword"), get_settings().ai_api_key):
        if secret:
            value = value.replace(str(secret), "[已隐藏]")
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value)


def _locator(page, spec: dict):
    by, value = spec.get("by"), spec.get("value")
    if not isinstance(value, str) or (not value and (by != "role" or not spec.get("role"))):
        raise MissingInput("操作步骤缺少可用的控件定位信息。")
    exact = bool(spec.get("exact", True))
    if by == "role" and spec.get("role"):
        return page.get_by_role(spec["role"], name=value, exact=exact) if value else page.get_by_role(spec["role"])
    if by == "label":
        return page.get_by_label(value, exact=exact)
    if by == "text":
        return page.get_by_text(value, exact=exact)
    if by == "testid":
        return page.get_by_test_id(value)
    raise MissingInput("仅支持角色、label、可见文字和 testid 定位。")


def _functional(task: dict, case: dict, folder: Path, cancelled, progress) -> dict:
    from playwright.sync_api import sync_playwright, expect

    steps = case.get("steps", [])
    if not steps or not any(step.get("action", "").startswith("assert_") for step in steps):
        raise MissingInput("功能用例缺少独立的预期结果断言。")
    if any(step.get("action") in {"click", "fill", "press", "select_option", "check", "uncheck"} for step in steps) and not task.get("dataScope", "").strip():
        raise MissingInput("执行交互前需填写允许修改的测试数据范围。")
    observed = {"status": "running", "steps": []}
    with sync_playwright() as playwright:
        try:
            browser = playwright.chromium.launch(headless=True)
        except Exception:
            raise MissingInput("Chromium 测试浏览器尚未安装；在服务器环境安装 Playwright Chromium 后重试。") from None
        context = browser.new_context(viewport={"width": 1920, "height": 1080}, service_workers="block")
        page = context.new_page()
        page.set_default_timeout(10000)
        allowed = {_origin(value) for value in (task.get("frontendUrl"), task.get("backendUrl")) if value}
        policy_error = ""
        def guard(route):
            nonlocal policy_error
            try:
                origin = _origin(route.request.url)
            except ValueError:
                route.abort()
                return
            if origin not in allowed and route.request.is_navigation_request():
                policy_error = "浏览器跳转超出任务配置的服务范围。"
            route.continue_() if origin in allowed else route.abort()
        context.route("**/*", guard)
        def websocket_guard(route):
            nonlocal policy_error
            parsed = urlsplit(route.url)
            scheme = {"ws": "http", "wss": "https"}.get(parsed.scheme)
            try:
                permitted = bool(scheme) and _origin(parsed._replace(scheme=scheme).geturl()) in allowed
            except ValueError:
                permitted = False
            if permitted:
                route.connect_to_server()
            else:
                # Routed sockets do not reach a server unless explicitly
                # connected. Avoid a nested synchronous close call from this
                # async-dispatched Playwright callback; end the failed context
                # after retaining the current step's screenshot instead.
                policy_error = "WebSocket 连接超出任务配置的服务范围，已阻止。"
        context.route_web_socket("**/*", websocket_guard)
        def close_popup(popup):
            nonlocal policy_error
            if popup != page:
                policy_error = "测试页面尝试打开新窗口，已阻止，请将用例限定在当前页面。"
                popup.close()
        context.on("page", close_popup)
        console_errors: list[str] = []
        page.on("pageerror", lambda error: console_errors.append(_redact(task, str(error))))
        try:
            for index, step in enumerate(steps):
                if cancelled():
                    observed["status"] = "stopped"
                    break
                action = step.get("action")
                locator = None
                observation = {"action": action, "description": step.get("description") or str(action), "status": "passed", "actual": ""}
                try:
                    if action == "open":
                        response = page.goto(_target(task, str(step.get("url", "")), frontend=True), wait_until="domcontentloaded", timeout=30000)
                        if response and response.status >= 400:
                            raise AssertionError(f"页面 HTTP {response.status}")
                        # DOMContentLoaded can precede Vue's first render. Wait for
                        # visible first-screen content, not generic business/network idle.
                        page.wait_for_function(_PAGE_READY_JS, timeout=_PAGE_READY_TIMEOUT_MS)
                        page.evaluate(_PAGE_PAINT_JS)
                        observation["actual"] = f"已打开 {page.url}"
                    elif action == "assert_url":
                        expected = _target(task, str(step.get("url", "")), frontend=True)
                        expect(page).to_have_url(expected)
                        observation["actual"] = f"页面地址符合预期：{page.url}"
                    elif action in {"click", "fill", "press", "select_option", "check", "uncheck", "assert_text", "assert_count", "assert_value", "assert_enabled"}:
                        locator = _locator(page, step.get("locator", {}))
                        if action == "click":
                            locator.click()
                            observation["actual"] = "控件点击完成"
                        elif action == "fill":
                            locator.fill(_value(task, step.get("value", "")))
                            observation["actual"] = "字段填写完成"  # Never persist password values.
                        elif action == "press":
                            key = str(step.get("value", "Enter"))
                            if key not in {"Enter", "Tab", "Escape", "ArrowDown", "ArrowUp"}:
                                raise MissingInput("此按键不在测试允许范围。")
                            locator.press(key)
                            observation["actual"] = f"已按 {key}"
                        elif action == "select_option":
                            locator.select_option(label=_value(task, step.get("value", "")))
                            observation["actual"] = "选项选择完成"
                        elif action in {"check", "uncheck"}:
                            locator.set_checked(action == "check")
                            observation["actual"] = "已勾选" if action == "check" else "已取消勾选"
                        elif action == "assert_value":
                            expect(locator).to_have_value(_value(task, step.get("value", "")))
                            observation["actual"] = "字段值符合预期"
                        elif action == "assert_enabled":
                            expect(locator).to_be_enabled() if step.get("enabled", True) else expect(locator).to_be_disabled()
                            observation["actual"] = "控件可用状态符合预期"
                        elif action == "assert_text":
                            expect(locator).to_be_visible()
                            if "value" in step:
                                expect(locator).to_have_text(_value(task, step["value"]))
                            observation["actual"] = "可见文字符合预期：" + _redact(task, locator.inner_text()[:1200])
                        else:
                            count = step.get("count")
                            if not isinstance(count, int) or not 0 <= count <= 10000:
                                raise MissingInput("数量断言需要有效 count。")
                            expect(locator).to_have_count(count)
                            observation["actual"] = f"实际控件数量：{count}"
                    else:
                        raise MissingInput(f"不支持的功能测试操作：{action}")
                except MissingInput as error:
                    observation.update(status="blocked", actual=str(error))
                    observed.update(status="blocked", error=str(error))
                except Exception as error:
                    message = _redact(task, str(error))[:4000]
                    observation.update(status="failed", actual=message)
                    observed.update(status="failed", error=message)
                image_name = f"{UUID(case['id'])}-{index + 1}.png"
                focus_rect = None
                screenshot_url = page.url
                if action in _FOCUS_ACTIONS and locator is not None:
                    try:
                        if locator.count() == 1 and locator.is_visible():
                            locator.scroll_into_view_if_needed(timeout=1000)
                            focus_rect = _focus_rect(locator.bounding_box(timeout=1000))
                    except Exception:
                        # Optional display metadata must not lose the full evidence.
                        logger.debug("focus metadata unavailable for screenshot", exc_info=True)
                try:
                    mask = [page.locator('input[type="password"],input[autocomplete="current-password"],input[autocomplete="new-password"]')]
                    page.screenshot(path=str(folder / image_name), full_page=False, animations="disabled", mask=mask, mask_color="gray")
                    observation["screenshot"] = image_name
                    if focus_rect and page.url == screenshot_url:
                        observation["focusRect"] = focus_rect
                except Exception:
                    observation["evidenceError"] = "无法取得该步骤截图。"
                    observed.update(status="failed", error="步骤缺少截图证据。")
                observed["steps"].append(observation)
                if policy_error:
                    observed.update(status="failed", error=policy_error)
                progress(observed)
                if observed["status"] in {"failed", "blocked"}:
                    break
            if observed["status"] == "running":
                observed["status"] = "passed"
            if console_errors:
                observed["consoleErrors"] = console_errors[:30]
                if observed["status"] == "passed":
                    observed.update(status="failed", error="页面出现未处理的 JavaScript 错误。")
        finally:
            context.close()
            browser.close()
    return observed


def _source(task: dict, side: str, relative: str | None = None) -> Path:
    if side not in {"frontend", "backend"} or not task.get(side + "Root"):
        raise MissingInput("尚未填写该用例的源码目录。")
    configured = Path(task[side + "Root"])
    root = configured.resolve()
    if not root.is_dir() or configured.is_symlink() or configured.is_junction():
        raise MissingInput("源码目录不存在。")
    if relative is None:
        return root
    from app.test_management import _source_path
    try:
        return _source_path(task, side, relative)
    except ValueError as error:
        raise MissingInput(str(error)) from None


MIN_NODE_MAJOR = 22


def _node_major(node: str) -> int:
    """Node 22 added the per-file test:summary event the result check relies on."""
    try:
        output = subprocess.run([node, "--version"], capture_output=True, text=True, timeout=10, check=False).stdout
        return int(output.strip().lstrip("v").split(".", 1)[0])
    except (OSError, subprocess.SubprocessError, ValueError):
        return 0


def _unit(task: dict, case: dict, folder: Path, cancelled, progress) -> dict:
    results = {"status": "running", "steps": []}
    if not case.get("steps"):
        raise MissingInput("单元用例没有可执行测试。")
    for index, step in enumerate(case["steps"]):
        if cancelled():
            results["status"] = "stopped"
            break
        root = _source(task, step.get("root", "backend"))
        generated = step.get("testCode")
        runner = step.get("runner")
        prefix = f"{UUID(case['id'])}-{index + 1}"
        junit = folder / (prefix + ".xml")
        node_result = folder / (prefix + "-node.jsonl")
        # Run folders are normally unique; retries must never consume old evidence.
        for previous in (junit, node_result):
            previous.unlink(missing_ok=True)
        if runner == "pytest":
            interpreter = root / ".venv" / ("Scripts/python.exe" if os.name == "nt" else "bin/python")
            if not interpreter.is_file():
                if getattr(sys, "frozen", False):
                    raise MissingInput("Python 单元测试需要目标项目的 .venv 环境，请先安装 Python、项目依赖和 pytest。")
                interpreter = Path(sys.executable)
            if generated:
                if not isinstance(generated, str) or len(generated) > 30000:
                    raise MissingInput("生成的单元测试代码超过允许长度。")
                tree = ast.parse(generated)
                aliases = set()
                reference = str(case.get("sourceReference", "")).replace("\\", "/")
                for node in ast.walk(tree):
                    imports = [(node.module, [name.asname or name.name for name in node.names])] if isinstance(node, ast.ImportFrom) and node.module and not node.level else [(name.name, [name.asname or name.name.split('.')[0]]) for name in node.names] if isinstance(node, ast.Import) else []
                    for module, names in imports:
                        module_path = module.replace(".", "/")
                        for relative in (module_path + ".py", module_path + "/__init__.py"):
                            if relative in reference:
                                try:
                                    _source(task, step.get("root", "backend"), relative)
                                except MissingInput:
                                    continue
                                aliases.update(names)
                def called_import(call, aliases=aliases):
                    target = call.func
                    while isinstance(target, ast.Attribute):
                        target = target.value
                    return isinstance(target, ast.Name) and target.id in aliases
                test_functions = [node for node in ast.walk(tree) if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)) and node.name.startswith("test_")]
                def has_assertion(function):
                    return any((isinstance(node, ast.Assert) and any(isinstance(expression, ast.Name) for expression in ast.walk(node.test)))
                               or (isinstance(node, ast.Call) and isinstance(node.func, ast.Attribute) and node.func.attr == "raises")
                               for node in ast.walk(function))
                if (not aliases or not test_functions or any(not has_assertion(function) or not any(
                    isinstance(node, ast.Call) and called_import(node) for node in ast.walk(function)
                ) for function in test_functions)):
                    results.update(status="blocked", error="生成的测试必须在每个 test_ 函数内调用源码依据，并包含独立断言。")
                    progress(results)
                    return results
                test_file = folder / ("test_" + prefix.replace("-", "_") + ".py")
                test_file.write_text(generated, encoding="utf-8")
            else:
                test_file = _source(task, step.get("root", "backend"), str(step.get("testFile", "")))
                if test_file.suffix != ".py":
                    raise MissingInput("pytest 只执行 Python 测试文件。")
            selector = str(step.get("selector") or "")
            if selector and not re.fullmatch(r"[A-Za-z_][\w]*(?:::[A-Za-z_][\w]*)*(?:\[[\w .-]+\])?", selector):
                raise MissingInput("单元测试选择器不正确。")
            target = str(test_file) + ("::" + selector if selector else "")
            command = [str(interpreter), "-m", "pytest", target, "-q", "--tb=short", f"--junitxml={junit}", "-p", "no:cacheprovider", "-c", "nul" if os.name == "nt" else "/dev/null"]
        elif runner == "node" and not generated:
            node = shutil.which("node")
            if not node:
                raise MissingInput("服务器未安装 Node.js。")
            if _node_major(node) < MIN_NODE_MAJOR:
                raise MissingInput(f"Node 内置测试需要 Node.js {MIN_NODE_MAJOR} 及以上版本，以取得逐文件测试汇总。")
            test_file = _source(task, step.get("root", "frontend"), str(step.get("testFile", "")))
            if test_file.suffix not in {".js", ".mjs", ".cjs"}:
                raise MissingInput("Node 内置测试只执行 JS/MJS/CJS 测试文件。")
            reporter = folder / (prefix + "-node-reporter.cjs")
            reporter.write_text(
                "module.exports = async function* (events) {\n"
                "  for await (const event of events) {\n"
                "    if (event.type === 'test:summary') yield JSON.stringify(event.data) + '\\n';\n"
                "  }\n};\n", encoding="utf-8")
            command = [node, "--test", "--test-reporter=spec", "--test-reporter-destination=stdout",
                       "--test-reporter=" + reporter.resolve().as_uri(), "--test-reporter-destination=" + str(node_result)]
            if step.get("selector"):
                command.append("--test-name-pattern=^" + re.escape(str(step["selector"])) + "$")
            command.append(str(test_file))
        else:
            raise MissingInput("当前支持 Python pytest 与 Node 内置测试；该项目测试框架尚未接入。")
        environment = {key: value for key, value in os.environ.items() if not any(word in key.upper() for word in ("KEY", "TOKEN", "PASSWORD", "SECRET"))}
        environment.update(PYTHONPATH=str(root), PYTHONDONTWRITEBYTECODE="1", PYTEST_DISABLE_PLUGIN_AUTOLOAD="1")
        log_file = folder / (prefix + ".log")
        started = time.monotonic()
        # ponytail: one subprocess per case, isolated containers when testing untrusted source.
        with log_file.open("wb") as output:
            process = subprocess.Popen(command, cwd=folder if generated else root, env=environment, stdout=output, stderr=subprocess.STDOUT, shell=False,
                                       creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0)
            status = "running"
            while process.poll() is None:
                if cancelled() or time.monotonic() - started > 120:
                    status = "stopped" if cancelled() else "blocked"
                    if os.name == "nt":
                        subprocess.run(["taskkill", "/PID", str(process.pid), "/T", "/F"], capture_output=True, timeout=10, creationflags=subprocess.CREATE_NO_WINDOW)
                    else:
                        process.kill()
                    process.wait(timeout=10)
                    break
                time.sleep(.1)
        log = _redact(task, log_file.read_text(encoding="utf-8", errors="replace"))
        log_file.write_text(log, encoding="utf-8")
        count = None
        if junit.exists():
            suites = ElementTree.parse(junit).getroot()
            count = sum(int(suite.get("tests", 0)) for suite in suites.iter("testsuite"))
            failures = sum(int(suite.get("failures", 0)) + int(suite.get("errors", 0)) for suite in suites.iter("testsuite"))
            skipped = sum(int(suite.get("skipped", 0)) for suite in suites.iter("testsuite"))
            if status == "running":
                status = "passed" if process.returncode == 0 and count > 0 and not skipped and not failures else "failed" if failures or process.returncode == 1 else "blocked"
        elif status == "running":
            summaries = [json.loads(line) for line in node_result.read_text(encoding="utf-8").splitlines() if line.strip()] if node_result.exists() else []
            # The parent summary treats an empty JS file as a passing file test.
            # Only node:test's file-scoped summary proves actual declared tests.
            actual_tests = [summary for summary in summaries if summary.get("file") and Path(summary["file"]).resolve() == test_file.resolve()]
            count = sum(summary.get("counts", {}).get("tests", 0) for summary in actual_tests)
            complete = actual_tests and all(summary.get("success") and summary.get("counts", {}).get("passed") == summary.get("counts", {}).get("tests")
                and not any(summary.get("counts", {}).get(key, 0) for key in ("failed", "cancelled", "skipped", "todo")) for summary in actual_tests)
            status = "passed" if runner == "node" and process.returncode == 0 and count > 0 and complete else "failed" if process.returncode == 1 else "blocked"
        actual = f"退出码 {process.returncode}；执行测试 {count if count is not None else '未取得'}；耗时 {time.monotonic() - started:.2f} 秒"
        observation = {"action": "unit", "description": step.get("description") or f"{runner} {step.get('testFile') or 'AI 生成测试'}", "status": status, "actual": actual, "log": log[-16000:]}
        results["steps"].append(observation)
        results["log"] = log[-16000:]
        if status != "passed":
            results.update(status=status, error="单元测试未通过；请查看实际执行日志。")
            progress(results)
            break
        progress(results)
    if results["status"] == "running":
        results["status"] = "passed"
    return results


async def _concurrency_async(task: dict, step: dict, cancelled, progress=None) -> dict:
    url = _target(task, str(step.get("url", "")))
    method = str(step.get("method", "GET")).upper()
    if method not in {"GET", "HEAD"}:
        raise MissingInput("并发测试当前只执行 GET/HEAD；写入接口需明确的专用测试数据方案。")
    count, concurrency = int(task.get("requests", 20)), int(task.get("concurrency", 5))
    expected = int(step.get("expectedStatus", 200))
    threshold = float(step.get("maxLatencyMs", 2000))
    if not 1 <= count <= 200 or not 1 <= concurrency <= 20 or not 100 <= expected <= 599 or not 0 < threshold <= 30000:
        raise MissingInput("并发参数或阈值不正确。")
    semaphore = asyncio.Semaphore(concurrency)
    samples: list[dict] = []
    started, last_progress = time.monotonic(), 0.0

    def observation(status):
        latencies = sorted(item["latencyMs"] for item in samples)
        failures = sum(not sample["passed"] for sample in samples)
        metrics = {"requested": count, "completed": len(samples), "concurrency": concurrency, "failed": failures,
                   "elapsedMs": round((time.monotonic() - started) * 1000, 2), "expectedStatus": expected, "maxLatencyMs": threshold,
                   "averageMs": round(sum(latencies) / len(latencies), 2) if latencies else None,
                   "p95Ms": latencies[max(0, int(len(latencies) * .95 + .999) - 1)] if latencies else None,
                   "samples": sorted(samples, key=lambda item: item["request"])}
        return {"status": status, "metrics": metrics, "steps": [{"action": "concurrency", "description": f"{concurrency} 并发 / {count} 请求 {method} {url}", "status": status,
                 "actual": f"完成 {len(samples)}，失败 {failures}，P95 {metrics['p95Ms']} ms"}]}

    async with httpx.AsyncClient(timeout=10, follow_redirects=False, trust_env=False) as client:
        async def request(index):
            nonlocal last_progress
            async with semaphore:
                if cancelled():
                    return
                started = time.monotonic()
                try:
                    response = await client.request(method, url)
                    sample = {"request": index + 1, "statusCode": response.status_code, "latencyMs": round((time.monotonic() - started) * 1000, 2)}
                    sample["passed"] = response.status_code == expected and sample["latencyMs"] < threshold
                except httpx.HTTPError as error:
                    sample = {"request": index + 1, "passed": False, "error": _redact(task, str(error))[:1000], "latencyMs": round((time.monotonic() - started) * 1000, 2)}
                samples.append(sample)
                if progress and time.monotonic() - last_progress >= .25:
                    last_progress = time.monotonic()
                    progress(observation("running"))
        batch = asyncio.gather(*(request(index) for index in range(count)))
        try:
            while not batch.done():
                await asyncio.wait([batch], timeout=.2)
                if cancelled():
                    batch.cancel()
                    break
            await asyncio.gather(batch, return_exceptions=True)
            if not cancelled():
                batch.result()  # Propagate a real executor/storage exception.
        finally:
            if not batch.done():
                batch.cancel()
                await asyncio.gather(batch, return_exceptions=True)
    status = "stopped" if cancelled() else "passed" if len(samples) == count and all(sample["passed"] for sample in samples) else "failed"
    return observation(status)


def initial_run_result(cases: list[dict]) -> dict:
    return {"status": "running", "startedAt": _now(), "cases": [
        {"caseId": case["id"], "title": case["title"], "kind": case["kind"], "expected": case.get("expected", ""),
         "planReference": case.get("planReference", ""), "sourceReference": case.get("sourceReference", ""),
         "status": "queued", "steps": []} for case in cases
    ]}


def execute_run(task: dict, cases: list[dict], run_id: str, cancelled_callback) -> dict:
    from app.test_management import save_run_progress
    folder = run_directory(run_id)
    result = initial_run_result(cases)
    save_run_progress(run_id, result)
    for case, current in zip(cases, result["cases"]):
        def progress(observed, current=current):
            current.update(deepcopy(observed))
            save_run_progress(run_id, result)
        try:
            if cancelled_callback():
                current.update(status="stopped", error="已停止，未执行该用例。")
            elif case.get("needsInput"):
                current.update(status="blocked", error=case["needsInput"])
            else:
                current["status"] = "running"
                save_run_progress(run_id, result)
                if case["kind"] == "functional":
                    if task.get("executionMode", "steps") == "agent":
                        from app.test_browser_agent import functional_agent
                        current.update(functional_agent(task, case, folder, cancelled_callback, progress))
                    else:
                        current.update(_functional(task, case, folder, cancelled_callback, progress))
                elif case["kind"] == "unit":
                    current.update(_unit(task, case, folder, cancelled_callback, progress))
                elif case["kind"] == "concurrency":
                    if len(case.get("steps", [])) != 1:
                        raise MissingInput("并发用例需要一个明确的请求目标。")
                    current.update(asyncio.run(_concurrency_async(task, case["steps"][0], cancelled_callback, progress)))
                else:
                    raise MissingInput("测试类型尚未接入。")
        except MissingInput as error:
            current.update(status="blocked", error=str(error))
        except Exception as error:
            current.update(status="failed", error=_redact(task, str(error))[:4000])
        save_run_progress(run_id, result)
    for case in result["cases"]:
        if case.get("status") not in {"passed", "failed", "blocked", "stopped"}:
            case.update(status="failed", error="执行器没有返回完整测试结论。")
    statuses = {case["status"] for case in result["cases"]}
    result["status"] = "stopped" if cancelled_callback() or "stopped" in statuses else "failed" if "failed" in statuses else "blocked" if "blocked" in statuses or not statuses else "completed"
    result["endedAt"] = _now()
    (folder / "result.json").write_text(_redact(task, json.dumps(result, ensure_ascii=False, indent=2)), encoding="utf-8")
    return result


_STATUS = {"passed": "通过", "failed": "失败", "blocked": "未执行（缺少条件）", "stopped": "已停止", "completed": "全部通过", "running": "执行中", "queued": "未执行"}


def _case_paragraphs(task: dict, case: dict) -> list[str]:
    values = ["类型：" + {"functional": "功能测试", "unit": "单元测试", "concurrency": "并发测试"}.get(case["kind"], case["kind"]),
              "计划依据：" + (case.get("planReference") or "未提供"),
              "源码依据：" + (case.get("sourceReference") or "未提供"),
              "预期结果：" + (case.get("expected") or "未提供"),
              "实际结果：" + _STATUS.get(case["status"], case["status"])]
    if case.get("error"):
        values.append("问题：" + str(case["error"]))
    return [_redact(task, value) for value in values]


def _step_paragraphs(task: dict, step: dict, index: int) -> tuple[str, str]:
    return (_redact(task, f"步骤 {index + 1}：{step.get('description') or step.get('action')}"),
            _STATUS.get(step["status"], step["status"]) + "；" + _redact(task, str(step.get("actual", ""))))


def _metric_rows(case: dict) -> list[tuple[str, str]]:
    metrics = case.get("metrics")
    return [(label, str(metrics.get(key))) for label, key in (
        ("并发数", "concurrency"), ("计划请求", "requested"), ("完成请求", "completed"), ("失败请求", "failed"),
        ("平均耗时（ms）", "averageMs"), ("P95（ms）", "p95Ms"), ("最大允许耗时（ms）", "maxLatencyMs"))] if metrics else []


def _report_image(folder: Path, case: dict, step: dict) -> Path | None:
    name = step.get("screenshot")
    path = folder / name if isinstance(name, str) and re.fullmatch(r"[A-Za-z0-9_-]+\.(?:png|jpg|jpeg)", name) else None
    if path is not None and (path.is_symlink() or not path.is_file() or not path.resolve().is_relative_to(folder.resolve())):
        path = None
    if case.get("kind") == "functional" and case.get("status") == "passed" and path is None:
        raise MissingInput("通过用例的完整截图证据缺失，无法生成完整测试报告。")
    return path


def result_markdown(task: dict, run: dict, cases: list[dict], *, heading_level: int = 2) -> tuple[str, list[dict]]:
    """Render stored observations into the existing Word-editing path, without AI."""
    folder = run_directory(run["id"])
    blocks, images = [], []

    def paragraph(value):
        blocks.append("<p>" + escape(_redact(task, str(value))).replace("\n", "<br />") + "</p>")

    level = min(6, heading_level)
    blocks.append("#" * level + " " + _markdown_text(_redact(task, task["name"]).replace("\n", " ") + " 实测记录"))
    paragraph(f"执行时间：{run.get('startedAt', '')} 至 {run.get('endedAt') or '尚未结束'}")
    paragraph("本次执行结果：" + _STATUS.get(run["status"], run["status"]))
    paragraph("测试数据范围：" + (task.get("dataScope") or "未配置"))
    for case in cases:
        if case.get("kind") == "functional" and case.get("status") == "passed" and not case.get("steps"):
            raise MissingInput("通过用例的完整截图证据缺失，无法生成完整测试报告。")
        child_level = min(6, level + 1)
        blocks.append("#" * child_level + " " + _markdown_text(_redact(task, str(case["title"])).replace("\n", " ")))
        for value in _case_paragraphs(task, case):
            paragraph(value)
        for index, step in enumerate(case.get("steps", [])):
            for value in _step_paragraphs(task, step, index):
                paragraph(value)
            picture = _report_image(folder, case, step)
            if picture is None:
                if step.get("screenshot"):
                    paragraph("该步骤截图缺失。")
                continue
            data = picture.read_bytes()
            if not data.startswith(b"\x89PNG\r\n\x1a\n") or len(data) > 20 * 1024 * 1024:
                raise MissingInput("测试截图不是有效的 PNG 图片。")
            ref = f"docx-image:test-{run['id']}-{case['caseId']}-{index}"
            rect = step.get("focusRect") if step.get("action") in _FOCUS_ACTIONS else None
            if not isinstance(rect, dict) or any(isinstance(rect.get(key), bool) or not isinstance(rect.get(key), (int, float))
                or not math.isfinite(rect[key]) or abs(rect[key]) > 1e9 for key in ("x", "y", "width", "height")):
                rect = None
            images.append({"ref": ref, "alt": f"步骤 {index + 1} 操作截图", "data": base64.b64encode(data).decode("ascii"),
                           "content_type": "image/png", "focusRect": rect})
            blocks.append(f"![步骤 {index + 1} 操作截图]({ref})")
        rows = _metric_rows(case)
        if rows:
            blocks.append("<table><tr><th>并发指标</th><th>实测值</th></tr>" + "".join(
                "<tr><td>" + escape(label) + "</td><td>" + escape(value) + "</td></tr>" for label, value in rows) + "</table>")
        if case.get("log"):
            paragraph("执行日志")
            for line in _redact(task, str(case["log"])).splitlines():
                paragraph(line)
    return "\n\n".join(blocks), images


def build_report(task: dict, run: dict) -> Path:
    folder = run_directory(run["id"])
    original = None
    if task.get("reportFileId"):
        with connect() as connection:
            row = connection.execute("SELECT storage_key FROM files WHERE id = %s", (task["reportFileId"],)).fetchone()
        if row is None:
            raise MissingInput("待补充的原 Word 文件不存在。")
        original = resolve_object(get_settings().file_data_dir, str(row["storage_key"]))
    document = Document(str(original)) if original else Document()
    if original:
        document.add_page_break()
    else:
        section = document.sections[0]
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.left_margin = section.right_margin = Mm(25)
        for name, size in (("Normal", 12), ("Heading 1", 18), ("Heading 2", 14)):
            style = document.styles[name]
            style.font.name, style.font.size, style.font.color.rgb = "宋体", Pt(size), RGBColor(0, 0, 0)
            style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
            style.paragraph_format.space_after = Pt(6)
    heading = "Heading 1" if "Heading 1" in document.styles else "Normal"
    subheading = "Heading 2" if "Heading 2" in document.styles else heading
    table_style = document.tables[0].style if document.tables else document.styles["Table Grid"] if "Table Grid" in document.styles else None
    document.add_paragraph(_redact(task, task["name"]) + " 实测记录", heading)
    document.add_paragraph(f"执行时间：{run.get('startedAt', '')} 至 {run.get('endedAt') or '尚未结束'}")
    document.add_paragraph("结果：" + _STATUS.get(run["status"], run["status"]))
    document.add_paragraph(_redact(task, f"前端地址：{task.get('frontendUrl') or '未配置'}\n后端地址：{task.get('backendUrl') or '未配置'}"))
    document.add_paragraph("测试数据范围：" + _redact(task, task.get("dataScope") or "未配置"))
    result = run.get("result") or {}
    results = result.get("cases", [])
    if not results:
        document.add_paragraph("未取得任何测试执行结果。")
    available_width = document.sections[-1].page_width - document.sections[-1].left_margin - document.sections[-1].right_margin
    for case in results:
        if case.get("status") == "passed" and case.get("kind") == "functional" and (
            not case.get("steps") or any(not step.get("screenshot") for step in case["steps"])
        ):
            raise MissingInput("通过用例的完整截图证据缺失，无法生成完整测试报告。")
        document.add_paragraph(_redact(task, case["title"]), subheading)
        for value in _case_paragraphs(task, case):
            document.add_paragraph(value)
        for index, step in enumerate(case.get("steps", [])):
            caption_text, actual_text = _step_paragraphs(task, step, index)
            caption = document.add_paragraph(caption_text)
            actual = document.add_paragraph(actual_text)
            if step.get("screenshot"):
                caption.paragraph_format.keep_with_next = actual.paragraph_format.keep_with_next = True
                full_image = _report_image(folder, case, step)
                # Old detail captures can show a different navigation frame.
                # A partial image must never replace the full operation evidence.
                image = full_image
                if image is None:
                    document.add_paragraph("该步骤截图缺失。")
                else:
                    shape = document.add_picture(str(image), width=available_width)
                    if step.get("action") in _FOCUS_ACTIONS:
                        _crop_report_picture(document, shape, step.get("focusRect"))
        metrics = case.get("metrics")
        if metrics:
            table = document.add_table(rows=1, cols=2)
            if table_style:
                table.style = table_style
            table.rows[0].cells[0].text, table.rows[0].cells[1].text = "并发指标", "实测值"
            for label, value in _metric_rows(case):
                cells = table.add_row().cells
                cells[0].text, cells[1].text = label, value
        if case.get("log"):
            document.add_paragraph("执行日志")
            for line in _redact(task, str(case["log"])).splitlines():
                document.add_paragraph(line)
    output = folder / "test-report.docx"
    temporary = folder / "test-report.tmp.docx"
    document.save(str(temporary))
    temporary.replace(output)
    return output
