"""内置文档模板与文档依赖链（30 §3.1、§3.6）。

内置目录（代码）给出阶段、章节与公共框架；模板上游存 template_document_upstreams（可编辑），
项目内文档上游存 document_upstreams（Q-1：模板预填，负责人可增删）。
"""

from __future__ import annotations

from io import BytesIO
import json
import re
from uuid import UUID, uuid4

from fastapi import APIRouter, Request

from app.db import connect
from app.document_access import document_principal, is_admin
from app.errors import AppError
from app.identity import require_permission
from app.markdown_blocks import heading_text, split

router = APIRouter()

BUILTIN_GROUP = "内置软件文档模板"
BUILTIN_CATEGORY = "软件研制"

STAGES = [
    {"key": "initiate", "name": "立项"},
    {"key": "requirement", "name": "需求"},
    {"key": "design", "name": "设计与计划"},
    {"key": "delivery", "name": "交付"},
]

# 公共框架章节：多份文档共用、单独分配负责人（30 §3.4）。key 只在数据里使用。
FRAMEWORKS = [
    {"key": "runtime", "name": "软件运行环境"},
    {"key": "glossary", "name": "术语和缩略语"},
    {"key": "references", "name": "引用文档"},
]

# 章节写成 (级别, 标题, 公共框架 key 或 None)
CATALOG = [
    {"key": "task-book", "name": "任务计划书", "stage": "initiate", "upstream": [],
     "description": "说明任务来源、目标、范围和进度要求。",
     "outline": [(1, "1 范围", None), (2, "1.1 标识", None), (2, "1.2 系统概述", None), (1, "2 引用文档", "references"),
                 (1, "3 任务目标", None), (1, "4 任务范围", None), (1, "5 进度要求", None), (1, "6 术语和缩略语", "glossary")]},
    {"key": "srs", "name": "软件需求规格说明", "stage": "requirement", "upstream": ["task-book"],
     "description": "逐项描述功能、性能、接口与质量需求。",
     "outline": [(1, "1 范围", None), (2, "1.1 标识", None), (1, "2 引用文档", "references"), (1, "3 需求", None),
                 (2, "3.1 运行环境", "runtime"), (2, "3.2 功能需求", None), (2, "3.3 性能需求", None), (2, "3.4 接口需求", None),
                 (2, "3.5 质量需求", None), (1, "4 术语和缩略语", "glossary")]},
    {"key": "dev-plan", "name": "软件开发计划", "stage": "design", "upstream": ["srs"],
     "description": "规定开发活动、任务分解、人员分工与进度。",
     "outline": [(1, "1 范围", None), (1, "2 引用文档", "references"), (1, "3 开发环境", "runtime"), (1, "4 任务分解", None),
                 (1, "5 公共框架分工", None), (1, "6 进度安排", None), (1, "7 术语和缩略语", "glossary")]},
    {"key": "design", "name": "软件概要设计说明", "stage": "design", "upstream": ["srs"],
     "description": "描述软件体系结构、部件与接口设计。",
     "outline": [(1, "1 范围", None), (1, "2 引用文档", "references"), (1, "3 运行环境", "runtime"), (1, "4 体系结构设计", None),
                 (1, "5 接口设计", None), (1, "6 术语和缩略语", "glossary")]},
    {"key": "test-plan", "name": "软件测试计划", "stage": "delivery", "upstream": ["dev-plan", "srs"],
     "description": "规定测试范围、测试环境、测试项与进度。",
     "outline": [(1, "1 范围", None), (1, "2 引用文档", "references"), (1, "3 测试环境", "runtime"), (1, "4 测试项", None),
                 (1, "5 进度安排", None), (1, "6 术语和缩略语", "glossary")]},
    {"key": "user-manual", "name": "软件用户手册", "stage": "delivery", "upstream": ["design"],
     "description": "指导用户安装、配置和使用软件。",
     "outline": [(1, "1 范围", None), (1, "2 引用文档", "references"), (1, "3 软件运行环境", "runtime"), (1, "4 安装与配置", None),
                 (1, "5 使用说明", None), (1, "6 术语和缩略语", "glossary")]},
]

PLAN_KEY = "dev-plan"
_BY_KEY = {item["key"]: item for item in CATALOG}
_BY_NAME = {item["name"]: item for item in CATALOG}
_FRAMEWORK_NAMES = {item["key"]: item["name"] for item in FRAMEWORKS}


def catalog_payload() -> dict:
    return {
        "stages": STAGES,
        "frameworks": FRAMEWORKS,
        "documents": [{
            "key": item["key"], "name": item["name"], "stage": item["stage"], "description": item["description"],
            "upstream": list(item["upstream"]), "chapterCount": len(item["outline"]),
            "outline": [{"level": level, "title": title, "framework": fw} for level, title, fw in item["outline"]],
        } for item in CATALOG],
    }


def match_catalog(name: str, template_name: str | None = None) -> dict | None:
    for candidate in (template_name, name):
        if candidate and candidate in _BY_NAME:
            return _BY_NAME[candidate]
    for item in CATALOG:  # 宽松匹配：项目里常加前缀，如“XX系统软件需求规格说明”
        if name and (item["name"] in name or item["name"].removeprefix("软件") in name):
            return item
    return None


def build_docx(item: dict) -> bytes:
    from docx import Document

    word = Document()
    for level, title, fw in item["outline"]:
        word.add_heading(title, level=level)
        hint = f"（公共框架：{_FRAMEWORK_NAMES[fw]}，各文档引用同一内容。）" if fw else "（在此编写本章内容。）"
        word.add_paragraph(hint)
    buffer = BytesIO()
    word.save(buffer)
    return buffer.getvalue()


def is_prompt_or_empty(block) -> bool:
    """占位提示语（“（在此编写本章内容。）”、来源注记等以“（”开头的段落）或空段：不算正文。全文统一用这一口径。"""
    text = (block.body or "").strip()
    return not text or text.startswith("（")


def chapter_progress(markdown: str) -> tuple[int, int, list[dict]]:
    """按标题划分章节，章节下有非提示正文即视为已写。返回 (已写, 总数, 章节列表)。"""
    _, blocks = split(markdown or "")
    chapters: list[dict] = []
    for block in blocks:
        if block.kind == "heading":
            chapters.append({"title": heading_text(block.body), "level": block.heading_level, "written": False})
        elif chapters and not is_prompt_or_empty(block):
            chapters[-1]["written"] = True
    return sum(1 for c in chapters if c["written"]), len(chapters), chapters


# ---------------------------------------------------------------- 接口

@router.get("/doc-chain/catalog")
def get_catalog(request: Request) -> dict:
    document_principal(request)
    return {"data": catalog_payload(), "requestId": request.state.request_id}


@router.post("/templates/builtin:seed")
def seed_builtin(request: Request) -> dict:
    """创建（或补齐）内置模板组；已存在的同名模板文档不覆盖。"""
    from app.template_versions import ensure_v1, find_or_store
    from app.template_import import parse_template, validate_structure
    from app.docx_parser import parse_docx
    from app.config import get_settings
    from app.files import resolve_object

    principal = require_permission(request, "templates.manage")
    created = []
    files = {}
    with connect() as conn:  # 文件先落盘（store_bytes 自带连接），再进写事务
        for item in CATALOG:
            files[item["key"]] = find_or_store(conn, build_docx(item), f"{item['name']}.docx")[0]
    with connect() as conn:
        conn.execute("BEGIN IMMEDIATE")
        row = conn.execute("SELECT id FROM templates WHERE name=%s", [BUILTIN_GROUP]).fetchone()
        if row is None:
            tid = uuid4()
            conn.execute("INSERT INTO templates(id,name,category,description,status) VALUES (%s,%s,%s,%s,'published')",
                         [tid, BUILTIN_GROUP, BUILTIN_CATEGORY, "系统内置的软件研制文档，按依赖链排列。"])
        else:
            tid = row["id"]
        existing = {r["name"] for r in conn.execute("SELECT name FROM template_documents WHERE template_id=%s", [tid]).fetchall()}
        for order, item in enumerate(CATALOG):
            if item["name"] in existing:
                continue
            file_id = files[item["key"]]
            storage = conn.execute("SELECT storage_key FROM files WHERE id=%s", [file_id]).fetchone()
            path = resolve_object(get_settings().file_data_dir, storage["storage_key"])
            parsed = parse_docx(str(path), outline_headings=True)
            structure = parse_template(path)  # 内置模板结构由我们生成，直接视为已核对
            structure = validate_structure(structure, {**structure, "reviewed": True})
            fmt = {**parsed.format_template, "importStructure": structure}
            td = uuid4()
            conn.execute("""INSERT INTO template_documents(id,template_id,name,description,display_order,word_file_id,format_template)
                            VALUES (%s,%s,%s,%s,%s,%s,%s)""",
                         [td, tid, item["name"], item["description"], order, file_id, json.dumps(fmt)])
            ensure_v1(conn, td, principal.user_id)
            created.append(item["name"])
        link_builtin_upstreams(conn, tid)
        conn.commit()
    return {"data": {"templateId": str(tid), "created": created}, "requestId": request.state.request_id}


def _project_documents(conn, request: Request, project_id: str):
    principal = document_principal(request)
    try:
        pid = UUID(project_id)
    except ValueError:
        raise AppError(404, "NOT_FOUND", "项目不存在或你无权访问。") from None
    from app.projects import _load, project_role
    admin = is_admin(principal) or "auditor" in principal.roles
    _load(conn, pid, principal.user_id, admin)
    role = "owner" if is_admin(principal) else project_role(conn, pid, principal.user_id)
    rows = conn.execute("""SELECT d.id,d.name,d.status,d.markdown_content,d.updated_at,td.name AS template_name,td.id AS td_id,
                                  u.name AS owner_name
                           FROM documents d LEFT JOIN template_documents td ON td.id=d.template_document_id
                           LEFT JOIN users u ON u.id=d.owner_id
                           WHERE d.project_id=%s ORDER BY d.created_at""", [pid]).fetchall()
    return principal, pid, role, rows


def _builtin_template_docs(conn) -> dict[str, str]:
    rows = conn.execute("""SELECT td.id,td.name FROM template_documents td JOIN templates t ON t.id=td.template_id
                           WHERE t.name=%s""", [BUILTIN_GROUP]).fetchall()
    return {r["name"]: str(r["id"]) for r in rows}


@router.get("/projects/{project_id}/doc-chain")
def get_project_chain(request: Request, project_id: str) -> dict:
    with connect() as conn:
        _, _, role, rows = _project_documents(conn, request, project_id)
        templates = _builtin_template_docs(conn)
    nodes, extras, by_key = [], [], {}
    for row in rows:
        item = match_catalog(row["name"], row["template_name"])
        written, total, _ = chapter_progress(row["markdown_content"])
        node = {"documentId": str(row["id"]), "name": row["name"], "status": row["status"], "created": True,
                "written": written, "total": total, "ownerName": row["owner_name"],
                "updatedAt": str(row["updated_at"])}
        if item and item["key"] not in by_key:
            by_key[item["key"]] = node
        else:
            extras.append({**node, "key": None, "stage": None, "upstream": []})
    with connect() as conn:
        links = {n["documentId"]: document_upstreams(conn, n["documentId"]) for n in [*by_key.values(), *extras]}
    key_of = {n["documentId"]: k for k, n in by_key.items()}
    with connect() as conn:
        status = {n["documentId"]: upstream_baseline_status(conn, n["documentId"]) for n in [*by_key.values(), *extras]}
        latest_pb = {n["documentId"]: (baseline_entries(conn, n["documentId"]) or [None])[-1] for n in [*by_key.values(), *extras]}
    for node in [*by_key.values(), *extras]:
        node["upstreamDocuments"] = links.get(node["documentId"], [])
        node["upstreamBaselines"] = status.get(node["documentId"], [])
        node["upstreamChanged"] = any(u["newBaseline"] for u in node["upstreamBaselines"])
        entry = latest_pb.get(node["documentId"])
        node["baselineSeq"] = entry["pbSeq"] if entry else None
    for item in CATALOG:
        node = by_key.get(item["key"]) or {"documentId": None, "name": item["name"], "status": "missing", "created": False,
                                          "written": 0, "total": len(item["outline"]), "ownerName": None, "updatedAt": None,
                                          "baselineSeq": None, "upstreamBaselines": [], "upstreamChanged": False}
        linked = [key_of[u["documentId"]] for u in node.get("upstreamDocuments", []) if u["documentId"] in key_of]
        nodes.append({**node, "key": item["key"], "stage": item["stage"],
                      "upstream": linked if node.get("upstreamDocuments") else list(item["upstream"]),
                      "templateDocumentId": templates.get(item["name"])})
    return {"data": {"stages": STAGES, "nodes": nodes, "others": extras,
                     "canCreate": role == "owner"}, "requestId": request.state.request_id}


# ---------------------------------------------------------------- 模板上游（可编辑）

from pydantic import BaseModel, ConfigDict, Field  # noqa: E402

class UpstreamBody(BaseModel):
    model_config = ConfigDict(extra="forbid")
    upstreamIds: list[str] = Field(default_factory=list, max_length=20)


def _uid(value: str, message: str = "对象不存在。") -> str:
    try:
        return str(UUID(value))
    except ValueError:
        raise AppError(404, "NOT_FOUND", message) from None


def template_upstreams(conn, td_id: str) -> list[dict]:
    rows = conn.execute("""SELECT u.upstream_id, td.name FROM template_document_upstreams u
                           JOIN template_documents td ON td.id=u.upstream_id WHERE u.template_document_id=%s ORDER BY td.name""",
                        [td_id]).fetchall()
    return [{"id": str(r["upstream_id"]), "name": r["name"]} for r in rows]


def _would_cycle(conn, table: str, key: str, up: str, start: str, targets: set[str]) -> bool:
    """从 targets 沿上游走，能回到 start 即成环。"""
    seen, stack = set(), list(targets)
    while stack:
        node = stack.pop()
        if node == start:
            return True
        if node in seen:
            continue
        seen.add(node)
        stack.extend(str(r[0]) for r in conn.execute(f"SELECT {up} FROM {table} WHERE {key}=%s", [node]).fetchall())
    return False


@router.get("/template-documents/{document_id}/upstreams")
def get_template_upstreams(request: Request, document_id: str) -> dict:
    require_permission(request, "templates.manage")
    td = _uid(document_id, "模板文档不存在。")
    with connect() as conn:
        if conn.execute("SELECT 1 FROM template_documents WHERE id=%s", [td]).fetchone() is None:
            raise AppError(404, "NOT_FOUND", "模板文档不存在。")
        candidates = conn.execute("""SELECT td.id, td.name, t.name AS group_name FROM template_documents td
                                     JOIN templates t ON t.id=td.template_id WHERE td.id<>%s ORDER BY t.name, td.display_order""",
                                  [td]).fetchall()
        return {"data": {"upstreams": template_upstreams(conn, td),
                         "candidates": [{"id": str(r["id"]), "name": r["name"], "groupName": r["group_name"]} for r in candidates]},
                "requestId": request.state.request_id}


@router.put("/template-documents/{document_id}/upstreams")
def put_template_upstreams(request: Request, document_id: str, body: UpstreamBody) -> dict:
    require_permission(request, "templates.manage")
    td = _uid(document_id, "模板文档不存在。")
    ids = list(dict.fromkeys(_uid(v, "上游模板不存在。") for v in body.upstreamIds))
    with connect() as conn:
        conn.execute("BEGIN IMMEDIATE")
        if conn.execute("SELECT 1 FROM template_documents WHERE id=%s", [td]).fetchone() is None:
            raise AppError(404, "NOT_FOUND", "模板文档不存在。")
        for value in ids:
            if value == td or conn.execute("SELECT 1 FROM template_documents WHERE id=%s", [value]).fetchone() is None:
                raise AppError(422, "UPSTREAM_INVALID", "上游模板不存在或是它自己。", field_errors=[{"field": "upstreamIds", "message": "请重新选择上游模板。"}])
        if _would_cycle(conn, "template_document_upstreams", "template_document_id", "upstream_id", td, set(ids)):
            raise AppError(422, "UPSTREAM_CYCLE", "这样设置会形成循环依赖。", field_errors=[{"field": "upstreamIds", "message": "上游模板不能依赖本模板。"}])
        conn.execute("DELETE FROM template_document_upstreams WHERE template_document_id=%s", [td])
        for value in ids:
            conn.execute("INSERT INTO template_document_upstreams(template_document_id,upstream_id) VALUES (%s,%s)", [td, value])
        conn.commit()
        return {"data": {"upstreams": template_upstreams(conn, td)}, "requestId": request.state.request_id}


def link_builtin_upstreams(conn, template_id) -> None:
    """内置模板组：按目录写入模板上游（已有关系不覆盖）。"""
    ids = {r["name"]: str(r["id"]) for r in conn.execute("SELECT id,name FROM template_documents WHERE template_id=%s", [template_id]).fetchall()}
    for item in CATALOG:
        me = ids.get(item["name"])
        if not me or conn.execute("SELECT 1 FROM template_document_upstreams WHERE template_document_id=%s", [me]).fetchone():
            continue
        for key in item["upstream"]:
            up = ids.get(_BY_KEY[key]["name"])
            if up:
                conn.execute("INSERT OR IGNORE INTO template_document_upstreams(template_document_id,upstream_id) VALUES (%s,%s)", [me, up])


# ---------------------------------------------------------------- 文档上游：创建时复制初稿 + 链接

_NUMBER = re.compile(r"^\s*(?:\d+(?:\.\d+)*|[一二三四五六七八九十]+、)\s*")


def chapter_key(title: str) -> str:
    return _NUMBER.sub("", title).strip().removeprefix("软件")


def default_upstreams(conn, project_id, template_document_id) -> list[str]:
    """模板上游 → 本项目中由这些模板创建的文档；没有配置时按内置目录名称匹配。"""
    if not template_document_id:
        return []
    ups = [u["id"] for u in template_upstreams(conn, str(template_document_id))]
    found: list[str] = []
    if ups:
        for up in ups:
            row = conn.execute("SELECT id FROM documents WHERE project_id=%s AND template_document_id=%s ORDER BY created_at LIMIT 1",
                               [str(project_id), up]).fetchone()
            if row:
                found.append(str(row["id"]))
        return found
    row = conn.execute("SELECT name FROM template_documents WHERE id=%s", [str(template_document_id)]).fetchone()
    item = match_catalog("", row["name"]) if row else None
    if not item:
        return []
    for doc in conn.execute("""SELECT d.id,d.name,td.name AS tname FROM documents d LEFT JOIN template_documents td
                               ON td.id=d.template_document_id WHERE d.project_id=%s ORDER BY d.created_at""", [str(project_id)]).fetchall():
        match = match_catalog(doc["name"], doc["tname"])
        if match and match["key"] in item["upstream"] and str(doc["id"]) not in found:
            found.append(str(doc["id"]))
    return found


def _sections(markdown: str) -> list[tuple[str, int, list]]:
    _, blocks = split(markdown or "")
    out: list[tuple[str, int, list]] = []
    for block in blocks:
        if block.kind == "heading":
            out.append((heading_text(block.body), block.heading_level, []))
        elif out:
            out[-1][2].append(block)
    return out


_FRAMEWORK_TITLES = {chapter_key(t) for item in CATALOG for _l, t, fw in item["outline"] if fw}


def seed_from_upstreams(conn, project_id, upstream_ids: list[str], markdown: str) -> tuple[str, list[dict]]:
    """把上游同名章节的正文复制进新文档对应章节（只在本章没有正文时），并加来源注记。返回 (新正文, 复制记录)。"""
    from app.markdown_blocks import assemble

    uppers = []
    for uid in upstream_ids:
        row = conn.execute("SELECT id,name,revision,markdown_content FROM documents WHERE id=%s AND project_id=%s",
                           [uid, str(project_id)]).fetchone()
        if row is None:
            raise AppError(422, "UPSTREAM_INVALID", "上游文档不存在或不在本项目中。", field_errors=[{"field": "upstreamDocumentIds", "message": "请重新选择上游文档。"}])
        # 上游章节只取真正的正文：只有提示语的章节不当作初稿复制（40 P2-3）
        uppers.append((row, {chapter_key(t): content for t, _lvl, body in _sections(row["markdown_content"])
                             if (content := [b for b in body if not is_prompt_or_empty(b)])}))
    if not uppers:
        return markdown, []
    prefix, blocks = split(markdown or "")
    out: list = []
    copied: list[dict] = []
    index = 0
    while index < len(blocks):
        block = blocks[index]
        out.append(block)
        index += 1
        if block.kind != "heading":
            continue
        title = heading_text(block.body)
        body_end = index
        while body_end < len(blocks) and blocks[body_end].kind != "heading":
            body_end += 1
        own = [b for b in blocks[index:body_end] if not is_prompt_or_empty(b)]
        key = chapter_key(title)
        if own or key in _FRAMEWORK_TITLES:
            continue
        for row, sections in uppers:
            source = sections.get(key)
            if not source:
                continue
            gap = block.gap_after or "\n\n"
            note = type(block)("paragraph", f"（来源：《{row['name']}》{title.split(' ', 1)[-1]}，创建时复制的初稿，可修改，不随上游更新）", gap, None)
            copies = [type(b)(b.kind, b.body, b.gap_after or gap, b.heading_level) for b in source]
            out.extend([note, *copies])
            index = body_end  # 本章只有提示语：被初稿替换
            copied.append({"chapter": title, "upstreamDocumentId": str(row["id"]), "blocks": len(source)})
            break
    if not copied:
        return markdown, []
    return assemble(prefix, out), copied


def record_upstreams(conn, document_id, upstream_ids: list[str], copied: list[dict], actor) -> None:
    for uid in upstream_ids:
        rev = conn.execute("SELECT revision FROM documents WHERE id=%s", [uid]).fetchone()
        conn.execute("""INSERT OR IGNORE INTO document_upstreams(document_id,upstream_document_id,source,upstream_revision,copied,created_by)
                        VALUES (%s,%s,'template',%s,%s,%s)""",
                     [str(document_id), uid, rev["revision"] if rev else None,
                      json.dumps([c for c in copied if c["upstreamDocumentId"] == uid], ensure_ascii=False), str(actor)])


def document_upstreams(conn, document_id) -> list[dict]:
    """供依赖链、AI 写作上下文使用：上游文档、链接时版本、当前版本与复制过的章节。"""
    rows = conn.execute("""SELECT u.upstream_document_id, u.source, u.upstream_revision, u.copied, d.name, d.revision
                           FROM document_upstreams u JOIN documents d ON d.id=u.upstream_document_id
                           WHERE u.document_id=%s ORDER BY d.created_at""", [str(document_id)]).fetchall()
    return [{"documentId": str(r["upstream_document_id"]), "name": r["name"], "source": r["source"],
             "linkedRevision": r["upstream_revision"], "currentRevision": r["revision"],
             "changedSinceLink": r["upstream_revision"] is not None and r["revision"] != r["upstream_revision"],
             "copiedChapters": json.loads(r["copied"] or "[]")} for r in rows]


class DocumentUpstreamBody(BaseModel):
    model_config = ConfigDict(extra="forbid")
    upstreamDocumentIds: list[str] = Field(default_factory=list, max_length=20)


@router.get("/documents/{document_id}/upstreams")
def get_document_upstreams(request: Request, document_id: str) -> dict:
    from app.document_access import load_document
    principal = document_principal(request)
    did = UUID(_uid(document_id, "文档不存在或你无权访问。"))
    with connect() as conn:
        load_document(conn, did, principal.user_id, is_admin(principal))
        return {"data": {"upstreams": document_upstreams(conn, did)}, "requestId": request.state.request_id}


@router.put("/documents/{document_id}/upstreams")
def put_document_upstreams(request: Request, document_id: str, body: DocumentUpstreamBody) -> dict:
    """Q-1：项目负责人可增删上游关系。只改链接，不改正文。"""
    from app.document_access import ensure_project_writable, load_document
    principal = document_principal(request)
    did = UUID(_uid(document_id, "文档不存在或你无权访问。"))
    ids = list(dict.fromkeys(_uid(v, "上游文档不存在。") for v in body.upstreamDocumentIds))
    with connect() as conn:
        doc = load_document(conn, did, principal.user_id, is_admin(principal))
        if not (doc.admin or doc.project_role == "owner"):
            raise AppError(403, "FORBIDDEN", "只有项目负责人可以调整文档上游。")
        ensure_project_writable(conn, doc.project_id)
        conn.execute("BEGIN IMMEDIATE")
        for value in ids:
            if value == str(did) or conn.execute("SELECT 1 FROM documents WHERE id=%s AND project_id=%s", [value, str(doc.project_id)]).fetchone() is None:
                raise AppError(422, "UPSTREAM_INVALID", "上游文档不存在或不在本项目中。", field_errors=[{"field": "upstreamDocumentIds", "message": "请重新选择上游文档。"}])
        if _would_cycle(conn, "document_upstreams", "document_id", "upstream_document_id", str(did), set(ids)):
            raise AppError(422, "UPSTREAM_CYCLE", "这样设置会形成循环依赖。", field_errors=[{"field": "upstreamDocumentIds", "message": "上游文档不能依赖本文档。"}])
        conn.execute("DELETE FROM document_upstreams WHERE document_id=%s AND upstream_document_id NOT IN (%s)" % ("%s", ",".join(["%s"] * len(ids)) or "''"),
                     [str(did), *ids])
        for value in ids:
            rev = conn.execute("SELECT revision FROM documents WHERE id=%s", [value]).fetchone()["revision"]
            conn.execute("""INSERT OR IGNORE INTO document_upstreams(document_id,upstream_document_id,source,upstream_revision,created_by)
                            VALUES (%s,%s,'manual',%s,%s)""", [str(did), value, rev, str(principal.user_id)])
        conn.commit()
        return {"data": {"upstreams": document_upstreams(conn, did)}, "requestId": request.state.request_id}


# ---------------------------------------------------------------- 上游有新基线（按项目基线判断，30 §3.1、§3.5）

def baseline_entries(conn, document_id) -> list[dict]:
    """文档进入过的项目基线，按基线号升序：pbSeq、文档版本号 versionSeq、版本对应的正文 revision。"""
    rows = conn.execute("""SELECT pb.seq AS pb_seq, pb.created_at, v.seq AS v_seq, v.revision
                           FROM project_baseline_items i JOIN project_baselines pb ON pb.id=i.project_baseline_id
                           JOIN document_baselines v ON v.id=i.document_baseline_id
                           WHERE i.document_id=%s ORDER BY pb.seq""", [str(document_id)]).fetchall()
    return [{"pbSeq": r["pb_seq"], "versionSeq": r["v_seq"], "revision": r["revision"]} for r in rows]


def upstream_baseline_status(conn, document_id) -> list[dict]:
    """每个上游：最近基线号、下游核对所依据的基线号、是否有新基线、对比用的文档版本号。

    下游所依据的基线 = 已“标记为已核对”的基线号；否则取下游首次进入的项目基线；
    下游从未进基线时取链接建立时已有的最新基线（没有则为 0）。
    上游在此之后进入了更新的项目基线、且正文版本不同，即“有新基线”。基线不冻结，也不修改下游内容（Q-4、Q-5）。
    """
    did = str(document_id)
    links = conn.execute("""SELECT u.upstream_document_id, u.acked_baseline_seq, u.created_at, d.name, d.project_id
                            FROM document_upstreams u JOIN documents d ON d.id=u.upstream_document_id
                            WHERE u.document_id=%s ORDER BY d.created_at""", [did]).fetchall()
    own = baseline_entries(conn, did)
    out = []
    for link in links:
        uid = str(link["upstream_document_id"])
        entries = baseline_entries(conn, uid)
        if link["acked_baseline_seq"] is not None:
            ref = int(link["acked_baseline_seq"])
        elif own:
            ref = own[0]["pbSeq"]
        else:
            row = conn.execute("SELECT max(seq) AS n FROM project_baselines WHERE project_id=%s AND created_at<=%s",
                               [str(link["project_id"]), link["created_at"]]).fetchone()
            ref = int(row["n"] or 0)
        latest = entries[-1] if entries else None
        based = next((e for e in reversed(entries) if e["pbSeq"] <= ref), None)
        new = bool(latest and latest["pbSeq"] > ref and (based is None or based["revision"] != latest["revision"]))
        out.append({"documentId": uid, "name": link["name"],
                    "latestBaselineSeq": latest["pbSeq"] if latest else None,
                    "baseBaselineSeq": based["pbSeq"] if based else (ref or None),
                    "newBaseline": new, "acked": link["acked_baseline_seq"],
                    "fromVersionSeq": based["versionSeq"] if based else None,
                    "toVersionSeq": latest["versionSeq"] if latest else None})
    return out


@router.get("/documents/{document_id}/upstream-notices")
def get_upstream_notices(request: Request, document_id: str) -> dict:
    """工作台提示条（UpstreamNotice）数据：只列出有新基线的上游。"""
    from app.document_access import load_document
    principal = document_principal(request)
    did = UUID(_uid(document_id, "文档不存在或你无权访问。"))
    with connect() as conn:
        doc = load_document(conn, did, principal.user_id, is_admin(principal))
        can_ack = not doc.archived and (doc.admin or doc.project_role in ("owner", "editor"))
        items = [s for s in upstream_baseline_status(conn, did) if s["newBaseline"]]
    return {"data": {"notices": items, "canAck": can_ack, "projectId": str(doc.project_id)}, "requestId": request.state.request_id}


@router.post("/documents/{document_id}/upstreams/{upstream_id}:ack")
def ack_upstream(request: Request, document_id: str, upstream_id: str) -> dict:
    """标记为已核对：下游此后以上游当前最新基线为依据（不强制，Q-5）。"""
    from app.document_access import ensure_project_writable, load_document
    principal = document_principal(request)
    did = UUID(_uid(document_id, "文档不存在或你无权访问。"))
    uid = _uid(upstream_id, "上游文档不存在。")
    with connect() as conn:
        doc = load_document(conn, did, principal.user_id, is_admin(principal))
        if not (doc.admin or doc.project_role in ("owner", "editor")):
            raise AppError(403, "FORBIDDEN", "只读成员不能标记核对。")
        ensure_project_writable(conn, doc.project_id)
        conn.execute("BEGIN IMMEDIATE")
        if conn.execute("SELECT 1 FROM document_upstreams WHERE document_id=%s AND upstream_document_id=%s", [str(did), uid]).fetchone() is None:
            raise AppError(404, "NOT_FOUND", "这份文档没有这个上游。")
        entries = baseline_entries(conn, uid)
        seq = entries[-1]["pbSeq"] if entries else 0
        conn.execute("UPDATE document_upstreams SET acked_baseline_seq=%s WHERE document_id=%s AND upstream_document_id=%s", [seq, str(did), uid])
        conn.commit()
        status = next(s for s in upstream_baseline_status(conn, did) if s["documentId"] == uid)
    return {"data": status, "requestId": request.state.request_id}
