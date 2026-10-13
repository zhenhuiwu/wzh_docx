"""A06 模板版本快照与完整编辑（17 §3）：版本列表、草稿、发布、恢复、变基、上传新 Word、对比、样式影响。

要点：
- 每个模板文档最多 1 份草稿、0 或 1 份已发布版本（部分唯一索引兜底）；已发布版本不可变；
- 草稿保存只做轻量校验 + revision CAS，不生成 Word；发布时在事务外基于草稿的基底 Word 增量生成并往返核对，
  事务内只做 CAS 与状态切换（§3.10）；
- 未审核的模板文档有一份“初始草稿”（parent 为空），审核确认后变成第 1 版；之后的编辑草稿 parent 必为某个版本；
- 原始 Word（source）从不修改；生成的 Word 按 sha256 + 逐字节比对去重（find_or_store）。
"""

from __future__ import annotations

import json
import tempfile
import threading
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from uuid import UUID, uuid4

from docx import Document
from fastapi import APIRouter, Request, Response
from pydantic import BaseModel, ConfigDict, Field

from app import template_merge as tm
from app import template_word as tw
from app.config import get_settings
from app.db import connect, write_transaction
from app.errors import AppError
from app.files import content_disposition, read_object_bytes, resolve_object, store_bytes
from app.identity import require_permission
from app.models import format_timestamp
from app.template_import import parse_template

router = APIRouter()
WORD = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_V1_NOTE = "导入时的模板（第 1 版）"

_locks: dict[UUID, threading.Lock] = {}
_locks_guard = threading.Lock()


def _publish_lock(td_id: UUID) -> threading.Lock:
    with _locks_guard:
        return _locks.setdefault(td_id, threading.Lock())


# ---------------------------------------------------------------- 请求体

class DraftBody(BaseModel):
    model_config = ConfigDict(extra="forbid")
    structure: dict
    format: dict | None = None
    expectedDraftRevision: int = Field(ge=0)
    # 草稿已被别人发布或放弃时，编辑者可以带上自己开始编辑的版本重新建草稿（发布时得到 TEMPLATE_DRAFT_STALE，再合并）
    parentVersionId: str | None = None


class PublishBody(BaseModel):
    model_config = ConfigDict(extra="forbid")
    draftRevision: int = Field(ge=1)
    note: str


class RestoreBody(BaseModel):
    model_config = ConfigDict(extra="forbid")
    versionId: str
    expectedDraftRevision: int = Field(ge=0)


class UploadBody(BaseModel):
    model_config = ConfigDict(extra="forbid")
    wordFileId: str
    expectedDraftRevision: int = Field(ge=0)


class Resolution(BaseModel):
    model_config = ConfigDict(extra="forbid")
    blockId: str | None = None
    field: str
    take: str


class RebaseBody(BaseModel):
    model_config = ConfigDict(extra="forbid")
    expectedDraftRevision: int = Field(ge=1)
    dryRun: bool = False
    resolutions: list[Resolution] = Field(default_factory=list)


class ImpactBody(BaseModel):
    model_config = ConfigDict(extra="forbid")
    changes: list[dict] | None = None


# ---------------------------------------------------------------- 读取工具

def _uuid(value: str, field: str) -> UUID:
    try:
        return UUID(str(value))
    except ValueError:
        raise AppError(422, "VALIDATION_FAILED", "标识不正确。", [{"field": field, "message": "需要有效标识。"}]) from None


def _json(value):
    return json.loads(value) if isinstance(value, str) else value


def _file(conn, file_id) -> dict:
    row = conn.execute("SELECT id, name, sha256, storage_key, status, media_type FROM files WHERE id=%s", [file_id]).fetchone()
    if row is None or row["status"] != "stored":
        raise AppError(404, "NOT_FOUND", "模板 Word 文件不存在。")
    return {"id": row["id"], "name": row["name"], "sha256": row["sha256"], "storageKey": row["storage_key"], "mediaType": row["media_type"]}


def _path(conn, file_id) -> Path:
    path = resolve_object(get_settings().file_data_dir, _file(conn, file_id)["storageKey"])
    if not path.is_file():
        raise AppError(404, "NOT_FOUND", "模板 Word 文件不存在，请联系管理员。")
    return path


def _td(conn, td_id: UUID):
    row = conn.execute("""SELECT td.id, td.name, td.template_id, td.word_file_id, td.format_template, td.current_version_id,
                                 td.updated_at, t.name AS template_name
                          FROM template_documents td JOIN templates t ON t.id = td.template_id WHERE td.id=%s""", [td_id]).fetchone()
    if row is None:
        raise AppError(404, "NOT_FOUND", "模板文档不存在。")
    return row


_COLUMNS = """v.id, v.template_document_id, v.version, v.status, v.source_file_id, v.source_sha256, v.edited_file_id, v.edited_sha256,
              v.structure_json, v.format_template_json, v.structure_sha256, v.parent_version_id, v.restored_from_version_id,
              v.note, v.revision, v.created_by, v.created_at, v.updated_by, v.updated_at, v.published_by, v.published_at,
              ub.name AS updated_by_name, pb.name AS published_by_name"""


def _row(row) -> dict | None:
    if row is None:
        return None
    data = {name: row[name] for name in row}
    data["structure"] = _json(data.pop("structure_json"))
    data["format"] = _json(data.pop("format_template_json"))
    return data


def _version(conn, version_id) -> dict | None:
    return _row(conn.execute(f"""SELECT {_COLUMNS} FROM template_document_versions v
                                LEFT JOIN users ub ON ub.id = v.updated_by LEFT JOIN users pb ON pb.id = v.published_by
                                WHERE v.id=%s""", [version_id]).fetchone())


def _draft(conn, td_id) -> dict | None:
    return _row(conn.execute(f"""SELECT {_COLUMNS} FROM template_document_versions v
                                LEFT JOIN users ub ON ub.id = v.updated_by LEFT JOIN users pb ON pb.id = v.published_by
                                WHERE v.template_document_id=%s AND v.status='draft'""", [td_id]).fetchone())


def word_file_id(version: dict) -> UUID:
    """版本对应的 Word：有编辑生成的用 edited，否则用 source（17 §3.2）。"""
    return version["edited_file_id"] or version["source_file_id"]


def _base_file_id(draft: dict) -> UUID:
    return UUID(draft["structure"]["baseFileId"])


def _seq(conn, version_id) -> int | None:
    if version_id is None:
        return None
    row = conn.execute("SELECT version FROM template_document_versions WHERE id=%s", [version_id]).fetchone()
    return None if row is None else row["version"]


def _ref(conn, version_id) -> dict | None:
    seq = _seq(conn, version_id)
    return None if version_id is None else {"versionId": str(version_id), "seq": seq}


def _person(user_id, name) -> dict | None:
    return None if user_id is None else {"id": str(user_id), "name": name or ""}


# ---------------------------------------------------------------- 第 1 版与回填（17 §3.4）

def _format_of(conn, file_id) -> dict:
    try:
        return tw.read_format(Document(_path(conn, file_id)))
    except Exception:
        return {"styles": {}}


def bind_documents(conn, td_id: UUID, v1_id: UUID, v1_hash: str) -> tuple[int, int]:
    """只绑定第 1 版：文档里的结构拷贝与第 1 版规范化哈希相同才绑定，否则保持 NULL（版本化之前），不伪造历史。"""
    bound = legacy = 0
    rows = conn.execute("""SELECT id, format_template FROM documents
                           WHERE template_document_id=%s AND template_version_id IS NULL""", [td_id]).fetchall()
    for row in rows:
        structure = (row["format_template"] or {}).get("importStructure") if isinstance(row["format_template"], dict) else None
        if structure and tm.canonical_hash(structure) == v1_hash:
            conn.execute("UPDATE documents SET template_version_id=%s WHERE id=%s AND template_version_id IS NULL", [v1_id, row["id"]])
            bound += 1
        else:
            legacy += 1
    return bound, legacy


def ensure_v1(conn, td_id: UUID, actor: UUID | None = None, bind: bool = True) -> UUID | None:
    """必须在写事务内调用。已审核 → 第 1 版（published），未审核 → 初始草稿；返回当前版本 ID（未审核为 None）。"""
    td = _td(conn, td_id)
    if td["current_version_id"] is not None:
        return td["current_version_id"]
    published = conn.execute("SELECT id FROM template_document_versions WHERE template_document_id=%s AND status='published'",
                             [td_id]).fetchone()
    if published is not None:
        conn.execute("UPDATE template_documents SET current_version_id=%s WHERE id=%s", [published["id"], td_id])
        return published["id"]
    format_template = td["format_template"] if isinstance(td["format_template"], dict) else {}
    structure = format_template.get("importStructure")
    source = _file(conn, td["word_file_id"])
    if structure and structure.get("reviewed"):
        conn.execute("DELETE FROM template_document_versions WHERE template_document_id=%s AND status='draft' AND parent_version_id IS NULL",
                     [td_id])
        version_id = uuid4()
        digest = tm.canonical_hash(structure)
        conn.execute("""INSERT INTO template_document_versions
            (id, template_document_id, version, status, source_file_id, source_sha256, edited_file_id, edited_sha256,
             structure_json, format_template_json, structure_sha256, note, revision, created_by, updated_by, published_by, published_at)
            VALUES (%s, %s, 1, 'published', %s, %s, NULL, NULL, %s, %s, %s, %s, 1, %s, %s, %s, %s)""",
                     [version_id, td_id, source["id"], source["sha256"], json.dumps(structure, ensure_ascii=False),
                      json.dumps(_format_of(conn, source["id"]), ensure_ascii=False), digest, _V1_NOTE,
                      actor, actor, actor, td["updated_at"]])
        conn.execute("UPDATE template_documents SET current_version_id=%s WHERE id=%s", [version_id, td_id])
        if bind:
            bind_documents(conn, td_id, version_id, digest)
        return version_id
    if _draft(conn, td_id) is None:
        if structure is None:
            try:
                structure = parse_template(_path(conn, source["id"]))
            except AppError:
                raise
            except Exception:
                raise AppError(422, "WORD_PARSE_FAILED", "无法读取 Word 文档，请确认文件有效后重新上传。") from None
        conn.execute("""INSERT INTO template_document_versions
            (id, template_document_id, version, status, source_file_id, source_sha256, structure_json, format_template_json,
             structure_sha256, revision, created_by, updated_by)
            VALUES (%s, %s, NULL, 'draft', %s, %s, %s, %s, %s, 1, %s, %s)""",
                     [uuid4(), td_id, source["id"], source["sha256"], json.dumps(structure, ensure_ascii=False),
                      json.dumps({"styles": {}}), tm.canonical_hash(structure), actor, actor])
    return None


def sync_initial_draft(conn, td_id: UUID, structure: dict, actor: UUID) -> None:
    """未审核阶段（PATCH /structure）同步初始草稿；审核确认时生成第 1 版。须在写事务内。"""
    if structure.get("reviewed"):
        ensure_v1(conn, td_id, actor)
        return
    ensure_v1(conn, td_id, actor)
    conn.execute("""UPDATE template_document_versions SET structure_json=%s, structure_sha256=%s, revision=revision+1,
                    updated_by=%s, updated_at=CURRENT_TIMESTAMP
                    WHERE template_document_id=%s AND status='draft' AND parent_version_id IS NULL""",
                 [json.dumps(structure, ensure_ascii=False), tm.canonical_hash(structure), actor, td_id])


def ensure(td_id: UUID, actor: UUID | None = None) -> None:
    """惰性兜底：还没有版本行的模板文档，在一个写事务里补上（17 §3.4）。"""
    with connect() as conn:
        row = conn.execute("""SELECT td.current_version_id,
                                     (SELECT count(*) FROM template_document_versions v WHERE v.template_document_id = td.id) AS n
                              FROM template_documents td WHERE td.id=%s""", [td_id]).fetchone()
        if row is None:
            raise AppError(404, "NOT_FOUND", "模板文档不存在。")
        if row["current_version_id"] is not None or row["n"]:
            return
        with write_transaction(conn):
            ensure_v1(conn, td_id, actor)


# ---------------------------------------------------------------- 文件去重（17 §3.3）

def find_or_store(conn, content: bytes, name: str) -> tuple[UUID, str]:
    digest = tw.sha256_bytes(content)
    rows = conn.execute("""SELECT id, storage_key FROM files WHERE sha256=%s AND purpose='template' AND status='stored'
                           AND size_bytes=%s ORDER BY created_at""", [digest, len(content)]).fetchall()
    for row in rows:
        try:
            if read_object_bytes(get_settings().file_data_dir, row["storage_key"]) == content:
                return row["id"], digest
        except Exception:
            continue
    stored = store_bytes(name=name, media_type=WORD, purpose="template", content=content)
    return stored.id, digest


# ---------------------------------------------------------------- 负载

def _analysis(conn, file_id) -> dict:
    return tw.analyze(_path(conn, file_id))


def _structure_out(structure: dict, analysis: dict) -> dict:
    result = tw.enrich(structure, analysis)
    result.pop("baseFileId", None)
    return result


def _version_item(conn, version: dict) -> dict:
    return {"versionId": str(version["id"]), "seq": version["version"], "status": version["status"],
            "publishedAt": format_timestamp(version["published_at"]) if version["published_at"] else None,
            "publishedBy": _person(version["published_by"], version["published_by_name"]),
            "note": version["note"], "restoredFrom": _ref(conn, version["restored_from_version_id"]),
            "parentVersion": _ref(conn, version["parent_version_id"]),
            "sourceSha256": version["source_sha256"], "editedSha256": version["edited_sha256"],
            "hasEditedFile": version["edited_file_id"] is not None}


def _draft_payload(conn, draft: dict, analysis: dict | None = None) -> dict:
    analysis = analysis or _analysis(conn, _base_file_id(draft))
    blocks = draft["structure"]["blocks"]
    return {"structure": _structure_out(draft["structure"], analysis), "format": draft["format"],
            "baseStyles": analysis["styles"], "draftRevision": draft["revision"],
            "parentVersion": _ref(conn, draft["parent_version_id"]), "restoredFrom": _ref(conn, draft["restored_from_version_id"]),
            "uploaded": draft["source_file_id"] != _version(conn, draft["parent_version_id"])["source_file_id"],
            "unplacedCount": sum(1 for block in blocks if block.get("unplaced")),
            "updatedAt": format_timestamp(draft["updated_at"]), "updatedBy": _person(draft["updated_by"], draft["updated_by_name"])}


def _conflict_details(draft: dict) -> dict:
    return {"currentRevision": draft["revision"], "lastActor": _person(draft["updated_by"], draft["updated_by_name"]),
            "lastUpdatedAt": format_timestamp(draft["updated_at"])}


def _current(conn, td_id: UUID) -> dict:
    td = _td(conn, td_id)
    if td["current_version_id"] is None:
        raise AppError(409, "TEMPLATE_REVIEW_REQUIRED", "请先在模板管理中核对并确认模板结构，再编辑模板。")
    return _version(conn, td["current_version_id"])


def _edit_error(error: tw.TemplateEditError) -> AppError:
    message = {"TEMPLATE_EDIT_UNSUPPORTED": "有内容不能在线修改，请按提示调整后再试。",
               "TEMPLATE_STRUCTURE_INVALID": "模板结构与 Word 不一致，请按提示调整后再试。"}.get(error.code, "模板内容不正确。")
    return AppError(422, error.code, message, details={"items": error.details})


def _usage(conn, td_id: UUID, principal) -> dict:
    admin = "system_admin" in principal.roles
    rows = conn.execute("""SELECT d.id, d.name, d.template_version_id, p.id AS project_id, p.name AS project_name, v.version,
                                  EXISTS (SELECT 1 FROM project_members m WHERE m.project_id = p.id AND m.user_id = %s) AS member
                           FROM documents d JOIN projects p ON p.id = d.project_id
                           LEFT JOIN template_document_versions v ON v.id = d.template_version_id
                           WHERE d.template_document_id=%s ORDER BY p.name, d.name""", [principal.user_id, td_id]).fetchall()
    visible = [{"projectId": str(row["project_id"]), "projectName": row["project_name"], "documentId": str(row["id"]),
                "documentName": row["name"], "seq": row["version"], "legacy": row["template_version_id"] is None}
               for row in rows if admin or row["member"]]
    return {"total": len(rows), "visible": visible[:200], "hiddenCount": len(rows) - len(visible)}


# ---------------------------------------------------------------- 只读接口

@router.get("/template-documents/{document_id}/versions")
def list_versions(request: Request, document_id: str) -> dict:
    principal = require_permission(request, "templates.read")
    td_id = _uuid(document_id, "documentId")
    ensure(td_id, principal.user_id)
    with connect() as conn:
        td = _td(conn, td_id)
        rows = conn.execute(f"""SELECT {_COLUMNS} FROM template_document_versions v
                                LEFT JOIN users ub ON ub.id = v.updated_by LEFT JOIN users pb ON pb.id = v.published_by
                                WHERE v.template_document_id=%s AND v.status <> 'draft' ORDER BY v.version DESC""", [td_id]).fetchall()
        versions = [_row(row) for row in rows]
        draft = _draft(conn, td_id)
        first = versions[-1] if versions else None
        original_id = first["source_file_id"] if first else td["word_file_id"]
        original = _file(conn, original_id)
        data = {"items": [_version_item(conn, version) for version in versions],
                "current": _ref(conn, td["current_version_id"]),
                "draft": None if draft is None or draft["parent_version_id"] is None else {
                    "draftRevision": draft["revision"], "updatedAt": format_timestamp(draft["updated_at"]),
                    "updatedBy": _person(draft["updated_by"], draft["updated_by_name"]),
                    "parentVersion": _ref(conn, draft["parent_version_id"])},
                "reviewRequired": td["current_version_id"] is None,
                "original": {"fileId": str(original["id"]), "name": original["name"], "sha256": original["sha256"]},
                "usage": _usage(conn, td_id, principal), "templateName": td["template_name"], "documentName": td["name"],
                "templateId": str(td["template_id"])}
    return {"data": data, "requestId": request.state.request_id}


def _owned_version(conn, td_id: UUID, version_id: str) -> dict:
    version = _version(conn, _uuid(version_id, "versionId"))
    if version is None or version["template_document_id"] != td_id or version["status"] == "draft":
        raise AppError(404, "NOT_FOUND", "模板版本不存在。")
    return version


@router.get("/template-documents/{document_id}/versions/{version_id}")
def get_version(request: Request, document_id: str, version_id: str) -> dict:
    require_permission(request, "templates.read")
    td_id = _uuid(document_id, "documentId")
    with connect() as conn:
        _td(conn, td_id)
        version = _owned_version(conn, td_id, version_id)
        analysis = _analysis(conn, word_file_id(version))
        data = {**_version_item(conn, version), "structure": _structure_out(version["structure"], analysis), "format": version["format"]}
    return {"data": data, "requestId": request.state.request_id}


@router.get("/template-documents/{document_id}/versions/{version_id}/file")
def download_version(request: Request, document_id: str, version_id: str, kind: str = "edited") -> Response:
    require_permission(request, "templates.read")
    if kind not in {"source", "edited"}:
        raise AppError(422, "VALIDATION_FAILED", "文件类型只能是 source 或 edited。")
    td_id = _uuid(document_id, "documentId")
    with connect() as conn:
        td = _td(conn, td_id)
        version = _owned_version(conn, td_id, version_id)
        file_id = version["source_file_id"] if kind == "source" else word_file_id(version)
        record = _file(conn, file_id)
    if request.headers.get("if-none-match", "").strip('"') == record["sha256"]:
        return Response(status_code=304, headers={"ETag": f'"{record["sha256"]}"'})
    payload = read_object_bytes(get_settings().file_data_dir, record["storageKey"])
    name = f"{td['name']}-原始文件.docx" if kind == "source" else f"{td['name']}-第{version['version']}版.docx"
    return Response(content=payload, media_type=WORD, headers={
        "ETag": f'"{record["sha256"]}"', "X-Content-Type-Options": "nosniff", "Cache-Control": "private, no-cache",
        "Content-Disposition": content_disposition("attachment", name)})


def _version_with_cells(conn, version: dict) -> tuple[dict, dict]:
    analysis = _analysis(conn, word_file_id(version))
    return tm.with_cells(version["structure"], analysis), analysis


@router.get("/template-documents/{document_id}/versions/{a}/diff/{b}")
def diff_versions(request: Request, document_id: str, a: str, b: str) -> dict:
    require_permission(request, "templates.read")
    td_id = _uuid(document_id, "documentId")
    with connect() as conn:
        _td(conn, td_id)
        if b == "draft":
            left = _owned_version(conn, td_id, a)
            draft = _draft(conn, td_id)
            if draft is None or draft["parent_version_id"] is None:
                raise AppError(404, "NOT_FOUND", "没有草稿。")
            right_structure = tm.with_cells(draft["structure"], _analysis(conn, _base_file_id(draft)))
            right_styles = draft["format"].get("styles", {})
        else:
            left = _owned_version(conn, td_id, a)
            right = _owned_version(conn, td_id, b)
            right_structure, _ = _version_with_cells(conn, right)
            right_styles = right["format"].get("styles", {})
        left_structure, _ = _version_with_cells(conn, left)
        left_styles = left["format"].get("styles", {})
        if b != "draft" and "defined" in left["format"] and "defined" in right["format"]:
            # 两个已发布版本：只比样式自身的定义，继承带来的连带变化不算一处修改
            left_styles, right_styles = left["format"]["defined"], right["format"]["defined"]
        data = tm.diff(left_structure, right_structure, left_styles, right_styles)
    return {"data": data, "requestId": request.state.request_id}


# ---------------------------------------------------------------- 草稿

@router.get("/template-documents/{document_id}/draft")
def get_draft(request: Request, document_id: str) -> dict:
    principal = require_permission(request, "templates.read")
    td_id = _uuid(document_id, "documentId")
    ensure(td_id, principal.user_id)
    with connect() as conn:
        td = _td(conn, td_id)
        if td["current_version_id"] is None:
            data = {"draft": None, "base": None, "reviewRequired": True}
        else:
            draft = _draft(conn, td_id)
            current = _version(conn, td["current_version_id"])
            if draft is not None and draft["parent_version_id"] is not None:
                data = {"draft": _draft_payload(conn, draft), "base": None, "reviewRequired": False}
            else:
                analysis = _analysis(conn, word_file_id(current))
                data = {"draft": None, "reviewRequired": False,
                        "base": {"structure": _structure_out(current["structure"], analysis), "format": current["format"],
                                 "baseStyles": analysis["styles"], "draftRevision": 0,
                                 "parentVersion": _ref(conn, current["id"]), "restoredFrom": None, "uploaded": False,
                                 "unplacedCount": 0, "updatedAt": None, "updatedBy": None}}
        data["current"] = _ref(conn, td["current_version_id"])
        data["templateName"] = td["template_name"]
        data["documentName"] = td["name"]
        data["templateId"] = str(td["template_id"])
    return {"data": data, "requestId": request.state.request_id}


def _reference(conn, draft: dict | None, current: dict) -> tuple[dict, UUID, dict]:
    """草稿校验的参照：已有草稿用草稿自身（path 指向 baseFileId），否则用当前版本。"""
    if draft is not None:
        base_file = _base_file_id(draft)
        reference = deepcopy(draft["structure"])
        parent = _version(conn, draft["parent_version_id"])
        if parent is not None and word_file_id(parent) == base_file:
            known = {block["id"] for block in reference["blocks"]}
            reference["blocks"] += [block for block in parent["structure"]["blocks"] if block["id"] not in known and block.get("path")]
        return reference, base_file, draft["format"]
    reference = deepcopy(current["structure"])
    return reference, word_file_id(current), current["format"]


def _same_draft(latest: dict | None, draft: dict | None) -> None:
    """事务内复核草稿 id（23 S-9，防 ABA：草稿被发布后又有人新建了同 revision 的草稿）。"""
    if latest is not None and draft is not None and latest["id"] != draft["id"]:
        raise AppError(409, "REVISION_CONFLICT", "草稿已被其他人替换，请刷新后再继续。", details=_conflict_details(latest))


def _cas(draft: dict | None, expected: int) -> None:
    if draft is None and expected != 0:
        raise AppError(409, "REVISION_CONFLICT", "草稿已被放弃或发布，请刷新后再继续。", details={"currentRevision": 0})
    if draft is not None and draft["revision"] != expected:
        raise AppError(409, "REVISION_CONFLICT", "草稿已被其他人修改，请刷新后再继续。", details=_conflict_details(draft))


def _write_draft(conn, td_id: UUID, draft: dict | None, current: dict, structure: dict, format_: dict, actor: UUID, *,
                 source_file: dict | None = None, restored_from=None, keep_restored: bool = True) -> None:
    text = json.dumps(structure, ensure_ascii=False)
    digest = tm.canonical_hash(structure)
    if draft is None:
        source = source_file or {"id": current["source_file_id"], "sha256": current["source_sha256"]}
        conn.execute("""INSERT INTO template_document_versions
            (id, template_document_id, version, status, source_file_id, source_sha256, structure_json, format_template_json,
             structure_sha256, parent_version_id, restored_from_version_id, revision, created_by, updated_by)
            VALUES (%s, %s, NULL, 'draft', %s, %s, %s, %s, %s, %s, %s, 1, %s, %s)""",
                     [uuid4(), td_id, source["id"], source["sha256"], text, json.dumps(format_, ensure_ascii=False), digest,
                      current["id"], restored_from, actor, actor])
        return
    source = source_file or {"id": draft["source_file_id"], "sha256": draft["source_sha256"]}
    restored = draft["restored_from_version_id"] if keep_restored and restored_from is None else restored_from
    conn.execute("""UPDATE template_document_versions SET structure_json=%s, format_template_json=%s, structure_sha256=%s,
                    source_file_id=%s, source_sha256=%s, parent_version_id=%s, restored_from_version_id=%s,
                    revision=revision+1, updated_by=%s, updated_at=CURRENT_TIMESTAMP WHERE id=%s AND revision=%s""",
                 [text, json.dumps(format_, ensure_ascii=False), digest, source["id"], source["sha256"], current["id"], restored, actor,
                  draft["id"], draft["revision"]])


def _saved(conn, td_id: UUID) -> dict:
    draft = _draft(conn, td_id)
    return {"draftRevision": draft["revision"], "updatedAt": format_timestamp(draft["updated_at"])}


@router.put("/template-documents/{document_id}/draft")
def put_draft(request: Request, document_id: str, body: DraftBody) -> dict:
    principal = require_permission(request, "templates.manage")
    td_id = _uuid(document_id, "documentId")
    ensure(td_id, principal.user_id)
    with connect() as conn:
        current = _current(conn, td_id)
        draft = _draft(conn, td_id)
        parent = current
        if draft is None and body.parentVersionId is not None:
            parent = _owned_version(conn, td_id, body.parentVersionId)
        elif draft is not None and body.parentVersionId is not None and str(draft["parent_version_id"]) != body.parentVersionId:
            raise AppError(409, "REVISION_CONFLICT", "草稿已被其他人替换，请刷新后再继续。", details=_conflict_details(draft))
        if draft is not None or body.parentVersionId is None:
            _cas(draft, body.expectedDraftRevision)
        reference, base_file, _format = _reference(conn, draft, parent)
        analysis = _analysis(conn, base_file)
        try:
            blocks = tm.validate_draft(body.structure, reference, analysis)
            format_ = tm.validate_format(body.format, analysis["styles"])
        except tm.TooManyBlocks:
            raise AppError(413, "TEMPLATE_TOO_LARGE", f"模板最多 {tw.MAX_BLOCKS} 块，请拆分后再保存。") from None
        except tw.TemplateEditError as error:
            raise _edit_error(error) from None
        revision = (draft["revision"] if draft is not None else 0) + 1
        structure = {"schemaVersion": 1, "revision": revision, "reviewed": True,
                     "sourceHash": reference.get("sourceHash"), "baseFileId": str(base_file), "blocks": blocks}
        with write_transaction(conn):
            latest = _draft(conn, td_id)
            if latest is None and draft is not None:
                _cas(None, body.expectedDraftRevision)
            if latest is not None:
                _same_draft(latest, draft)
                _cas(latest, draft["revision"] if draft is not None else 0)
            if latest is None and body.parentVersionId is None and _td(conn, td_id)["current_version_id"] != current["id"]:
                raise AppError(409, "REVISION_CONFLICT", "模板刚发布了新版本，请刷新后再编辑。", details={"currentRevision": 0})
            _write_draft(conn, td_id, latest, parent if latest is None else _version(conn, latest["parent_version_id"]),
                         structure, format_, principal.user_id)
        return {"data": _saved(conn, td_id), "requestId": request.state.request_id}


@router.delete("/template-documents/{document_id}/draft", status_code=204)
def delete_draft(request: Request, document_id: str, expectedDraftRevision: int) -> Response:
    require_permission(request, "templates.manage")
    td_id = _uuid(document_id, "documentId")
    with connect() as conn:
        _td(conn, td_id)
        with write_transaction(conn):
            draft = _draft(conn, td_id)
            if draft is None or draft["parent_version_id"] is None:
                return Response(status_code=204)
            _cas(draft, expectedDraftRevision)
            conn.execute("DELETE FROM template_document_versions WHERE id=%s AND status='draft'", [draft["id"]])
    return Response(status_code=204)


# ---------------------------------------------------------------- 发布

def _unchanged(draft: dict, parent: dict) -> bool:
    return (str(word_file_id(parent)) == draft["structure"]["baseFileId"]
            and tm.semantic(draft["structure"]) == tm.semantic(parent["structure"])
            and (draft["format"].get("styles") or {}) == (parent["format"].get("styles") or {}))


@router.post("/template-documents/{document_id}/versions", status_code=201)
def publish(request: Request, document_id: str, body: PublishBody) -> dict:
    principal = require_permission(request, "templates.publish")
    td_id = _uuid(document_id, "documentId")
    note = body.note.strip()
    if not 1 <= len(note) <= 200:
        raise AppError(422, "VALIDATION_FAILED", "变更说明需要 1–200 字。", [{"field": "note", "message": "需要 1–200 字。"}])
    with _publish_lock(td_id), connect() as conn:
        td = _td(conn, td_id)
        draft = _draft(conn, td_id)
        if draft is None or draft["parent_version_id"] is None:
            raise AppError(409, "REVISION_CONFLICT", "没有可发布的草稿，请刷新后再继续。", details={"currentRevision": 0})
        _cas(draft, body.draftRevision)
        current = _version(conn, td["current_version_id"])
        if draft["parent_version_id"] != current["id"]:
            raise AppError(409, "TEMPLATE_DRAFT_STALE", f"{current['published_by_name'] or '其他人'} 已发布第 {current['version']} 版，请先把草稿合并到最新版本。",
                           details={"currentVersion": _version_item(conn, current)})
        unplaced = [block for block in draft["structure"]["blocks"] if block.get("unplaced")]
        if unplaced:
            raise AppError(422, "TEMPLATE_STRUCTURE_INVALID", f"还有 {len(unplaced)} 条编写指导没有指定位置。",
                           details={"items": [{"blockId": block["id"], "message": "这条编写指导还没有指定位置。"} for block in unplaced]})
        if _unchanged(draft, current):
            raise AppError(409, "TEMPLATE_UNCHANGED", "与当前版本相比没有变化，无需发布。")
        base_file = _base_file_id(draft)
        base_path = _path(conn, base_file)
        base_styles = _analysis(conn, base_file)["styles"]
        source_bytes = _path(conn, draft["source_file_id"]).read_bytes()
        # 事务外：生成 Word 并往返核对（17 §3.5-1）
        with tempfile.TemporaryDirectory(prefix="tpl-publish-") as folder:
            out = Path(folder) / "edited.docx"
            try:
                tw.apply_draft(base_path, draft["structure"], base_styles, draft["format"].get("styles") or {}, out)
                structure = tw.build_version(out, draft["structure"], current["version"] + 1)
                tw.clean_version(out, structure)
            except tw.TemplateEditError as error:
                raise _edit_error(error) from None
            except ValueError as error:
                raise AppError(422, "TEMPLATE_STRUCTURE_INVALID", str(error), details={"items": [{"blockId": None, "message": str(error)}]}) from None
            content = out.read_bytes()
            version_format = tw.read_format(Document(BytesIO(content)))
        edited_id = edited_sha = None
        if content != source_bytes:
            edited_id, edited_sha = find_or_store(conn, content, f"{td['name']}.docx")
            conn.commit()
        with write_transaction(conn):
            latest = _draft(conn, td_id)
            now_current = _td(conn, td_id)["current_version_id"]
            if latest is None or latest["id"] != draft["id"] or latest["revision"] != draft["revision"]:
                raise AppError(409, "REVISION_CONFLICT", "草稿已被其他人修改，请刷新后再继续。",
                               details=_conflict_details(latest) if latest else {"currentRevision": 0})
            if now_current != current["id"]:
                newer = _version(conn, now_current)
                raise AppError(409, "TEMPLATE_DRAFT_STALE", f"{newer['published_by_name'] or '其他人'} 已发布第 {newer['version']} 版，请先把草稿合并到最新版本。",
                               details={"currentVersion": _version_item(conn, newer)})
            seq = conn.execute("SELECT coalesce(max(version), 0) + 1 AS n FROM template_document_versions WHERE template_document_id=%s",
                               [td_id]).fetchone()["n"]
            structure["revision"] = seq
            conn.execute("UPDATE template_document_versions SET status='superseded' WHERE id=%s AND status='published'", [current["id"]])
            conn.execute("""UPDATE template_document_versions SET status='published', version=%s, edited_file_id=%s, edited_sha256=%s,
                            structure_json=%s, structure_sha256=%s, format_template_json=%s, note=%s, published_by=%s,
                            published_at=CURRENT_TIMESTAMP, updated_by=%s, updated_at=CURRENT_TIMESTAMP WHERE id=%s""",
                         [seq, edited_id, edited_sha, json.dumps(structure, ensure_ascii=False), tm.canonical_hash(structure),
                          json.dumps(version_format, ensure_ascii=False), note, principal.user_id, principal.user_id, draft["id"]])
            # 双写过渡（17 §3.4）：尚未改造的读者仍读 importStructure
            format_template = dict(td["format_template"] or {})
            format_template["importStructure"] = structure
            conn.execute("UPDATE template_documents SET current_version_id=%s, format_template=%s, updated_at=now() WHERE id=%s",
                         [draft["id"], json.dumps(format_template, ensure_ascii=False), td_id])
        version = _version(conn, draft["id"])
        return {"data": {"version": _version_item(conn, version)}, "requestId": request.state.request_id}


# ---------------------------------------------------------------- 恢复、变基、上传

def _undeletable(structure: dict, analysis: dict) -> set[str]:
    result = set()
    for block in structure["blocks"]:
        if block.get("path") and not block.get("tableId"):
            info = analysis["blocks"].get(block["path"]) or {}
            if block.get("isTable") or info.get("deleteLock"):
                result.add(block["id"])
    return result


def _merged_structure(merged_blocks: list[dict], current: dict, revision: int) -> dict:
    return {"schemaVersion": 1, "revision": revision, "reviewed": True, "sourceHash": current["structure"].get("sourceHash"),
            "baseFileId": str(word_file_id(current)), "blocks": merged_blocks}


def _validated(merged: dict, current: dict, analysis: dict) -> list[dict]:
    try:
        return tm.validate_draft(merged, current["structure"], analysis)
    except tw.TemplateEditError as error:
        raise _edit_error(error) from None


def _restore_styles(base: dict, current: dict, target: dict) -> dict:
    """恢复只改“两版定义不同”的样式属性：目标版本自身有定义的取它的有效值，目标版本没有定义的写 None（未设置，25 M-2）。"""
    result = deepcopy(base)
    current_defined, target_defined = current.get("defined"), target.get("defined")
    for key, props in result.items():
        for prop in tw.STYLE_PROPS:
            wanted = (target.get("styles", {}).get(key) or {}).get(prop)
            if current_defined is not None and target_defined is not None:
                mine, theirs = (current_defined.get(key) or {}).get(prop), (target_defined.get(key) or {}).get(prop)
                if mine == theirs:
                    continue
                if theirs is None:
                    props[prop] = None
                    continue
            if wanted is not None and wanted != props.get(prop) and tw.validate_value(prop, wanted):
                props[prop] = wanted
    return result


@router.post("/template-documents/{document_id}/draft:restore")
def restore(request: Request, document_id: str, body: RestoreBody) -> dict:
    principal = require_permission(request, "templates.manage")
    td_id = _uuid(document_id, "documentId")
    with connect() as conn:
        current = _current(conn, td_id)
        target = _owned_version(conn, td_id, body.versionId)
        draft = _draft(conn, td_id)
        _cas(draft, body.expectedDraftRevision)
        current_structure, analysis = _version_with_cells(conn, current)
        target_structure, _ = _version_with_cells(conn, target)
        merged = tm.merge3(current_structure, current_structure, target_structure,
                           (current["format"].get("styles", {}), current["format"].get("styles", {}), target["format"].get("styles", {})),
                           undeletable=_undeletable(current["structure"], analysis))
        blocks = _validated({"blocks": merged["blocks"]}, current, analysis)
        structure = _merged_structure(blocks, current, body.expectedDraftRevision + 1)
        format_ = {"styles": _restore_styles(analysis["styles"], current["format"], target["format"])}
        with write_transaction(conn):
            latest = _draft(conn, td_id)
            _cas(latest, body.expectedDraftRevision)
            source = {"id": current["source_file_id"], "sha256": current["source_sha256"]}
            _write_draft(conn, td_id, latest, current, structure, format_, principal.user_id, source_file=source,
                         restored_from=target["id"], keep_restored=False)
        draft = _draft(conn, td_id)
        data = _draft_payload(conn, draft, analysis)
        data["notes"] = merged["notes"]
        data["notePrefill"] = f"恢复自第 {target['version']} 版"
    return {"data": data, "requestId": request.state.request_id}


@router.post("/template-documents/{document_id}/draft:rebase")
def rebase(request: Request, document_id: str, body: RebaseBody) -> dict:
    principal = require_permission(request, "templates.manage")
    td_id = _uuid(document_id, "documentId")
    for item in body.resolutions:
        if item.take not in {"mine", "theirs"}:
            raise AppError(422, "VALIDATION_FAILED", "合并选择只能是 mine 或 theirs。")
    resolutions = {(item.blockId, item.field): item.take for item in body.resolutions}
    with connect() as conn:
        current = _current(conn, td_id)
        draft = _draft(conn, td_id)
        if draft is None or draft["parent_version_id"] is None:
            raise AppError(409, "REVISION_CONFLICT", "没有草稿，请刷新后再继续。", details={"currentRevision": 0})
        _cas(draft, body.expectedDraftRevision)
        parent = _version(conn, draft["parent_version_id"])
        if str(word_file_id(parent)) != draft["structure"]["baseFileId"]:
            raise AppError(409, "TEMPLATE_REBASE_UNSUPPORTED", "这份草稿来自上传的新 Word，不能自动合并；请基于最新版本重新上传。")
        base_structure, _ = _version_with_cells(conn, parent)
        theirs_structure, analysis = _version_with_cells(conn, current)
        mine_structure = tm.with_cells(draft["structure"], _analysis(conn, _base_file_id(draft)))
        base_styles, theirs_styles = parent["format"].get("styles", {}), current["format"].get("styles", {})
        merged = tm.merge3(base_structure, theirs_structure, mine_structure,
                           (base_styles, theirs_styles, draft["format"].get("styles", {})), resolutions,
                           undeletable=_undeletable(current["structure"], analysis))
        theirs_changes = tm.diff(base_structure, theirs_structure, base_styles, theirs_styles)
        unresolved = [item for item in merged["conflicts"] if item["resolved"] not in {"mine", "theirs"}]
        details = {"merged": None, "conflicts": merged["conflicts"], "theirsChanges": theirs_changes,
                   "currentVersion": _version_item(conn, current), "notes": merged["notes"]}
        if body.dryRun:
            details["merged"] = {"blocks": merged["blocks"], "styles": merged["styles"]}
            return {"data": details, "requestId": request.state.request_id}
        if unresolved:
            raise AppError(409, "TEMPLATE_REBASE_CONFLICT", f"还有 {len(unresolved)} 处需要选择。", details=details)
        blocks = _validated({"blocks": merged["blocks"]}, current, analysis)
        structure = _merged_structure(blocks, current, draft["revision"] + 1)
        format_ = {"styles": merged["styles"]}
        with write_transaction(conn):
            latest = _draft(conn, td_id)
            _same_draft(latest, draft)
            _cas(latest, body.expectedDraftRevision)
            if _td(conn, td_id)["current_version_id"] != current["id"]:
                raise AppError(409, "TEMPLATE_REBASE_CONFLICT", "又有新版本发布，请重新检查。", details=details)
            source = {"id": current["source_file_id"], "sha256": current["source_sha256"]}
            _write_draft(conn, td_id, latest, current, structure, format_, principal.user_id, source_file=source)
        data = _draft_payload(conn, _draft(conn, td_id), analysis)
        data["notes"] = merged["notes"]
    return {"data": data, "requestId": request.state.request_id}


@router.post("/template-documents/{document_id}/draft:upload")
def upload(request: Request, document_id: str, body: UploadBody) -> dict:
    principal = require_permission(request, "templates.manage")
    td_id = _uuid(document_id, "documentId")
    file_id = _uuid(body.wordFileId, "wordFileId")
    with connect() as conn:
        current = _current(conn, td_id)
        draft = _draft(conn, td_id)
        _cas(draft, body.expectedDraftRevision)
        record = _file(conn, file_id)
        if record["mediaType"] != WORD:
            raise AppError(422, "VALIDATION_FAILED", "请选择 Word .docx 文件。")
        if record["sha256"] in {current["source_sha256"], current["edited_sha256"]}:
            raise AppError(409, "TEMPLATE_UNCHANGED", "与当前版本相比没有变化，无需上传。")
        try:
            parsed = parse_template(_path(conn, file_id))
            analysis = _analysis(conn, file_id)
        except AppError:
            raise
        except Exception:
            raise AppError(422, "WORD_PARSE_FAILED", "无法读取 Word 文档，请确认文件有效后重新上传。") from None
        structure, alignment = tw.align_upload(parsed, current["structure"])
        structure.update({"reviewed": True, "revision": body.expectedDraftRevision + 1, "baseFileId": str(file_id)})
        format_ = {"styles": analysis["styles"]}
        with write_transaction(conn):
            latest = _draft(conn, td_id)
            _cas(latest, body.expectedDraftRevision)
            _write_draft(conn, td_id, latest, current, structure, format_, principal.user_id,
                         source_file={"id": file_id, "sha256": record["sha256"]}, keep_restored=False)
        data = _draft_payload(conn, _draft(conn, td_id), analysis)
        data["alignment"] = alignment
    return {"data": data, "requestId": request.state.request_id}


@router.post("/template-documents/{document_id}/draft:format-impact")
def format_impact(request: Request, document_id: str, body: ImpactBody) -> dict:
    require_permission(request, "templates.manage")
    td_id = _uuid(document_id, "documentId")
    with connect() as conn:
        current = _current(conn, td_id)
        draft = _draft(conn, td_id)
        if draft is not None and draft["parent_version_id"] is not None:
            base_file, structure, styles = _base_file_id(draft), draft["structure"], draft["format"].get("styles", {})
        else:
            base_file, structure, styles = word_file_id(current), current["structure"], current["format"].get("styles", {})
        path = _path(conn, base_file)
    analysis = tw.analyze(path)
    if body.changes is None:
        changes = tw.style_changes(analysis["styles"], styles)
    else:
        changes = []
        for change in body.changes[:50]:
            if not isinstance(change, dict) or change.get("style") not in analysis["styles"] or change.get("prop") not in tw.STYLE_PROPS:
                raise AppError(422, "VALIDATION_FAILED", "样式或属性不正确。")
            changes.append({"style": change["style"], "prop": change["prop"], "value": change.get("value")})
    path_to_id = {block["path"]: block["id"] for block in structure["blocks"] if block.get("path")}
    items = tw.impact_on(Document(path), changes, path_to_id)
    return {"data": {"items": items}, "requestId": request.state.request_id}
