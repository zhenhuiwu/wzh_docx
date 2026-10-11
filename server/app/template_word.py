"""A06 模板 Word 引擎（17 §3.5–§3.7、§10.8–§10.9）：在上一版 Word 上按块增量编辑，不从头生成。

- analyze：按 docx 静态计算每块的可编辑性（改文字/删除）与表格 tableEditable/lockedCells/cells、命名样式白名单的取值；
- apply_draft：把草稿（带块 ID、指向基底 Word 的 path）落到基底 Word：删除 → 改字/改级别 → 单元格改字 → 调整顺序 → 新增 → 样式（L-S，含 U-15 清直接格式）；
- build_version：再解析生成的 Word，与草稿逐块核对（往返一致），并把块 ID 按草稿顺序改写回来（§3.7）；
- clean_version：从某版本新建文档时清掉编写指导（与 clean_template 同一做法，按 path 定位）；
- format_impact：样式修改影响段数（K-10），发布摘要复用。
纯 Python，不占 Word COM 槽。
"""

from __future__ import annotations

import hashlib
import re
from copy import deepcopy
from difflib import SequenceMatcher
from functools import lru_cache
from io import BytesIO
from pathlib import Path
from uuid import uuid4
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from lxml import etree, html

from app.docx_export import _PROTECTED_EDIT, _edit_paragraph, _insert_block, _removal_reason
from app.docx_parser import heading_level
from app.template_import import _elements, parse_template

# ---------------------------------------------------------------- 常量

NEW_ID = re.compile(r"^block-[0-9a-f]{16}$")
MAX_BLOCKS = 4000
NEW_KINDS = {"heading", "body", "guidance", "placeholder", "example"}
# 删除时允许随段落一起删掉的书签：目录/光标/OLE 链接书签不承载引用内容（目录需在 Word 中更新域）。
_FREE_BOOKMARKS = ("_Toc", "_GoBack", "OLE_LINK")
_DELETE_GUARD = (".//w:drawing|.//w:pict|.//w:object|.//w:hyperlink|.//w:fldChar|.//w:fldSimple"
                 "|.//w:footnoteReference|.//w:endnoteReference|.//m:oMath|.//w:sdt"
                 "|.//w:br[@w:type='page' or @w:type='column']|.//w:sectPr|.//w:gridSpan|.//w:vMerge")
_TABLE_GUARD = ".//w:drawing|.//w:pict|.//w:object|.//w:fldChar|.//w:fldSimple|.//w:sdt|.//m:oMath"

# 样式白名单（16 §3.2）：键是界面名称，值是 Word 样式名候选（按顺序取第一个存在的）。
STYLE_WHITELIST: dict[str, tuple[str, ...]] = {
    "正文": ("Normal",),
    "标题 1": ("Heading 1",), "标题 2": ("Heading 2",), "标题 3": ("Heading 3",),
    "标题 4": ("Heading 4",), "标题 5": ("Heading 5",), "标题 6": ("Heading 6",),
    "表格正文": ("表格文字", "表格内容", "表格正文", "Table Text"),
    "题注": ("Caption",),
}
# 可调属性：类型 char 看 run 的 rPr，para 看段落 pPr（17 §10.8）。
STYLE_PROPS: dict[str, str] = {
    "fontEastAsia": "char", "fontAscii": "char", "size": "char",
    "firstLineIndent": "para", "leftIndent": "para", "rightIndent": "para",
    "lineSpacing": "para", "spaceBefore": "para", "spaceAfter": "para",
}
PROP_LABELS = {"fontEastAsia": "中文字体", "fontAscii": "西文字体", "size": "字号", "firstLineIndent": "首行缩进",
               "leftIndent": "左缩进", "rightIndent": "右缩进", "lineSpacing": "行距",
               "spaceBefore": "段前", "spaceAfter": "段后"}


class TemplateEditError(ValueError):
    """发布/预检失败：code 为 TEMPLATE_EDIT_UNSUPPORTED 或 TEMPLATE_STRUCTURE_INVALID，details=[{blockId,message}]。"""

    def __init__(self, code: str, details: list[dict]):
        super().__init__(details[0]["message"] if details else code)
        self.code = code
        self.details = details


# ---------------------------------------------------------------- 工具

def sha256_bytes(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def new_block_id() -> str:
    return "block-" + uuid4().hex[:16]


def html_paragraph(text: str, tag: str = "p"):
    element = html.Element(tag)
    parts = text.split("\n")
    element.text = parts[0]
    for part in parts[1:]:
        br = etree.SubElement(element, "br")
        br.tail = part
    return element


def _paragraph_text(doc, element) -> str:
    return Paragraph(element, doc).text.strip()


def _bookmark_names(element) -> list[str]:
    ids = {node.get(qn("w:id")) for node in element.iter(qn("w:bookmarkEnd"))}
    names = [node.get(qn("w:name")) or "" for node in element.iter(qn("w:bookmarkStart"))]
    if ids and not names:
        # 只有结束标记：按全文同 id 的开始标记取名
        root = element.getroottree().getroot()
        names = [node.get(qn("w:name")) or "" for node in root.iter(qn("w:bookmarkStart")) if node.get(qn("w:id")) in ids]
    return names


def _ref_targets(root) -> set[str]:
    """REF/NOTEREF（不含目录用的 PAGEREF/超链接）引用的书签：这类书签不能随段删除。"""
    names: set[str] = set()
    texts = [node.text or "" for node in root.iter(qn("w:instrText"))]
    texts += [node.get(qn("w:instr")) or "" for node in root.iter(qn("w:fldSimple"))]
    for text in texts:
        names.update(re.findall(r"(?<![A-Z])(?:REF|NOTEREF)\s+(\S+)", text))
    return names


def delete_lock(element, refs: set[str] | None = None) -> str | None:
    """None = 可以删除；否则返回原因（与 docx_export._removal_reason 同文案）。"""
    if element.xpath(_DELETE_GUARD):
        return _removal_reason(element)
    names = _bookmark_names(element)
    if names:
        refs = _ref_targets(element.getroottree().getroot()) if refs is None else refs
        if any(not name.startswith(_FREE_BOOKMARKS) or name in refs for name in names):
            return _removal_reason(element)
    return None


def text_lock(element) -> str | None:
    if element.xpath(_PROTECTED_EDIT):
        return "此段包含图片、链接、域或复杂 Word 内容，无法安全替换；请在 Word 中修改后上传为新草稿。"
    return None


def _remove_with_bookmarks(element) -> None:
    root = element.getroottree().getroot()
    starts = {node.get(qn("w:id")) for node in element.iter(qn("w:bookmarkStart"))}
    ends = {node.get(qn("w:id")) for node in element.iter(qn("w:bookmarkEnd"))}
    element.getparent().remove(element)
    for node in list(root.iter(qn("w:bookmarkEnd"))):
        if node.get(qn("w:id")) in starts:
            node.getparent().remove(node)
    for node in list(root.iter(qn("w:bookmarkStart"))):
        if node.get(qn("w:id")) in ends:
            node.getparent().remove(node)


# ---------------------------------------------------------------- 表格

def _cell_paragraphs(tc) -> list:
    return tc.xpath("./w:p")


def table_info(doc, tbl) -> dict:
    """tableEditable/lockedCells/tableEditReasons/cells（按原始 w:tc，合并格只出现一次）。"""
    cells, locked, reasons = [], [], []
    table_protected = bool(tbl.xpath(_TABLE_GUARD))
    merged_any = bool(tbl.xpath(".//w:gridSpan|.//w:vMerge"))
    for r, row in enumerate(tbl.xpath("./w:tr")):
        texts = []
        for c, tc in enumerate(row.xpath("./w:tc")):
            paragraphs = [p for p in _cell_paragraphs(tc) if Paragraph(p, doc).text.strip()]
            texts.append("\n".join(Paragraph(p, doc).text for p in paragraphs) if len(paragraphs) > 1
                         else (Paragraph(paragraphs[0], doc).text if paragraphs else ""))
            code = message = None
            if tc.xpath("./w:tcPr/w:gridSpan|./w:tcPr/w:vMerge"):
                code, message = "MERGED_CELLS", "不能改写合并单元格，请保留原表头和表格结构。"
            elif len(paragraphs) > 1:
                code, message = "MULTI_PARAGRAPH_CELL", f"表格第 {r + 1} 行第 {c + 1} 列包含多段文字，无法安全替换。"
            elif tc.xpath(".//w:tbl") or tc.xpath(_PROTECTED_EDIT):
                code, message = "PROTECTED_CONTENT", "此单元格包含图片、链接、域或复杂 Word 内容，无法安全替换；请在原稿中修改该段。"
            if code:
                locked.append([r, c])
                found = next((item for item in reasons if item["code"] == code and item["message"] == message), None)
                if found:
                    found["cells"].append([r, c])
                else:
                    reasons.append({"code": code, "message": message, "cells": [[r, c]]})
        cells.append(texts)
    total = sum(len(row) for row in cells)
    if table_protected:
        state = "none"
        reasons.insert(0, {"code": "PROTECTED_CONTENT", "message": "此表含 Word 域、图片或内容控件，不支持在线编辑；请在 Word 中修改后上传为新草稿。", "cells": []})
    elif not locked:
        state = "cells"
    elif len(locked) >= total:
        state = "none"
    else:
        state = "partial"
    rows = tbl.xpath("./w:tr")
    rows_editable = not merged_any and not table_protected and any(
        not row.xpath(".//w:gridSpan|.//w:vMerge|.//w:drawing|.//w:fldChar") for row in rows[1:])
    if not rows_editable and not merged_any and not table_protected:
        reasons.append({"code": "NO_PROTOTYPE_ROW", "message": "表格没有可复制的数据行，无法保留表格版式。", "cells": []})
    return {"tableEditable": state, "rowsEditable": rows_editable,
            "lockedCells": [] if state == "none" else locked, "tableEditReasons": reasons, "cells": cells}


def _edit_cell(doc, tc, old: str, new: str) -> None:
    paragraphs = [p for p in _cell_paragraphs(tc) if Paragraph(p, doc).text.strip()]
    target = paragraphs[0] if paragraphs else _cell_paragraphs(tc)[0]
    _edit_paragraph(Paragraph(target, doc), html_paragraph(new), html_paragraph(old))


# ---------------------------------------------------------------- 样式（L-S）

def _style_by_name(doc) -> dict[str, object]:
    found = {}
    for style in doc.styles.element.iter(qn("w:style")):
        if style.get(qn("w:type")) != "paragraph":
            continue
        name = style.find(qn("w:name"))
        if name is not None:
            found.setdefault(_normal_name(name.get(qn("w:val")) or ""), style)
    return found


def _normal_name(name: str) -> str:
    lowered = name.lower()
    if lowered == "normal":
        return "Normal"
    if lowered.startswith("heading ") and lowered[8:].isdigit():
        return "Heading " + lowered[8:]
    if lowered == "caption":
        return "Caption"
    return name


def whitelist_styles(doc) -> dict[str, object]:
    by_name = _style_by_name(doc)
    result = {}
    for key, names in STYLE_WHITELIST.items():
        for name in names:
            if name in by_name:
                result[key] = by_name[name]
                break
    return result


def _styles_by_id(doc) -> dict[str, object]:
    return {style.get(qn("w:styleId")): style for style in doc.styles.element.iter(qn("w:style"))}


def _default_paragraph_style(doc):
    for style in doc.styles.element.iter(qn("w:style")):
        if style.get(qn("w:type")) == "paragraph" and style.get(qn("w:default")) in ("1", "true", "on"):
            return style
    return None


def _own_value(container, prop: str):
    """容器（w:style / w:p 的 pPr、rPr 父）自身对某属性的取值；没有返回 None。container 是 w:style 或 w:pPr/w:rPr 的父元素。"""
    if container is None:
        return None
    ppr = container.find(qn("w:pPr"))
    rpr = container.find(qn("w:rPr"))
    if prop == "fontEastAsia":
        node = rpr.find(qn("w:rFonts")) if rpr is not None else None
        if node is None:
            return None
        return node.get(qn("w:eastAsia")) or ("（主题字体）" if node.get(qn("w:eastAsiaTheme")) else None)
    if prop == "fontAscii":
        node = rpr.find(qn("w:rFonts")) if rpr is not None else None
        if node is None:
            return None
        return node.get(qn("w:ascii")) or ("（主题字体）" if node.get(qn("w:asciiTheme")) else None)
    if prop == "size":
        node = rpr.find(qn("w:sz")) if rpr is not None else None
        value = node.get(qn("w:val")) if node is not None else None
        return float(value) / 2 if value and value.isdigit() else None
    if ppr is None:
        return None
    if prop in {"firstLineIndent", "leftIndent", "rightIndent"}:
        node = ppr.find(qn("w:ind"))
        if node is None:
            return None
        chars, twips = {"firstLineIndent": ("firstLineChars", "firstLine"),
                        "leftIndent": ("leftChars", "left"), "rightIndent": ("rightChars", "right")}[prop]
        value = node.get(qn("w:" + chars))
        if value is not None and value.lstrip("-").isdigit():
            return {"unit": "chars", "value": int(value) / 100}
        value = node.get(qn("w:" + twips)) or (node.get(qn("w:start")) if twips == "left" else node.get(qn("w:end")) if twips == "right" else None)
        if prop == "firstLineIndent" and value is None and node.get(qn("w:hanging")):
            return {"unit": "pt", "value": -int(node.get(qn("w:hanging"))) / 20}
        if value is not None and value.lstrip("-").isdigit():
            return {"unit": "pt", "value": int(value) / 20}
        return None
    node = ppr.find(qn("w:spacing"))
    if node is None:
        return None
    if prop == "lineSpacing":
        line, rule = node.get(qn("w:line")), node.get(qn("w:lineRule")) or "auto"
        if line is None or not line.lstrip("-").isdigit():
            return None
        return f"{int(line) / 240:g}" if rule == "auto" else f"{rule}:{int(line) / 20:g}"
    attribute = "before" if prop == "spaceBefore" else "after"
    lines = node.get(qn("w:" + attribute + "Lines"))
    if lines is not None and lines.isdigit():
        return {"unit": "lines", "value": int(lines) / 100}
    value = node.get(qn("w:" + attribute))
    return {"unit": "pt", "value": int(value) / 20} if value is not None and value.isdigit() else None


def _doc_defaults(doc):
    return doc.styles.element.find(qn("w:docDefaults"))


def effective_style_value(doc, style, prop: str):
    by_id, seen = _styles_by_id(doc), set()
    while style is not None and style.get(qn("w:styleId")) not in seen:
        seen.add(style.get(qn("w:styleId")))
        value = _own_value(style, prop)
        if value is not None:
            return value
        based = style.find(qn("w:basedOn"))
        style = by_id.get(based.get(qn("w:val"))) if based is not None else None
    defaults = _doc_defaults(doc)
    if defaults is not None:
        holder = etree.Element("holder")
        for tag in ("w:pPrDefault", "w:rPrDefault"):
            node = defaults.find(qn(tag))
            if node is not None:
                for child in node:
                    holder.append(deepcopy(child))
        value = _own_value(holder, prop)
        if value is not None:
            return value
    return 10.5 if prop == "size" else None


def read_styles(doc) -> dict[str, dict]:
    return {key: {prop: effective_style_value(doc, style, prop) for prop in STYLE_PROPS}
            for key, style in whitelist_styles(doc).items()}


def read_defined(doc) -> dict[str, dict]:
    """样式自身定义的取值（不含继承），用于版本对比只列真正改过的“样式 · 属性”。"""
    return {key: {prop: _own_value(style, prop) for prop in STYLE_PROPS} for key, style in whitelist_styles(doc).items()}


def read_format(doc) -> dict:
    return {"styles": read_styles(doc), "defined": read_defined(doc)}


def _value_ok(prop: str, value) -> bool:
    if value is None:
        return False
    if prop in {"fontEastAsia", "fontAscii"}:
        return isinstance(value, str) and 0 < len(value.strip()) <= 64 and "（" not in value
    if prop == "size":
        return isinstance(value, (int, float)) and not isinstance(value, bool) and 5 <= value <= 72
    if prop == "lineSpacing":
        return isinstance(value, str) and bool(re.fullmatch(r"(?:[0-9](?:\.[0-9]{1,2})?|exact:[0-9]{1,2}(?:\.[05])?|atLeast:[0-9]{1,2}(?:\.[05])?)", value))
    if not isinstance(value, dict) or set(value) != {"unit", "value"}:
        return False
    number = value["value"]
    if isinstance(number, bool) or not isinstance(number, (int, float)):
        return False
    if prop in {"firstLineIndent", "leftIndent", "rightIndent"}:
        return value["unit"] in {"chars", "pt"} and -20 <= number <= (40 if value["unit"] == "chars" else 400)
    return value["unit"] in {"pt", "lines"} and 0 <= number <= (100 if value["unit"] == "pt" else 10)


def validate_value(prop: str, value) -> bool:
    return prop in STYLE_PROPS and _value_ok(prop, value)


def _child(parent, tag: str, first: bool = False):
    node = parent.find(qn(tag))
    if node is None:
        node = etree.SubElement(parent, qn(tag))
        if first:
            parent.remove(node)
            parent.insert(0, node)
    return node


def _set_style_value(doc, style, prop: str, value) -> None:
    if STYLE_PROPS[prop] == "char":
        rpr = _child(style, "w:rPr")
        if prop == "size":
            for tag in ("w:sz", "w:szCs"):
                _child(rpr, tag).set(qn("w:val"), str(round(value * 2)))
            return
        fonts = _child(rpr, "w:rFonts")
        if prop == "fontEastAsia":
            fonts.set(qn("w:eastAsia"), value)
            fonts.attrib.pop(qn("w:eastAsiaTheme"), None)
        else:
            for name in ("ascii", "hAnsi"):
                fonts.set(qn("w:" + name), value)
                fonts.attrib.pop(qn(f"w:{name}Theme"), None)
        return
    ppr = _child(style, "w:pPr")
    if prop in {"firstLineIndent", "leftIndent", "rightIndent"}:
        ind = _child(ppr, "w:ind")
        size = effective_style_value(doc, style, "size") or 10.5
        twips = round(value["value"] * (size if value["unit"] == "chars" else 1) * 20)
        if prop == "firstLineIndent":
            for name in ("firstLine", "firstLineChars", "hanging", "hangingChars"):
                ind.attrib.pop(qn("w:" + name), None)
            if value["value"] < 0:
                ind.set(qn("w:hanging"), str(-twips))
                if value["unit"] == "chars":
                    ind.set(qn("w:hangingChars"), str(round(-value["value"] * 100)))
            else:
                ind.set(qn("w:firstLine"), str(twips))
                if value["unit"] == "chars":
                    ind.set(qn("w:firstLineChars"), str(round(value["value"] * 100)))
            return
        side = "left" if prop == "leftIndent" else "right"
        for name in (side, side + "Chars", "start" if side == "left" else "end", ("start" if side == "left" else "end") + "Chars"):
            ind.attrib.pop(qn("w:" + name), None)
        ind.set(qn("w:" + side), str(twips))
        if value["unit"] == "chars":
            ind.set(qn("w:" + side + "Chars"), str(round(value["value"] * 100)))
        return
    spacing = _child(ppr, "w:spacing")
    if prop == "lineSpacing":
        if ":" in value:
            rule, points = value.split(":")
            spacing.set(qn("w:line"), str(round(float(points) * 20)))
            spacing.set(qn("w:lineRule"), rule)
        else:
            spacing.set(qn("w:line"), str(round(float(value) * 240)))
            spacing.set(qn("w:lineRule"), "auto")
        return
    attribute = "before" if prop == "spaceBefore" else "after"
    for name in (attribute, attribute + "Lines", attribute + "Autospacing"):
        spacing.attrib.pop(qn("w:" + name), None)
    if value["unit"] == "lines":
        spacing.set(qn("w:" + attribute + "Lines"), str(round(value["value"] * 100)))
        spacing.set(qn("w:" + attribute), str(round(value["value"] * 240)))
    else:
        spacing.set(qn("w:" + attribute), str(round(value["value"] * 20)))


_DIRECT_ATTRS = {
    "firstLineIndent": ("ind", ("firstLine", "firstLineChars", "hanging", "hangingChars")),
    "leftIndent": ("ind", ("left", "leftChars", "start", "startChars")),
    "rightIndent": ("ind", ("right", "rightChars", "end", "endChars")),
    "lineSpacing": ("spacing", ("line", "lineRule")),
    "spaceBefore": ("spacing", ("before", "beforeLines", "beforeAutospacing")),
    "spaceAfter": ("spacing", ("after", "afterLines", "afterAutospacing")),
}
_CHAR_ATTRS = {"fontEastAsia": ("eastAsia", "eastAsiaTheme"), "fontAscii": ("ascii", "hAnsi", "asciiTheme", "hAnsiTheme")}


def _run_properties(paragraph_element) -> list:
    found = [node for node in paragraph_element.iter(qn("w:rPr"))
             if node.getparent().tag == qn("w:r")]
    mark = paragraph_element.find(qn("w:pPr") + "/" + qn("w:rPr"))
    return found + ([mark] if mark is not None else [])


def has_direct(paragraph_element, prop: str) -> bool:
    if STYLE_PROPS[prop] == "para":
        tag, names = _DIRECT_ATTRS[prop]
        node = paragraph_element.find(qn("w:pPr") + "/" + qn("w:" + tag))
        return node is not None and any(node.get(qn("w:" + name)) is not None for name in names)
    for rpr in _run_properties(paragraph_element):
        if prop == "size" and (rpr.find(qn("w:sz")) is not None or rpr.find(qn("w:szCs")) is not None):
            return True
        if prop in _CHAR_ATTRS:
            fonts = rpr.find(qn("w:rFonts"))
            if fonts is not None and any(fonts.get(qn("w:" + name)) is not None for name in _CHAR_ATTRS[prop]):
                return True
    return False


def clear_direct(paragraph_element, prop: str) -> None:
    if STYLE_PROPS[prop] == "para":
        tag, names = _DIRECT_ATTRS[prop]
        node = paragraph_element.find(qn("w:pPr") + "/" + qn("w:" + tag))
        if node is not None:
            for name in names:
                node.attrib.pop(qn("w:" + name), None)
            if not node.attrib:
                node.getparent().remove(node)
        return
    for rpr in _run_properties(paragraph_element):
        if prop == "size":
            for tag in ("w:sz", "w:szCs"):
                for node in rpr.findall(qn(tag)):
                    rpr.remove(node)
        else:
            fonts = rpr.find(qn("w:rFonts"))
            if fonts is not None:
                for name in _CHAR_ATTRS[prop]:
                    fonts.attrib.pop(qn("w:" + name), None)
                if not fonts.attrib:
                    rpr.remove(fonts)


def _affected(paragraph_element, target_style, prop: str, by_id: dict, default_style) -> bool:
    style_id = paragraph_element.xpath("./w:pPr/w:pStyle/@w:val")
    style = by_id.get(style_id[0]) if style_id else default_style
    seen = set()
    while style is not None and id(style) not in seen:
        seen.add(id(style))
        if style is target_style:
            return True
        if _own_value(style, prop) is not None:
            return False
        based = style.find(qn("w:basedOn"))
        style = by_id.get(based.get(qn("w:val"))) if based is not None else None
    return False


def _block_paths(doc) -> dict:
    return {element: path for path, element, _ in _elements(doc)}


def impact_on(doc, changes: list[dict], path_to_id: dict[str, str], limit: int = 500) -> list[dict]:
    """K-10：受影响段数、带直接格式段数及对应块 ID（不改文档）。"""
    styles, by_id, default = whitelist_styles(doc), _styles_by_id(doc), _default_paragraph_style(doc)
    paths = _block_paths(doc)
    paragraphs = [p for p in doc.element.body.iter(qn("w:p")) if "".join(p.xpath(".//w:t/text()")).strip()]
    items = []
    for change in changes:
        style = styles.get(change["style"])
        if style is None:
            continue
        prop = change["prop"]
        affected, direct = [], []
        for paragraph in paragraphs:
            if _affected(paragraph, style, prop, by_id, default):
                affected.append(paragraph)
                if has_direct(paragraph, prop):
                    direct.append(paragraph)

        def ids(elements):
            found = []
            for element in elements:
                block = path_to_id.get(paths.get(element))
                if block and block not in found:
                    found.append(block)
            return found
        block_ids, direct_ids = ids(affected), ids(direct)
        items.append({"style": change["style"], "prop": prop, "before": effective_style_value(doc, style, prop),
                      "after": change.get("value"), "affectedParagraphs": len(affected),
                      "directFormatParagraphs": len(direct), "willClear": len(direct),
                      "blockIds": block_ids[:limit], "directBlockIds": direct_ids[:limit],
                      "truncated": len(block_ids) > limit or len(direct_ids) > limit})
    return items


def apply_style_changes(doc, changes: list[dict]) -> None:
    """先算受影响段落（基于修改前的样式链），再改样式定义，最后按 U-15 只清被改那一项的直接格式。"""
    styles, by_id, default = whitelist_styles(doc), _styles_by_id(doc), _default_paragraph_style(doc)
    paragraphs = [p for p in doc.element.body.iter(qn("w:p")) if "".join(p.xpath(".//w:t/text()")).strip()]
    plan = []
    for change in changes:
        style = styles.get(change["style"])
        if style is None:
            raise TemplateEditError("TEMPLATE_EDIT_UNSUPPORTED", [{"blockId": None, "message": f"模板中没有“{change['style']}”样式。"}])
        plan.append((style, change, [p for p in paragraphs if _affected(p, style, change["prop"], by_id, default)]))
    for style, change, affected in plan:
        _set_style_value(doc, style, change["prop"], change["value"])
        for paragraph in affected:
            clear_direct(paragraph, change["prop"])


def style_changes(parent_styles: dict, draft_styles: dict) -> list[dict]:
    changes = []
    for key, props in (draft_styles or {}).items():
        for prop, value in props.items():
            if prop in STYLE_PROPS and value != (parent_styles.get(key) or {}).get(prop):
                changes.append({"style": key, "prop": prop, "value": value})
    return changes


# ---------------------------------------------------------------- 分析

def _analysis(content: bytes) -> dict:
    doc = Document(BytesIO(content))
    refs = _ref_targets(doc.element.body)
    by_path, tables = {}, {}
    names = {style.get(qn("w:styleId")): (style.find(qn("w:name")).get(qn("w:val")) if style.find(qn("w:name")) is not None else "")
             for style in doc.styles.element.iter(qn("w:style"))}
    key_of = {id(style): key for key, style in whitelist_styles(doc).items()}
    by_id = _styles_by_id(doc)
    default = _default_paragraph_style(doc)
    for path, element, table_path in _elements(doc):
        if element.tag == qn("w:tbl"):
            tables[path] = table_info(doc, element)
            by_path[path] = {"textLock": None, "deleteLock": "表格不能在线删除；请在 Word 中修改后上传为新草稿。"}
            continue
        style_id = element.xpath("./w:pPr/w:pStyle/@w:val")
        style = by_id.get(style_id[0]) if style_id else default
        by_path[path] = {"textLock": text_lock(element), "deleteLock": delete_lock(element, refs) if table_path is None else "单元格段落随表格保留。",
                         "styleName": _normal_name(names.get(style.get(qn("w:styleId")), "") if style is not None else ""),
                         "styleKey": key_of.get(id(style)) if style is not None else None,
                         "text": Paragraph(element, doc).text.strip(),
                         "level": heading_level(Paragraph(element, doc)) if table_path is None else None}
    return {"blocks": by_path, "tables": tables, "styles": read_styles(doc)}


@lru_cache(maxsize=16)
def _cached_analysis(digest: str, path: str) -> dict:
    return _analysis(Path(path).read_bytes())


def analyze(path: str | Path) -> dict:
    content = Path(path).read_bytes()
    return deepcopy(_cached_analysis(sha256_bytes(content), str(path)))


def enrich(structure: dict, analysis: dict) -> dict:
    """给块加上编辑能力标注（只读派生字段，保存时服务端忽略）。"""
    result = deepcopy(structure)
    for block in result["blocks"]:
        info = analysis["blocks"].get(block.get("path")) if block.get("path") else None
        if info is None:
            block["lock"] = {"text": None, "delete": None}
            block["styleKey"] = None
            block["styleName"] = None
            continue
        block["lock"] = {"text": info["textLock"], "delete": info["deleteLock"]}
        block["styleKey"] = info.get("styleKey")
        block["styleName"] = info.get("styleName")
        if block.get("isTable"):
            table = analysis["tables"][block["path"]]
            for key in ("tableEditable", "rowsEditable", "lockedCells", "tableEditReasons"):
                block[key] = table[key]
            block["baseCells"] = table["cells"]
            if "cells" not in block:
                block["cells"] = deepcopy(table["cells"])
    return result


# ---------------------------------------------------------------- 发布：应用草稿

def _sample_paragraph(doc, anchor, level: int | None):
    nearby = list(anchor.itersiblings(preceding=True)) + list(anchor.itersiblings())
    for element in nearby:
        if element.tag != qn("w:p"):
            continue
        paragraph = Paragraph(element, doc)
        if not paragraph.text.strip():
            continue
        found = heading_level(paragraph)
        if level is not None and found == level:
            return paragraph
        if level is None and found is None and paragraph.style is not None \
                and not paragraph.style.name.startswith(("List", "toc", "TOC", "Caption", "Title", "Subtitle")) \
                and paragraph.alignment != WD_ALIGN_PARAGRAPH.CENTER \
                and not re.match(r"^[图表]\s*\d+(?:[.－—-]\d+)*\s*", paragraph.text.strip()) \
                and not element.xpath(_DELETE_GUARD):
            return paragraph
    return None


def _set_level(doc, element, level: int | None) -> None:
    sample = _sample_paragraph(doc, element, level)
    paragraph = Paragraph(element, doc)
    if sample is not None and sample._p.pPr is not None:
        properties = deepcopy(sample._p.pPr)
        for boundary in properties.xpath("./w:sectPr"):
            properties.remove(boundary)
        if level is None:
            for numbering in properties.xpath("./w:numPr"):
                properties.remove(numbering)
        old = element.pPr
        if old is not None:
            for boundary in old.xpath("./w:sectPr"):
                properties.append(deepcopy(boundary))
            element.remove(old)
        element.insert(0, properties)
        runs = [run for run in sample.runs if run.text]
        model = max(runs, key=lambda run: len(run.text)) if runs else None
        if model is not None:
            for run in paragraph.runs:
                if run._r.rPr is not None:
                    run._r.remove(run._r.rPr)
                if model._r.rPr is not None:
                    run._r.insert(0, deepcopy(model._r.rPr))
    if heading_level(paragraph) != level:
        properties = element.get_or_add_pPr()
        for item in properties.xpath("./w:outlineLvl"):
            properties.remove(item)
        outline = etree.SubElement(properties, qn("w:outlineLvl"))
        outline.set(qn("w:val"), str(level - 1 if level is not None else 9))


def _word_level(block: dict) -> int | None:
    return block["headingLevel"] if block["kind"] == "heading" else None


def apply_draft(base_path: str | Path, draft: dict, base_styles: dict, draft_styles: dict, out_path: str | Path) -> int:
    """把草稿落到基底 Word，返回改动数（0 表示 Word 无变化，仍会写出与基底相同的文件）。失败抛 TemplateEditError。"""
    content = Path(base_path).read_bytes()
    doc = Document(BytesIO(content))
    body = doc.element.body
    located = {path: element for path, element, _ in _elements(doc)}
    refs = _ref_targets(body)
    errors: list[dict] = []
    ops = 0
    blocks = draft["blocks"]
    top = [block for block in blocks if not block.get("tableId")]
    real = [block for block in top if block.get("path")]
    for block in real:
        if block["path"] not in located:
            raise TemplateEditError("TEMPLATE_STRUCTURE_INVALID", [{"blockId": block["id"], "message": "模板块与 Word 位置对不上，请刷新后重试。"}])
    element_of = {block["id"]: located[block["path"]] for block in real}
    referenced = {block["path"] for block in blocks if block.get("path")}
    base_top = [path for path in located if path.count("/") == 1]
    parsed_top = {block["path"] for block in parse_template(content)["blocks"] if not block["tableId"]}

    # 1 删除（只删基底中有块、草稿里没有的顶层元素）
    for path in base_top:
        if path in referenced or path not in parsed_top:
            continue
        element = located[path]
        if element.tag == qn("w:tbl"):
            errors.append({"blockId": None, "message": "表格不能在线删除；请在 Word 中修改后上传为新草稿。", "path": path})
            continue
        reason = delete_lock(element, refs)
        if reason:
            errors.append({"blockId": None, "message": reason, "path": path})
            continue
        _remove_with_bookmarks(element)
        ops += 1

    # 2 改字、改级别、单元格
    for block in real:
        element = element_of[block["id"]]
        if element.tag == qn("w:tbl"):
            if "cells" not in block:
                continue
            info = table_info(doc, element)
            rows = element.xpath("./w:tr")
            if len(block["cells"]) != len(rows) or any(len(new) != len(row.xpath("./w:tc")) for new, row in zip(block["cells"], rows)):
                errors.append({"blockId": block["id"], "message": "表格行列与模板不一致；增删行列请在 Word 中修改后上传为新草稿。"})
                continue
            locked = {tuple(cell) for cell in info["lockedCells"]}
            for r, (new_row, row) in enumerate(zip(block["cells"], rows)):
                for c, (new, tc) in enumerate(zip(new_row, row.xpath("./w:tc"))):
                    old = info["cells"][r][c]
                    if new == old:
                        continue
                    if info["tableEditable"] == "none" or (r, c) in locked:
                        reason = next((item["message"] for item in info["tableEditReasons"] if [r, c] in item["cells"]), None) \
                            or (info["tableEditReasons"][0]["message"] if info["tableEditReasons"] else "此单元格不能在线修改。")
                        errors.append({"blockId": block["id"], "message": reason, "cell": [r, c]})
                        continue
                    try:
                        _edit_cell(doc, tc, old, new)
                        ops += 1
                    except ValueError as exc:
                        errors.append({"blockId": block["id"], "message": str(exc), "cell": [r, c]})
            continue
        paragraph = Paragraph(element, doc)
        if block["kind"] != "toc":
            current = heading_level(paragraph)
            wanted = _word_level(block)
            if current != wanted:
                try:
                    _set_level(doc, element, wanted)
                    ops += 1
                except ValueError as exc:
                    errors.append({"blockId": block["id"], "message": str(exc)})
        if block["kind"] == "guidance":
            continue
        old_text = paragraph.text.strip()
        new_text = block["text"].strip()
        if new_text != old_text:
            reason = text_lock(element)
            if reason:
                errors.append({"blockId": block["id"], "message": reason})
                continue
            try:
                _edit_paragraph(paragraph, html_paragraph(new_text), html_paragraph(paragraph.text))
                ops += 1
            except ValueError as exc:
                errors.append({"blockId": block["id"], "message": str(exc)})
    if errors:
        raise TemplateEditError("TEMPLATE_EDIT_UNSUPPORTED", errors)

    # 3 调整顺序（不跨分节符）
    ops += _reorder(doc, [block["id"] for block in real], element_of)

    # 4 新增（顶层、非指导；指导是虚拟块，不进 Word）
    previous = None
    for block in top:
        if block.get("path"):
            previous = element_of[block["id"]]
            continue
        if block["kind"] == "guidance":
            continue
        tag = f"h{block['headingLevel']}" if block["kind"] == "heading" else "p"
        anchor = previous.getnext() if previous is not None else next(iter(element_of.values()), body.sectPr)
        if anchor is None:
            anchor = body.sectPr
        before = anchor.getprevious()
        try:
            _insert_block(doc, html_paragraph(block["text"].strip(), tag), anchor, outline_headings=True)
        except (ValueError, KeyError) as exc:
            errors.append({"blockId": block["id"], "message": str(exc)})
            continue
        inserted = anchor.getprevious()
        if inserted is None or inserted is before:
            errors.append({"blockId": block["id"], "message": "新增内容没有写入 Word。"})
            continue
        if block["kind"] == "heading" and heading_level(Paragraph(inserted, doc)) != block["headingLevel"]:
            _set_level(doc, inserted, block["headingLevel"])
        elif block["kind"] != "heading" and heading_level(Paragraph(inserted, doc)) is not None:
            _set_level(doc, inserted, None)
        element_of[block["id"]] = inserted
        previous = inserted
        ops += 1
    if errors:
        raise TemplateEditError("TEMPLATE_EDIT_UNSUPPORTED", errors)

    # 5 样式（新增段落之后做：新段落复制来的直接格式同样按 U-15 清除）
    changes = style_changes(base_styles, draft_styles)
    if changes:
        apply_style_changes(doc, changes)
        ops += len(changes)

    if ops == 0:
        Path(out_path).write_bytes(content)
        return 0
    output = BytesIO()
    doc.save(output)
    Path(out_path).write_bytes(output.getvalue())
    return ops


def _reorder(doc, order: list[str], element_of: dict) -> int:
    body = doc.element.body
    final = body.sectPr
    kids = [kid for kid in body if kid is not final]
    owner = {id(element): block_id for block_id, element in element_of.items()}
    groups: list[dict] = [{"id": None, "elements": [], "barrier": False}]
    last_block = None
    for kid in kids:
        block_id = owner.get(id(kid))
        barrier = next(kid.iter(qn("w:sectPr")), None) is not None
        if block_id is not None:
            groups.append({"id": block_id, "elements": [kid], "barrier": barrier})
            last_block = block_id
        elif barrier:
            groups.append({"id": None, "elements": [kid], "barrier": True, "after": last_block})
        else:
            groups[-1]["elements"].append(kid)
    head = groups[0]
    current = [group["id"] for group in groups[1:] if group["id"] is not None]
    if current == order:
        return 0
    by_id = {group["id"]: group for group in groups if group["id"] is not None}
    followers: dict[str | None, list] = {}
    for group in groups[1:]:
        if group["id"] is None:
            followers.setdefault(group["after"], []).append(group)

    def segments(sequence: list[dict]) -> dict[str, int]:
        count, result = 0, {}
        for group in sequence:
            if group["id"] is not None:
                result[group["id"]] = count
            if group["barrier"]:
                count += 1
        return result

    def chain(block_id):
        result = []
        for follower in followers.get(block_id, []):
            result.append(follower)
        return result

    old_sequence = [head] + chain(None) + [g for group in groups[1:] if group["id"] is not None for g in [group] + chain(group["id"])]
    new_sequence = [head] + chain(None) + [g for block_id in order for g in [by_id[block_id]] + chain(block_id)]
    before, after = segments(old_sequence), segments(new_sequence)
    moved = [block_id for block_id in order if before[block_id] != after[block_id]]
    if moved:
        raise TemplateEditError("TEMPLATE_EDIT_UNSUPPORTED", [{"blockId": moved[0], "message": "调整顺序跨越了 Word 分节符（页面方向、页眉页脚等按节设置），不能在线移动；请在 Word 中修改后上传为新草稿。"}])
    for group in new_sequence:
        for element in group["elements"]:
            final.addprevious(element) if final is not None else body.append(element)
    return 1


# ---------------------------------------------------------------- 发布：往返校验与块 ID 延续

def _suffix(path: str, table_path: str) -> str:
    return path[len(table_path):]


def build_version(out_path: str | Path, draft: dict, revision: int) -> dict:
    content = Path(out_path).read_bytes()
    parsed = parse_template(content)
    doc = Document(BytesIO(content))
    tables = {path: table_info(doc, element) for path, element, _ in _elements(doc) if element.tag == qn("w:tbl")}
    parsed_top = [block for block in parsed["blocks"] if not block["tableId"]]
    children: dict[str, list] = {}
    for block in parsed["blocks"]:
        if block["tableId"]:
            table_path = block["path"].split("/row/")[0]
            children.setdefault(table_path, []).append(block)
    expected = [block for block in draft["blocks"] if not block.get("tableId") and not _is_virtual(block)]
    draft_children: dict[str, dict] = {}
    for block in draft["blocks"]:
        if block.get("tableId") and block.get("path"):
            draft_children.setdefault(block["tableId"], {})[block["path"].split("/row/", 1)[1]] = block
    problems = []
    for index, block in enumerate(expected):
        if index >= len(parsed_top):
            problems.append({"blockId": block["id"], "message": "生成的 Word 中缺少这一块。"})
            break
        found = parsed_top[index]
        if found["isTable"] != bool(block.get("isTable")):
            problems.append({"blockId": block["id"], "message": "生成的 Word 中这一块的类型不一致。"})
            break
        if block.get("isTable"):
            if "cells" in block and tables[found["path"]]["cells"] != block["cells"]:
                problems.append({"blockId": block["id"], "message": "生成的 Word 中表格文字与草稿不一致。"})
            continue
        want = block["sourceText"] if block["kind"] == "guidance" and block.get("path") else block["text"]
        if found["text"] != want.strip():
            problems.append({"blockId": block["id"], "message": "生成的 Word 中这一段的文字与草稿不一致。"})
        if block["kind"] == "heading" and (found["kind"] != "heading" or found["headingLevel"] != block["headingLevel"]):
            problems.append({"blockId": block["id"], "message": "生成的 Word 中标题级别与草稿不一致。"})
    if len(parsed_top) != len(expected) and not problems:
        extra = parsed_top[len(expected)] if len(parsed_top) > len(expected) else None
        problems.append({"blockId": None, "message": "生成的 Word 多出了内容：" + (extra["text"][:20] if extra else "")})
    if problems:
        raise TemplateEditError("TEMPLATE_STRUCTURE_INVALID", problems)
    result, pointer = [], 0
    for block in draft["blocks"]:
        if block.get("tableId"):
            continue
        if _is_virtual(block):
            result.append({"id": block["id"], "path": None, "kind": "guidance", "text": block["text"], "sourceText": block["text"],
                           "headingLevel": None, "headingId": block.get("headingId"), "tableId": None,
                           "protected": False, "isTable": False, "virtual": True})
            continue
        found = parsed_top[pointer]
        pointer += 1
        item = {key: found[key] for key in ("path", "sourceText", "protected", "isTable")}
        item.update({"id": block["id"], "kind": block["kind"],
                     "text": block["text"] if block["kind"] == "guidance" and block.get("path") else found["text"],
                     "headingLevel": block["headingLevel"] if block["kind"] == "heading" else None,
                     "headingId": block.get("headingId"), "tableId": None})
        if found.get("warning") and found["kind"] == "body" and block["kind"] == "body" and "书签或域" in found["warning"]:
            item["warning"] = found["warning"]
        result.append(item)
        if found["isTable"]:
            previous = draft_children.get(block["id"], {})
            for child in children.get(found["path"], []):
                suffix = child["path"].split("/row/", 1)[1]
                old = previous.get(suffix)
                entry = {key: child[key] for key in ("path", "sourceText", "protected", "isTable")}
                kind = old["kind"] if old else child["kind"]
                text = child["text"]
                if old and kind == "guidance" and old.get("sourceText", "").strip() == child["text"]:
                    text = old["text"]
                entry.update({"id": old["id"] if old else new_block_id(), "kind": kind, "text": text,
                              "headingLevel": None, "headingId": None, "tableId": block["id"]})
                result.append(entry)
    _rebind_headings(result)
    return {"schemaVersion": 1, "revision": revision, "reviewed": True, "sourceHash": sha256_bytes(content), "blocks": result}


def _is_virtual(block: dict) -> bool:
    return not block.get("path") and block.get("kind") == "guidance"


def _rebind_headings(blocks: list[dict]) -> None:
    headings = {block["id"] for block in blocks if block["kind"] == "heading"}
    current = None
    for block in blocks:
        if block["kind"] == "heading":
            current = block["id"]
            block["headingId"] = None
        elif block.get("unplaced"):
            block["headingId"] = None
        elif block["kind"] == "guidance":
            if block.get("headingId") not in headings:
                block["headingId"] = current
        else:
            block["headingId"] = current


# ---------------------------------------------------------------- 新建文档：清掉指导

def clean_version(path: str | Path, structure: dict) -> bytes:
    """与 clean_template 同一做法（只清编写指导/示例/填写位置，其余 ZIP 原样），按版本结构里的 path 定位。"""
    path = Path(path)
    document = Document(path)
    locations = {location: element for location, element, _ in _elements(document)}
    samples, current_heading = set(), None
    for block in structure["blocks"]:
        if block["kind"] == "heading":
            current_heading = block["id"]
        if not block.get("path"):
            continue
        element = locations.get(block["path"])
        if element is None:
            raise ValueError("模板版本与 Word 文件对不上，请联系模板维护者重新发布。")
        if element.tag == qn("w:p") and Paragraph(element, document).text.strip() != block["sourceText"].strip():
            raise ValueError("模板版本与 Word 文件对不上，请联系模板维护者重新发布。")
        if block["kind"] in {"guidance", "example", "placeholder"}:
            if block["tableId"] is None:
                if current_heading in samples:
                    element.getparent().remove(element)
                    continue
                samples.add(current_heading)
            runs = element.xpath("./w:r")
            formatting = deepcopy(runs[0].find(qn("w:rPr"))) if runs else None
            for child in list(element):
                if child.tag != qn("w:pPr"):
                    element.remove(child)
            if formatting is not None:
                empty_run = etree.SubElement(element, qn("w:r"))
                empty_run.append(formatting)
            for numbering in element.xpath("./w:pPr/w:numPr"):
                numbering.getparent().remove(numbering)
        if element.tag == qn("w:p"):
            paragraph = Paragraph(element, document)
            old_level = heading_level(paragraph)
            level = block["headingLevel"] if block["kind"] == "heading" else None
            if old_level != level and block["kind"] != "toc":
                properties = element.get_or_add_pPr()
                for item in properties.xpath("./w:outlineLvl"):
                    properties.remove(item)
                outline = etree.SubElement(properties, qn("w:outlineLvl"))
                outline.set(qn("w:val"), str(level - 1 if level is not None else 9))
    output = BytesIO()
    with ZipFile(path) as source, ZipFile(output, "w") as target:
        for entry in source.infolist():
            content = etree.tostring(document.element, xml_declaration=True, encoding="UTF-8", standalone=True) if entry.filename == "word/document.xml" else source.read(entry.filename)
            target.writestr(entry, content)
    return output.getvalue()


# ---------------------------------------------------------------- 上传新 Word（L4）：对齐继承块 ID

def _class(block: dict) -> str:
    return "table" if block["isTable"] else "heading" if block["kind"] == "heading" else "p"


def _norm(text: str) -> str:
    return re.sub(r"\s+", "", text)


def align_upload(new_structure: dict, parent: dict) -> tuple[dict, dict]:
    """17 §3.7 L4：按 (类别, 规范化文字) 序列对齐，标题路径不同的同文字块不配；对上的继承 ID、分类和指导；
    对不上的旧指导变成“未落位”的虚拟指导（headingId=None, unplaced=True），需人工指定位置。"""
    old_top = [block for block in parent["blocks"] if not block.get("tableId") and block.get("path")]
    new_top = [block for block in new_structure["blocks"] if not block["tableId"]]
    old_keys = [(_class(block), _norm(block["sourceText"])) for block in old_top]
    new_keys = [(_class(block), _norm(block["text"])) for block in new_top]
    pairs: dict[int, int] = {}
    for op, a0, a1, b0, b1 in SequenceMatcher(None, old_keys, new_keys, autojunk=False).get_opcodes():
        if op == "equal":
            pairs.update({b0 + offset: a0 + offset for offset in range(a1 - a0)})
        elif op == "replace":
            # 改过的段落：在这一段里按顺序找同类、文字相似度 ≥0.6 的旧块（保持单调，不交叉配对）
            pointer = a0
            for j in range(b0, b1):
                new = new_keys[j]
                for i in range(pointer, a1):
                    old = old_keys[i]
                    if old[0] == new[0] and SequenceMatcher(None, old[1], new[1], autojunk=False).ratio() >= 0.6:
                        pairs[j] = i
                        pointer = i + 1
                        break
    id_map: dict[str, str] = {}
    result = deepcopy(new_structure)
    old_children: dict[str, dict] = {}
    for block in parent["blocks"]:
        if block.get("tableId") and block.get("path"):
            old_children.setdefault(block["tableId"], {})[block["path"].split("/row/", 1)[1]] = block
    new_index = {block["id"]: index for index, block in enumerate(new_top)}
    out = []
    matched_old: set[str] = set()
    for block in result["blocks"]:
        if block["tableId"]:
            continue
        index = new_index[block["id"]]
        original_id = block["id"]
        if index in pairs:
            old = old_top[pairs[index]]
            matched_old.add(old["id"])
            block["id"] = old["id"]
            if _class(old) == _class(block):
                block["kind"] = old["kind"]
                if old["kind"] == "heading":
                    block["headingLevel"] = block["headingLevel"] or old["headingLevel"]
                if old["kind"] == "guidance" and _norm(old["sourceText"]) == _norm(block["text"]):
                    block["text"] = old["text"]
        else:
            block["id"] = new_block_id()
        id_map[original_id] = block["id"]
        out.append(block)
        for child in [item for item in result["blocks"] if item["tableId"] == original_id]:
            suffix = child["path"].split("/row/", 1)[1]
            previous = old_children.get(block["id"], {}).get(suffix)
            child["tableId"] = block["id"]
            if previous and _norm(previous["sourceText"]) == _norm(child["text"]):
                child["id"] = previous["id"]
                child["kind"] = previous["kind"]
                if previous["kind"] == "guidance":
                    child["text"] = previous["text"]
            else:
                child["id"] = new_block_id()
            out.append(child)
    headings_now = {block["id"] for block in out if block["kind"] == "heading"}
    unplaced = []
    for block in parent["blocks"]:
        if block["kind"] != "guidance" or block.get("tableId"):
            continue
        if block.get("path") and block["id"] in matched_old:
            continue
        if not block.get("path") and block.get("headingId") in headings_now:
            # 虚拟指导：所属标题还在，放回该章末尾
            position = max(i for i, item in enumerate(out) if item["id"] == block["headingId"] or item.get("headingId") == block["headingId"])
            out.insert(position + 1, {**deepcopy(block)})
            continue
        unplaced.append({"id": block["id"], "path": None, "kind": "guidance", "text": block["text"], "sourceText": block["text"],
                         "headingLevel": None, "headingId": None, "tableId": None, "protected": False, "isTable": False,
                         "virtual": True, "unplaced": True})
    _rebind_headings(out)
    result["blocks"] = unplaced + out
    alignment = {"matched": len(pairs), "total": len(new_top),
                 "unmatched": [{"blockId": block["id"], "text": block["text"][:200]} for block in unplaced]}
    return result, alignment
