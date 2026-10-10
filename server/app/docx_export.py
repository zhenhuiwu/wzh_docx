"""Apply Markdown edits to the uploaded DOCX without rebuilding its layout."""

from __future__ import annotations

from copy import deepcopy
from dataclasses import dataclass
from difflib import SequenceMatcher
from pathlib import Path
from shutil import copyfile
from io import BytesIO
import base64
import math
import re
from typing import Sequence

import markdown
from lxml import etree
from docx import Document
from docx.oxml.table import CT_Tbl
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.text.paragraph import CT_P
from docx.table import Table, _Row
from docx.text.paragraph import Paragraph
from docx.section import Section
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.image.image import Image
from docx.shared import Inches
from docx.text.run import Run
from lxml import html

from app.docx_parser import parse_docx, _list_info, heading_level


_PROTECTED_INLINE = (
    ".//w:drawing|.//w:pict|.//w:object|.//w:hyperlink|.//w:fldChar|.//w:fldSimple"
    "|.//w:footnoteReference|.//w:endnoteReference|.//m:oMath|.//w:sdt"
    "|.//w:bookmarkStart|.//w:bookmarkEnd|.//w:br[@w:type='page' or @w:type='column']"
)
# 改写段落文字时书签不承载内容，原样保留（界面验收 M1：438C“文档标识号”行只含一个 OLE_LINK 书签）。
# 删除整段时仍按 _PROTECTED_INLINE 保守拒绝（书签可能被 REF 域引用）。
_PROTECTED_EDIT = (
    ".//w:drawing|.//w:pict|.//w:object|.//w:hyperlink|.//w:fldChar|.//w:fldSimple"
    "|.//w:footnoteReference|.//w:endnoteReference|.//m:oMath|.//w:sdt"
    "|.//w:br[@w:type='page' or @w:type='column']"
)


class WordEditError(ValueError):
    """A Word write-back failure tied to one Markdown top-level block (01 §4.10)."""

    def __init__(self, message: str, block_index: int | None, removed_text: str | None = None):
        super().__init__(message)
        self.block_index = block_index
        # 界面复验 N1：被拒的是“删除整段”时，被删段落在原稿里的文字（编辑稿里已没有它），供“撤销这一段”按原内容找回。
        self.removed_text = removed_text


class _RemoveRefused(ValueError):
    def __init__(self, message: str, removed_text: str):
        super().__init__(message)
        self.removed_text = removed_text


def _remove_element(element) -> None:
    if element.xpath(_PROTECTED_INLINE + "|.//w:sectPr|.//w:gridSpan|.//w:vMerge"):
        raise ValueError("此内容包含图片、域、书签、合并单元格或分节信息，无法安全删除；请在原稿中修改。")
    element.getparent().remove(element)


@dataclass
class _Block:
    tag: str
    elements: list[CT_P | CT_Tbl]


def _html_blocks(content: str) -> list:
    root = html.fragment_fromstring(
        markdown.markdown(content, extensions=["tables", "fenced_code"]),
        create_parent="div",
    )
    # 工作台重塑轮：正文里的 HTML 注释在 Word 中没有对应内容，整体去掉（保留其后的文字），
    # 不再把 lxml 的英文报错（“Input object is not an XML element: HtmlComment”）抛给用户，也不把注释文字写进 Word。
    for comment in list(root.iter(etree.Comment)):
        comment.drop_tree()
    # 只写“#”的空标题：Word 里不会产生可对应的标题段（预览因此 422），按不存在处理；编辑区重绘后这一行随之消失。
    return [child for child in root if isinstance(child.tag, str)
            and not (child.tag in _HEADING_TAGS and not "".join(child.itertext()).strip() and not child.xpath(".//img"))]


_HEADING_TAGS = frozenset({"h1", "h2", "h3", "h4", "h5", "h6"})


def _word_blocks(doc, *, flat_lists: bool = False, outline_headings: bool = False) -> list[_Block]:
    blocks: list[_Block] = []
    for element in doc.element.body:
        if isinstance(element, CT_Tbl):
            blocks.append(_Block("table", [element]))
        elif isinstance(element, CT_P):
            paragraph = Paragraph(element, doc)
            text = paragraph.text.strip()
            listing = _list_info(paragraph)
            if (level := heading_level(paragraph, outline_headings=outline_headings)) is not None and text:
                blocks.append(_Block(f"h{level}", [element]))
            elif listing and text:
                tag, depth = listing
                if flat_lists:
                    depth = 0
                if blocks and blocks[-1].tag in {"ul", "ol"} and (flat_lists or blocks[-1].tag == tag or depth > 0):
                    blocks[-1].elements.append(element)
                else:
                    blocks.append(_Block(tag, [element]))
            elif text:
                blocks.append(_Block("p", [element]))
            for _ in element.xpath(".//w:drawing//a:blip[@r:embed]"):
                blocks.append(_Block("p", [element]))
    return blocks


def _text(element) -> str:
    return " ".join("".join(element.itertext()).split())


def _list_items(element, depth=0):
    """Yield independent list items without folding child lists into their text."""
    if depth > 8:
        raise ValueError("Word 列表最多支持九级，请减少列表嵌套层级。")
    for item in element.xpath("./li"):
        if len(item.xpath("./p")) > 1:
            raise ValueError("列表项目包含多个正文段落，无法安全保持列表版式。")
        own = html.Element("li")
        own.text = item.text if item.text and item.text.strip() else None
        for child in item:
            if child.tag in {"ul", "ol"}:
                if child.tail and child.tail.strip():
                    if len(own):
                        own[-1].tail = (own[-1].tail or "") + child.tail
                    else:
                        own.text = (own.text or "") + child.tail
                continue
            copied = deepcopy(child)
            if copied.tail and not copied.tail.strip():
                copied.tail = None
            own.append(copied)
        yield own, depth, element.tag, element
        for child in item.xpath("./ul|./ol"):
            yield from _list_items(child, depth + 1)


def _new_list_numbering(doc, kind: str) -> int:
    """Allocate a native multilevel definition for one newly inserted HTML list."""
    numbering = doc.part.numbering_part.element
    abstract_id = max([int(value) for value in numbering.xpath("./w:abstractNum/@w:abstractNumId")] + [-1]) + 1
    number_id = max([int(value) for value in numbering.xpath("./w:num/@w:numId")] + [0]) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    mode = OxmlElement("w:multiLevelType")
    mode.set(qn("w:val"), "multilevel")
    abstract.append(mode)
    for depth in range(9):
        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), str(depth))
        for tag, value in (("start", "1"), ("numFmt", "decimal" if kind == "ol" else "bullet"),
                           ("lvlText", f"%{depth + 1}." if kind == "ol" else ("•", "○", "▪")[depth % 3]),
                           ("lvlJc", "left")):
            child = OxmlElement("w:" + tag)
            child.set(qn("w:val"), value)
            level.append(child)
        properties, tabs, tab, indent = (OxmlElement("w:" + tag) for tag in ("pPr", "tabs", "tab", "ind"))
        left = 720 + depth * 360
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), str(left))
        tabs.append(tab)
        properties.append(tabs)
        indent.set(qn("w:left"), str(left))
        indent.set(qn("w:hanging"), "360")
        properties.append(indent)
        level.append(properties)
        abstract.append(level)
    existing_numbers = numbering.xpath("./w:num")
    if existing_numbers:
        existing_numbers[0].addprevious(abstract)
    else:
        numbering.append(abstract)
    number = OxmlElement("w:num")
    number.set(qn("w:numId"), str(number_id))
    reference = OxmlElement("w:abstractNumId")
    reference.set(qn("w:val"), str(abstract_id))
    number.append(reference)
    numbering.append(number)
    return number_id


def _set_list_level(paragraph: Paragraph, kind: str, depth: int, number_id: int) -> None:
    paragraph.style = "List Number" if kind == "ol" else "List Bullet"
    properties = paragraph._p.get_or_add_pPr()
    for old in properties.xpath("./w:numPr|./w:ind|./w:tabs"):
        properties.remove(old)
    numbering = properties.get_or_add_numPr()
    numbering.get_or_add_ilvl().val = depth
    numbering.get_or_add_numId().val = number_id


def _key(element) -> tuple:
    tag = element.tag
    if tag == "table":
        rows = element.xpath(".//tr")
        content = tuple(tuple(_text(cell) for cell in row.xpath("./th|./td")) for row in rows)
    elif tag in {"ul", "ol"}:
        content = tuple((depth, kind, _key(item)) for item, depth, kind, _ in _list_items(element))
    else:
        content = _text(element)
    images = tuple(image.get("src") for image in element.xpath(".//img"))
    inline = tuple((node.tag, _text(node)) for node in element.iterdescendants()
                   if node.tag in {"strong", "em", "code", "a"})
    return tag, content, images, inline


def _inline_chunks(element, bold=False, italic=False):
    bold = bold or element.tag == "strong"
    italic = italic or element.tag == "em"
    if element.text:
        yield element.text, bold, italic
    for child in element:
        if child.tag == "br":
            yield "\n", bold, italic
        else:
            yield from _inline_chunks(child, bold, italic)
        if child.tail:
            # Python-Markdown emits a serialization newline after <br />.
            # The break already supplied that newline; keeping both creates
            # an extra blank line in Word after every Markdown hard break.
            tail = child.tail.removeprefix("\n") if child.tag == "br" else child.tail
            if tail:
                yield tail, bold, italic


def _anchored_operations(old: Sequence, new: Sequence):
    """Anchor unchanged boundaries so repeated words cannot steal their run style."""
    prefix = 0
    while prefix < min(len(old), len(new)) and old[prefix] == new[prefix]:
        prefix += 1
    suffix = 0
    while suffix < min(len(old), len(new)) - prefix and old[-1 - suffix] == new[-1 - suffix]:
        suffix += 1
    if prefix:
        yield "equal", 0, prefix, 0, prefix
    for operation, start, end, new_start, new_end in SequenceMatcher(
        None, old[prefix:len(old) - suffix], new[prefix:len(new) - suffix], autojunk=False
    ).get_opcodes():
        yield operation, start + prefix, end + prefix, new_start + prefix, new_end + prefix
    if suffix:
        yield "equal", len(old) - suffix, len(old), len(new) - suffix, len(new)


def _edit_paragraph(paragraph: Paragraph, target, source=None) -> None:
    new_text = "".join(text for text, _, _ in _inline_chunks(target))
    source_inline = _key(source)[-1] if source is not None else ()
    target_inline = _key(target)[-1]
    if paragraph.text == new_text and source_inline == target_inline:
        return
    if paragraph._p.xpath(_PROTECTED_EDIT):
        raise ValueError("此段包含图片、链接、域或复杂 Word 内容，无法安全替换；请在原稿中修改该段。")

    runs = paragraph.runs
    # 复审：段内有书签时也走按偏移重建的路径，书签保持原范围（段首/书签边界处加字不再被并进书签）。
    if source_inline != target_inline or paragraph._p.xpath("./w:bookmarkStart|./w:bookmarkEnd"):
        _edit_inline_runs(paragraph, target)
        return

    if not runs:
        paragraph.add_run(new_text)
        return
    old_text = "".join(run.text for run in runs)
    if old_text != paragraph.text:
        raise ValueError("此段包含复杂 Word 内容，无法安全替换；请在原稿中修改该段。")
    owners = [index for index, run in enumerate(runs) for _ in run.text]
    pieces = [""] * len(runs)
    for operation, start, end, new_start, new_end in _anchored_operations(old_text, new_text):
        if operation == "equal":
            for old_index in range(start, end):
                pieces[owners[old_index]] += old_text[old_index]
        elif operation in {"insert", "replace"}:
            owner = owners[start] if start < len(owners) else (owners[-1] if owners else 0)
            pieces[owner] += new_text[new_start:new_end]
    for run, value in zip(runs, pieces):
        if not value and run.text:
            run._r.getparent().remove(run._r)
        else:
            run.text = value


def _edit_inline_runs(paragraph: Paragraph, target) -> None:
    """Retain each text span's original run properties when adding emphasis."""
    runs = paragraph.runs
    old_text = "".join(run.text for run in runs)
    if old_text != paragraph.text:
        raise ValueError("此段包含复杂 Word 内容，无法安全替换；请在原稿中修改该段。")
    target_chars = [(char, bold, italic) for text, bold, italic in _inline_chunks(target) for char in text]
    new_text = "".join(char for char, _, _ in target_chars)
    owners = [index for index, run in enumerate(runs) for _ in run.text]
    new_owners = [0] * len(new_text)
    for operation, start, end, new_start, new_end in _anchored_operations(old_text, new_text):
        if operation == "equal":
            new_owners[new_start:new_end] = owners[start:end]
        elif operation in {"insert", "replace"}:
            owner = owners[start] if start < len(owners) else owners[-1] if owners else 0
            new_owners[new_start:new_end] = [owner] * (new_end - new_start)
    marks = _bookmark_marks(paragraph, old_text, new_text)
    cuts = {offset for offset, _ in marks}
    groups = []
    for position, ((char, bold, italic), owner) in enumerate(zip(target_chars, new_owners)):
        key = (owner, bold, italic)
        if groups and groups[-1][0] == key and position not in cuts:
            groups[-1][1] += char
        else:
            groups.append([key, char])
    paragraph.clear()
    pending = list(marks)
    offset = 0

    def flush(upto: int) -> None:
        while pending and pending[0][0] <= upto:
            paragraph._p.append(pending.pop(0)[1])

    for (owner, bold, italic), text in groups:
        flush(offset)
        offset += len(text)
        original = runs[owner] if runs else None
        if original is not None and text == original.text and (not bold or original.bold is True) and (not italic or original.italic is True):
            paragraph._p.append(deepcopy(original._r))
            continue
        run = paragraph.add_run(text)
        if original is not None and original._r.rPr is not None:
            run._r.insert(0, deepcopy(original._r.rPr))
        if bold:
            run.bold = True
        if italic:
            run.italic = True
    flush(len(new_text))


_BOOKMARK_FREE = ("_Toc", "_GoBack", "OLE_LINK")


def _referenced_bookmarks(paragraph: Paragraph) -> set[str]:
    root = paragraph._p.getroottree().getroot()
    texts = [node.text or "" for node in root.iter(qn("w:instrText"))]
    texts += [node.get(qn("w:instr")) or "" for node in root.iter(qn("w:fldSimple"))]
    names = {node.get(qn("w:anchor")) for node in root.iter(qn("w:hyperlink")) if node.get(qn("w:anchor"))}
    for text in texts:
        names.update(re.findall(r"\b(?:REF|PAGEREF|NOTEREF)\s+(\S+)", text))
        names.update(re.findall(r"\bHYPERLINK\s+\\l\s+\"?([^\s\"]+)", text))
    return names


def _bookmark_marks(paragraph: Paragraph, old_text: str, new_text: str) -> list[tuple[int, object]]:
    """复审：改格式（整段重建 run）时书签保持原范围。

    记下每个书签开始/结束标记在原文字里的偏移，按新旧文字的差异映射到新文字里，重建时插回对应位置：
    落在未改动的文字上就精确对应；落在被替换/删除的文字里时，开始往前、结束往后取（范围只扩不缩）。
    被交叉引用（REF/PAGEREF/NOTEREF/超链接锚点）的书签若边界落在被改的文字里、范围说不清，按受保护处理；
    嵌在其他元素里的书签（重建会丢）同样按受保护处理。_Toc/_GoBack/OLE_LINK 不承载引用内容，不受此限。
    """
    p = paragraph._p
    if p.xpath("./*[not(self::w:bookmarkStart or self::w:bookmarkEnd)][.//w:bookmarkStart or .//w:bookmarkEnd]"):
        raise ValueError("此段的书签嵌在其他 Word 结构里，改格式会丢失书签；请在原稿中修改该段。")
    found = []
    offset = 0
    for child in p:
        if child.tag == qn("w:r"):
            offset += len(Run(child, paragraph).text)
        elif child.tag in (qn("w:bookmarkStart"), qn("w:bookmarkEnd")):
            found.append((offset, child))
    if not found:
        return []
    operations = list(_anchored_operations(old_text, new_text))
    names_by_id = {node.get(qn("w:id")): node.get(qn("w:name")) or "" for node in p.getroottree().getroot().iter(qn("w:bookmarkStart"))}
    # 测试复测 N3：整段文字全被换掉（例如行首退格并入上一段后又回车新写一段）时，书签会落到无关的新内容上。
    # 标题目录书签（_Toc/_GoBack）随标题改名是正常的；其他书签按受保护处理。
    # 判定“整段换掉”：首尾都没有保留原文字，且保留下来的零散字符不到原文一半。
    kept = sum(end - start for operation, start, end, _, _ in operations if operation == "equal")
    anchored = operations and (operations[0][0] == "equal" or operations[-1][0] == "equal")
    if old_text.strip() and not anchored and kept * 2 < len(old_text):
        names = [node.get(qn("w:name")) if node.tag == qn("w:bookmarkStart") else names_by_id.get(node.get(qn("w:id")), "")
                 for _, node in found]
        if any(not (name or "").startswith(("_Toc", "_GoBack")) for name in names):
            raise ValueError("此段带书签，整段文字都被替换后书签会落到新内容上；请只修改其中的文字，或在原稿中修改该段。")
    referenced = None
    marks = []
    for order, (old_offset, node) in enumerate(found):
        is_start = node.tag == qn("w:bookmarkStart")
        mapped, exact = None, False
        for operation, start, end, new_start, new_end in operations:
            if operation == "equal" and start <= old_offset <= end:
                mapped, exact = new_start + old_offset - start, True
                break
        if mapped is None:
            for operation, start, end, new_start, new_end in operations:
                if start <= old_offset <= end:
                    mapped = new_start if is_start else new_end
                    break
        if mapped is None:
            mapped = len(new_text)
        if not exact and old_offset not in (0, len(old_text)):
            name = node.get(qn("w:name")) if is_start else names_by_id.get(node.get(qn("w:id")), "")
            if not (name or "").startswith(_BOOKMARK_FREE):
                referenced = _referenced_bookmarks(paragraph) if referenced is None else referenced
                if name in referenced:
                    raise ValueError("此段被交叉引用的书签边界处文字已改动，无法确定书签范围；请在原稿中修改该段。")
        marks.append((mapped, order, deepcopy(node)))
    marks.sort(key=lambda item: (item[0], item[1]))
    return [(mapped, node) for mapped, _, node in marks]


def _edit_table(table: Table, target, source) -> None:
    target_rows = target.xpath(".//tr")
    source_rows = source.xpath(".//tr")
    word_rows = list(table.rows)
    if len(word_rows) != len(source_rows):
        raise ValueError("原表格结构无法对应，未修改表格。")
    columns = len(source_rows[0].xpath("./th|./td"))
    if any(len(row.xpath("./th|./td")) != columns for row in target_rows):
        raise ValueError("表格列数已改变，无法保留原表格版式。")

    def edit_row(word_row, old_row, new_row, row_index):
        old_cells = old_row.xpath("./th|./td")
        new_cells = new_row.xpath("./th|./td")
        if len(word_row.cells) != len(new_cells) or len(old_cells) != len(new_cells):
            raise ValueError("表格列数已改变，无法保留原表格版式。")
        for cell_index, (word_cell, old_cell, new_cell) in enumerate(zip(word_row.cells, old_cells, new_cells)):
            if _key(old_cell) == _key(new_cell):
                continue
            if word_cell._tc.xpath("./w:tcPr/w:gridSpan|./w:tcPr/w:vMerge"):
                raise ValueError("不能改写合并单元格，请保留原表头和表格结构。")
            paragraphs = [p for p in word_cell.paragraphs if p.text.strip()]
            if len(paragraphs) > 1:
                raise ValueError(f"表格第 {row_index + 1} 行第 {cell_index + 1} 列包含多段文字，无法安全替换。")
            _edit_paragraph(paragraphs[0] if paragraphs else word_cell.paragraphs[0], new_cell, old_cell)

    def row_key(row):
        return tuple(_key(cell) for cell in row.xpath("./th|./td"))

    for operation, start, end, new_start, new_end in reversed(SequenceMatcher(
        None, [row_key(row) for row in source_rows], [row_key(row) for row in target_rows], autojunk=False
    ).get_opcodes()):
        if operation == "equal":
            continue
        old, new = word_rows[start:end], target_rows[new_start:new_end]
        if len(old) != len(new) and any(row._tr.xpath(".//w:gridSpan|.//w:vMerge") for row in old):
            raise ValueError("不能增删含合并单元格的行，请保留原表头和表格结构。")
        for offset in range(min(len(old), len(new))):
            edit_row(old[offset], source_rows[start + offset], new[offset], new_start + offset)
        for row in old[len(new):]:
            _remove_element(row._tr)
        if len(new) > len(old):
            candidates = list(range(start, end)) + list(range(start - 1, 0, -1)) + list(range(end, len(word_rows)))
            prototype_index = next((index for index in candidates
                if not word_rows[index]._tr.xpath(".//w:gridSpan|.//w:vMerge|.//w:drawing|.//w:fldChar")), None)
            if prototype_index is None:
                raise ValueError("表格没有可复制的数据行，无法保留表格版式。")
            prototype = word_rows[prototype_index]
            anchor = next((row._tr for row in word_rows[end:] if row._tr.getparent() is table._tbl), None)
            for offset, new_row in enumerate(new[len(old):], len(old)):
                element = deepcopy(prototype._tr)
                if anchor is None:
                    table._tbl.append(element)
                else:
                    anchor.addprevious(element)
                # The prototype may have just been edited; compare against its current text.
                old_html = deepcopy(source_rows[prototype_index])
                for cell, html_cell in zip(prototype.cells, old_html.xpath("./th|./td")):
                    html_cell.clear()
                    html_cell.text = cell.text
                edit_row(_Row(element, table), old_html, new_row, new_start + offset)


def _edit_block(doc, block: _Block, target, source) -> None:
    if block.tag == "table":
        _edit_table(Table(block.elements[0], doc), target, source)
    elif block.tag in {"ul", "ol"}:
        items = list(_list_items(target))
        old_items = list(_list_items(source))
        if len(old_items) != len(block.elements):
            raise ValueError("原列表结构无法对应，未修改列表。")
        number_ids = {}
        def number_id(container, kind):
            if container not in number_ids:
                number_ids[container] = _new_list_numbering(doc, kind)
            return number_ids[container]
        for element, (new_item, depth, kind, container), (old_item, old_depth, old_kind, _) in zip(block.elements, items, old_items):
            paragraph = Paragraph(element, doc)
            _edit_paragraph(paragraph, new_item, old_item)
            if (depth, kind) != (old_depth, old_kind):
                _set_list_level(paragraph, kind, depth, number_id(container, kind))
        for element in block.elements[len(items):]:
            _remove_element(element)
        if len(items) > len(block.elements):
            prototype = block.elements[-1]
            current = prototype
            for item, depth, kind, container in items[len(block.elements):]:
                element = deepcopy(prototype)
                current.addnext(element)
                paragraph = Paragraph(element, doc)
                old_html = html.Element("li")
                old_html.text = paragraph.text
                _edit_paragraph(paragraph, item, old_html)
                # Continue an unchanged flat list's original numbering; nested
                # items receive their own real level rather than soft line breaks.
                if (depth, kind) != (_list_info(Paragraph(prototype, doc)) or (block.tag, 0)):
                    _set_list_level(paragraph, kind, depth, number_id(container, kind))
                current = element
    elif len(block.elements) == 1:
        _edit_paragraph(Paragraph(block.elements[0], doc), target, source)
    else:
        raise ValueError("此段包含复合内容，无法安全替换。")


def _resize_plain_paragraphs(doc, old_blocks: list[_Block], old_sources: list, count: int) -> tuple[list[_Block], list]:
    """Keep native runs available to the text diff when paragraphs merge or split."""
    if any(element.xpath(_PROTECTED_INLINE + "|.//w:sectPr")
           for block in old_blocks for element in block.elements):
        raise ValueError("选区包含图片、链接、域或分节信息，无法安全跨段替换；请缩小选区。")
    if len(old_blocks) > count:
        merged = old_blocks[count - 1:]
        paragraph = Paragraph(merged[0].elements[0], doc)
        if count > 1:
            properties = merged[-1].elements[0].pPr
            if paragraph._p.pPr is not None:
                paragraph._p.remove(paragraph._p.pPr)
            if properties is not None:
                paragraph._p.insert(0, deepcopy(properties))
        combined = deepcopy(old_sources[count - 1])
        for block, source in zip(merged[1:], old_sources[count:]):
            paragraph.add_run("\n\n")
            for child in block.elements[0]:
                if child.tag != qn("w:pPr"):
                    paragraph._p.append(deepcopy(child))
            if len(combined):
                combined[-1].tail = (combined[-1].tail or "") + "\n\n" + (source.text or "")
            else:
                combined.text = (combined.text or "") + "\n\n" + (source.text or "")
            for child in source:
                combined.append(deepcopy(child))
            _remove_element(block.elements[0])
        old_blocks, old_sources = old_blocks[:count], old_sources[:count - 1] + [combined]
    elif len(old_blocks) < count:
        prototype = old_blocks[-1].elements[0]
        current = prototype
        while len(old_blocks) < count:
            cloned = deepcopy(prototype)
            current.addnext(cloned)
            old_blocks.append(_Block("p", [cloned]))
            old_sources.append(deepcopy(old_sources[-1]))
            current = cloned
    return old_blocks, old_sources


def crop_picture(document, shape, rect: dict | None) -> None:
    """Crop only Word's display; keep the original complete image part."""
    if not isinstance(rect, dict):
        return
    values = [rect.get(key) for key in ("x", "y", "width", "height")]
    if any(isinstance(value, bool) or not isinstance(value, (int, float)) or abs(value) > 1e9 or not math.isfinite(value) for value in values):
        return
    x, y, width, height = values
    blip_fill = shape._inline.graphic.graphicData.pic.blipFill
    image = document.part.related_parts[blip_fill.blip.embed].image
    pixel_width, pixel_height = image.px_width, image.px_height
    if x < 0 or y < 0 or width <= 0 or height <= 0 or x + width > pixel_width or y + height > pixel_height:
        return
    left, top, right, bottom = (round(x / pixel_width * 100000), round(y / pixel_height * 100000),
                              round((pixel_width - x - width) / pixel_width * 100000),
                              round((pixel_height - y - height) / pixel_height * 100000))
    if left + right >= 100000 or top + bottom >= 100000:
        return
    crop = OxmlElement("a:srcRect")
    for key, value in (("l", left), ("t", top), ("r", right), ("b", bottom)):
        crop.set(key, str(value))
    blip_fill.insert(1, crop)
    shape.height = round(shape.width * height / width)


def diagram_images(images: list[dict]) -> dict[str, bytes | dict]:
    """Only server-owned diagram/test PNGs may be inserted; never fetch URLs."""
    sources = {}
    for image in images:
        ref = image.get("ref", "")
        if not ref.startswith(("docx-image:diagram-", "docx-image:test-")):
            continue
        try:
            data = base64.b64decode(image["data"], validate=True)
        except (ValueError, KeyError):
            raise ValueError("文档图片数据不正确。") from None
        if not data.startswith(b"\x89PNG\r\n\x1a\n") or len(data) > 20 * 1024 * 1024:
            raise ValueError("文档图片必须是有效的 PNG 图片。")
        sources[ref] = {"data": data, "focusRect": image.get("focusRect")} if ref.startswith("docx-image:test-") else data
    return sources


def _insert_block(doc, target, anchor, style_override: str | None = None, image_sources: dict[str, bytes | dict] | None = None,
                  *, list_depth: int | None = None, list_kind: str | None = None, list_number_id: int | None = None,
                  outline_headings: bool = False) -> None:
    tag = target.tag
    if target.xpath(".//img"):
        images = target.xpath(".//img")
        if tag != "p" or len(images) != 1 or _text(target) or images[0].get("src") not in (image_sources or {}):
            raise ValueError("只能插入此文档的图表或测试截图，不能使用外部图片地址。")
        source = image_sources[images[0].get("src")]
        data = source["data"] if isinstance(source, dict) else source
        image = Image.from_blob(data)
        section = doc.sections[-1]
        for element in [anchor, *anchor.itersiblings()]:
            boundaries = element.xpath("./w:pPr/w:sectPr")
            if boundaries:
                section = Section(boundaries[0], doc.part)
                break
        width = section.page_width - section.left_margin - section.right_margin
        height = section.page_height - section.top_margin - section.bottom_margin - Inches(1)
        width = min(width, int(height * image.px_width / image.px_height))
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture = paragraph.add_run().add_picture(BytesIO(data), width=width)
        picture._inline.docPr.set("descr", images[0].get("alt") or "文档图表")
        if isinstance(source, dict) and images[0].get("src", "").startswith("docx-image:test-"):
            crop_picture(doc, picture, source.get("focusRect"))
        anchor.addprevious(paragraph._p)
        return
    if tag == "table":
        rows = target.xpath(".//tr")
        columns = len(rows[0].xpath("./th|./td")) if rows else 0
        if not rows or not columns:
            return
        if any(len(row.xpath("./th|./td")) != columns for row in rows):
            raise ValueError("表格列数不一致，无法安全写入 Word。")
        table = doc.add_table(rows=len(rows), cols=columns)
        nearby = [element for element in [*anchor.itersiblings(preceding=True), *anchor.itersiblings()]
                  if isinstance(element, CT_Tbl) and element is not table._tbl]
        prototype = next((Table(element, doc) for element in nearby
                          if len(Table(element, doc).columns) == columns), None)
        table.style = prototype.style if prototype else "Table Grid"
        if prototype is not None:
            table._tbl.replace(table._tbl.tblPr, deepcopy(prototype._tbl.tblPr))
            table._tbl.replace(table._tbl.tblGrid, deepcopy(prototype._tbl.tblGrid))
        for row_index, (word_row, row) in enumerate(zip(table.rows, rows)):
            for column_index, (cell, html_cell) in enumerate(zip(word_row.cells, row.xpath("./th|./td"))):
                if prototype is not None:
                    original = prototype.cell(0 if row_index == 0 else min(1, len(prototype.rows) - 1), column_index)
                    if original._tc.tcPr is not None:
                        cell._tc.replace(cell._tc.tcPr, deepcopy(original._tc.tcPr))
                        for merge in cell._tc.tcPr.xpath("./w:gridSpan|./w:vMerge"):
                            cell._tc.tcPr.remove(merge)
                    cell.width = prototype.columns[column_index].width
                    paragraph = cell.paragraphs[0]
                    sample = original.paragraphs[0]
                    if sample._p.pPr is not None:
                        paragraph._p.insert(0, deepcopy(sample._p.pPr))
                    run = paragraph.add_run()
                    if sample.runs and sample.runs[0]._r.rPr is not None:
                        run._r.insert(0, deepcopy(sample.runs[0]._r.rPr))
                _edit_paragraph(cell.paragraphs[0], html_cell)
                if html_cell.tag == "th":
                    for run in cell.paragraphs[0].runs:
                        run.bold = True
        element = table._tbl
    elif tag in {"p", "h1", "h2", "h3", "h4", "h5", "h6", "ul", "ol"}:
        if tag in {"ul", "ol"}:
            number_ids = {}
            for item, depth, kind, container in _list_items(target):
                if container not in number_ids:
                    number_ids[container] = _new_list_numbering(doc, kind)
                paragraph_html = html.Element("p")
                paragraph_html.text = item.text
                for child in item:
                    paragraph_html.append(deepcopy(child))
                _insert_block(doc, paragraph_html, anchor,
                              "List Number" if kind == "ol" else "List Bullet", image_sources,
                              list_depth=depth, list_kind=kind, list_number_id=number_ids[container],
                              outline_headings=outline_headings)
            return
        style = style_override or (f"Heading {tag[1:]}" if tag.startswith("h") else "Normal")
        # Inherit from this part of the document, rather than its cover's first Normal paragraph.
        nearby = list(anchor.itersiblings(preceding=True)) + list(anchor.itersiblings())
        paragraphs = [Paragraph(element, doc) for element in nearby if isinstance(element, CT_P)]
        sample = next((p for p in paragraphs if (p.text or p.runs) and
                       (p.style.name == style if style_override else
                        heading_level(p, outline_headings=outline_headings) == int(tag[1:]) if tag.startswith("h")
                        else heading_level(p, outline_headings=outline_headings) is None and not p.style.name.startswith(("List", "toc", "TOC", "Caption", "Title", "Subtitle"))
                        and p.alignment != WD_ALIGN_PARAGRAPH.CENTER
                        and not re.match(r"^[图表]\s*\d+(?:[.－—-]\d+)*\s*", p.text.strip())
                        and not p._p.xpath(_PROTECTED_INLINE))), None)
        if sample is not None and not style_override:
            style = sample.style.name
        paragraph = doc.add_paragraph(style=style)
        if sample is not None and sample._p.pPr is not None:
            if paragraph._p.pPr is not None:
                paragraph._p.remove(paragraph._p.pPr)
            properties = deepcopy(sample._p.pPr)
            # A new paragraph inherits formatting, never a new section boundary.
            for boundary in properties.xpath("./w:sectPr"):
                properties.remove(boundary)
            if list_depth is None and tag == "p" and not style_override:
                for numbering in properties.xpath("./w:numPr"):
                    properties.remove(numbering)
            paragraph._p.insert(0, properties)
        if list_depth is not None:
            _set_list_level(paragraph, list_kind, list_depth, list_number_id)
        # Character emphasis belongs to spans, not to the whole body prototype.
        sample_run = next((run for run in sorted(sample.runs, key=lambda item: len(item.text), reverse=True)
                           if run.text and run.bold is not True and run.italic is not True), sample.runs[0] if sample.runs else None) if sample is not None else None
        for text, bold, italic in _inline_chunks(target):
            run = paragraph.add_run(text)
            if sample_run is not None and sample_run._r.rPr is not None:
                properties = deepcopy(sample_run._r.rPr)
                if not tag.startswith("h"):
                    for emphasis in properties.xpath("./w:b|./w:bCs|./w:i|./w:iCs"):
                        properties.remove(emphasis)
                run._r.insert(0, properties)
            if bold:
                run.bold = True
            if italic:
                run.italic = True
        element = paragraph._p
    else:
        raise ValueError("文档包含暂不支持的排版块，未导出以避免破坏格式。")
    anchor.addprevious(element)


def edit_original_docx(markdown_content: str, original_path: str | Path, output_path: str | Path, *, image_sources: dict[str, bytes | dict] | None = None, outline_headings: bool = False) -> None:
    """Retain untouched OOXML and patch only edited blocks."""
    original_path = Path(original_path)
    doc = Document(original_path)
    source = _html_blocks(parse_docx(str(original_path), outline_headings=outline_headings).markdown)
    target = _html_blocks(markdown_content)
    blocks = _word_blocks(doc, outline_headings=outline_headings)
    source_levels = [depth for element in source if element.tag in {"ul", "ol"}
                     for _, depth, _, _ in _list_items(element)]
    target_levels = [depth for element in target if element.tag in {"ul", "ol"}
                     for _, depth, _, _ in _list_items(element)]
    if any(source_levels) and target_levels and not any(target_levels):
        # Documents imported by the older parser stored every list level flat.
        # Treat missing nesting as that legacy view, not permission to remove
        # the original numbering. Explicit nested target markup uses the current
        # parser. This changes only the matching baseline, never stored content.
        source = _html_blocks(parse_docx(str(original_path), flat_lists=True, outline_headings=outline_headings).markdown)
        blocks = _word_blocks(doc, flat_lists=True, outline_headings=outline_headings)
    if [block.tag for block in blocks] != [element.tag for element in source]:
        raise ValueError("原稿结构无法与编辑稿对应，未导出以避免破坏格式。")

    source_keys = [_key(element) for element in source]
    target_keys = [_key(element) for element in target]
    if source_keys == target_keys:
        copyfile(original_path, output_path)
        return

    test_pictures = []
    for operation, old_start, old_end, new_start, new_end in reversed(
        list(_anchored_operations(source_keys, target_keys))
    ):
        if operation == "equal":
            continue
        try:
            _apply_operation(doc, blocks, source, target, old_start, old_end, new_start, new_end,
                             image_sources, outline_headings, test_pictures)
        except WordEditError:
            raise
        except ValueError as exc:
            raise WordEditError(str(exc), _failing_index(exc, new_start, new_end, len(target)),
                                getattr(exc, "removed_text", None)) from None
    _finish_document(doc, test_pictures)
    doc.save(output_path)


def _failing_index(exc: ValueError, new_start: int, new_end: int, total: int) -> int | None:
    index = getattr(exc, "block_index", None)
    if index is not None:
        return index
    if new_end > new_start:
        return new_start
    return min(new_start, total - 1) if total else None


def _apply_operation(doc, blocks, source, target, old_start, old_end, new_start, new_end,
                     image_sources, outline_headings, test_pictures) -> None:
    old_blocks = blocks[old_start:old_end]
    new_blocks = target[new_start:new_end]
    if (old_blocks and new_blocks and len(old_blocks) != len(new_blocks)
            and all(block.tag == "p" and len(block.elements) == 1 for block in old_blocks)
            and all(element.tag == "p" and not _key(element)[2]
                    for element in source[old_start:old_end] + new_blocks)):
        if len(old_blocks) > len(new_blocks):
            # 界面复验 N1：相邻段刚被改过时，删掉的受保护段会和它落在同一个片段里走“跨段替换”。
            # 受保护段的文字在新内容里已经找不到 → 按“删除”拒绝并带上原文字，界面才能把它插回去。
            kept = "".join("".join("".join(element.itertext()).split()) for element in new_blocks)
            for position, block in enumerate(old_blocks):
                if not any(element.xpath(_PROTECTED_INLINE + "|.//w:sectPr") for element in block.elements):
                    continue
                text = " ".join("".join(source[old_start + position].itertext()).split())
                if text and "".join(text.split()) not in kept:
                    raise _RemoveRefused("此内容包含图片、域、书签、合并单元格或分节信息，无法安全删除；请在原稿中修改。", text)
        resized, previous = _resize_plain_paragraphs(doc, old_blocks, source[old_start:old_end], len(new_blocks))
        for offset, (block, old, new) in enumerate(zip(resized, previous, new_blocks)):
            _located(new_start + offset, _edit_block, doc, block, new, old)
        return
    prefix = 0
    while (prefix < min(len(old_blocks), len(new_blocks))
           and old_blocks[prefix].tag == new_blocks[prefix].tag
           and _key(source[old_start + prefix])[2] == _key(new_blocks[prefix])[2]):
        _located(new_start + prefix, _edit_block, doc, old_blocks[prefix], new_blocks[prefix], source[old_start + prefix])
        prefix += 1
    suffix = 0
    while (suffix < min(len(old_blocks), len(new_blocks)) - prefix
           and old_blocks[-1 - suffix].tag == new_blocks[-1 - suffix].tag
           and _key(source[old_end - 1 - suffix])[2] == _key(new_blocks[-1 - suffix])[2]):
        _located(new_end - 1 - suffix, _edit_block, doc, old_blocks[-1 - suffix], new_blocks[-1 - suffix],
                 source[old_end - 1 - suffix])
        suffix += 1
    old_blocks = old_blocks[prefix:len(old_blocks) - suffix if suffix else len(old_blocks)]
    new_blocks = new_blocks[prefix:len(new_blocks) - suffix if suffix else len(new_blocks)]
    anchor = (blocks[old_end - suffix].elements[0] if suffix else next(
        (element for block in blocks[old_end:] for element in block.elements
         if element.getparent() is not None),
        doc.element.body.sectPr,
    ))
    for position, block in enumerate(old_blocks):
        for element in block.elements:
            if element.getparent() is not None:
                try:
                    _remove_element(element)
                except ValueError as exc:
                    removed = " ".join("".join(source[old_start + prefix + position].itertext()).split())
                    raise _RemoveRefused(str(exc), removed) from None
    for offset, element in enumerate(new_blocks):
        _located(new_start + prefix + offset, _insert_block, doc, element, anchor,
                 image_sources=image_sources, outline_headings=outline_headings)
        if element.tag == "p" and element.xpath('./img[starts-with(@src,"docx-image:test-")]'):
            test_pictures.append(anchor.getprevious())


def _located(index: int, action, *args, **kwargs) -> None:
    try:
        action(*args, **kwargs)
    except WordEditError:
        raise
    except ValueError as exc:
        raise WordEditError(str(exc), index) from None


def _finish_document(doc, test_pictures) -> None:
    # Apply after insertion so later body paragraphs do not inherit this binding.
    for picture in test_pictures:
        preceding = list(picture.itersiblings(preceding=True))[:2]
        if len(preceding) == 2 and all(isinstance(element, CT_P) for element in preceding):
            actual, caption = (Paragraph(element, doc) for element in preceding)
            if re.match(r"^步骤 \d+：", caption.text):
                caption.paragraph_format.keep_with_next = actual.paragraph_format.keep_with_next = True
    for table in doc.tables:
        if len(table.rows) < 2 or len(table.columns) < 2 or table._tbl.xpath(".//w:gridSpan|.//w:vMerge"):
            continue
        header = table.rows[0]
        runs = [run for cell in header.cells for paragraph in cell.paragraphs for run in paragraph.runs if run.text.strip()]
        if not runs or not all(run.bold is True for run in runs):
            continue
        properties = header._tr.get_or_add_trPr()
        if not properties.xpath("./w:tblHeader"):
            properties.append(OxmlElement("w:tblHeader"))
        for cell in header.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True
