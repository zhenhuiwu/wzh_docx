"""Project the edited DOCX's typography onto its editable Markdown structure.

This is a continuous body preview, not a Word pagination engine.  The Markdown
tags remain the save contract: native Word emphasis is CSS, while emphasis the
user explicitly stored in Markdown remains ``strong``/``em``.
"""

from __future__ import annotations

import base64
import json
import re
from pathlib import Path
from tempfile import TemporaryDirectory
from urllib.parse import urlsplit

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from lxml import etree, html

from app.fast_diff import bounded_opcodes
from app.docx_export import (
    _html_blocks, _list_items, _word_blocks,
    diagram_images, edit_original_docx,
)
from app.docx_parser import _list_info, parse_docx

_ALLOWED_TAGS = {
    "p", "h1", "h2", "h3", "h4", "h5", "h6", "ul", "ol", "li", "table",
    "thead", "tbody", "tr", "th", "td", "br", "strong", "em", "code", "pre",
    "blockquote", "a", "img", "hr", "del", "s", "span",
}
_FIELD_LIKE = (
    ".//w:fldChar|.//w:fldSimple|.//w:hyperlink|.//w:footnoteReference|.//w:endnoteReference"
    "|.//m:oMath|.//w:sdt|.//w:br[@w:type='page' or @w:type='column']"
)
# 测试 L4：Python-Markdown 会把制表符展开成空格，编辑区再存回去时 \t 就变成了空格。
# 行内（行首非空白之后）的制表符先换成非字符占位，渲染后在文本里还原成 \t（行首缩进不动）。
# 复审：正文里本来就有的非字符不能被当成 tab 还原——每次从候选里挑一个正文没出现过的；
# 候选全被占用（几乎不可能）时干脆不替换，退回旧行为（tab 展开为空格），不会改动正文原有字符。
_TAB_MARK_CANDIDATES = tuple(chr(code) for code in range(0xFDEF, 0xFDCF, -1)) + tuple(
    chr(plane * 0x10000 + 0xFFFE + low) for plane in range(1, 17) for low in (0, 1)
)


def _pick_tab_mark(markdown_text: str) -> str | None:
    return next((mark for mark in _TAB_MARK_CANDIDATES if mark not in markdown_text), None)


def _hide_inline_tabs(markdown_text: str, mark: str | None) -> str:
    if mark is None:
        return markdown_text

    def line(text: str) -> str:
        stripped = text.lstrip(" \t")
        lead = text[:len(text) - len(stripped)]
        return lead + stripped.replace("\t", mark)
    return "".join(line(part) for part in markdown_text.splitlines(keepends=True))


def _restore_inline_tabs(root, mark: str | None) -> None:
    if mark is None:
        return
    for node in root.iter():
        if isinstance(node.tag, str):
            if node.text and mark in node.text:
                node.text = node.text.replace(mark, "\t")
            if node.tail and mark in node.tail:
                node.tail = node.tail.replace(mark, "\t")


_DROP_TAGS = {"script", "style", "iframe", "object", "embed", "svg", "math", "form", "input", "button"}
_COLOR = re.compile(r"^[0-9A-Fa-f]{6}$")
_IMAGE_TYPES = {"image/png", "image/jpeg", "image/gif", "image/webp"}
_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _sanitize(root) -> None:
    """Strip raw Markdown HTML attributes before adding trusted Word styles."""
    for node in list(root.iterdescendants()):
        tag = node.tag.lower() if isinstance(node.tag, str) else ""
        if tag in _DROP_TAGS or not tag:
            node.drop_tree()
            continue
        if tag not in _ALLOWED_TAGS:
            node.drop_tag()
            continue
        allowed = {"href", "title"} if tag == "a" else {"src", "alt", "title"} if tag == "img" else set()
        for key in list(node.attrib):
            if key not in allowed:
                del node.attrib[key]
        if tag == "a":
            address = node.get("href", "").strip()
            if urlsplit(address).scheme.lower() not in {"", "http", "https", "mailto"}:
                node.attrib.pop("href", None)


def _style_chain(style) -> list:
    result, visited = [], set()
    while style is not None and style.style_id not in visited:
        visited.add(style.style_id)
        result.append(style.element)
        style = style.base_style
    return list(reversed(result))


def _table_style_sources(table, paragraph, kind) -> list:
    styles = _style_chain(table.style) if table.style else []
    result = [child for style in styles for child in style.findall(_W + kind)]
    cell = paragraph._p.getparent()
    row = cell.getparent() if cell.tag == qn("w:tc") else None
    if row is None or row not in table._tbl.tr_lst:
        return result
    look = {}
    for source in styles + [table._tbl]:
        for element in source.findall(f"./{_W}tblPr/{_W}tblLook"):
            look.update({etree.QName(key).localname: value for key, value in element.attrib.items()})
    row_index, column_index = table._tbl.tr_lst.index(row), row.tc_lst.index(cell)
    conditions = []
    if look.get("noHBand") == "0":
        conditions.append("band1Horz" if row_index % 2 == 0 else "band2Horz")
    for condition, flag, active in (
        ("firstRow", "firstRow", row_index == 0),
        ("lastRow", "lastRow", row_index == len(table.rows) - 1),
        ("firstCol", "firstColumn", column_index == 0),
        ("lastCol", "lastColumn", column_index == len(row.tc_lst) - 1),
    ):
        if active and look.get(flag) == "1":
            conditions.append(condition)
    for condition in conditions:
        for style in styles:
            for block in style.findall(_W + "tblStylePr"):
                if block.get(qn("w:type")) == condition:
                    result.extend(block.findall(_W + kind))
    return result


def _properties(paragraph: Paragraph, kind: str, run: Run | None = None, table: Table | None = None) -> dict[str, dict]:
    """Resolve doc defaults, base styles, character style, then direct OOXML."""
    doc = paragraph.part.document
    sources = doc.styles.element.xpath(f"./w:docDefaults/w:{kind}Default/w:{kind}")
    if table is not None and table.style:
        sources += _table_style_sources(table, paragraph, kind)
    sources += [child for style in _style_chain(paragraph.style) for child in style.findall(_W + kind)]
    if run is not None and run.style:
        sources += [child for style in _style_chain(run.style) for child in style.findall(_W + kind)]
    direct = paragraph._p.pPr if kind == "pPr" else run._r.rPr if run is not None else None
    if direct is not None:
        sources.append(direct)
    result: dict[str, dict] = {}
    for source in sources:
        for child in source:
            tag = etree.QName(child).localname
            attrs = {etree.QName(key).localname: value for key, value in child.attrib.items()}
            if tag in {"rFonts", "ind", "spacing"}:
                resolved = result.setdefault(tag, {})
                if tag == "rFonts":
                    for category in ("ascii", "hAnsi", "eastAsia", "cs"):
                        if category in attrs and category + "Theme" not in attrs:
                            resolved.pop(category + "Theme", None)
                        elif category + "Theme" in attrs and category not in attrs:
                            resolved.pop(category, None)
                resolved.update(attrs)
            else:
                result[tag] = attrs
    return result


def _theme(doc) -> tuple[dict[str, str], dict[str, str]]:
    colors, fonts = {}, {}
    language = doc.settings.element.xpath("./w:themeFontLang")
    east_asia = language[0].get(qn("w:eastAsia"), "zh-CN").lower() if language else "zh-cn"
    script = "Jpan" if east_asia.startswith("ja") else "Hang" if east_asia.startswith("ko") else "Hant" if east_asia in {"zh-tw", "zh-hk", "zh-mo"} else "Hans"
    for relation in doc.part.rels.values():
        if not relation.reltype.endswith("/theme") or relation.is_external:
            continue
        tree = etree.fromstring(relation.target_part.blob, parser=etree.XMLParser(resolve_entities=False, no_network=True))
        for node in tree.findall(f".//{_A}clrScheme/*"):
            if len(node):
                value = node[0].get("val") if etree.QName(node[0]).localname == "srgbClr" else node[0].get("lastClr")
                if value and _COLOR.fullmatch(value):
                    colors[etree.QName(node).localname] = value
        for category in ("major", "minor"):
            font = tree.find(f".//{_A}{category}Font")
            if font is not None:
                for name in ("latin", "ea", "cs"):
                    element = font.find(_A + name)
                    if element is not None and element.get("typeface"):
                        fonts[category + name] = element.get("typeface")
                for element in font.findall(_A + "font"):
                    if element.get("script") == script:
                        fonts.setdefault(category + "ea", element.get("typeface", ""))
    for name, alias in (("text1", "dk1"), ("text2", "dk2"), ("background1", "lt1"), ("background2", "lt2")):
        if alias in colors:
            colors[name] = colors[alias]
    return colors, fonts


def _number(value, divisor=1) -> float | None:
    try:
        number = float(value) / divisor
        return number if abs(number) <= 10000 else None
    except (TypeError, ValueError):
        return None


def _measure(value, divisor=20) -> str | None:
    number = _number(value, divisor)
    return f"{number:g}pt" if number is not None else None


def _color(attrs: dict, theme: dict[str, str]) -> str | None:
    fallback = attrs.get("val")
    value = theme.get(attrs.get("themeColor", "")) or (fallback if fallback and _COLOR.fullmatch(fallback) else None)
    if not value:
        return None
    components = [int(value[index:index + 2], 16) for index in (0, 2, 4)]
    for modifier in ("themeShade", "themeTint"):
        try:
            factor = int(attrs[modifier], 16) / 255
        except (KeyError, ValueError):
            continue
        components = [round(channel * factor if modifier == "themeShade" else channel + (255 - channel) * (1 - factor)) for channel in components]
    return "#" + "".join(f"{channel:02X}" for channel in components)


def _on(attrs: dict) -> bool:
    return attrs.get("val", "1").lower() not in {"0", "false", "off", "none"}


def _run_css(paragraph: Paragraph, run: Run | None, theme, table=None, *, east_asian: bool | None = None) -> dict[str, str]:
    properties = _properties(paragraph, "rPr", run, table)
    colors, fonts = theme
    css = {"white-space": "pre-wrap"}
    names = properties.get("rFonts", {})
    cjk = east_asian if east_asian is not None else run is not None and bool(re.search(r"[\u2e80-\uffff]", run.text))
    category = "eastAsia" if cjk else "ascii"
    themed = names.get(category + "Theme")
    if not themed and not names.get(category):
        themed = names.get("asciiTheme", "")
    font = fonts.get(("major" if themed.startswith("major") else "minor") + ("ea" if cjk else "latin")) if themed else None
    font = font or names.get(category) or names.get("ascii")
    if font and re.fullmatch(r"[\w\s\-]+", font, re.UNICODE):
        css["font-family"] = json.dumps(font, ensure_ascii=False)
    size = _measure(properties.get("sz", {}).get("val"), 2)
    if size:
        css["font-size"] = size
    for property_name, css_name, enabled, disabled in (("b", "font-weight", "700", "400"), ("i", "font-style", "italic", "normal")):
        if property_name in properties:
            css[css_name] = enabled if _on(properties[property_name]) else disabled
    decorations = []
    if "u" in properties and _on(properties["u"]):
        decorations.append("underline")
    if "strike" in properties and _on(properties["strike"]):
        decorations.append("line-through")
    if decorations:
        css["text-decoration"] = " ".join(decorations)
    color = _color(properties.get("color", {}), colors)
    if color:
        css["color"] = color
    shading = properties.get("shd", {}).get("fill")
    if shading and _COLOR.fullmatch(shading):
        css["background-color"] = "#" + shading
    position = properties.get("vertAlign", {}).get("val")
    if position in {"superscript", "subscript"}:
        css["vertical-align"] = "super" if position == "superscript" else "sub"
    return css


def _paragraph_css(paragraph: Paragraph, table=None) -> dict[str, str]:
    properties = _properties(paragraph, "pPr", table=table)
    css = {}
    alignment = properties.get("jc", {}).get("val", "left")
    css["text-align"] = {"both": "justify", "distribute": "justify", "start": "left", "end": "right"}.get(alignment, alignment) if alignment in {"left", "right", "center", "both", "distribute", "start", "end"} else "left"
    indent = properties.get("ind", {})
    for name, target in (("left", "margin-left"), ("right", "margin-right"), ("firstLine", "text-indent")):
        if (value := _measure(indent.get(name))) is not None:
            css[target] = value
        if (value := _number(indent.get(name + "Chars"), 100)) is not None:
            css[target] = f"{value:g}em"
    if "hanging" in indent and (value := _number(indent["hanging"], 20)) is not None:
        css["text-indent"] = f"{-value:g}pt"
    if (value := _number(indent.get("hangingChars"), 100)) is not None:
        css["text-indent"] = f"{-value:g}em"
    spacing = properties.get("spacing", {})
    if "line" in spacing:
        css["line-height"] = (_measure(spacing["line"]) if spacing.get("lineRule") in {"exact", "atLeast"}
                              else f"{_number(spacing['line'], 240) or 1:g}")
        if spacing.get("lineRule") == "atLeast" and css["line-height"]:
            css["line-height"] = f"max({css['line-height']}, 1em)"
    for name, target in (("before", "margin-top"), ("after", "margin-bottom")):
        css[target] = _measure(spacing.get(name, "0")) or "0pt"
        if not _on(spacing.get(name + "Autospacing", {"val": "0"})):
            if (value := _number(spacing.get(name + "Lines"), 100)) is not None:
                line = css.get("line-height", "1.2")
                unit = line if not re.fullmatch(r"[\d.]+", line) else f"{line}em"
                css[target] = f"calc({value:g} * {unit})"
    return css


def _set_css(node, css: dict[str, str]) -> None:
    if css:
        node.set("style", ";".join(f"{name}:{value}" for name, value in css.items() if value is not None))


def _text_nodes(node):
    if node.text:
        yield node, "text", node.text
    for child in node:
        if child.tag not in {"ul", "ol"}:
            yield from _text_nodes(child)
        if child.tail:
            yield child, "tail", child.tail


def _project_runs(node, paragraphs: list[Paragraph], theme, table=None) -> None:
    """Decorate text nodes in place; existing semantic tags are never rewritten."""
    owners, source = [], ""
    for index, paragraph in enumerate(paragraphs):
        if index:
            source += "\n"
            owners.append(_run_css(paragraph, None, theme, table))
        # Include hyperlink/field runs for display, but leave edit protection to
        # the existing exporter.  Never execute a relationship or field.
        for element in paragraph._p.xpath(".//w:r"):
            run = Run(element, paragraph)
            latin = _run_css(paragraph, run, theme, table, east_asian=False)
            cjk = _run_css(paragraph, run, theme, table, east_asian=True)
            source += run.text
            owners += [cjk if '\u2e80' <= char <= '\uffff' else latin for char in run.text]
    text_nodes = list(_text_nodes(node))
    target = "".join(value for _, _, value in text_nodes)
    fallback = (_run_css(paragraphs[0], None, theme, table) if paragraphs else {}) | {"white-space": "normal"}
    mapped = [fallback] * len(target)
    # 测试工程师 N7：不再整段 SequenceMatcher（二次方），见 app/fast_diff.py
    for operation, start, end, new_start, new_end in bounded_opcodes(source, target):
        if operation == "equal":
            mapped[new_start:new_end] = owners[start:end]
        elif operation in {"insert", "replace"}:
            css = owners[start] if start < len(owners) else owners[-1] if owners else fallback
            # HTML-only list/table serialization whitespace has no Word run.
            mapped[new_start:new_end] = [css | {"white-space": "normal"}] * (new_end - new_start)
    offset = 0
    for owner, attribute, value in text_nodes:
        groups = []
        for index, (char, css) in enumerate(zip(value, mapped[offset:offset + len(value)])):
            if attribute == "tail" and owner.tag == "br" and index == 0 and char == "\n":
                # This character is Python-Markdown's HTML serialization
                # whitespace; <br> already displays the native Word break.
                css = css | {"white-space": "normal"}
            # 测试工程师 N7：同一样式的字收集成列表再拼接（逐字拼字符串在长段落上是二次方）
            if groups and (groups[-1][1] is css or groups[-1][1] == css):
                groups[-1][0].append(char)
            else:
                groups.append(([char], css))
        offset += len(value)
        setattr(owner, attribute, None)
        if attribute == "text":
            position, parent = 0, owner
        else:
            parent = owner.getparent()
            position = parent.index(owner) + 1
        for chars, css in groups:
            span = html.Element("span")
            span.text = "".join(chars)
            _set_css(span, css)
            parent.insert(position, span)
            position += 1


def _numbering(paragraph: Paragraph, counters: dict) -> tuple[dict, str | None]:
    """Read native numbering, including styles, without inserting body text."""
    sources = [style for style in _style_chain(paragraph.style)] + [paragraph._p]
    number_id, depth = None, None
    for source in sources:
        for tag in ("numId", "ilvl"):
            values = source.xpath(f"./w:pPr/w:numPr/w:{tag}/@w:val")
            if values and values[0].isdigit():
                if tag == "numId":
                    number_id = int(values[0])
                else:
                    depth = int(values[0])
    if number_id is None or number_id == 0:
        return {}, None
    depth = depth or 0
    numbering = paragraph.part.numbering_part.element
    numbers = numbering.xpath(f"./w:num[@w:numId='{number_id}']")
    if not numbers:
        return {}, None
    abstracts = numbers[0].xpath("./w:abstractNumId/@w:val")
    if not abstracts:
        return {}, None
    levels = numbering.xpath(f"./w:abstractNum[@w:abstractNumId='{abstracts[0]}']/w:lvl[@w:ilvl='{depth}']")
    if not levels:
        return {}, None
    level = levels[0]
    override = numbers[0].xpath(f"./w:lvlOverride[@w:ilvl='{depth}']/w:lvl")
    if override:
        level = override[0]
    def level_value(tag):
        element = level.find(qn("w:" + tag))
        value = element.get(qn("w:val")) if element is not None else None
        return [value] if value is not None else []

    starts = numbers[0].xpath(f"./w:lvlOverride[@w:ilvl='{depth}']/w:startOverride/@w:val") or level_value("start")
    start = int(starts[0]) if starts and starts[0].isdigit() else 1
    key = number_id, depth
    counters[key] = counters.get(key, start - 1) + 1
    for key_to_reset in list(counters):
        if key_to_reset[0] == number_id and key_to_reset[1] > depth:
            del counters[key_to_reset]
    formats = level_value("numFmt")
    texts = level_value("lvlText")
    marker = texts[0] if texts else "%1."
    marker = re.sub(r"%([1-9])", lambda match: str(counters.get((number_id, int(match[1]) - 1), 1)), marker)
    css = {}
    indents = level.findall(f"./{qn('w:pPr')}/{qn('w:ind')}")
    if indents:
        indent = indents[0]
        if (left := _measure(indent.get(qn("w:left")))):
            css["margin-left"] = left
        if (hanging := _number(indent.get(qn("w:hanging")), 20)) is not None:
            css["text-indent"] = f"{-hanging:g}pt"
    kind = formats[0] if formats else "decimal"
    css["list-style-type"] = {"bullet": "disc", "decimal": "decimal", "lowerLetter": "lower-alpha", "upperLetter": "upper-alpha", "lowerRoman": "lower-roman", "upperRoman": "upper-roman", "none": "none"}.get(kind, "decimal")
    return css, marker


def _decorate_paragraph(node, paragraphs, theme, table=None) -> None:
    if not paragraphs:
        return
    paragraph = paragraphs[0]
    sample = next((run for run in paragraph.runs if run.text.strip()), None)
    _set_css(node, _run_css(paragraph, sample, theme, table) | _paragraph_css(paragraph, table))
    _project_runs(node, paragraphs, theme, table)


def _decorate_table(node, table: Table, theme) -> None:
    css = {"border-collapse": "collapse", "table-layout": "fixed"}
    widths = [column.width for column in table.columns]
    if widths and all(width is not None for width in widths):
        css["width"] = f"{sum(width.pt for width in widths):g}pt"
        css["max-width"] = "100%"
    _set_css(node, css)
    borders = {}
    for source in (_style_chain(table.style) if table.style else []) + [table._tbl]:
        for properties in source.findall(f"./{_W}tblPr/{_W}tblBorders"):
            for border in properties:
                borders[etree.QName(border).localname] = border
    for row_index, (target_row, source_row) in enumerate(zip(node.xpath(".//tr"), table.rows)):
        for column_index, (target_cell, source_cell) in enumerate(zip(target_row.xpath("./th|./td"), source_row.cells)):
            _decorate_paragraph(target_cell, source_cell.paragraphs, theme, table)
            cell_css = dict(item.split(":", 1) for item in target_cell.get("style", "").split(";") if ":" in item)
            cell_css.update({"vertical-align": "top", "font-weight": "400", "border": "none"})
            cell_sources = _table_style_sources(table, source_cell.paragraphs[0], "tcPr")
            if source_cell._tc.tcPr is not None:
                cell_sources.append(source_cell._tc.tcPr)
            for properties in cell_sources:
                for shading in properties.findall(_W + "shd"):
                    fill = shading.get(qn("w:fill"), "")
                    if _COLOR.fullmatch(fill):
                        cell_css["background-color"] = "#" + fill
                for valign in properties.findall(_W + "vAlign"):
                    value = valign.get(qn("w:val"))
                    cell_css["vertical-align"] = {"center": "middle"}.get(value, value) if value in {"top", "center", "bottom"} else "top"
            cell_borders = dict(borders)
            cell_edges = set()
            for properties in cell_sources:
                for border_group in properties.findall(_W + "tcBorders"):
                    for border in border_group:
                        side = etree.QName(border).localname
                        cell_borders[side] = border
                        cell_edges.add(side)
            for side, inside in (("top", "insideH" if row_index else "top"), ("bottom", "insideH" if row_index < len(table.rows) - 1 else "bottom"), ("left", "insideV" if column_index else "left"), ("right", "insideV" if column_index < len(source_row.cells) - 1 else "right")):
                border = cell_borders.get(side) if side in cell_edges else cell_borders.get(inside)
                if border is None:
                    continue
                mode = border.get(qn("w:val"), "single")
                width = _measure(border.get(qn("w:sz"), "4"), 8) or "0.5pt"
                color = _color({"val": border.get(qn("w:color"), "000000")}, theme[0]) or "#000000"
                cell_css["border-" + side] = "none" if mode in {"nil", "none"} else f"{width} { {'double': 'double', 'dotted': 'dotted', 'dashed': 'dashed'}.get(mode, 'solid')} {color}"
            # Markdown already represents merged cells as grid entries.  Use
            # grid widths here rather than duplicating a merged cell's width.
            if column_index < len(widths) and widths[column_index] is not None:
                cell_css["width"] = f"{widths[column_index].pt:g}pt"
            _set_css(target_cell, cell_css)


def _images(root, supplied: list[dict]) -> None:
    safe_images = {}
    for image in supplied:
        if not isinstance(image, dict) or image.get("content_type") not in _IMAGE_TYPES:
            continue
        try:
            data = base64.b64decode(image.get("data", ""), validate=True)
        except (ValueError, TypeError):
            continue
        if data and len(data) <= 20 * 1024 * 1024 and isinstance(image.get("ref"), str):
            safe_images[image["ref"]] = f"data:{image['content_type']};base64,{base64.b64encode(data).decode('ascii')}"
    for node in root.xpath(".//img"):
        source = node.get("src", "")
        if source in safe_images:
            node.set("src", safe_images[source])
            node.set("data-word-image-ref", source)
            _set_css(node, {"max-width": "100%", "height": "auto"})
        else:
            node.attrib.pop("src", None)


def memoize_styles(doc) -> None:
    """终验 S-7：python-docx 每次取段落/字符样式都要线性扫描全部样式找默认样式（styles.default_for），
    SRS 这样的文档一次预览约 7 s，其中约 6 s 花在这里。一次渲染内样式不变，按 (style_id, type) 缓存。"""
    part = doc.part
    original = part.get_style
    cache: dict[tuple, object] = {}

    def get_style(style_id, style_type):
        key = (style_id, style_type)
        if key not in cache:
            cache[key] = original(style_id, style_type)
        return cache[key]

    part.get_style = get_style


def render_document_html(source_path: str | Path, current_markdown: str, images: list[dict] | None = None, *, outline_headings: bool = False) -> str:
    """Return safe editable body HTML styled from the actual edited Word file.

    Raises the same safe-edit errors as DOCX export.  No original file, stored
    Markdown, or database is changed; temporary exports are always removed.
    """
    supplied = images or []
    with TemporaryDirectory(prefix="docwb-preview-") as folder:
        edited = Path(folder) / "edited.docx"
        edit_original_docx(current_markdown, source_path, edited, image_sources=diagram_images(supplied), outline_headings=outline_headings)
        doc = Document(edited)
        memoize_styles(doc)
        root = html.Element("div")
        tab_mark = _pick_tab_mark(current_markdown)
        for block in _html_blocks(_hide_inline_tabs(current_markdown, tab_mark)):
            root.append(block)
        _restore_inline_tabs(root, tab_mark)
        _sanitize(root)
        target = list(root)
        flat = any((_list_info(paragraph) or ("", 0))[1] for paragraph in doc.paragraphs) and not any(
            depth for element in target if element.tag in {"ul", "ol"} for _, depth, _, _ in _list_items(element)
        )
        blocks = _word_blocks(doc, flat_lists=flat, outline_headings=outline_headings)
        if [element.tag for element in target] != [block.tag for block in blocks]:
            raise ValueError("编辑预览结构无法与 Word 对应，请使用 Word 预览检查完整排版。")
        theme, counters = _theme(doc), {}
        for node, block in zip(target, blocks):
            if node.tag == "table":
                _decorate_table(node, Table(block.elements[0], doc), theme)
            elif node.tag in {"ul", "ol"}:
                items = list(_list_items(node))
                actual_items = list(node.xpath(".//li"))
                if len(actual_items) != len(block.elements):
                    raise ValueError("编辑预览列表无法与 Word 对应，请使用 Word 预览检查完整排版。")
                for container in [node] + list(node.xpath(".//ul|.//ol")):
                    _set_css(container, {"margin": "0", "padding-left": "0"})
                margins = {}
                for item, (_, depth, _, _), element in zip(actual_items, items, block.elements):
                    paragraph = Paragraph(element, doc)
                    _decorate_paragraph(item, [paragraph], theme)
                    css = _run_css(paragraph, None, theme) | _paragraph_css(paragraph)
                    numbering_css, marker = _numbering(paragraph, counters)
                    css.update(numbering_css)
                    left = _number(css.get("margin-left", "0pt").removesuffix("pt")) or 0
                    css["margin-left"] = f"{left - margins.get(depth - 1, 0):g}pt"
                    margins[depth] = left
                    css["text-indent"] = "0pt"
                    _set_css(item, css)
                    if item.getparent().tag == "ol" and marker and (ordinal := re.search(r"\d+", marker)):
                        item.set("value", ordinal.group())
            else:
                paragraph = Paragraph(block.elements[0], doc)
                _decorate_paragraph(node, [paragraph], theme)
                style_name = (paragraph.style.name if paragraph.style is not None else "") or ""
                # 测试复测 N4：[02][03] 的目录行是静态文字（不含域），但仍是目录（样式 toc N），同样只读
                if block.elements[0].xpath(_FIELD_LIKE) or re.match(r"(?i)^toc\s*\d", style_name):
                    # 测试 L3：目录行、域、超链接等写回 Word 会被拒绝，编辑区里直接设为只读并说明。
                    node.set("contenteditable", "false")
                    node.set("data-word-protected", "1")
                    node.set("title", "此段含目录、域或链接等 Word 内容，不能在这里修改；请在原稿中修改。")
                # Legacy Normal+numPr stays a p in Markdown.  A CSS pseudo
                # marker can show its native number without changing textContent.
                numbering_css, marker = _numbering(paragraph, counters)
                if marker and not node.xpath(".//img"):
                    css = _run_css(paragraph, None, theme) | _paragraph_css(paragraph)
                    css.update({key: value for key, value in numbering_css.items() if key != "list-style-type"})
                    css.update({"position": "relative", "text-indent": "0pt"})
                    _set_css(node, css)
                    node.set("data-word-marker", marker)
        # Source image refs are stable; generated diagram refs come from the
        # caller.  Do not resolve external Markdown image URLs.
        _images(root, parse_docx(str(source_path), outline_headings=outline_headings).images + supplied)
        if len(root):
            # shortcut: continuous editing uses the final body section; inspect mixed sections in Word layout preview.
            section = doc.sections[-1]
            if all(value is not None for value in (section.page_width, section.page_height, section.left_margin,
                                                   section.right_margin, section.top_margin, section.bottom_margin)):
                root[0].set("data-word-page", json.dumps({
                    "width": f"{(section.page_width - section.left_margin - section.right_margin) / 12700:g}pt",
                    "minHeight": f"{(section.page_height - section.top_margin - section.bottom_margin) / 12700:g}pt",
                    "paddingTop": f"{section.top_margin.pt:g}pt", "paddingBottom": f"{section.bottom_margin.pt:g}pt",
                    "paddingLeft": f"{section.left_margin.pt:g}pt", "paddingRight": f"{section.right_margin.pt:g}pt",
                }))
        return "".join(html.tostring(node, encoding="unicode") for node in root)
