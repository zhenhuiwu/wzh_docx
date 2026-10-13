"""A07 段落 / 章节负责人（硬限制，17 §4.5、§4.6、§4.10、§10）：归属计算、检查点规则 1–6、blockIds、被拒保存、
接口权限、成员联动、新建文档整篇归创建者、分工模式。"""

from __future__ import annotations

from io import BytesIO
from uuid import uuid4

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app import document_assignments as da
from app.db import connect, migrate
from app.files import store_bytes
from app.identity import create_user
from app.markdown_blocks import Block, assemble, split

pytestmark = pytest.mark.enforced_new_documents

PASSWORD = "a07-assign-pw"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


# ---------------------------------------------------------------- 纯函数：归属计算

def N(id, kind="paragraph", level=None, body=None):
    return da.Node(id, kind, level, body if body is not None else f"{id} 正文")


def H(id, level, title=None):
    return da.Node(id, "heading", level, "#" * level + " " + (title or id))


def owners(nodes, explicit):
    return [sorted(o.owners) for o in da.resolve(nodes, {k: frozenset(v) for k, v in explicit.items()})]


def test_resolve_nearest_layer_wins_and_document_fallback():
    nodes = [N("cover"), H("h1", 1), N("p1"), H("h3", 1), N("p3"), H("h32", 2), N("p32"), N("p32b"), H("h33", 2), N("p33"), H("h4", 1), N("p4")]
    explicit = {"h3": {"A"}, "h32": {"B"}, "p32b": {"C"}}
    got = owners(nodes, explicit)
    assert got[0] == [] and got[1] == [] and got[2] == []          # 封面、未分配章节
    assert got[3] == ["A"] and got[4] == ["A"]
    assert got[5] == ["B"] and got[6] == ["B"]                     # 最近一层为准，不取并集
    assert got[7] == ["C"]                                         # 段落级覆盖章节
    assert got[8] == ["A"] and got[9] == ["A"]                     # 3.3 继承 3
    assert got[10] == [] and got[11] == []
    # 文档级分配：封面、未分配章节都随整篇
    got = owners(nodes, {**explicit, da.DOCUMENT: {"W"}})
    assert got[0] == ["W"] and got[1] == ["W"] and got[11] == ["W"] and got[6] == ["B"]
    # 多人
    assert owners([H("h", 1), N("p")], {"h": {"A", "B"}})[1] == ["A", "B"]


def test_heading_level_change_recomputes_and_move_follows_section():
    before = [H("h1", 1), N("a"), H("h2", 1), N("b")]
    assert owners(before, {"h1": {"A"}})[3] == []
    after = [H("h1", 1), N("a"), H("h2", 2), N("b")]               # h2 降级为 h1 的子标题
    assert owners(after, {"h1": {"A"}})[3] == ["A"]
    moved = [H("h1", 1), H("h2", 1), N("b"), N("a")]
    assert owners(moved, {"h1": {"A"}})[3] == []                   # 移到 2 章后随新章节


def kinds(found):
    return sorted((item[0], item[1].id) for item in found)


BASE = [N("cover"), H("h31", 2, "3.1"), N("p31"), H("h32", 2, "3.2"), N("p32"), N("p32b"), H("h33", 2, "3.3"), N("p33")]
EXPLICIT = {"h31": frozenset({"Z"}), "h32": frozenset({"A"})}


def test_rules_allow_own_edits_and_reject_others():
    edit = [n if n.id != "p32" else N("p32", body="新内容") for n in BASE]
    assert da.check(BASE, edit, EXPLICIT, "A") == []                           # 规则 1 允许
    edit = [n if n.id != "p31" else N("p31", body="改别人") for n in BASE]
    assert kinds(da.check(BASE, edit, EXPLICIT, "A")) == [("modify", "p31")]   # 规则 1 违规
    deleted = [n for n in BASE if n.id != "p32b"]
    assert da.check(BASE, deleted, EXPLICIT, "A") == []                        # 规则 2 允许
    deleted = [n for n in BASE if n.id != "p33"]
    assert kinds(da.check(BASE, deleted, EXPLICIT, "A")) == [("delete", "p33")]
    deleted = [n for n in BASE if n.id != "h32"]                                 # 删除有显式分配的标题
    assert ("heading_structure", "h32") in kinds(da.check(BASE, deleted, EXPLICIT, "A"))


def test_insert_boundary_before_and_after_next_heading():
    at_end = BASE[:6] + [N("new")] + BASE[6:]                                   # 3.3 标题之前：属于 3.2
    assert da.check(BASE, at_end, EXPLICIT, "A") == []
    after_next = BASE[:7] + [N("new")] + BASE[7:]                               # 3.3 标题之后：属于 3.3（未分配）
    assert kinds(da.check(BASE, after_next, EXPLICIT, "A")) == [("insert", "new")]


def test_new_subheading_deeper_allowed_same_level_reassigns():
    deeper = BASE[:5] + [H("sub", 3, "3.2.1")] + BASE[5:]
    assert da.check(BASE, deeper, EXPLICIT, "A") == []
    same = BASE[:5] + [H("sib", 2, "3.2b")] + BASE[5:]
    found = kinds(da.check(BASE, same, EXPLICIT, "A"))
    assert ("insert", "sib") in found and ("reassign", "p32b") in found


def test_move_rules_and_explicit_heading_level():
    moved = [BASE[0], BASE[1], BASE[2], BASE[3], BASE[5], BASE[4], BASE[6], BASE[7]]   # 本章内调换
    assert da.check(BASE, moved, EXPLICIT, "A") == []
    out = [BASE[0], BASE[1], BASE[2], BASE[3], BASE[5], BASE[6], BASE[7], BASE[4]]     # 移出到 3.3
    assert ("move", "p32") in kinds(da.check(BASE, out, EXPLICIT, "A"))
    leveled = [n if n.id != "h32" else H("h32", 3, "3.2") for n in BASE]
    assert ("heading_structure", "h32") in kinds(da.check(BASE, leveled, EXPLICIT, "A"))
    renamed = [n if n.id != "h32" else H("h32", 2, "3.2 改名") for n in BASE]
    assert da.check(BASE, renamed, EXPLICIT, "A") == []                                # 改标题文字允许（U-11）


def test_merge_across_others_block_is_rejected():
    merged = [n for n in BASE if n.id not in ("h33", "p33")]
    merged = [n if n.id != "p32b" else N("p32b", body="p32b 正文 p33 正文") for n in merged]
    found = kinds(da.check(BASE, merged, EXPLICIT, "A"))
    assert ("delete", "p33") in found


def test_adjacent_inheritance_k3():
    explicit = {"h32": frozenset({"A"}), "p32": frozenset({"B"})}             # p32 是 B 在 A 章里的例外段
    enter_end = BASE[:5] + [N("new")] + BASE[5:]
    assert da.check(BASE, enter_end, explicit, "B") == []                     # 例外段末尾回车：归 B
    inherited = da.plan_inheritance({n.id for n in BASE}, enter_end, explicit, "B")
    assert inherited["new"][0] == {"B"} and inherited["new"][2] is False      # 本人不标待确认
    assert kinds(da.check(BASE, enter_end, explicit, "A")) == [("insert", "new")]   # A 在 B 的例外段后插入：被拒
    owner_split = da.plan_inheritance({n.id for n in BASE}, enter_end, explicit, "Z")
    assert owner_split["new"][2] is True                                      # 负责人拆分：待确认
    start = BASE[:4] + [N("new")] + BASE[4:]                                  # 例外段开头回车（章节开头看后一块）
    assert da.check(BASE, start, explicit, "B") == []


def test_reassign_detects_boundary_change_by_deleting_unassigned_heading():
    base = [H("h1", 1), N("a"), H("h2", 1), N("b")]
    explicit = {"h1": frozenset({"A"})}
    after = [H("h1", 1), N("a"), N("b")]                                       # 删掉未分配的 h2：b 被划进 A 的章
    found = kinds(da.check(base, after, explicit, "A"))
    assert ("delete", "h2") in found and ("reassign", "b") in found


# ---------------------------------------------------------------- 接口与检查点

def _login(client, account):
    client.cookies.clear()
    token = client.get("/api/v1/auth/csrf").json()["data"]["csrfToken"]
    response = client.post("/api/v1/auth/login", json={"account": account, "password": PASSWORD}, headers={"X-CSRF-Token": token})
    assert response.status_code == 200, response.text
    return {"X-CSRF-Token": response.json()["data"]["csrfToken"]}


def _docx() -> bytes:
    word = Document()
    word.add_paragraph("封面文字")
    word.add_heading("范围", level=1)
    word.add_paragraph("范围正文。")
    word.add_heading("需求", level=1)
    word.add_heading("功能概述", level=2)
    word.add_paragraph("概述正文。")
    word.add_heading("功能需求", level=2)
    word.add_paragraph("功能正文甲。")
    word.add_paragraph("功能正文乙。")
    word.add_heading("性能需求", level=2)
    word.add_paragraph("性能正文。")
    buffer = BytesIO()
    word.save(buffer)
    return buffer.getvalue()


class Ctx:
    def __init__(self, client, users, project, file_id):
        self.client, self.users, self.project, self.file_id = client, users, project, file_id
        self.headers: dict = {}
        self.who = None

    def as_(self, name):
        self.headers = _login(self.client, f"a07-{name}")
        self.who = name
        return self

    def req(self, method, path, body=None):
        return self.client.request(method, f"/api/v1{path}", json=body, headers=self.headers)

    def ok(self, method, path, body=None, status=200):
        response = self.req(method, path, body)
        assert response.status_code == status, response.text
        return response.json()["data"] if response.content else None

    def create(self, name="需求规格"):
        return self.ok("POST", f"/projects/{self.project}/documents",
                       {"name": name, "sourceType": "upload", "wordFileId": str(self.file_id)})

    def state(self, did):
        """(markdown, revision, [(id, Block)])：对象与正文块一一对应。"""
        data = self.ok("GET", f"/documents/{did}/objects")
        doc = self.ok("GET", f"/documents/{did}")
        prefix, blocks = split(doc["markdownContent"])
        ids = [item["id"] for item in data["objects"]]
        assert len(ids) == len(blocks)
        return prefix, doc["revision"], list(zip(ids, blocks)), data

    def heading(self, did, title):
        _prefix, _rev, pairs, _data = self.state(did)
        return next(object_id for object_id, block in pairs if block.kind == "heading" and title in block.body)

    def assign(self, did, object_ids, names, status=200):
        rev = self.ok("GET", f"/documents/{did}/assignments")["assignmentsRevision"]
        response = self.req("PUT", f"/documents/{did}/assignments", {
            "objectIds": object_ids, "assigneeIds": [str(self.users[n]) for n in names], "expectedAssignmentsRevision": rev})
        assert response.status_code == status, response.text
        return response.json()

    def save(self, did, edit, *, block_ids=True, merge=True, expected=None):
        prefix, revision, pairs, _ = self.state(did)
        pairs = edit(list(pairs))
        markdown = assemble(prefix, [block for _id, block in pairs])
        body = {"expectedRevision": revision if expected is None else expected, "markdownContent": markdown,
                "allowMerge": merge, "clientSaveKey": uuid4().hex}
        if block_ids:
            body["blockIds"] = [object_id for object_id, _block in pairs]
        return self.req("PATCH", f"/documents/{did}/content", body)


def replace_text(pairs, needle, text):
    out = []
    for object_id, block in pairs:
        if needle in block.body:
            block = Block(block.kind, block.body.replace(needle, text), block.gap_after, block.heading_level)
        out.append((object_id, block))
    return out


@pytest.fixture
def ctx():
    migrate()
    run = uuid4().hex[:6]
    users = {}
    for key, name, roles in (("admin", "管理员", ["system_admin"]), ("owner", "张工", []), ("wang", "王测", []),
                             ("sun", "孙敏", []), ("liu", "刘洋", []), ("viewer", "只读甲", []), ("out", "外人", [])):
        users[key] = create_user(account=f"a07-{key}", name=name, password=PASSWORD, roles=roles, allow_no_roles=True)
    project = uuid4()
    with connect() as connection:
        connection.execute("INSERT INTO projects(id,name,owner_id) VALUES(%s,%s,%s)", (project, f"A07 {run}", users["owner"]))
        for key, role in (("owner", "owner"), ("wang", "editor"), ("sun", "editor"), ("liu", "editor"), ("viewer", "viewer")):
            connection.execute("INSERT INTO project_members(project_id,user_id,role,created_at) VALUES(%s,%s,%s,now())",
                               (project, users[key], role))
        connection.commit()
    stored = store_bytes(name="需求.docx", media_type=DOCX, purpose="attachment", project_id=project, content=_docx())
    from app.main import create_app
    with TestClient(create_app()) as client:
        yield Ctx(client, users, project, stored.id)


def _counts(did):
    with connect() as connection:
        return tuple(connection.execute(sql, (str(did),)).fetchone()[0] for sql in (
            "SELECT revision FROM documents WHERE id=%s",
            "SELECT COUNT(*) FROM document_objects WHERE document_id=%s",
            "SELECT COUNT(*) FROM document_changes WHERE document_id=%s",
            "SELECT COUNT(*) FROM document_conflicts WHERE document_id=%s",
            "SELECT COUNT(*) FROM document_rejected_saves WHERE document_id=%s"))


def test_owner_created_document_is_enforced_and_unassigned(ctx):
    doc = ctx.as_("owner").create()
    view = ctx.ok("GET", f"/documents/{doc['id']}/assignments?resolved=1")
    assert view["assignmentMode"] == "enforced" and view["explicit"] == []
    assert view["projectOwners"] == [{"id": str(ctx.users["owner"]), "name": "张工"}]
    assert view["summary"]["unassignedRoots"] == 2 and view["summary"]["unassignedSections"] == 5
    assert all(item["editable"] for item in view["resolved"])                     # 负责人全部可改
    ctx.as_("wang")
    view = ctx.ok("GET", f"/documents/{doc['id']}/assignments?resolved=1")
    assert not any(item["editable"] for item in view["resolved"])
    assert {item["readonlyReason"] for item in view["resolved"]} == {"assignment"}
    detail = ctx.ok("GET", f"/documents/{doc['id']}")
    assert detail["permissions"]["readonlyReason"] == "assignment" and detail["permissions"]["canAssign"] is False
    assert detail["permissions"]["projectOwners"][0]["name"] == "张工"


def test_writer_created_document_belongs_to_creator(ctx):
    doc = ctx.as_("wang").create("王测的文档")
    view = ctx.ok("GET", f"/documents/{doc['id']}/assignments?resolved=1")
    assert [item["objectId"] for item in view["explicit"]] == ["@document"]
    assert all(item["editable"] and item["via"] == "document" for item in view["resolved"])
    assert ctx.save(doc["id"], lambda p: replace_text(p, "封面文字", "封面改了")).status_code == 200
    # 文末新增一级章节
    def add_chapter(pairs):
        return pairs + [(None, Block("heading", "# 新章", "\n\n", 1)), (None, Block("paragraph", "新章正文。", "", None))]
    response = ctx.save(doc["id"], lambda p: add_chapter([(i, Block(b.kind, b.body, b.gap_after or "\n\n", b.heading_level)) for i, b in p]))
    assert response.status_code == 200, response.text
    # 负责人把“功能需求”分给孙敏：覆盖整篇
    ctx.as_("owner")
    ctx.assign(doc["id"], [ctx.heading(doc["id"], "功能需求")], ["sun"])
    ctx.as_("wang")
    rejected = ctx.save(doc["id"], lambda p: replace_text(p, "功能正文甲", "王测越权"))
    assert rejected.status_code == 409 and rejected.json()["error"]["code"] == "SECTION_NOT_ASSIGNED"
    assert ctx.save(doc["id"], lambda p: replace_text(p, "范围正文", "王测仍可改")).status_code == 200


def test_writer_hard_limit_end_to_end(ctx):
    doc = ctx.as_("owner").create()
    did = doc["id"]
    feature = ctx.heading(did, "功能需求")
    ctx.assign(did, [feature], ["wang"])
    ctx.as_("wang")
    # 规则 1 允许：改自己的
    assert ctx.save(did, lambda p: replace_text(p, "功能正文甲", "王测改甲")).status_code == 200
    # 越权：改别人的（概述，未分配）+ 自己的，整体拒绝，什么都不变
    before = _counts(did)
    response = ctx.save(did, lambda p: replace_text(replace_text(p, "概述正文", "越权"), "功能正文乙", "王测改乙"))
    assert response.status_code == 409, response.text
    error = response.json()["error"]
    assert error["code"] == "SECTION_NOT_ASSIGNED"
    assert error["details"]["total"] == 1 and error["details"]["violations"][0]["type"] == "modify"
    assert error["details"]["violations"][0]["headingPath"] == ["需求", "功能概述"]
    after = _counts(did)
    assert after[:4] == before[:4] and after[4] == before[4] + 1                      # 只多一份被拒保存
    rid = error["details"]["rejectedSaveId"]
    saved = ctx.ok("GET", f"/documents/{did}/rejected-saves/{rid}")
    assert "越权" in saved["markdown"] and "王测改乙" in saved["markdown"]
    # 暂存同步：只更新最新一份，份数不变
    for index in range(3):
        ctx.ok("PUT", f"/documents/{did}/rejected-saves/{rid}", {"markdown": f"第{index}次", "blockIds": None})
    assert ctx.ok("GET", f"/documents/{did}/rejected-saves/{rid}")["markdown"] == "第2次"
    assert _counts(did)[4] == after[4]
    ctx.as_("sun")
    assert ctx.req("GET", f"/documents/{did}/rejected-saves/{rid}").status_code == 404
    assert ctx.req("PUT", f"/documents/{did}/rejected-saves/{rid}", {"markdown": "x"}).status_code == 404
    ctx.as_("wang")
    # 不带 blockIds：422
    response = ctx.save(did, lambda p: replace_text(p, "功能正文乙", "x"), block_ids=False)
    assert response.status_code == 422 and response.json()["error"]["code"] == "BLOCK_IDS_REQUIRED"
    # 严格路径同样经过检查点
    response = ctx.save(did, lambda p: replace_text(p, "性能正文", "越权"), merge=False)
    assert response.status_code == 409 and response.json()["error"]["code"] == "SECTION_NOT_ASSIGNED"
    # 章节末尾新增段：允许；插在下一个标题之后：拒绝
    def at_end(pairs):
        index = next(i for i, (_id, b) in enumerate(pairs) if "功能正文乙" in b.body)
        return pairs[:index + 1] + [(None, Block("paragraph", "王测新增。", "\n\n", None))] + pairs[index + 1:]
    assert ctx.save(did, at_end).status_code == 200
    def after_next(pairs):
        index = next(i for i, (_id, b) in enumerate(pairs) if "性能需求" in b.body)
        return pairs[:index + 1] + [(None, Block("paragraph", "插到性能。", "\n\n", None))] + pairs[index + 1:]
    response = ctx.save(did, after_next)
    assert response.status_code == 409 and response.json()["error"]["details"]["violations"][0]["type"] == "insert"
    # 负责人不受限
    ctx.as_("owner")
    assert ctx.save(did, lambda p: replace_text(p, "性能正文", "负责人改"), block_ids=False).status_code == 200


def test_reassign_between_load_and_save_is_rejected_and_conflict_not_stored(ctx):
    doc = ctx.as_("owner").create()
    did = doc["id"]
    feature, perf = ctx.heading(did, "功能需求"), ctx.heading(did, "性能需求")
    ctx.assign(did, [feature, perf], ["wang"])
    ctx.as_("wang")
    prefix, revision, pairs, _ = ctx.state(did)
    ctx.as_("owner")
    ctx.assign(did, [perf], ["sun"])
    ctx.as_("wang")
    edited = replace_text(replace_text(pairs, "性能正文", "王测"), "功能正文甲", "王测甲")
    body = {"expectedRevision": revision, "markdownContent": assemble(prefix, [b for _i, b in edited]), "allowMerge": True,
            "clientSaveKey": uuid4().hex, "blockIds": [i for i, _b in edited]}
    response = ctx.req("PATCH", f"/documents/{did}/content", body)
    assert response.status_code == 409 and response.json()["error"]["code"] == "SECTION_NOT_ASSIGNED"
    # 冲突段触及别人的块：不生成冲突卡
    ctx.as_("sun")
    assert ctx.save(did, lambda p: replace_text(p, "性能正文", "孙敏先改")).status_code == 200
    ctx.as_("wang")
    before = _counts(did)
    response = ctx.req("PATCH", f"/documents/{did}/content", {**body, "clientSaveKey": uuid4().hex})
    assert response.status_code == 409 and response.json()["error"]["code"] == "SECTION_NOT_ASSIGNED"
    assert _counts(did)[3] == before[3]


def test_assignment_api_permissions_and_validation(ctx):
    doc = ctx.as_("owner").create()
    did = doc["id"]
    feature = ctx.heading(did, "功能需求")
    rev = ctx.ok("GET", f"/documents/{did}/assignments")["assignmentsRevision"]
    for who in ("wang", "viewer"):
        ctx.as_(who)
        response = ctx.req("PUT", f"/documents/{did}/assignments", {"objectIds": [feature], "assigneeIds": [], "expectedAssignmentsRevision": rev})
        assert response.status_code == 403
        assert ctx.req("PUT", f"/documents/{did}/assignment-mode", {"mode": "open", "expectedAssignmentsRevision": rev}).status_code == 403
    ctx.as_("out")
    assert ctx.req("GET", f"/documents/{did}/assignments").status_code == 404
    ctx.as_("owner")
    for name, reason in (("viewer", "readonly"), ("out", "not_member")):
        error = ctx.assign(did, [feature], [name], status=422)["error"]
        assert error["code"] == "ASSIGNEE_NOT_ELIGIBLE" and error["details"]["reason"] == reason
    with connect() as connection:
        connection.execute("UPDATE users SET status='disabled' WHERE id=%s", (ctx.users["liu"],))
        connection.commit()
    assert ctx.assign(did, [feature], ["liu"], status=422)["error"]["details"]["reason"] == "disabled"
    assert ctx.assign(did, [str(uuid4())], ["wang"], status=409)["error"]["code"] == "OBJECT_STALE"
    stale = ctx.req("PUT", f"/documents/{did}/assignments", {"objectIds": [feature], "assigneeIds": [], "expectedAssignmentsRevision": rev + 9})
    assert stale.status_code == 409 and stale.json()["error"]["code"] == "REVISION_CONFLICT"
    data = ctx.assign(did, [feature, "@document"], ["wang", "sun"])["data"]
    assert {item["objectId"] for item in data["explicit"]} == {feature, "@document"}
    # 停用后保留，userActive=false
    with connect() as connection:
        connection.execute("UPDATE users SET status='disabled' WHERE id=%s", (ctx.users["sun"],))
        connection.commit()
    view = ctx.ok("GET", f"/documents/{did}/assignments")
    sun = [a for item in view["explicit"] for a in item["assignees"] if a["name"] == "孙敏"]
    assert sun and not any(a["userActive"] for a in sun)
    # 分工模式
    rev = view["assignmentsRevision"]
    mode = ctx.ok("PUT", f"/documents/{did}/assignment-mode", {"mode": "open", "expectedAssignmentsRevision": rev})
    assert mode == {"assignmentMode": "open", "assignmentsRevision": rev + 1}
    assert ctx.req("PUT", f"/documents/{did}/assignment-mode", {"mode": "enforced", "expectedAssignmentsRevision": rev}).status_code == 409


def test_open_mode_writers_edit_everything(ctx):
    doc = ctx.as_("owner").create()
    did = doc["id"]
    rev = ctx.ok("GET", f"/documents/{did}/assignments")["assignmentsRevision"]
    ctx.ok("PUT", f"/documents/{did}/assignment-mode", {"mode": "open", "expectedAssignmentsRevision": rev})
    ctx.as_("wang")
    assert ctx.save(did, lambda p: replace_text(p, "概述正文", "开放编辑"), block_ids=False).status_code == 200
    view = ctx.ok("GET", f"/documents/{did}/assignments?resolved=1")
    assert all(item["editable"] for item in view["resolved"])


def test_member_demote_and_remove_clear_assignments_and_counts(ctx):
    doc = ctx.as_("owner").create()
    did = doc["id"]
    ctx.assign(did, [ctx.heading(did, "功能需求"), ctx.heading(did, "性能需求")], ["wang"])
    ctx.assign(did, [next(i for i, b in ctx.state(did)[2] if "概述正文" in b.body)], ["wang", "sun"])
    members = ctx.ok("GET", f"/projects/{ctx.project}/members")
    wang = next(item for item in members["items"] if item["name"] == "王测")
    assert wang["assignments"] == {"sections": 2, "blocks": 1, "documents": 1}
    rev = members["membersRevision"]
    response = ctx.ok("PUT", f"/projects/{ctx.project}/members/{ctx.users['wang']}", {"role": "viewer", "expectedMembersRevision": rev})
    assert response["assignmentsCleared"] == {"sections": 2, "blocks": 1, "documents": 1}
    view = ctx.ok("GET", f"/documents/{did}/assignments")
    assert all(a["name"] != "王测" for item in view["explicit"] for a in item["assignees"])
    rev = ctx.ok("GET", f"/projects/{ctx.project}/members")["membersRevision"]
    response = ctx.req("DELETE", f"/projects/{ctx.project}/members/{ctx.users['sun']}?expectedMembersRevision={rev}")
    assert response.status_code == 200 and response.json()["data"]["assignmentsCleared"]["blocks"] == 1
    assert ctx.ok("GET", f"/documents/{did}/assignments")["explicit"] == []


def test_summary_and_version_restore_and_objects_flags(ctx):
    doc = ctx.as_("owner").create()
    did = doc["id"]
    ctx.assign(did, [ctx.heading(did, "需求")], ["wang"])
    summary = ctx.ok("GET", f"/projects/{ctx.project}/assignment-summary")
    item = summary["byDocument"][0]
    assert item["sectionsTotal"] == 5 and item["sectionsAssigned"] == 4 and item["unassignedRoots"] == 1
    assert item["unassignedFrontMatter"] == 1
    assert [p["name"] for p in summary["byPerson"]] == ["王测"]
    assert {e["name"] for e in summary["unassignedEditors"]} == {"孙敏", "刘洋"}
    ctx.as_("wang")
    data = ctx.ok("GET", f"/documents/{did}/objects")
    flags = {o["markdown"]: o["editable"] for o in data["objects"]}
    assert flags["概述正文。"] is True and flags["范围正文。"] is False and "assignmentsRevision" in data
    assert ctx.save(did, lambda p: replace_text(p, "概述正文", "王测改")).status_code == 200
    response = ctx.req("POST", f"/documents/{did}/versions/1/restore", {"expectedRevision": 1})
    assert response.status_code == 403
    ctx.as_("viewer")
    assert ctx.req("GET", f"/projects/{ctx.project}/assignment-summary").status_code == 200


def test_split_inherits_block_assignment_and_deleted_object_assignment_removed(ctx):
    doc = ctx.as_("owner").create()
    did = doc["id"]
    ctx.assign(did, [ctx.heading(did, "功能需求")], ["wang"])
    exception = next(i for i, b in ctx.state(did)[2] if "功能正文甲" in b.body)
    ctx.assign(did, [exception], ["sun"])
    ctx.as_("sun")
    def enter(pairs):
        index = next(i for i, (oid, _b) in enumerate(pairs) if oid == exception)
        return pairs[:index + 1] + [(None, Block("paragraph", "孙敏回车新段。", "\n\n", None))] + pairs[index + 1:]
    assert ctx.save(did, enter).status_code == 200
    view = ctx.ok("GET", f"/documents/{did}/assignments?resolved=1")
    new = next(item for item in view["explicit"] if item["source"] == "split")
    assert new["needsConfirm"] is False and new["assignees"][0]["name"] == "孙敏"
    ctx.as_("wang")
    response = ctx.save(did, lambda p: replace_text(p, "孙敏回车新段", "王测改"))
    assert response.status_code == 409
    # 孙敏删除自己的例外段：分配随对象删除
    ctx.as_("sun")
    assert ctx.save(did, lambda p: [(i, b) for i, b in p if "孙敏回车新段" not in b.body]).status_code == 200
    view = ctx.ok("GET", f"/documents/{did}/assignments")
    assert all(item["source"] != "split" for item in view["explicit"])


def test_rejected_saves_keep_latest_five(ctx):
    doc = ctx.as_("owner").create()
    did = doc["id"]
    ctx.as_("wang")
    for index in range(7):
        response = ctx.save(did, lambda p, index=index: replace_text(p, "概述正文", f"越权{index}"))
        assert response.status_code == 409
    with connect() as connection:
        rows = connection.execute("SELECT markdown FROM document_rejected_saves WHERE document_id=%s", (did,)).fetchall()
    assert len(rows) == 5 and any("越权6" in row[0] for row in rows) and not any("越权0" in row[0] for row in rows)
