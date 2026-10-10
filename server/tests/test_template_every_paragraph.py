"""界面验收 M1：20 个 438C 模板里每个可编辑段落都改一次，整篇写回 Word 不再 422。

可编辑 = 原稿段落里没有图片、域（目录）、超链接、脚注、公式等受保护内容；只含书签的段落（如“文档标识号”行）可编辑。
所有可编辑段落一次全部追加标记，edit_original_docx 必须成功，导出稿里每个标记都在，书签仍在。
"""

from __future__ import annotations

from dataclasses import replace
from pathlib import Path

import pytest
from docx import Document

from app.docx_export import _PROTECTED_EDIT, _html_blocks, _word_blocks, edit_original_docx
from app.docx_parser import parse_docx
from app.markdown_blocks import assemble, plain_text, split

ROOT = Path(__file__).resolve().parents[2]
TEMPLATES = sorted((ROOT / "438C-2021全套模板").glob("*.docx"))


def _text(element) -> str:
    return " ".join("".join(element.itertext()).split())


@pytest.mark.skipif(not TEMPLATES, reason="仓库中没有 438C 模板")
@pytest.mark.parametrize("template", TEMPLATES, ids=[path.name[:10] for path in TEMPLATES])
def test_every_editable_paragraph_can_be_saved(template: Path, tmp_path: Path):
    markdown = parse_docx(str(template)).markdown
    source = _html_blocks(markdown)
    words = _word_blocks(Document(template))
    assert [block.tag for block in words] == [element.tag for element in source]
    editable_texts = []
    protected = 0
    for element, block in zip(source, words):
        if element.tag != "p" or not _text(element) or element.xpath(".//img"):
            continue
        if any(item.xpath(_PROTECTED_EDIT) for item in block.elements):
            protected += 1
            continue
        editable_texts.append(_text(element))
    prefix, blocks = split(markdown)
    pending = list(editable_texts)
    edited = 0
    for index, block in enumerate(blocks):
        if not pending or block.kind in {"heading", "table", "code"}:
            continue
        if plain_text(block.body) == pending[0]:
            pending.pop(0)
            blocks[index] = replace(block, body=f"{block.body}【改{index}】")
            edited += 1
    assert not pending, f"{template.name}: {len(pending)} 个可编辑段落没对上 Markdown 块，例如 {pending[0]!r}"
    target = assemble(prefix, blocks)
    output = tmp_path / "out.docx"
    edit_original_docx(target, template, output)
    exported = Document(output)
    text = "\n".join(paragraph.text for paragraph in exported.paragraphs)
    for table in exported.tables:
        text += "\n".join(cell.text for row in table.rows for cell in row.cells)
    assert text.count("【改") == edited, template.name
    for paragraph in exported.paragraphs:
        if paragraph.text.startswith("文档标识号") and "【改" in paragraph.text:
            original = next(p for p in Document(template).paragraphs if p.text.startswith("文档标识号"))
            assert len(paragraph._p.xpath("./w:bookmarkStart")) == len(original._p.xpath("./w:bookmarkStart"))
    print(f"{template.name}: 可编辑 {edited}，受保护 {protected}")
    assert edited > 10


@pytest.mark.skipif(not TEMPLATES, reason="仓库中没有 438C 模板")
def test_preview_marks_field_paragraphs_readonly():
    """测试 L3：目录行等含域段落在编辑区只读（contenteditable=false），“文档标识号”行可编辑。"""
    from lxml import html as html_dom

    from app.docx_preview import render_document_html

    template = next(path for path in TEMPLATES if "[SDP]" in path.name)
    markdown = parse_docx(str(template)).markdown
    root = html_dom.fragment_fromstring(render_document_html(template, markdown), create_parent="div")
    protected = root.xpath("//*[@data-word-protected]")
    assert protected and all(node.get("contenteditable") == "false" for node in protected)
    identifier = next(node for node in root if "文档标识号" in "".join(node.itertext()))
    assert identifier.get("contenteditable") is None


@pytest.mark.skipif(not TEMPLATES, reason="仓库中没有 438C 模板")
@pytest.mark.parametrize("name", ["[SDP]", "[SRS]"])
def test_preview_keeps_inline_tabs(name):
    """测试 L4：编辑区 HTML 保留行内制表符（Python-Markdown 默认展开成空格），首次保存不再把 \\t 改成空格。"""
    from app.docx_preview import render_document_html

    template = next(path for path in TEMPLATES if name in path.name)
    markdown = parse_docx(str(template)).markdown
    inline_tabs = sum(line.lstrip(" \t").count("\t") for line in markdown.splitlines())
    assert inline_tabs and render_document_html(template, markdown).count("\t") == inline_tabs
