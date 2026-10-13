"""A reviewed template must separate instructions without changing its Word package."""

from copy import deepcopy
from io import BytesIO
from pathlib import Path
from uuid import uuid4
from zipfile import ZipFile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from fastapi.testclient import TestClient
import pytest

from app.template_import import clean_template, guidance_for_section, parse_template, validate_structure
from app.docx_parser import parse_docx, heading_level
from app.docx_export import edit_original_docx
from app.db import connect, migrate
from app.files import store_bytes, resolve_object
from app.identity import create_user

MEDIA = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def word(tmp_path: Path, *, protected=False) -> Path:
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "固定页眉"
    doc.sections[0].footer.paragraphs[0].text = "固定页脚"
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "文档封面"
    doc.add_paragraph("目录")
    doc.add_paragraph("1 概述\t1", style="TOC Heading")
    doc.add_heading("概述", 1)
    doc.add_paragraph("本章应说明整体依据。")
    custom = doc.styles.add_style("自定义二级", WD_STYLE_TYPE.PARAGRAPH)
    custom.base_style = doc.styles["Heading 2"]
    doc.add_paragraph("系统用途", style=custom)
    guide = doc.add_paragraph("本条应描述真实系统用途：")
    if protected:
        field = OxmlElement("w:fldChar")
        field.set(qn("w:fldCharType"), "begin")
        guide.add_run()._r.append(field)
    doc.add_paragraph("a）描述运行环境。")
    doc.add_paragraph("1）列出使用对象。")
    doc.add_paragraph("a. 必须有实际资料依据。")
    p = doc.add_paragraph("系统应响应实际业务请求，不能仅凭‘应’字删掉本段。")
    p.runs[0].font.name, p.runs[0].font.size = "Microsoft YaHei", Pt(11)
    doc.add_paragraph("标题：")
    doc.add_paragraph("表1 功能列表")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "功能", "说明"
    table.cell(1, 0).text, table.cell(1, 1).text = "固定功能", "固定说明"
    doc.add_heading("接口", 2)
    doc.add_paragraph("本条应列出接口依据。")
    path = tmp_path / "template.docx"
    doc.save(path)
    return path


def test_inherited_headings_guidance_chain_and_word_package_preservation(tmp_path):
    path = word(tmp_path)
    source = path.read_bytes()
    structure = parse_template(path)
    blocks = structure["blocks"]
    assert next(item for item in blocks if item["text"] == "系统用途")["headingLevel"] == 2
    assert len([item for item in blocks if item["kind"] == "guidance"]) == 6
    assert next(item for item in blocks if item["text"].startswith("系统应响应"))["kind"] == "body"
    assert next(item for item in blocks if item["text"] == "文档封面")["kind"] == "cover"
    assert len([item for item in blocks if item["kind"] == "table"]) == 1
    assert next(item for item in blocks if item["text"] == "表1 功能列表")["kind"] == "caption"
    structure["reviewed"] = True
    cleaned = clean_template(path, structure)
    result = Document(BytesIO(cleaned))
    texts = [p.text for p in result.paragraphs]
    assert not any(text.startswith(("本条应", "本章应", "a）", "1）", "a.")) for text in texts)
    assert any(text.startswith("系统应响应") for text in texts)
    assert "标题：" not in texts
    assert texts.count("") == 3  # One reusable format sample per actual heading.
    assert result.tables[1].cell(1, 1).text == "固定说明"
    with ZipFile(BytesIO(source)) as before, ZipFile(BytesIO(cleaned)) as after:
        assert before.namelist() == after.namelist()
        for name in before.namelist():
            if name != "word/document.xml":
                assert before.read(name) == after.read(name), name
    assert path.read_bytes() == source
    base = tmp_path / "clean.docx"
    base.write_bytes(cleaned)
    parsed = parse_docx(str(base), outline_headings=True)
    assert "## 系统用途" in parsed.markdown
    actual_guides = guidance_for_section(structure, parsed.markdown, 1)
    assert any("整体依据" in rule["text"] for rule in actual_guides)
    assert any("真实系统用途" in rule["text"] for rule in actual_guides)
    assert not any("接口依据" in rule["text"] for rule in actual_guides)
    unmatched = guidance_for_section(structure, parsed.markdown.replace("系统用途", "改名用途"), 1)
    assert any(item["kind"] == "notice" for item in unmatched)
    assert not any("真实系统用途" in item["text"] for item in unmatched)
    assert guidance_for_section(structure, parsed.markdown, None) == []
    edited = tmp_path / "edited.docx"
    edit_original_docx(parsed.markdown.replace("系统应响应实际业务请求", "系统应按既定流程响应业务请求"), base, edited, outline_headings=True)
    exported = Document(edited)
    assert exported.sections[0].header.paragraphs[0].text == "固定页眉"
    actual = next(p for p in exported.paragraphs if p.text.startswith("系统应按"))
    assert actual.runs[0].font.name == "Microsoft YaHei" and actual.runs[0].font.size == Pt(11)
    from app.docx_preview import render_document_html
    rendered = render_document_html(base, parsed.markdown, outline_headings=True)
    assert "系统用途" in rendered and "本条应" not in rendered


def test_traceability_table_instruction_is_guidance_and_keeps_table_layout(tmp_path):
    path = word(tmp_path)
    document = Document(path)
    document.tables[1].cell(1, 1).text = "若存在一对多则此形式表达"
    document.save(path)
    structure = parse_template(path)
    block = next(item for item in structure["blocks"] if not item["isTable"] and item["text"] == "若存在一对多则此形式表达")
    assert block["kind"] == "guidance" and block["headingId"]
    cleaned = Document(BytesIO(clean_template(path, {**structure, "reviewed": True})))
    assert len(cleaned.tables) == 2
    assert cleaned.tables[1].cell(1, 1).text == ""
    assert cleaned.tables[1].cell(0, 0).text == document.tables[1].cell(0, 0).text
    assert len(cleaned.tables[1].cell(1, 1).paragraphs) == 1


def test_long_guidance_chain_leaves_one_body_format_sample_and_new_text_stays_body(tmp_path):
    document = Document()
    custom = document.styles.add_style("模板二级标题", WD_STYLE_TYPE.PARAGRAPH)
    custom.base_style = document.styles["Heading 2"]
    document.add_heading("需求", 1)
    document.add_paragraph("接口", style=custom)
    guide = document.add_paragraph("本条应列出接口要求：")
    guide.paragraph_format.first_line_indent = Pt(22)
    guide.paragraph_format.space_after = Pt(6)
    guide.runs[0].font.name, guide.runs[0].font.size = "Microsoft YaHei", Pt(11)
    for index in range(45):
        document.add_paragraph(f"a）第{index + 1}项编写要求。")
    document.add_paragraph("字段：")
    document.add_heading("其他", 2)
    path = tmp_path / "long-guidance.docx"
    document.save(path)
    structure = parse_template(path)
    field = next(item for item in structure["blocks"] if item["text"] == "字段：")
    field["kind"] = "placeholder"
    cleaned = clean_template(path, {**structure, "reviewed": True})
    result = Document(BytesIO(cleaned))
    assert [p.text for p in result.paragraphs] == ["需求", "接口", "", "其他"]
    sample = result.paragraphs[2]
    assert sample.paragraph_format.first_line_indent == Pt(22)
    assert sample.runs[0].font.name == "Microsoft YaHei" and sample.runs[0].font.size == Pt(11)
    base = tmp_path / "clean.docx"
    base.write_bytes(cleaned)
    markdown = parse_docx(str(base), outline_headings=True).markdown.replace("## 接口", "## 接口\n\n实际接口正文。")
    target = tmp_path / "written.docx"
    edit_original_docx(markdown, base, target, outline_headings=True)
    added = next(p for p in Document(target).paragraphs if p.text == "实际接口正文。")
    assert heading_level(added) is None
    assert added.paragraph_format.first_line_indent == Pt(22)
    assert added.runs[0].font.name == "Microsoft YaHei" and added.runs[0].font.size == Pt(11)
    from app.docx_preview import render_document_html
    assert "实际接口正文" in render_document_html(base, markdown, outline_headings=True)


def test_validation_prevents_source_tampering_missing_blocks_and_protected_cleanup(tmp_path):
    source = parse_template(word(tmp_path))
    for mutation in (lambda x: x["blocks"].pop(), lambda x: x["blocks"].append(x["blocks"][0]),
                     lambda x: x["blocks"][0].update(sourceText="篡改"), lambda x: x.update(sourceHash="bad"),
                     lambda x: x["blocks"][0].update(kind="guidance"), lambda x: x["blocks"][0].update(protected=False)):
        bad = deepcopy(source)
        mutation(bad)
        # A same-valued false is not tampering; force a protected-marker mismatch.
        if bad == source:
            bad["blocks"][0]["protected"] = not source["blocks"][0]["protected"]
        with pytest.raises(ValueError):
            validate_structure(source, bad)
    guide = next(x for x in source["blocks"] if x["kind"] == "guidance")
    edited = deepcopy(source)
    next(x for x in edited["blocks"] if x["id"] == guide["id"])["text"] += " 用户补充约束。"
    assert "用户补充" in validate_structure(source, edited)["blocks"][source["blocks"].index(guide)]["text"]
    protected_path = word(tmp_path, protected=True)
    protected = parse_template(protected_path)
    protected["reviewed"] = True
    with pytest.raises(ValueError, match="安全清理"):
        clean_template(protected_path, protected)


def test_reviewed_level_and_manual_guidance_binding_take_effect(tmp_path):
    path = word(tmp_path)
    source = parse_template(path)
    candidate = deepcopy(source)
    target = next(item for item in candidate["blocks"] if item["text"] == "系统用途")
    target["headingLevel"] = 3
    interface = next(item for item in candidate["blocks"] if item["text"] == "接口")
    guide = next(item for item in candidate["blocks"] if item["text"] == "本条应描述真实系统用途：")
    guide["headingId"] = interface["id"]
    candidate["reviewed"] = True
    cleaned = clean_template(path, candidate)
    document = Document(BytesIO(cleaned))
    assert heading_level(next(p for p in document.paragraphs if p.text == "系统用途")) == 3
    assert next(p for p in document.paragraphs if p.text == "系统用途").style.name == "自定义二级"
    markdown = parse_docx(BytesIO(cleaned), outline_headings=True).markdown
    assert "### 系统用途" in markdown
    assert not any("真实系统用途" in x["text"] for x in guidance_for_section(candidate, markdown, 1))
    assert any("真实系统用途" in x["text"] for x in guidance_for_section(candidate, markdown, 2))
    downgraded = deepcopy(candidate)
    next(x for x in downgraded["blocks"] if x["id"] == target["id"]).update(kind="body", headingLevel=None)
    root = next(x for x in downgraded["blocks"] if x["text"] == "概述")
    for item in downgraded["blocks"]:
        if item["kind"] == "guidance" and item["headingId"] == target["id"]:
            item["headingId"] = root["id"]
    normalized = validate_structure(source, downgraded)
    assert next(x for x in normalized["blocks"] if x["text"].startswith("系统应响应"))["headingId"] == root["id"]
    assert heading_level(next(p for p in Document(BytesIO(clean_template(path, normalized))).paragraphs if p.text == "系统用途")) is None
    promoted = deepcopy(source)
    new_heading = next(x for x in promoted["blocks"] if x["text"] == "标题：")
    new_heading.update(kind="heading", headingLevel=3, headingId=None)
    assert heading_level(next(p for p in Document(BytesIO(clean_template(path, {**promoted, "reviewed": True}))).paragraphs if p.text == "标题：")) == 3


def test_new_subheading_inherits_unique_ancestors_and_ambiguous_paths_are_rejected(tmp_path):
    path = word(tmp_path)
    source = parse_template(path)
    source["reviewed"] = True
    markdown = parse_docx(BytesIO(clean_template(path, source)), outline_headings=True).markdown
    added = markdown.replace("## 系统用途", "## 系统用途\n\n### 设备登记")
    inherited = guidance_for_section(source, added, 2)
    assert any("整体依据" in item["text"] for item in inherited)
    assert any("真实系统用途" in item["text"] for item in inherited)
    assert any(item["kind"] == "notice" for item in inherited)
    duplicate = deepcopy(source)
    root = next(item for item in source["blocks"] if item["kind"] == "heading")
    duplicate["blocks"].append({**root, "id": "duplicate-heading"})
    with pytest.raises(ValueError, match="重复"):
        guidance_for_section(duplicate, markdown, 1)
    with pytest.raises(ValueError, match="无法对应"):
        guidance_for_section(source, markdown.replace("# 概述", "# 改名总则"), 1)


def test_legacy_custom_heading_markdown_export_is_unchanged(tmp_path):
    path = word(tmp_path)
    parsed = parse_docx(str(path))
    assert "## 系统用途" not in parsed.markdown
    target = tmp_path / "legacy.docx"
    edit_original_docx(parsed.markdown, path, target)
    assert target.read_bytes() == path.read_bytes()
    edited = tmp_path / "legacy-edited.docx"
    edit_original_docx(parsed.markdown.replace("固定说明", "修改说明"), path, edited)
    assert Document(edited).tables[1].cell(1, 1).text == "修改说明"


@pytest.fixture()
def api():
    migrate()
    from app.main import create_app
    with TestClient(create_app()) as client:
        user = create_user(account="template-import-admin", name="模板负责人", password="correct-horse-1", roles=["system_admin"])
        csrf = client.get("/api/v1/auth/csrf").json()["data"]["csrfToken"]
        login = client.post("/api/v1/auth/login", json={"account": "template-import-admin", "password": "correct-horse-1"}, headers={"X-CSRF-Token": csrf})
        headers = {"X-CSRF-Token": login.json()["data"]["csrfToken"]}
        project = uuid4()
        with connect() as connection:
            connection.execute("INSERT INTO projects (id,name,owner_id) VALUES (%s,%s,%s)", (project, "独立测试项目", user))
            connection.commit()
        yield client, headers, project


def test_api_review_cas_creation_and_guidance_free_export(api, tmp_path):
    client, headers, project = api
    file = store_bytes(name="模板.docx", media_type=MEDIA, purpose="template", content=word(tmp_path).read_bytes())
    group = client.post("/api/v1/templates", json={"name": "模板组", "category": "438C"}, headers=headers).json()["data"]
    imported = client.post(f"/api/v1/templates/{group['id']}/documents", json={"name": "需求说明", "wordFileId": str(file.id)}, headers=headers)
    assert imported.status_code == 200, imported.text
    template = imported.json()["data"]
    endpoint = f"/api/v1/template-documents/{template['id']}/structure"
    new_document = {"name": "项目需求", "sourceType": "template", "templateDocumentId": template["id"]}
    assert client.post(f"/api/v1/projects/{project}/documents", json=new_document, headers=headers).status_code == 409
    structure = client.get(f"/api/v1/template-documents/{template['id']}").json()["data"]["importStructure"]
    structure["reviewed"] = True
    saved = client.patch(endpoint, json={"expectedRevision": 0, "importStructure": structure}, headers=headers)
    assert saved.status_code == 200, saved.text
    assert saved.json()["data"]["importStructure"]["revision"] == 1
    assert client.patch(endpoint, json={"expectedRevision": 0, "importStructure": structure}, headers=headers).status_code == 409
    assert client.get(f"/api/v1/templates/{group['id']}/documents").json()["data"]["items"][0]["importStructure"]["reviewed"]
    created = client.post(f"/api/v1/projects/{project}/documents", json=new_document, headers=headers)
    assert created.status_code == 200, created.text
    document = created.json()["data"]
    assert "本条应" not in document["markdownContent"]
    assert document["originalDocxFileId"] != str(file.id)
    assert document["formatTemplate"]["importStructure"]["revision"] == 1
    preview = client.get(f"/api/v1/documents/{document['id']}/preview")
    assert preview.status_code == 200, preview.text
    assert "本条应" not in preview.json()["data"]["html"]
    structure = saved.json()["data"]["importStructure"]
    next(x for x in structure["blocks"] if x["kind"] == "guidance")["text"] += " 后续模板变化。"
    # A06：确认后的结构是第 1 版，之后只能经草稿 → 发布新版本修改；已有文档仍用自己的拷贝
    refused = client.patch(endpoint, json={"expectedRevision": 1, "importStructure": structure}, headers=headers)
    assert refused.status_code == 409 and refused.json()["error"]["code"] == "TEMPLATE_VERSIONED"
    draft = client.put(f"/api/v1/template-documents/{template['id']}/draft", json={"structure": structure, "expectedDraftRevision": 0}, headers=headers)
    assert draft.status_code == 200, draft.text
    published = client.post(f"/api/v1/template-documents/{template['id']}/versions", json={"draftRevision": 1, "note": "改指导"}, headers=headers)
    assert published.status_code == 201, published.text
    assert published.json()["data"]["version"]["seq"] == 2
    assert client.get(f"/api/v1/documents/{document['id']}").json()["data"]["formatTemplate"]["importStructure"]["revision"] == 1
    exported = client.post(f"/api/v1/documents/{document['id']}/export", json={"expectedRevision": 0}, headers=headers)
    assert exported.status_code == 200, exported.text
    downloaded = client.get(f"/api/v1/files/{exported.json()['data']['fileId']}/content")
    assert downloaded.status_code == 200
    assert not any("本条应" in p.text for p in Document(BytesIO(downloaded.content)).paragraphs)


def test_parse_api_source_permissions_corruption_and_immutable_fields(api, tmp_path):
    client, headers, _ = api
    broken = store_bytes(name="坏.docx", media_type=MEDIA, purpose="template", content=b"bad")
    response = client.post("/api/v1/template-documents/parse", json={"wordFileId": str(broken.id)}, headers=headers)
    assert response.status_code == 422
    file = store_bytes(name="模板.docx", media_type=MEDIA, purpose="template", content=word(tmp_path).read_bytes())
    parsed = client.post("/api/v1/template-documents/parse", json={"wordFileId": str(file.id)}, headers=headers)
    assert parsed.status_code == 200, parsed.text
    structure = parsed.json()["data"]
    structure["blocks"][0]["sourceText"] = "伪造来源"
    group = client.post("/api/v1/templates", json={"name": "模板组", "category": "438C"}, headers=headers).json()["data"]
    refused = client.post(f"/api/v1/templates/{group['id']}/documents", json={"name": "伪造", "wordFileId": str(file.id), "importStructure": structure}, headers=headers)
    assert refused.status_code == 422
    from app.config import get_settings
    resolve_object(get_settings().file_data_dir, file.storage_key).unlink()
    assert client.post("/api/v1/template-documents/parse", json={"wordFileId": str(file.id)}, headers=headers).status_code == 404


def test_parse_and_import_refuse_another_projects_private_file(api, tmp_path):
    client, _, _ = api
    owner = create_user(account="private-owner", name="资料负责人", password="correct-horse-1", roles=["project_creator"])
    private_project = uuid4()
    with connect() as connection:
        connection.execute("INSERT INTO projects (id,name,owner_id) VALUES (%s,%s,%s)", (private_project, "私有项目", owner))
        connection.commit()
    file = store_bytes(name="私有.docx", media_type=MEDIA, purpose="attachment", project_id=private_project, content=word(tmp_path).read_bytes())
    create_user(account="limited-manager", name="模板管理员", password="correct-horse-1", roles=["template_manager"])
    with TestClient(client.app) as other:
        csrf = other.get("/api/v1/auth/csrf").json()["data"]["csrfToken"]
        login = other.post("/api/v1/auth/login", json={"account": "limited-manager", "password": "correct-horse-1"}, headers={"X-CSRF-Token": csrf})
        headers = {"X-CSRF-Token": login.json()["data"]["csrfToken"]}
        assert other.post("/api/v1/template-documents/parse", json={"wordFileId": str(file.id)}, headers=headers).status_code == 404
        group = other.post("/api/v1/templates", json={"name": "无权限模板", "category": "438C"}, headers=headers).json()["data"]
        assert other.post(f"/api/v1/templates/{group['id']}/documents", json={"name": "私有", "wordFileId": str(file.id)}, headers=headers).status_code == 404
