"""Independent reports reuse the editing path's original Word formatting."""

from contextlib import nullcontext
from types import SimpleNamespace
import struct
from uuid import uuid4
from zipfile import ZipFile
import zlib

import pytest

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

from app import test_executor


def test_report_inherits_direct_paragraph_heading_table_and_picture_format(tmp_path, monkeypatch):
    source = Document()
    source.sections[0].left_margin = Mm(31)
    source.sections[0].right_margin = Mm(23)
    source.sections[0].header.paragraphs[0].text = "原稿页眉"
    source.sections[0].footer.paragraphs[0].text = "原稿页脚"

    def format_paragraph(paragraph, font, size, indent, spacing):
        paragraph.paragraph_format.first_line_indent = Pt(indent)
        paragraph.paragraph_format.line_spacing = Pt(spacing)
        paragraph.paragraph_format.space_after = Pt(7)
        for run in paragraph.runs:
            run.font.name, run.font.size = font, Pt(size)
            run._r.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)

    heading = source.add_heading("原稿一级标题", 1)
    format_paragraph(heading, "黑体", 18, 0, 30)
    subheading = source.add_heading("原稿二级标题", 2)
    format_paragraph(subheading, "仿宋", 16, 14, 28)
    body = source.add_paragraph("原稿正文格式必须保留，不采用默认样式中的字号。")
    body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    format_paragraph(body, "仿宋", 14, 28, 24)
    body.paragraph_format.left_indent = Pt(8)
    table = source.add_table(rows=2, cols=2)
    table.style = "Light Shading Accent 1"
    table.autofit = False
    table.columns[0].width, table.columns[1].width = Mm(55), Mm(105)
    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            cell.text = f"原表{row_index}{column_index}"
            format_paragraph(cell.paragraphs[0], "宋体", 11, 0, 18)
    source_path = tmp_path / "original.docx"
    source.save(source_path)
    original_bytes = source_path.read_bytes()
    original_xml = [p._p.xml for p in source.paragraphs]

    def chunk(kind, data):
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))

    png = (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 1920, 1080, 8, 2, 0, 0, 0))
           + chunk(b"IDAT", zlib.compress((b"\0" + b"\x80\x80\x80" * 1920) * 1080)) + chunk(b"IEND", b""))
    (tmp_path / "evidence.png").write_bytes(png)
    connection = SimpleNamespace(execute=lambda *args: SimpleNamespace(fetchone=lambda: {"storage_key": "fixture"}))
    monkeypatch.setattr(test_executor, "connect", lambda: nullcontext(connection))
    monkeypatch.setattr(test_executor, "resolve_object", lambda *args: source_path)
    monkeypatch.setattr(test_executor, "run_directory", lambda _: tmp_path)
    cases = [
        {"caseId": str(uuid4()), "title": "功能测试", "kind": "functional", "status": "passed", "expected": "显示提示", "steps": [
            {"action": "assert_text", "description": "核对提示", "status": "passed", "actual": "提示可见", "screenshot": "evidence.png"},
        ]},
        {"caseId": str(uuid4()), "title": "并发测试", "kind": "concurrency", "status": "failed", "expected": "全部通过", "steps": [],
         "metrics": {"requested": 20, "completed": 20, "concurrency": 5, "failed": 1, "averageMs": 10, "p95Ms": 15, "maxLatencyMs": 2000}},
    ]
    output = test_executor.build_report({"name": "报告格式验收", "reportFileId": "fixture",
                                        "frontendUrl": "https://frontend.invalid", "backendUrl": "https://backend.invalid/api"},
                                       {"id": str(uuid4()), "status": "failed", "result": {"cases": cases}})
    result = Document(output)
    assert source_path.read_bytes() == original_bytes
    assert [p._p.xml for p in result.paragraphs[:len(original_xml)]] == original_xml
    assert result.sections[0].left_margin == source.sections[0].left_margin
    assert result.sections[0].right_margin == source.sections[0].right_margin
    assert result.sections[0].header.paragraphs[0].text == "原稿页眉"
    assert result.sections[0].footer.paragraphs[0].text == "原稿页脚"
    for address in ("前端地址：https://frontend.invalid", "后端地址：https://backend.invalid/api"):
        paragraphs = [p for p in result.paragraphs if p.text == address]
        assert len(paragraphs) == 1
        assert "\n" not in paragraphs[0].text and not paragraphs[0]._p.xpath('.//w:br')

    def check_format(paragraph, sample):
        assert paragraph.runs[0].font.name == sample.runs[0].font.name
        assert paragraph.runs[0].font.size == sample.runs[0].font.size
        assert paragraph.runs[0]._r.rPr.rFonts.get(qn("w:eastAsia")) == sample.runs[0]._r.rPr.rFonts.get(qn("w:eastAsia"))
        assert paragraph.paragraph_format.first_line_indent == sample.paragraph_format.first_line_indent
        assert paragraph.paragraph_format.line_spacing == sample.paragraph_format.line_spacing
        assert paragraph.paragraph_format.space_after == sample.paragraph_format.space_after

    title = next(p for p in result.paragraphs if p.text == "报告格式验收 实测记录")
    check_format(title, heading)
    assert title.paragraph_format.page_break_before is True
    check_format(next(p for p in result.paragraphs if p.text == "功能测试"), subheading)
    for text in ("类型：功能测试", "步骤 1：核对提示", "实际结果：失败"):
        paragraph = next(p for p in result.paragraphs if p.text == text)
        check_format(paragraph, body)
        assert paragraph.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
        assert paragraph.paragraph_format.left_indent == Pt(8)
    assert len(result.tables) == 2
    metrics = result.tables[1]
    assert metrics.style == result.tables[0].style and metrics.autofit is False
    assert metrics._tbl.tblGrid.xml == result.tables[0]._tbl.tblGrid.xml
    check_format(metrics.cell(1, 0).paragraphs[0], table.cell(1, 0).paragraphs[0])
    assert metrics.cell(0, 0).paragraphs[0].runs[0].font.size == Pt(11)
    assert metrics.rows[0]._tr.xpath('./w:trPr/w:tblHeader')
    assert len(result.inline_shapes) == 1
    shape = result.inline_shapes[0]
    assert shape.width <= result.sections[0].page_width - result.sections[0].left_margin - result.sections[0].right_margin
    picture = shape._inline.xpath('ancestor::w:p')[0]
    actual, caption = list(picture.itersiblings(preceding=True))[:2]
    assert picture.xpath('./w:pPr/w:jc')[0].get(qn("w:val")) == "center"
    assert caption.xpath('./w:pPr/w:keepNext') and actual.xpath('./w:pPr/w:keepNext')
    with ZipFile(output) as package, ZipFile(source_path) as original:
        assert package.read("word/styles.xml") == original.read("word/styles.xml")
        assert [package.read(name) for name in package.namelist() if name.startswith("word/media/")] == [png]


@pytest.mark.parametrize("case_id", [None, "", "../../bad", "<bad>"])
def test_report_refuses_invalid_case_ids_before_writing_output(tmp_path, monkeypatch, case_id):
    monkeypatch.setattr(test_executor, "run_directory", lambda _: tmp_path)
    # Evidence validation reaches the identifier without needing an image decoder.
    (tmp_path / "evidence.png").write_bytes(b"\x89PNG\r\n\x1a\nfixture")
    case = {"caseId": case_id, "title": "异常标识", "kind": "functional", "status": "passed", "steps": [
        {"action": "assert_text", "status": "passed", "screenshot": "evidence.png"},
    ]}
    with pytest.raises(test_executor.MissingInput, match="测试结果标识不正确"):
        test_executor.build_report({"name": "异常标识"}, {"id": str(uuid4()), "status": "completed", "result": {"cases": [case]}})
    assert not (tmp_path / "test-report.docx").exists()
