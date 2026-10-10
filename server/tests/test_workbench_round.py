"""工作台重塑轮：上轮发现的两个保存/预览缺陷。"""

import pytest

from app.docx_export import _html_blocks
from tests.test_documents import create_document, document_context  # noqa: F401


def test_html_comments_are_dropped_and_empty_headings_ignored():
    blocks = _html_blocks("# 范围\n\n<!-- 注释 -->\n\n正文<!-- 行内注释 -->继续。\n\n#\n\n## \n\n### 有字的标题")
    assert [block.tag for block in blocks] == ["h1", "p", "h3"]
    assert "".join(blocks[1].itertext()) == "正文继续。"
    assert all("注释" not in "".join(block.itertext()) for block in blocks)


@pytest.mark.parametrize("addition", ["<!-- 注释 -->", "正文<!-- 注释 -->继续。", "#", "## "])
def test_comment_and_empty_heading_save_and_preview(document_context, addition):  # noqa: F811
    client, headers, *_ = document_context
    document = create_document(document_context, name=f"缺陷-{addition}")
    content = document["markdownContent"].replace("原始正文。", f"原始正文。\n\n{addition}\n\n")
    saved = client.patch(f"/api/v1/documents/{document['id']}/content",
                         json={"expectedRevision": 0, "markdownContent": content}, headers=headers)
    # 以前：注释保存返回英文 lxml 报错（422）；只写“#”时保存成功但预览 422
    assert saved.status_code == 200, saved.text
    preview = client.get(f"/api/v1/documents/{document['id']}/preview")
    assert preview.status_code == 200, preview.text
    assert "注释" not in preview.json()["data"]["html"]


def _bookmarked(paragraph, name="OLE_LINK1"):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), "1")
    start.set(qn("w:name"), name)
    paragraph._p.insert(0, start)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), "1")
    paragraph._p.append(end)


def test_refusal_reason_names_the_actual_cause(tmp_path):
    """界面复验 N8：拒绝删除时写明实际原因（书签），删掉的是标题行时单独说明。"""
    from docx import Document

    from app.docx_export import WordEditError, edit_original_docx
    from app.docx_parser import parse_docx

    word = Document()
    word.add_heading("1 范围", level=1)
    _bookmarked(word.add_heading("1.1 标识", level=2), "OLE_LINK2")
    word.add_paragraph("本文档适用于系统。")
    _bookmarked(word.add_paragraph("文档标识号：TN/x"))
    source = tmp_path / "src.docx"
    word.save(source)
    markdown = parse_docx(str(source)).markdown
    reasons = []
    for edited in (markdown.replace("\n\n文档标识号：TN/x", ""), markdown.replace("## 1.1 标识\n\n", "")):
        assert edited != markdown
        with pytest.raises(WordEditError) as error:
            edit_original_docx(edited, source, tmp_path / "out.docx")
        reasons.append(str(error.value))
    assert "书签" in reasons[0] and "图片" not in reasons[0] and "标题行" not in reasons[0]
    assert "标题行" in reasons[1] and "书签" in reasons[1]


def test_preview_object_ids_match_tables_ignoring_whitespace(monkeypatch):
    """工作台重塑轮：表格单元格在预览里有空白分隔，不能因此判为对不上（否则几乎所有模板文档都显示“整篇保存”）。"""
    from uuid import uuid4

    import app.document_objects as objects

    rows = [{"id": uuid4(), "revision": 1, "body": "正文一段。"},
            {"id": uuid4(), "revision": 2, "body": "| 密 级： | 内部 |\n|---|---|\n| 阶 段： | A 版 |"}]
    monkeypatch.setattr(objects, "live_objects", lambda connection, document_id: rows)
    html = "<p>正文一段。</p><table><tr><td>密 级：</td><td>内部</td></tr><tr><td>阶 段：</td><td>A 版</td></tr></table>"
    decorated, ok = objects.decorate_object_ids(None, uuid4(), html)
    assert ok and decorated.count("data-object-id") == 2
    # 嵌套列表：预览顶层 li 与对象数对不上，仍按整篇保存
    nested = [{"id": uuid4(), "revision": 1, "body": "- 一\n  - 一点一\n- 二"}]
    monkeypatch.setattr(objects, "live_objects", lambda connection, document_id: nested)
    _, ok = objects.decorate_object_ids(None, uuid4(), "<ul><li>一<ul><li>一点一</li></ul></li><li>二</li></ul>")
    assert ok is False
