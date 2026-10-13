from copy import deepcopy
from io import BytesIO
from uuid import uuid4

import pytest
from docx import Document

from app import template_merge as tm
from app import template_word as tw
from app.docx_parser import parse_docx
from app.markdown_blocks import split
from app.template_guidance import bind_guidance_targets
from app.template_import import clean_template, parse_template, validate_guidance_targets, validate_structure


@pytest.fixture
def source(tmp_path):
    document = Document()
    document.add_heading("需求", level=1)
    document.add_paragraph("编写说明：第一项约束。")
    document.add_paragraph("编写说明：第二项约束。")
    document.add_paragraph("重复正文。")
    document.add_paragraph("重复正文。")
    document.add_paragraph("条目正文。", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    for row in table.rows:
        for cell in row.cells:
            cell.text = "数据"
    path = tmp_path / "template.docx"
    document.save(path)
    structure = parse_template(path)
    structure["reviewed"] = True
    guidance = [block for block in structure["blocks"] if block["kind"] == "guidance"]
    bodies = [block for block in structure["blocks"] if block["kind"] == "body" and not block["tableId"]]
    table_block = next(block for block in structure["blocks"] if block["isTable"])
    guidance[0]["targetBlockIds"] = [bodies[1]["id"]]
    guidance[1]["targetBlockIds"] = [table_block["id"], bodies[2]["id"]]
    return path, structure, guidance, bodies, table_block


@pytest.mark.parametrize("clean", [clean_template, tw.clean_version])
def test_binding_survives_cleanup_and_duplicate_text(source, clean):
    path, structure, guidance, bodies, table = source
    paths = {}
    cleaned = clean(path, structure, target_paths=paths)
    markdown = parse_docx(BytesIO(cleaned), outline_headings=True).markdown
    _, blocks = split(markdown)
    objects = [{"id": uuid4(), "kind": block.kind, "body": block.body} for block in blocks]
    bindings = bind_guidance_targets(cleaned, structure, markdown, objects, target_paths=paths)
    duplicates = [row for row in objects if row["body"] == "重复正文。"]
    assert bindings[guidance[0]["id"]] == [str(duplicates[1]["id"])]
    table_row = next(row for row in objects if row["kind"] == "table")
    list_row = next(row for row in objects if row["kind"] == "list_item")
    assert bindings[guidance[1]["id"]] == [str(table_row["id"]), str(list_row["id"])]
    assert paths[bodies[1]["id"]] != bodies[1]["path"]
    assert b"GUIDANCETARGET" not in cleaned


@pytest.mark.parametrize("failure", ["deleted", "empty", "duplicate", "cell", "guidance", "protected", "wrong_type"])
def test_invalid_binding_fails_explicitly(source, failure):
    _, structure, guidance, bodies, _ = source
    if failure == "deleted":
        guidance[0]["targetBlockIds"] = ["block-0000000000000000"]
    elif failure == "empty":
        bodies[1]["text"] = " "
    elif failure == "duplicate":
        guidance[0]["targetBlockIds"] *= 2
    elif failure == "cell":
        guidance[0]["targetBlockIds"] = [next(b["id"] for b in structure["blocks"] if b["tableId"])]
    elif failure == "guidance":
        guidance[0]["targetBlockIds"] = [guidance[1]["id"]]
    elif failure == "protected":
        bodies[1]["protected"] = True
    else:
        guidance[0]["targetBlockIds"] = "wrong"
    with pytest.raises(ValueError, match="指导"):
        validate_guidance_targets(structure["blocks"])


def test_create_refuses_missing_paths_and_changed_object_snapshot(source):
    path, structure, _, _, _ = source
    paths = {}
    cleaned = clean_template(path, structure, target_paths=paths)
    markdown = parse_docx(BytesIO(cleaned), outline_headings=True).markdown
    _, blocks = split(markdown)
    objects = [{"id": uuid4(), "kind": block.kind, "body": block.body} for block in blocks]
    with pytest.raises(ValueError, match="目标"):
        bind_guidance_targets(cleaned, structure, markdown, objects, target_paths={})
    with pytest.raises(ValueError, match="未对齐"):
        bind_guidance_targets(cleaned, structure, markdown, objects[::-1], target_paths=paths)
    with pytest.raises(ValueError, match="正文已变化"):
        bind_guidance_targets(cleaned, structure, markdown + "新增", objects, target_paths=paths)


@pytest.mark.parametrize("clean", [clean_template, tw.clean_version])
def test_table_with_only_instructions_is_not_a_writable_target(source, clean):
    path, structure, _, _, table = source
    for block in structure["blocks"]:
        if block.get("tableId") == table["id"]:
            block["kind"] = "guidance"
    with pytest.raises(ValueError, match="清理提示词后没有正文"):
        clean(path, structure)


def test_schema_draft_publish_and_merge_preserve_bindings(source, tmp_path):
    path, structure, guidance, bodies, _ = source
    raw = parse_template(path)
    validated = validate_structure(raw, structure)
    assert validated["blocks"] == structure["blocks"]
    analysis = tw.analyze(path)
    draft = {**structure, "blocks": tm.validate_draft(structure, structure, analysis)}
    assert draft["blocks"][1]["targetBlockIds"] == guidance[0]["targetBlockIds"]
    output = tmp_path / "published.docx"
    tw.apply_draft(path, draft, analysis["styles"], analysis["styles"], output)
    published = tw.build_version(output, draft, 2)
    assert published["blocks"][1]["targetBlockIds"] == guidance[0]["targetBlockIds"]
    mine = deepcopy(draft)
    mine["blocks"][1]["targetBlockIds"] = [bodies[0]["id"]]
    assert tm.semantic(mine) != tm.semantic(draft)
    merged = tm.merge3(draft, draft, mine, styles=({}, {}, {}))
    assert merged["blocks"][1]["targetBlockIds"] == [bodies[0]["id"]]
    changes = tm.diff(draft, mine, {}, {})
    change = next(item for item in changes["blocks"] if item["blockId"] == guidance[0]["id"])
    assert change["before"]["targetTexts"] != change["after"]["targetTexts"]
    assert "重复正文。" in change["after"]["targetTexts"][0]


def test_reupload_keeps_binding_and_missing_target_is_unplaced(source):
    path, structure, guidance, _, _ = source
    result, _ = tw.align_upload(parse_template(path), structure)
    assert result["blocks"][1]["targetBlockIds"] == guidance[0]["targetBlockIds"]
    new = parse_template(path)
    target = guidance[0]["targetBlockIds"][0]
    new["blocks"] = [block for block in new["blocks"] if block["id"] != target]
    result, _ = tw.align_upload(new, structure)
    instruction = next(block for block in result["blocks"] if block["id"] == guidance[0]["id"])
    assert instruction.get("unplaced") is True
    with pytest.raises(ValueError, match="已删除或未匹配"):
        validate_guidance_targets(result["blocks"])


def test_reupload_does_not_borrow_another_identical_paragraph(source):
    path, structure, guidance, bodies, _ = source
    # Deleting the first identical paragraph used to reassign the second one's position.
    guidance[0]["targetBlockIds"] = [bodies[0]["id"]]
    new = parse_template(path)
    new["blocks"] = [block for block in new["blocks"] if block["id"] != bodies[0]["id"]]
    result, _ = tw.align_upload(new, structure)
    instruction = next(block for block in result["blocks"] if block["id"] == guidance[0]["id"])
    assert instruction.get("unplaced") is True


def test_reupload_does_not_fuzzy_match_changed_target(source):
    path, structure, guidance, bodies, _ = source
    guidance[0]["targetBlockIds"] = [bodies[2]["id"]]
    new = parse_template(path)
    target = next(block for block in new["blocks"] if block["id"] == bodies[2]["id"])
    target["text"] = "条目正文变动。"
    result, _ = tw.align_upload(new, structure)
    instruction = next(block for block in result["blocks"] if block["id"] == guidance[0]["id"])
    assert instruction.get("unplaced") is True


def test_cell_guidance_binding_conflict_is_not_silently_overwritten(source):
    _, structure, _, bodies, table = source
    base = deepcopy(structure)
    cell = next(block for block in base["blocks"] if block.get("tableId") == table["id"])
    cell["kind"] = "guidance"
    cell["targetBlockIds"] = []
    theirs, mine = deepcopy(base), deepcopy(base)
    next(b for b in theirs["blocks"] if b["id"] == cell["id"])["targetBlockIds"] = [bodies[0]["id"]]
    next(b for b in mine["blocks"] if b["id"] == cell["id"])["targetBlockIds"] = [bodies[1]["id"]]
    merged = tm.merge3(base, theirs, mine, styles=({}, {}, {}))
    assert any(item["blockId"] == cell["id"] and item["field"] == "targetBlockIds" for item in merged["conflicts"])
    conflict = next(item for item in merged["conflicts"] if item["field"] == "targetBlockIds")
    assert conflict["theirsTargetTexts"] != conflict["mineTargetTexts"]
    assert next(b for b in merged["blocks"] if b["id"] == cell["id"])["targetBlockIds"] == [bodies[0]["id"]]
