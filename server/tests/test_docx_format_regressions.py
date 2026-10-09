"""Neutral DOCX regressions for run styles, numbering and nested list exports."""

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
import pytest

from app.docx_export import edit_original_docx
from app.docx_parser import parse_docx, _list_info


def test_added_emphasis_preserves_unmodified_run_properties_and_plain_body(tmp_path):
    source, output = tmp_path / "source.docx", tmp_path / "edited.docx"
    doc = Document()
    doc.add_heading("格式检查", 1)
    paragraph = doc.add_paragraph()
    first = paragraph.add_run("原有重点")
    first.bold, first.font.size, first.font.color.rgb = True, Pt(14), RGBColor(160, 0, 0)
    last = paragraph.add_run("原有普通")
    last.bold, last.font.size, last.font.color.rgb = False, Pt(12), RGBColor(0, 0, 0)
    last_xml = last._r.xml
    doc.sections[0].header.paragraphs[0].text = "保留页眉"
    doc.sections[0].footer.paragraphs[0].text = "保留页脚"
    doc.save(source)
    target = parse_docx(str(source)).markdown.replace("原有重点原有普通", "**更新重点**原有普通")
    edit_original_docx(target + "\n\n普通新增段落", source, output)
    result = Document(output)
    ordinary = next(run for run in result.paragraphs[1].runs if run.text == "原有普通")
    assert ordinary._r.xml == last_xml
    assert (ordinary.bold, ordinary.font.size, ordinary.font.color.rgb) == (False, Pt(12), RGBColor(0, 0, 0))
    added = result.paragraphs[-1]
    assert added.text == "普通新增段落" and all(run.bold is not True for run in added.runs)
    assert all(run.font.size == Pt(12) and run.font.color.rgb == RGBColor(0, 0, 0) for run in added.runs)
    assert result.sections[0].header.paragraphs[0].text == "保留页眉"
    assert result.sections[0].footer.paragraphs[0].text == "保留页脚"


def test_plain_paragraph_does_not_inherit_direct_numbering(tmp_path):
    source, output = tmp_path / "source.docx", tmp_path / "edited.docx"
    doc = Document()
    doc.add_heading("流程检查", 1)
    paragraph = doc.add_paragraph("步骤正文", style="Normal")
    properties = OxmlElement("w:numPr")
    for tag, value in (("ilvl", "0"), ("numId", "5")):
        element = OxmlElement("w:" + tag)
        element.set(qn("w:val"), value)
        properties.append(element)
    paragraph._p.get_or_add_pPr().append(properties)
    original_xml = paragraph._p.xml
    doc.save(source)
    edit_original_docx(parse_docx(str(source)).markdown + "\n\n新增说明段落", source, output)
    result = Document(output)
    assert result.paragraphs[1]._p.xml == original_xml
    assert result.paragraphs[1]._p.xpath("./w:pPr/w:numPr")
    assert result.paragraphs[2].text == "新增说明段落"
    assert not result.paragraphs[2]._p.xpath("./w:pPr/w:numPr")


def test_legacy_plain_markdown_keeps_normal_direct_numbering_during_text_edits(tmp_path):
    source, output = tmp_path / "source.docx", tmp_path / "edited.docx"
    doc = Document()
    doc.add_heading("旧稿兼容", 1)
    paragraph = doc.add_paragraph("原步骤文字", style="Normal")
    numbering = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    numbering.get_or_add_numId().val = 5
    numbering.get_or_add_ilvl().val = 0
    original_properties = paragraph._p.pPr.xml
    ordinary = doc.add_paragraph("保留普通正文")
    ordinary_xml = ordinary._p.xml
    doc.save(source)
    legacy = "# 旧稿兼容\n\n原步骤文字\n\n保留普通正文"
    assert _list_info(paragraph) is None
    edit_original_docx(legacy, source, output)
    assert output.read_bytes() == source.read_bytes()
    edit_original_docx(legacy.replace("原步骤文字", "修改后的步骤文字"), source, output)
    result = Document(output)
    assert result.paragraphs[1].text == "修改后的步骤文字"
    assert result.paragraphs[1]._p.pPr.xml == original_properties
    assert result.paragraphs[2]._p.xml == ordinary_xml


@pytest.mark.parametrize("styles", [
    ["List Bullet", "List Bullet 2", "List Bullet 3"],
    ["List Number", "List Bullet 2", "List Number 3"],
])
def test_legacy_flat_list_markdown_does_not_remove_original_nesting(tmp_path, styles):
    source, output = tmp_path / "source.docx", tmp_path / "edited.docx"
    doc = Document()
    doc.add_heading("旧稿层级", 1)
    for text, style in zip(("主项", "子项", "孙项"), styles):
        doc.add_paragraph(text, style=style)
    doc.add_paragraph("保留普通正文")
    original_properties = [paragraph._p.pPr.xml for paragraph in doc.paragraphs[1:4]]
    original_parents = [doc.paragraphs[index]._p.xml for index in (1, 3)]
    doc.save(source)
    # This is the exact old parser representation, not a database migration.
    legacy = parse_docx(str(source), flat_lists=True).markdown
    assert legacy != parse_docx(str(source)).markdown and "\n    " not in legacy
    edit_original_docx(legacy, source, output)
    assert output.read_bytes() == source.read_bytes()
    edit_original_docx(legacy.replace("子项", "子项更新"), source, output)
    result = Document(output)
    assert [paragraph.text for paragraph in result.paragraphs[1:4]] == ["主项", "子项更新", "孙项"]
    assert [_list_info(paragraph)[1] for paragraph in result.paragraphs[1:4]] == [0, 1, 2]
    assert [paragraph._p.pPr.xml for paragraph in result.paragraphs[1:4]] == original_properties
    assert [result.paragraphs[index]._p.xml for index in (1, 3)] == original_parents


@pytest.mark.parametrize("listing,kinds", [
    ("- 主项\n    - 子项\n        - **孙项**", ["ul", "ul", "ul"]),
    ("1. 主项\n    - 子项\n        1. **孙项**", ["ol", "ul", "ol"]),
])
def test_nested_lists_export_as_independent_native_levels_and_can_be_edited_again(tmp_path, listing, kinds):
    source, output, again = tmp_path / "source.docx", tmp_path / "edited.docx", tmp_path / "again.docx"
    doc = Document()
    doc.add_heading("层级检查", 1)
    doc.add_paragraph("原有正文")
    doc.save(source)
    target = "# 层级检查\n\n" + listing + "\n\n独立段落"
    edit_original_docx(target, source, output)
    result = Document(output)
    items = result.paragraphs[1:4]
    assert [paragraph.text for paragraph in items] == ["主项", "子项", "孙项"]
    assert [_list_info(paragraph) for paragraph in items] == list(zip(kinds, [0, 1, 2]))
    assert next(run for run in items[-1].runs if run.text == "孙项").bold is True
    numbers = result.part.numbering_part.element
    for paragraph in items:
        number_id = paragraph._p.xpath("./w:pPr/w:numPr/w:numId/@w:val")[0]
        abstract_id = numbers.xpath(f"./w:num[@w:numId='{number_id}']/w:abstractNumId/@w:val")[0]
        assert len(numbers.xpath(f"./w:abstractNum[@w:abstractNumId='{abstract_id}']/w:lvl")) == 9
    assert not result.paragraphs[-1]._p.xpath("./w:pPr/w:numPr")
    roundtrip = parse_docx(str(output)).markdown
    assert "\n    - 子项\n        " in roundtrip
    unchanged = tmp_path / "unchanged.docx"
    edit_original_docx(roundtrip, output, unchanged)
    assert unchanged.read_bytes() == output.read_bytes()
    parent_xml, grandchild_xml = items[0]._p.xml, items[-1]._p.xml
    edit_original_docx(roundtrip.replace("子项", "子项更新"), output, again)
    edited = Document(again)
    assert [paragraph.text for paragraph in edited.paragraphs[1:4]] == ["主项", "子项更新", "孙项"]
    assert edited.paragraphs[1]._p.xml == parent_xml and edited.paragraphs[3]._p.xml == grandchild_xml
    assert [_list_info(paragraph)[1] for paragraph in edited.paragraphs[1:4]] == [0, 1, 2]


def test_adding_child_to_existing_list_preserves_parent_xml(tmp_path):
    source, output = tmp_path / "source.docx", tmp_path / "edited.docx"
    doc = Document()
    doc.add_heading("层级检查", 1)
    parent = doc.add_paragraph("主项", style="List Bullet")
    parent_xml = parent._p.xml
    doc.save(source)
    edit_original_docx(parse_docx(str(source)).markdown + "\n    - 新子项\n", source, output)
    result = Document(output)
    assert result.paragraphs[1]._p.xml == parent_xml
    assert [paragraph.text for paragraph in result.paragraphs[1:]] == ["主项", "新子项"]
    assert [_list_info(paragraph) for paragraph in result.paragraphs[1:]] == [("ul", 0), ("ul", 1)]


def test_list_nesting_beyond_word_limit_is_rejected_before_writing(tmp_path):
    source, output = tmp_path / "source.docx", tmp_path / "edited.docx"
    doc = Document()
    doc.add_heading("层级检查", 1)
    doc.save(source)
    target = "# 层级检查\n\n" + "\n".join("    " * depth + "- 项目" for depth in range(10))
    with pytest.raises(ValueError, match="九级"):
        edit_original_docx(target, source, output)
    assert not output.exists()
