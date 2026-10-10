"""S5：整篇入口 allowMerge 的多人并发（01 §8.2 用例 12–16、19、21、22）。

A=editor、B=editor2 都是非管理员项目编写人员；C=viewer；D=outsider。
"""

import threading
from io import BytesIO
from uuid import UUID
from zipfile import ZipFile

import pytest
from docx import Document

from app.db import connect
from app.document_objects import check_invariants
from app.identity import create_user
from tests.test_document_access import PASSWORD, _as, members  # noqa: F401
from tests.test_documents import document_context  # noqa: F401

TEXT = "# 1 范围\n\n第一段。\n\n第二段。\n\n第三段。\n\n# 2 引用\n\n引用段。\n"


@pytest.fixture
def pair(members):  # noqa: F811
    client, project, document, users = members
    with connect() as connection:
        users["editor2"] = create_user(account="m-editor2", name="成员editor2", password=PASSWORD,
                                       roles=["template_manager"], connection=connection)
        connection.execute("INSERT INTO project_members(project_id,user_id,role) VALUES(%s,%s,'editor')", (project, users["editor2"]))
        connection.commit()
    return client, document, users


@pytest.fixture
def team(pair, monkeypatch):
    monkeypatch.setattr("app.documents._check_word", lambda *args, **kwargs: None)
    client, document, users = pair
    headers = _as(client, "editor")
    assert save(client, document["id"], 0, TEXT, headers).status_code == 200
    return client, document["id"], users


def save(client, did, base, text, headers, key=None, merge=True):
    body = {"expectedRevision": base, "markdownContent": text}
    if merge:
        body["allowMerge"] = True
    if key:
        body["clientSaveKey"] = key
    return client.patch(f"/api/v1/documents/{did}/content", json=body, headers=headers)


def current(client, did):
    return client.get(f"/api/v1/documents/{did}").json()["data"]


@pytest.mark.parametrize("order", ["AB", "BA"])
def test_different_paragraphs_merge_in_either_order(team, order):
    # 用例 12（2a）
    client, did, users = team
    edits = {"editor": TEXT.replace("第一段。", "A 改第一段。"), "editor2": TEXT.replace("第三段。", "B 改第三段。")}
    sequence = ["editor", "editor2"] if order == "AB" else ["editor2", "editor"]
    first = save(client, did, 1, edits[sequence[0]], _as(client, sequence[0]))
    assert first.status_code == 200 and first.json()["data"]["merged"] is False
    second = save(client, did, 1, edits[sequence[1]], _as(client, sequence[1]))
    assert second.status_code == 200, second.text
    data = second.json()["data"]
    assert data["merged"] is True and data["revision"] == 3
    assert "A 改第一段。" in data["markdownContent"] and "B 改第三段。" in data["markdownContent"]
    assert data["changedBy"] == [{"id": str(users[sequence[0]]), "name": f"成员{sequence[0]}", "objects": 1}]
    with connect() as connection:
        actors = [str(row["actor_id"]) for row in connection.execute(
            "SELECT actor_id FROM document_changes WHERE document_id=%s AND doc_revision>1 ORDER BY doc_revision", (did,)).fetchall()]
        assert actors == [str(users[sequence[0]]), str(users[sequence[1]])]
        assert check_invariants(connection, UUID(did))


def test_same_paragraph_conflict_keeps_copy_and_visibility(team):
    # 用例 13
    client, did, users = team
    assert save(client, did, 1, TEXT.replace("第二段。", "A 的第二段。"), _as(client, "editor")).status_code == 200
    b_headers = _as(client, "editor2")
    mine = TEXT.replace("第二段。", "B 的第二段。").replace("引用段。", "B 改引用段。")
    rejected = save(client, did, 1, mine, b_headers)
    assert rejected.status_code == 409, rejected.text
    error = rejected.json()["error"]
    assert error["code"] == "OBJECT_CONFLICT" and error["currentRevision"] == 2
    details = error["details"]
    [item] = details["conflicts"]
    assert item["reason"] == "MODIFIED" and item["base"]["markdown"] == "第二段。"
    assert item["theirs"]["markdown"] == "A 的第二段。" and item["theirs"]["updatedBy"]["name"] == "成员editor"
    assert item["mine"]["markdown"] == "B 的第二段。"
    assert item["location"]["headingPath"] == ["1 范围"] and item["location"]["ordinal"] == 2
    # 可部分保存：无冲突的修改在 appliableMarkdown 里
    assert "B 改引用段。" in details["appliableMarkdown"] and "A 的第二段。" in details["appliableMarkdown"]
    assert current(client, did)["markdownContent"] == TEXT.replace("第二段。", "A 的第二段。")
    listed = client.get(f"/api/v1/documents/{did}/conflicts").json()["data"]["items"]
    assert len(listed) == 1 and listed[0]["mine"]["markdown"] == "B 的第二段。"
    assert client.get(f"/api/v1/documents/{did}/conflicts/summary").json()["data"] == {"mineOpen": 1, "mineHandled": 0, "allOpen": None}
    with connect() as connection:
        row = connection.execute("SELECT payload_json FROM document_conflicts WHERE document_id=%s", (did,)).fetchone()
        assert mine in row["payload_json"].replace("\\n", "\n")
    # A（editor）看所有人的冲突 → 403；负责人、管理员能看到但不能处理（决策3）
    _as(client, "editor")
    assert client.get(f"/api/v1/documents/{did}/conflicts", params={"actor": "all"}).status_code == 403
    assert client.get(f"/api/v1/documents/{did}/conflicts/{details['conflictId']}").status_code == 403
    owner = _as(client, "owner")
    everyone = client.get(f"/api/v1/documents/{did}/conflicts", params={"actor": "all"}).json()["data"]["items"]
    assert [entry["actor"]["name"] for entry in everyone] == ["成员editor2"]
    assert client.get(f"/api/v1/documents/{did}/conflicts/summary").json()["data"] == {"mineOpen": 0, "mineHandled": 0, "allOpen": 1}
    assert client.get(f"/api/v1/documents/{did}/conflicts/{details['conflictId']}").status_code == 200
    denied = client.post(f"/api/v1/documents/{did}/conflicts/{details['conflictId']}/resolve", headers=owner,
                         json={"clientSaveKey": "owner-key-1", "baseDocumentRevision": 2,
                               "items": [{"itemIndex": 0, "action": "use_theirs"}]})
    assert denied.status_code == 403
    # 部分保存：B 自动把 appliableMarkdown 用新键再提交
    b_headers = _as(client, "editor2")
    partial = save(client, did, details["appliableBaseRevision"], details["appliableMarkdown"], b_headers, key="partial-key-1")
    assert partial.status_code == 200 and "B 改引用段。" in partial.json()["data"]["markdownContent"]


def test_list_counts_open_conflicts(team):
    # 用例 45
    client, did, _ = team
    assert save(client, did, 1, TEXT.replace("第二段。", "A。"), _as(client, "editor")).status_code == 200
    assert save(client, did, 1, TEXT.replace("第二段。", "B。"), _as(client, "editor2")).status_code == 409
    project = current(client, did)["projectId"] if "projectId" in current(client, did) else None
    with connect() as connection:
        project = project or str(connection.execute("SELECT project_id FROM documents WHERE id=%s", (did,)).fetchone()["project_id"])
    item = client.get(f"/api/v1/projects/{project}/documents").json()["data"]["items"][0]
    assert item["mineOpenConflicts"] == 1 and item["memberOpenConflicts"] is None
    _as(client, "owner")
    item = client.get(f"/api/v1/projects/{project}/documents").json()["data"]["items"][0]
    assert item["mineOpenConflicts"] == 0 and item["memberOpenConflicts"] == 1
    _as(client, "editor")
    assert client.get(f"/api/v1/projects/{project}/documents").json()["data"]["items"][0]["mineOpenConflicts"] == 0


def test_concurrent_requests_never_lose_updates(team, monkeypatch):
    # 用例 14：两请求读取同一快照后再提交（Word 预检处汇合）
    client, did, users = team
    barrier = threading.Barrier(2, timeout=10)
    calls = {"n": 0}

    def slow_check(*args, **kwargs):
        calls["n"] += 1
        if calls["n"] <= 2:
            barrier.wait()

    monkeypatch.setattr("app.documents._check_word", slow_check)
    from fastapi.testclient import TestClient
    from app.main import create_app

    results: dict[str, object] = {}

    def run(role, text, base=1):
        with TestClient(create_app()) as own:
            headers = _as(own, role)
            results[role] = save(own, did, base, text, headers)

    threads = [threading.Thread(target=run, args=("editor", TEXT.replace("第一段。", "A1。"))),
               threading.Thread(target=run, args=("editor2", TEXT.replace("第三段。", "B3。")))]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join()
    assert {response.status_code for response in results.values()} == {200}
    final = current(client, did)
    assert "A1。" in final["markdownContent"] and "B3。" in final["markdownContent"] and final["revision"] == 3
    # 同段：一个 200 一个 409
    calls["n"] = 0
    barrier.reset()
    results.clear()
    threads = [threading.Thread(target=run, args=("editor", final["markdownContent"].replace("第二段。", "A2。"), 3)),
               threading.Thread(target=run, args=("editor2", final["markdownContent"].replace("第二段。", "B2。"), 3))]
    monkeypatch.setattr("app.documents._check_word", lambda *args, **kwargs: None)
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join()
    # 两个线程基线都是 3：先提交的 200，后提交的看到同段被改 → 409
    statuses = sorted(response.status_code for response in results.values())
    assert statuses == [200, 409]
    text = current(client, did)["markdownContent"]
    assert ("A2。" in text) != ("B2。" in text)
    with connect() as connection:
        assert check_invariants(connection, UUID(did))


def test_same_key_replays_and_detects_reuse(team):
    # 用例 15、16
    client, did, _ = team
    headers = _as(client, "editor")
    text = TEXT.replace("第三段。", "第三段。\n\n新增段落。")
    first = save(client, did, 1, text, headers, key="save-key-0001")
    again = save(client, did, 1, text, headers, key="save-key-0001")
    assert first.status_code == again.status_code == 200
    assert again.json()["data"]["revision"] == first.json()["data"]["revision"] == 2
    assert again.json()["data"]["markdownContent"].count("新增段落。") == 1
    with connect() as connection:
        assert connection.execute("SELECT COUNT(*) AS n FROM document_changes WHERE client_save_key='save-key-0001'").fetchone()["n"] == 1
    reused = save(client, did, 1, text + "\n另一份。\n", headers, key="save-key-0001")
    assert reused.status_code == 422 and reused.json()["error"]["code"] == "IDEMPOTENCY_KEY_REUSED"
    # 响应丢失后（期间他人又改了别处）同键重发：不重复插入
    assert save(client, did, 2, text.replace("引用段。", "他人改引用。"), _as(client, "editor2")).status_code == 200
    headers = _as(client, "editor")
    replay = save(client, did, 1, text, headers, key="save-key-0001")
    assert replay.status_code == 200 and current(client, did)["markdownContent"].count("新增段落。") == 1


def test_without_allow_merge_stays_strict(team):
    # 用例 19
    client, did, _ = team
    assert save(client, did, 1, TEXT.replace("第一段。", "A。"), _as(client, "editor")).status_code == 200
    stale = save(client, did, 1, TEXT.replace("第三段。", "B。"), _as(client, "editor2"), merge=False)
    assert stale.status_code == 409 and stale.json()["error"]["code"] == "REVISION_CONFLICT"


def test_merged_export_keeps_untouched_word_parts(pair, document_context):  # noqa: F811
    # 用例 21：真实 Word 预检；A 改正文、B 改表格单元格
    client, document, _ = pair
    did = document["id"]
    base = document["markdownContent"]
    assert "原始正文。" in base and "通过" in base
    assert save(client, did, 0, base.replace("原始正文。", "A 改正文。"), _as(client, "editor")).status_code == 200
    merged = save(client, did, 0, base.replace("| 固定测试 | 通过 |", "| 固定测试 | 未通过 |"), _as(client, "editor2"))
    assert merged.status_code == 200, merged.text
    assert merged.json()["data"]["merged"] is True
    exported = client.post(f"/api/v1/documents/{did}/export", json={"expectedRevision": 2}, headers=_as(client, "editor2"))
    assert exported.status_code == 200, exported.text
    content = client.get(f"/api/v1/files/{exported.json()['data']['fileId']}/content").content
    edited = Document(BytesIO(content))
    assert [p.text for p in edited.paragraphs] == ["范围", "A 改正文。"]
    assert [[cell.text for cell in row.cells] for row in edited.tables[0].rows] == [["项目", "结论"], ["固定测试", "未通过"]]
    source = document_context[5]
    with ZipFile(BytesIO(source)) as original, ZipFile(BytesIO(content)) as result:
        for name in original.namelist():
            if name != "word/document.xml":
                assert result.read(name) == original.read(name), name


def test_merge_that_breaks_word_becomes_format_conflict(team, monkeypatch):
    # 用例 22
    client, did, _ = team
    assert save(client, did, 1, TEXT.replace("第一段。", "A。"), _as(client, "editor")).status_code == 200
    from app.errors import AppError

    def failing(*args, **kwargs):
        raise AppError(422, "AI_FORMAT_INVALID", "合并后无法写回。")

    monkeypatch.setattr("app.documents._check_word", failing)
    rejected = save(client, did, 1, TEXT.replace("第三段。", "B。"), _as(client, "editor2"))
    assert rejected.status_code == 409
    [item] = rejected.json()["error"]["details"]["conflicts"]
    assert item["reason"] == "FORMAT" and "B。" in item["mine"]["markdown"]
    with connect() as connection:
        assert connection.execute("SELECT COUNT(*) AS n FROM document_conflicts WHERE document_id=%s", (did,)).fetchone()["n"] == 1
    # 没有他人修改时（不合并）仍是 422
    own = save(client, did, 2, TEXT.replace("第三段。", "C。"), _as(client, "editor"))
    assert own.status_code == 422
