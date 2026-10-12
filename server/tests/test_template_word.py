"""A06 模板 Word 引擎：20 份真实 438C 模板逐份“改字/改级别/新增/删除/单元格/移章/改样式”→ 生成 Word → 往返核对。"""

from __future__ import annotations

import copy
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

from app import template_word as tw
from app.template_import import parse_template

TEMPLATES = sorted((Path(__file__).resolve().parents[2] / "438C-2021全套模板").glob("*.docx"))


def _edit(content: bytes) -> tuple[dict, dict, list[str], dict]:
    structure = parse_template(content)
    structure["reviewed"] = True
    analysis = tw._analysis(content)
    draft = copy.deepcopy(structure)
    blocks = draft["blocks"]
    top = [block for block in blocks if not block["tableId"]]
    heads = [block for block in top if block["kind"] == "heading"]
    expect: dict = {}
    actions = []
    body = next((b for b in top if b["kind"] == "body" and not analysis["blocks"][b["path"]]["textLock"]), None)
    if body:
        body["text"] += "（已修改）"
        expect["edit"] = (body["id"], body["text"])
        actions.append("edit")
    heading = next((b for b in heads[1:] if not analysis["blocks"][b["path"]]["textLock"] and b["headingLevel"] < 6), None)
    if heading:
        heading["text"] += "甲"
        heading["headingLevel"] += 1
        expect["hlevel"] = (heading["id"], heading["text"], heading["headingLevel"])
        actions.append("hlevel")
    new = {"id": tw.new_block_id(), "path": None, "kind": "body", "text": "新增正文段落。", "sourceText": "", "headingLevel": None,
           "headingId": heads[0]["id"], "tableId": None, "protected": False, "isTable": False}
    blocks.insert(blocks.index(heads[0]) + 1, new)
    expect["insert"] = new["id"]
    actions.append("insert")
    removed = next((b for b in top if b["kind"] == "body" and b is not body and not analysis["blocks"][b["path"]]["deleteLock"]), None)
    if removed:
        blocks.remove(removed)
        expect["delete"] = removed["id"]
        actions.append("delete")
    for table in (b for b in top if b["isTable"] and analysis["tables"][b["path"]]["tableEditable"] in ("cells", "partial")):
        info = analysis["tables"][table["path"]]
        locked = {tuple(cell) for cell in info["lockedCells"]}
        free = [(r, c) for r, row in enumerate(info["cells"]) for c in range(len(row)) if (r, c) not in locked]
        if free:
            cells = copy.deepcopy(info["cells"])
            r, c = free[-1]
            cells[r][c] = (cells[r][c] + "改").strip()
            table["cells"] = cells
            expect["cell"] = (table["id"], cells)
            actions.append("cell")
            break
    level = heads[0]["headingLevel"]
    firsts = [b for b in blocks if not b["tableId"] and b["kind"] == "heading" and b["headingLevel"] == level]
    if len(firsts) >= 3:
        topb = [b for b in blocks if not b["tableId"]]
        a, b_ = topb.index(firsts[-2]), topb.index(firsts[-1])
        order = topb[:a] + topb[b_:] + topb[a:b_]
        children: dict[str, list] = {}
        for block in blocks:
            if block["tableId"]:
                children.setdefault(block["tableId"], []).append(block)
        draft["blocks"] = [x for block in order for x in [block] + children.get(block["id"], [])]
        expect["move"] = [b["id"] for b in order if not tw._is_virtual(b)]
        actions.append("move")
    styles = copy.deepcopy(analysis["styles"])
    styles["正文"]["size"] = 12
    return draft, styles, actions, expect


@pytest.mark.parametrize("source", TEMPLATES, ids=lambda p: p.name[:5])
def test_real_template_round_trip(source: Path, tmp_path: Path) -> None:
    content = source.read_bytes()
    analysis = tw._analysis(content)
    draft, styles, actions, expect = _edit(content)
    out = tmp_path / "v2.docx"
    ops = tw.apply_draft(source, draft, analysis["styles"], styles, out)
    assert ops > 0
    version = tw.build_version(out, draft, 2)
    by_id = {block["id"]: block for block in version["blocks"]}
    if "edit" in expect:
        assert by_id[expect["edit"][0]]["text"] == expect["edit"][1]
    if "hlevel" in expect:
        block = by_id[expect["hlevel"][0]]
        assert (block["text"], block["headingLevel"]) == expect["hlevel"][1:]
    assert by_id[expect["insert"]]["text"] == "新增正文段落。" and by_id[expect["insert"]]["path"]
    if "delete" in expect:
        assert expect["delete"] not in by_id
    if "cell" in expect:
        info = tw._analysis(out.read_bytes())["tables"][by_id[expect["cell"][0]]["path"]]
        assert info["cells"] == expect["cell"][1]
    if "move" in expect:
        assert [b["id"] for b in version["blocks"] if not b["tableId"] and not b.get("virtual")] == expect["move"]
    assert tw.read_styles(Document(out))["正文"]["size"] == 12
    # 版本 Word 可再次作为基底：未改动时零操作
    again = tmp_path / "v3.docx"
    assert tw.apply_draft(out, version, tw.read_styles(Document(out)), tw.read_styles(Document(out)), again) == 0
    # 新建文档：清掉指导后仍是合法 Word
    Document(BytesIO(tw.clean_version(out, version)))



# ---------------------------------------------------------------- 23 M-3 / M-4、25 M-1

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _order_violations(xml: bytes) -> int:
    from lxml import etree
    order = {_W + name: tw._ORDER[name] for name in ("style", "pPr", "rPr")}
    count = 0
    for element in etree.fromstring(xml).iter(*order):
        sequence = order[element.tag]
        ranks = [sequence.index(child.tag[len(_W):]) for child in element
                 if isinstance(child.tag, str) and child.tag.startswith(_W) and child.tag[len(_W):] in sequence]
        count += any(a > b for a, b in zip(ranks, ranks[1:]))
    return count


_PROP_VALUES = {"fontEastAsia": "黑体", "fontAscii": "Arial", "size": 15, "firstLineIndent": {"unit": "chars", "value": 2},
                "leftIndent": {"unit": "chars", "value": 2}, "rightIndent": {"unit": "pt", "value": 6}, "lineSpacing": "1.5",
                "spaceBefore": {"unit": "pt", "value": 6}, "spaceAfter": {"unit": "lines", "value": 0.5}}


@pytest.mark.parametrize("source", TEMPLATES, ids=lambda p: p.name[:5])
def test_style_edits_keep_schema_order_and_untouched_parts(source: Path, tmp_path: Path) -> None:
    """每种样式属性（每个白名单样式）写出后 style/pPr/rPr 子元素仍按 schema 顺序；未改部件逐字节相同、目录条目不丢。"""
    from zipfile import ZipFile
    content = source.read_bytes()
    structure = parse_template(content)
    analysis = tw._analysis(content)
    styles = copy.deepcopy(analysis["styles"])
    for name in styles:
        styles[name].update(_PROP_VALUES)
    out = tmp_path / "out.docx"
    assert tw.apply_draft(source, structure, analysis["styles"], styles, out) > 0
    with ZipFile(source) as before, ZipFile(out) as after:
        for part in ("word/styles.xml", "word/document.xml"):
            assert _order_violations(after.read(part)) <= _order_violations(before.read(part)), part
        assert [info.filename for info in after.infolist()] == [info.filename for info in before.infolist()]
        changed = [name for name in before.namelist() if before.read(name) != after.read(name)]
        assert set(changed) <= {"word/styles.xml", "word/document.xml"}, changed
    written = tw.read_format(Document(out))["styles"]
    for name in styles:
        assert written[name]["size"] == 15 and written[name]["fontEastAsia"] == "黑体", name


@pytest.mark.parametrize("source", TEMPLATES, ids=lambda p: p.name[:5])
def test_text_edit_keeps_untouched_parts_byte_identical(source: Path, tmp_path: Path) -> None:
    from zipfile import ZipFile
    content = source.read_bytes()
    structure = parse_template(content)
    analysis = tw._analysis(content)
    draft = copy.deepcopy(structure)
    block = next((b for b in draft["blocks"] if not b["tableId"] and b["kind"] == "body"
                  and not analysis["blocks"][b["path"]]["textLock"]), None)
    if block is None:
        pytest.skip("没有可改正文")
    block["text"] += "（改）"
    out = tmp_path / "out.docx"
    assert tw.apply_draft(source, draft, analysis["styles"], analysis["styles"], out) > 0
    with ZipFile(source) as before, ZipFile(out) as after:
        assert [info.filename for info in after.infolist()] == [info.filename for info in before.infolist()]
        changed = [name for name in before.namelist() if before.read(name) != after.read(name)]
        assert changed == ["word/document.xml"]
        assert _order_violations(after.read("word/document.xml")) <= _order_violations(before.read("word/document.xml"))


@pytest.mark.parametrize("source", TEMPLATES, ids=lambda p: p.name[:5])
def test_style_key_matches_paragraph_style(source: Path) -> None:
    """25 M-1：Normal 段落的 styleKey 是“正文”，Heading n 是“标题 n”；非白名单样式为 None。多跑两次结果稳定。"""
    content = source.read_bytes()
    first = tw._analysis(content)["blocks"]
    tw._cached_analysis.cache_clear() if hasattr(tw._cached_analysis, "cache_clear") else None
    second = tw._analysis(content)["blocks"]
    checked = 0
    for path, info in first.items():
        if "styleName" not in info:
            continue
        assert second[path]["styleKey"] == info["styleKey"]
        name = info["styleName"]
        if name == "Normal":
            assert info["styleKey"] == "正文", (path, name)
            checked += 1
        elif name.startswith("Heading ") and name[8:].isdigit() and int(name[8:]) <= 6:
            assert info["styleKey"] == f"标题 {name[8:]}", (path, name)
            checked += 1
        elif name in {"Normal Indent", "toc 1", "toc 2", "toc 3", "header", "footer"}:
            assert info["styleKey"] is None, (path, name)
    assert checked > 0
