"""A07 编写分工（硬限制）浏览器验收（16 §A07、17 §4.5，1920×1080 与 1366×768）。

环境变量：E2E_API、E2E_WEB、DATABASE_URL、FILE_DATA_DIR 必须设置（脚本直接在被测库里建账号）；
E2E_VIEWPORT（如 1366x768，默认 1920x1080）、E2E_SHOTS、E2E_OUT。admin/admin123 登录、不改密。
覆盖：G1 负责人在工作台“编写分工”面板按章节分配、模式标签；G2 编写人员只能改分给自己的章节（其余只读，点了有提示）；
G3 保存过程中被改分工 → 409 SECTION_NOT_ASSIGNED 提示条、不自动重试、可下载；G4 接口：越权 409、缺 blockIds 422、负责人不受限；
G5 编写人员新建的文档整篇归创建人；G6 编写人员不能恢复版本；G7 成员表“编写分工”列、点数字到概述按人视图；
G8 S-5 降为只读前确认（取消不改、确认后分工清除）；G9 概述“编写分工”按文档/按人。
修复第 1 轮（27/28/29）：G1 抽屉 400 宽、焦点移入、Esc 关闭还焦点、不遮正文、筛选、多人分配；G1b 模式切换确认；
G2 只读样式/gutter/我的部分目录/首次提示、只读块输入进暂存并放到我的一处；G3 被拒横幅与面板（请联系、定位、下载、只保存我负责的部分）；
G7/G9 概述链接带筛选；G10 3000 块文档编写人员保存不错位；G11 没有分工的编写人员只读，分到后提示。
"""

from __future__ import annotations

import os
import re
import sys
import tempfile
import uuid
from io import BytesIO
from pathlib import Path

import httpx
from playwright.sync_api import Page, expect, sync_playwright

ROOT = Path(__file__).resolve().parents[2]
API = os.environ.get("E2E_API", "")
WEB = os.environ.get("E2E_WEB", "")
OUT = Path(os.environ.get("E2E_OUT", str(Path(tempfile.gettempdir()) / "e2e_a07")))
SHOTS = Path(os.environ.get("E2E_SHOTS", str(OUT / "shots")))
PASSWORD = "E2e-a07-pass-1"
ADMIN, ADMIN_PASSWORD = "admin", "admin123"
RUN = uuid.uuid4().hex[:6]
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
results: list[tuple[str, bool, str]] = []
open_pages: list[Page] = []


def step(name: str):
    def wrap(fn):
        def run(*args, **kwargs):
            if os.environ.get("E2E_ONLY") and os.environ["E2E_ONLY"] not in name:
                print(f"SKIP {name}", flush=True)
                return
            try:
                fn(*args, **kwargs)
                results.append((name, True, ""))
                print(f"PASS {name}", flush=True)
            except Exception as error:  # noqa: BLE001 - 汇总后统一失败
                results.append((name, False, f"{type(error).__name__}: {error}"))
                print(f"FAIL {name}: {error}", flush=True)
                OUT.mkdir(parents=True, exist_ok=True)
                for index, page in enumerate(open_pages):
                    if not page.is_closed():
                        page.screenshot(path=str(OUT / f"fail-{len(results)}-{index}.png"))
                if os.environ.get("E2E_STOP_ON_FAIL") == "1":
                    raise
        return run
    return wrap


class Api:
    def __init__(self, account: str, password: str = PASSWORD):
        self.client = httpx.Client(base_url=API, timeout=120)
        csrf = self.client.get("/auth/csrf").json()["data"]["csrfToken"]
        response = self.client.post("/auth/login", json={"account": account, "password": password}, headers={"X-CSRF-Token": csrf})
        response.raise_for_status()
        self.csrf = response.json()["data"]["csrfToken"]

    def raw(self, method: str, path: str, **kwargs) -> httpx.Response:
        headers = {"X-CSRF-Token": self.csrf, **kwargs.pop("headers", {})}
        return self.client.request(method, path, headers=headers, **kwargs)

    def call(self, method: str, path: str, **kwargs):
        response = self.raw(method, path, **kwargs)
        if response.status_code >= 400:
            raise AssertionError(f"{method} {path} → {response.status_code} {response.text[:300]}")
        return response.json()["data"] if response.content else None

    def members_revision(self, project: str) -> int:
        return self.call("GET", f"/projects/{project}/members")["membersRevision"]

    def set_role(self, project: str, user, role: str):
        """已是成员 → PUT 改角色；不是成员 → members:batch 加入（PUT 只改已有成员，A05 20 I-3）。"""
        snapshot = self.call("GET", f"/projects/{project}/members")
        if any(item["userId"] == str(user) for item in snapshot["items"]):
            return self.call("PUT", f"/projects/{project}/members/{user}", json={"role": role, "expectedMembersRevision": snapshot["membersRevision"]})
        return self.call("POST", f"/projects/{project}/members:batch",
                         json={"userIds": [str(user)], "role": role, "expectedMembersRevision": snapshot["membersRevision"]})

    def project(self, name: str, owner) -> dict:
        return self.call("POST", "/projects", json={"name": name, "ownerId": str(owner), "status": "active"},
                         headers={"Idempotency-Key": f"a07-{uuid.uuid4().hex}"})





VIEW = os.environ.get("E2E_VIEWPORT", "1920x1080")
WIDTH, HEIGHT = (int(value) for value in VIEW.split("x"))
TAG = "" if VIEW == "1920x1080" else f"-{VIEW}"


def make_word(title_paragraphs: list[tuple[str, list[str]]]) -> bytes:
    from docx import Document as WordDocument
    word = WordDocument()
    for title, paragraphs in title_paragraphs:
        word.add_heading(title, 1)
        for text in paragraphs:
            word.add_paragraph(text)
    buffer = BytesIO()
    word.save(buffer)
    return buffer.getvalue()


def upload_doc(api: "Api", project_id: str, name: str, content: bytes) -> dict:
    upload = api.client.post("/files", files={"file": (f"{name}.docx", content, DOCX)},
                             data={"purpose": "attachment", "projectId": project_id},
                             headers={"X-CSRF-Token": api.csrf, "Idempotency-Key": f"a07-upload-{uuid.uuid4().hex}"})
    assert upload.status_code == 201, upload.text
    return api.call("POST", f"/projects/{project_id}/documents",
                    json={"name": name, "sourceType": "upload", "wordFileId": upload.json()["data"]["id"]})


CHAPTERS = [("概述", ["概述第一段：系统的建设背景。", "概述第二段：系统的建设目标。"]),
            ("功能需求", ["功能需求第一段：用户管理。", "功能需求第二段：权限管理。"]),
            ("附录", ["附录第一段：术语表。"])]


def seed() -> dict:
    sys.path.insert(0, str(ROOT / "server"))
    from app.db import connect
    from app.identity import create_user
    users = {}
    with connect() as connection:
        row = connection.execute("select id from users where account_key = %s", (ADMIN,)).fetchone()
        users["admin"] = row[0] if row else create_user(account=ADMIN, name="系统管理员", password=ADMIN_PASSWORD, roles=["system_admin"], connection=connection)
        for key, name in (("lisi", "李四"), ("wangce", "王测"), ("liuyang", "刘洋")):
            users[key] = create_user(account=f"{key}{RUN}", name=name, password=PASSWORD, roles=[], department="软件一部",
                                     allow_no_roles=True, connection=connection)
        connection.commit()
    admin = Api(ADMIN, ADMIN_PASSWORD)
    project = admin.project(f"分工验收项目{RUN}", users["admin"])
    for key in ("lisi", "wangce", "liuyang"):
        admin.set_role(project["id"], users[key], "editor")
    document = upload_doc(admin, project["id"], "软件需求规格说明", make_word(CHAPTERS))
    return {"users": users, "project": project, "document": document, "admin": admin}


def login(page: Page, account: str, password: str = PASSWORD):
    page.context.clear_cookies()
    page.goto(f"{WEB}/login")
    page.fill("#staff-id", account)
    page.fill("#password", password)
    page.locator("form button[type=submit]").click()
    page.wait_for_url(re.compile(r"/projects"), timeout=30000)


def shot(page: Page, name: str):
    SHOTS.mkdir(parents=True, exist_ok=True)
    page.screenshot(path=str(SHOTS / f"{name}{TAG}.png"))


def open_doc(page: Page, pid: str, did: str, query: str = ""):
    page.goto(f"{WEB}/projects/{pid}/documents/{did}{query}")
    expect(page.locator(".paper-content")).to_be_visible(timeout=60000)
    expect(page.get_by_test_id("save-status")).to_contain_text(re.compile("保存"), timeout=30000)


def para(page: Page, text: str):
    return page.locator(".paper-content p", has_text=text).first


def wait_saved(page: Page, timeout: int = 30000):
    expect(page.get_by_test_id("save-status")).to_contain_text(re.compile("刚刚保存|已保存"), timeout=timeout)


def pick(panel, row_text: str, user_ids: list):
    """多人分配：行内［分配］→ 勾选 → 确定（空列表 = 取消分配）。"""
    row = panel.get_by_test_id("assign-heading").filter(has_text=row_text)
    row.get_by_test_id("assign-edit").click()
    box = row.get_by_test_id("assign-pick")
    for item in box.locator("input[type=checkbox]").all():
        if item.is_checked():
            item.uncheck()
    for user in user_ids:
        box.locator(f'input[data-user-id="{user}"]').check()
    box.get_by_test_id("assign-pick-ok").click()


def heading_ids(api: "Api", did: str) -> dict[str, str]:
    objects = api.call("GET", f"/documents/{did}/objects")["objects"]
    return {item["markdown"].lstrip("# ").strip(): item["id"] for item in objects if item["kind"] == "heading"}


def require_targets() -> None:
    missing = [name for name, value in (("E2E_API", API), ("E2E_WEB", WEB), ("DATABASE_URL", os.environ.get("DATABASE_URL", "")),
                                        ("FILE_DATA_DIR", os.environ.get("FILE_DATA_DIR", ""))) if not value]
    if missing:
        raise SystemExit(f"请设置环境变量 {', '.join(missing)}（被测服务的接口、前端地址和同一个数据库），脚本不使用默认值。")


def main() -> int:  # noqa: C901 - 验收脚本按用例顺序展开
    require_targets()
    data = seed()
    users, admin = data["users"], data["admin"]
    pid, did = data["project"]["id"], data["document"]["id"]
    lisi, wangce = Api(f"lisi{RUN}"), Api(f"wangce{RUN}")

    with sync_playwright() as playwright:
        browser = playwright.chromium.launch()
        ctx_a = browser.new_context(viewport={"width": WIDTH, "height": HEIGHT}, accept_downloads=True)
        ctx_b = browser.new_context(viewport={"width": WIDTH, "height": HEIGHT}, accept_downloads=True)
        page, page2 = ctx_a.new_page(), ctx_b.new_page()
        open_pages.extend([page, page2])
        login(page, ADMIN, ADMIN_PASSWORD)
        login(page2, f"lisi{RUN}")

        @step("G1 负责人在工作台“编写分工”面板按章节分配，文档显示“按分工编辑”")
        def g1():
            open_doc(page, pid, did)
            expect(page.get_by_test_id("doc-assign-mode-chip")).to_have_text("按分工编辑")
            page.get_by_test_id("assign-open").click()
            panel = page.get_by_test_id("assign-panel")
            expect(panel).to_be_visible()
            # 29 抽屉：宽 400、焦点在面板里、不压在正文上
            box, paper = panel.bounding_box(), page.locator(".paper").bounding_box()
            assert box and 396 <= box["width"] <= 404, box
            assert paper and paper["x"] + paper["width"] <= box["x"] + 1, (paper, box)
            assert page.evaluate("() => !!document.activeElement?.closest('[data-testid=assign-panel]')")
            expect(panel.get_by_test_id("assign-summary")).to_contain_text("未分配章节 3 个")
            rows = panel.get_by_test_id("assign-heading")
            expect(rows).to_have_count(3)
            pick(panel, "概述", [users["lisi"]])
            expect(panel.get_by_test_id("assign-summary")).to_contain_text("未分配章节 2 个", timeout=15000)
            pick(panel, "功能需求", [users["wangce"]])
            expect(panel.get_by_test_id("assign-summary")).to_contain_text("未分配章节 1 个", timeout=15000)
            expect(rows.filter(has_text="概述").get_by_test_id("assign-who")).to_have_text("李四")
            # 多人分配：附录给王测和刘洋，再取消
            pick(panel, "附录", [users["wangce"], users["liuyang"]])
            expect(rows.filter(has_text="附录").get_by_test_id("assign-who")).to_have_text("王测、刘洋", timeout=15000)
            view = admin.call("GET", f"/documents/{did}/assignments")
            appendix = [item for item in view["explicit"] if {person["name"] for person in item["assignees"]} == {"王测", "刘洋"}]
            assert len(appendix) == 1, view["explicit"]
            pick(panel, "附录", [])
            expect(rows.filter(has_text="附录").get_by_test_id("assign-who")).to_have_text("未分配", timeout=15000)
            # 筛选［全部｜未分配｜某人▾］
            panel.get_by_test_id("assign-filter-unassigned").click()
            expect(rows).to_have_count(1)
            expect(rows.first).to_contain_text("附录")
            panel.get_by_test_id("assign-filter-person").select_option(str(users["wangce"]))
            expect(rows).to_have_count(1)
            expect(rows.first).to_contain_text("功能需求")
            panel.get_by_test_id("assign-filter-all").click()
            expect(rows).to_have_count(3)
            shot(page, "A07-01-分工面板")
            view = admin.call("GET", f"/documents/{did}/assignments")
            assert sorted(item["assignees"][0]["name"] for item in view["explicit"]) == ["李四", "王测"], view["explicit"]
            # Esc 关闭，焦点回到“编写分工”按钮
            panel.get_by_test_id("assign-close").focus()
            page.keyboard.press("Escape")
            expect(panel).to_have_count(0)
            assert page.evaluate("() => document.activeElement?.getAttribute('data-testid')") == "assign-open"
        g1()

        @step("G1b 切换分工模式先确认（16 §4.3），开放编辑时标签为“开放编辑”")
        def g1b():
            page.get_by_test_id("assign-open").click()
            panel = page.get_by_test_id("assign-panel")
            panel.get_by_test_id("assign-mode-toggle").click()
            dialog = page.get_by_test_id("assign-mode-confirm")
            expect(dialog).to_contain_text("改为开放编辑？")
            dialog.get_by_role("button", name="取消").click()
            expect(dialog).to_have_count(0)
            assert admin.call("GET", f"/documents/{did}/assignments")["assignmentMode"] == "enforced"
            panel.get_by_test_id("assign-mode-toggle").click()
            page.get_by_test_id("assign-mode-confirm-ok").click()
            expect(page.get_by_test_id("doc-assign-mode-chip")).to_have_text("开放编辑", timeout=15000)
            panel.get_by_test_id("assign-mode-toggle").click()
            dialog = page.get_by_test_id("assign-mode-confirm")
            expect(dialog).to_contain_text("开启按分工编辑？")
            expect(dialog.get_by_test_id("assign-mode-unassigned")).to_contain_text("还有 1 章未分配")
            shot(page, "A07-01b-开启确认")
            page.get_by_test_id("assign-mode-confirm-ok").click()
            expect(page.get_by_test_id("doc-assign-mode-chip")).to_have_text("按分工编辑", timeout=15000)
            assert admin.call("GET", f"/documents/{did}/assignments")["assignmentMode"] == "enforced"
            page.get_by_test_id("assign-close").click()
        g1b()

        @step("G2 编写人员只能改分给自己的章节：其余只读，点了有提示；自己的章节保存成功")
        def g2():
            open_doc(page2, pid, did)
            expect(page2.get_by_test_id("doc-assign-mode-chip")).to_be_visible()
            mine, theirs, free = para(page2, "概述第一段"), para(page2, "功能需求第一段"), para(page2, "附录第一段")
            expect(theirs).to_have_attribute("data-editable", "false", timeout=15000)
            expect(theirs).to_have_attribute("data-assignee", "王测")
            expect(free).to_have_attribute("data-editable", "false")
            expect(mine).not_to_have_attribute("data-editable", "false")
            expect(page2.locator(".paper-content")).to_have_attribute("contenteditable", "true", timeout=30000)
            # 29 M-1：只读样式、gutter 姓名片 / 未分配、我的部分竖条、目录“我的部分”、首次提示
            expect(theirs).to_have_attribute("data-assignee-label", "王测")
            expect(free).to_have_attribute("data-assignee-label", "未分配")
            expect(mine).to_have_attribute("data-mine", "")
            assert theirs.evaluate("n => getComputedStyle(n).color") != mine.evaluate("n => getComputedStyle(n).color")
            assert "王测" in theirs.evaluate("n => getComputedStyle(n, '::before').content")
            expect(page2.get_by_test_id("skill-settings-open")).to_be_disabled()   # 27 复验 N-3 / Q-5
            expect(page2.get_by_test_id("mine-intro")).to_contain_text("你负责本文档的 1 章，其他部分只读。")
            expect(page2.get_by_test_id("outline-filter-mine")).to_have_text("我的部分 1")
            expect(page2.locator(".outline-item")).to_have_count(1)
            page2.get_by_test_id("outline-filter-all").click()
            expect(page2.locator(".outline-item")).to_have_count(3)
            expect(page2.get_by_test_id("outline-lock")).to_have_count(2)
            mine.click()
            page2.keyboard.press("End")
            page2.keyboard.type("（李四补充）")
            expect(mine).to_contain_text("（李四补充）")
            page2.keyboard.press("Control+s")
            saved = ""
            for _ in range(60):
                saved = admin.call("GET", f"/documents/{did}")["markdownContent"]
                if "（李四补充）" in saved:
                    break
                page2.wait_for_timeout(500)
            wait_saved(page2)
            assert "概述第一段：系统的建设背景。（李四补充）" in saved, saved[:400]
            # 点只读段：不能输入，出现归属提示
            before = admin.call("GET", f"/documents/{did}")["markdownContent"]
            theirs.click()
            page2.keyboard.type("越权")
            page2.keyboard.press("Control+s")
            expect(theirs).not_to_contain_text("越权")
            free.click()
            page2.keyboard.press("Backspace")
            shot(page2, "A07-02-编写人员只读章节")
            page2.wait_for_timeout(1500)
            after = admin.call("GET", f"/documents/{did}")["markdownContent"]
            assert after == before, (before[:400], after[:400])
            # 29 M-2：只读块里打的字进了暂存（16 §4.8 文案），可以放到我负责的一处
            # （Playwright 输入中文走 insertText、不产生按键；这里用英文按键）
            theirs.click()
            page2.keyboard.type("ABC")
            hint = page2.get_by_test_id("locked-hint")
            expect(hint).to_contain_text("该段由 王测 负责，你不能修改。如需修改请联系项目负责人")
            expect(hint.get_by_test_id("locked-stash")).to_contain_text("ABC")  # 前面输入的“越权”（insertText）也已暂存
            expect(theirs).not_to_contain_text("ABC")
            hint.get_by_test_id("locked-stash-place").click()
            expect(para(page2, "概述第二段")).to_contain_text("ABC")
            for _ in range(60):
                if "ABC" in admin.call("GET", f"/documents/{did}")["markdownContent"]:
                    break
                page2.wait_for_timeout(500)
            else:
                raise AssertionError("暂存的字没有保存到我的段落")
            wait_saved(page2)
            assert "# 概述\n" in admin.call("GET", f"/documents/{did}")["markdownContent"], "暂存不能写进标题"
            # 复验 N-2：同一章先点标题再点段落，每次都有提示；暂存只保留这次的输入
            page2.locator(".paper-content h1", has_text="功能需求").first.click()
            expect(hint).to_be_visible()
            para(page2, "功能需求第二段").click()
            expect(hint).to_be_visible()
            page2.keyboard.type("XY")
            expect(hint.get_by_test_id("locked-stash")).to_have_text(re.compile("刚才输入的“XY”已暂存"))
            para(page2, "附录第一段").click()
            expect(hint).to_contain_text("该段尚未分配")
            page2.keyboard.type("Z")
            expect(hint.get_by_test_id("locked-stash")).to_contain_text("刚才输入的“Z”已暂存")
            hint.get_by_role("button", name="复制").click()
        g2()

        @step("G2b 在自己段首按退格不会并入别人的段（前端拦下并提示）")
        def g2b():
            first_theirs = para(page2, "功能需求第一段")
            heading = page2.locator(".paper-content h1", has_text="功能需求").first
            expect(heading).to_have_attribute("data-editable", "false")
            para(page2, "概述第二段").click()
            page2.keyboard.press("End")
            page2.keyboard.press("Delete")
            expect(page2.get_by_test_id("locked-hint")).to_contain_text("王测", timeout=5000)
            expect(first_theirs).to_be_visible()
        g2b()

        @step("G3 正在编辑时被改了分工：保存被拒（409），提示条不自动重试、可下载，服务器正文不变")
        def g3():
            heads = heading_ids(admin, did)
            view = admin.call("GET", f"/documents/{did}/assignments")
            mine = para(page2, "概述第二段")
            mine.click()
            page2.keyboard.press("End")
            # 管理员先把“概述”改给王测（李四的页面 45 秒内才会知道），李四紧接着保存
            admin.call("PUT", f"/documents/{did}/assignments",
                       json={"objectIds": [heads["概述"]], "assigneeIds": [str(users["wangce"])], "expectedAssignmentsRevision": view["assignmentsRevision"]})
            before = admin.call("GET", f"/documents/{did}")["markdownContent"]
            page2.keyboard.type("（被拒）")
            page2.keyboard.press("Control+s")
            banner = page2.get_by_test_id("save-rejected-banner")
            expect(banner).to_contain_text("这次修改没有保存：其中 1 处属于他人负责的部分。", timeout=30000)
            expect(banner.get_by_role("button", name="重试保存")).to_have_count(0)
            expect(page2.get_by_test_id("save-status")).to_contain_text("未保存")
            assert admin.call("GET", f"/documents/{did}")["markdownContent"] == before
            # 归属已刷新：概述也锁上了
            expect(para(page2, "概述第一段")).to_have_attribute("data-editable", "false", timeout=15000)
            banner.get_by_test_id("save-rejected-open").click()
            panel = page2.get_by_test_id("save-rejected-panel")
            item = panel.get_by_test_id("save-rejected-item").first
            expect(item).to_contain_text("修改了 王测 负责的内容")
            expect(item).to_contain_text("请联系 王测")
            item.click()
            shot(page2, "A07-03-保存被拒")
            with page2.expect_download() as info:
                panel.get_by_test_id("save-rejected-download").click()
            assert "（被拒）" in Path(info.value.path()).read_text(encoding="utf-8")
            # 只保存我负责的部分：这里已没有我的部分 → 服务器内容不变，越权的字进暂存
            panel.get_by_test_id("save-rejected-save-mine").click()
            expect(banner).to_have_count(0, timeout=30000)
            expect(page2.get_by_test_id("locked-stash-bar")).to_contain_text("（被拒）", timeout=15000)
            expect(para(page2, "概述第二段")).not_to_contain_text("（被拒）")
            assert "（被拒）" not in admin.call("GET", f"/documents/{did}")["markdownContent"]
            page2.get_by_test_id("locked-stash-bar").get_by_role("button", name="清除").click()
            # 还给李四
            view = admin.call("GET", f"/documents/{did}/assignments")
            admin.call("PUT", f"/documents/{did}/assignments",
                       json={"objectIds": [heads["概述"]], "assigneeIds": [str(users["lisi"])], "expectedAssignmentsRevision": view["assignmentsRevision"]})
        g3()

        @step("G4 接口：越权 409 SECTION_NOT_ASSIGNED、缺 blockIds 422、负责人不受限")
        def g4():
            record = lisi.call("GET", f"/documents/{did}")
            objects = lisi.call("GET", f"/documents/{did}/objects")["objects"]
            ids = [item["id"] for item in objects]
            body = "\n\n".join(item["markdown"] for item in objects) + "\n"
            changed = body.replace("功能需求第二段：权限管理。", "功能需求第二段：李四改的。")
            response = lisi.raw("PATCH", f"/documents/{did}/content", json={"expectedRevision": record["revision"], "markdownContent": changed, "allowMerge": True, "blockIds": ids})
            assert response.status_code == 409 and response.json()["error"]["code"] == "SECTION_NOT_ASSIGNED", response.text
            details = response.json()["error"]["details"]
            assert details["total"] >= 1 and details["violations"][0]["owners"][0]["name"] == "王测", details
            response = lisi.raw("PATCH", f"/documents/{did}/content", json={"expectedRevision": record["revision"], "markdownContent": changed, "allowMerge": True})
            assert response.status_code == 422 and response.json()["error"]["code"] == "BLOCK_IDS_REQUIRED", response.text
            response = admin.raw("PATCH", f"/documents/{did}/content", json={"expectedRevision": record["revision"], "markdownContent": changed.replace("李四改的", "负责人改的"), "allowMerge": True, "blockIds": ids})
            assert response.status_code == 200, response.text
        g4()

        @step("G5 编写人员新建的文档整篇归创建人，其他编写人员只读")
        def g5():
            own = upload_doc(wangce, pid, "测试计划", make_word([("测试范围", ["王测写的测试范围。"])]))
            view = wangce.call("GET", f"/documents/{own['id']}/assignments")
            assert view["assignmentMode"] == "enforced"
            assert [(item["objectId"], item["assignees"][0]["name"]) for item in view["explicit"]] == [("@document", "王测")], view["explicit"]
            page2.goto(f"{WEB}/projects/{pid}/documents/{own['id']}")
            expect(para(page2, "王测写的测试范围")).to_have_attribute("data-assignee", "王测", timeout=60000)
            expect(page2.get_by_test_id("no-assign-notice")).to_contain_text("还没有分配给你的部分")
            detail = lisi.call("GET", f"/documents/{own['id']}")
            assert detail["permissions"]["readonlyReason"] == "assignment", detail["permissions"]
        g5()

        @step("G6 编写人员在按分工编辑的文档上不能恢复版本")
        def g6():
            open_doc(page2, pid, did)
            page2.get_by_role("button", name="版本记录", exact=True).click()
            page2.locator(".history-panel .vitem").nth(1).click()
            restore = page2.get_by_role("button", name="恢复此版本")
            expect(restore).to_be_disabled(timeout=15000)
            expect(page2.locator(".history-hint")).to_contain_text("只有项目负责人能恢复版本")
            record = lisi.call("GET", f"/documents/{did}")
            response = lisi.raw("POST", f"/documents/{did}/versions/1/restore", json={"expectedRevision": record["revision"]})
            assert response.status_code == 403, response.text
            page2.get_by_role("button", name="版本记录", exact=True).click()
        g6()

        @step("G7 成员表“编写分工”列；点数字打开概述按人视图并定位")
        def g7():
            page.goto(f"{WEB}/projects/{pid}")
            row = page.locator(f'[data-testid=member-row][data-user-id="{users["lisi"]}"]')
            expect(row.get_by_test_id("member-assign-count")).to_contain_text("1 章", timeout=30000)
            expect(row.get_by_test_id("member-assign-count")).to_contain_text("1 篇文档")
            expect(page.locator(f'[data-testid=member-row][data-user-id="{users["liuyang"]}"]').get_by_test_id("member-assign-count")).to_have_text("—")
            row.get_by_test_id("member-assign-count").click()
            person = page.locator(f'[data-testid=overview-assign-person-row][data-user-id="{users["lisi"]}"]')
            expect(person).to_be_visible(timeout=15000)
            expect(person).to_contain_text("软件需求规格说明 · 1 章")
            expect(page.get_by_test_id("overview-assign-no-work")).to_contain_text("刘洋")
            shot(page, "A07-04-概述按人")
            person.get_by_test_id("overview-assign-person-link").first.click()
            expect(page.get_by_test_id("assign-filter-person")).to_have_value(str(users["lisi"]), timeout=60000)
            expect(page.get_by_test_id("assign-heading")).to_have_count(1)
            page.goto(f"{WEB}/projects/{pid}")
        g7()

        @step("G8 S-5：降为只读前确认（取消不改；确认后分工清除）")
        def g8():
            row = page.locator(f'[data-testid=member-row][data-user-id="{users["lisi"]}"]')
            row.get_by_test_id("member-role-select").select_option("viewer")
            dialog = page.get_by_test_id("member-demote-dialog")
            expect(dialog).to_be_visible()
            expect(dialog).to_contain_text("1 章将变为未分配")
            shot(page, "A07-05-降级确认")
            dialog.get_by_role("button", name="取消").click()
            expect(row.get_by_test_id("member-role-select")).to_have_value("editor")
            assert any(item["userId"] == str(users["lisi"]) and item["role"] == "editor" for item in admin.call("GET", f"/projects/{pid}/members")["items"])
            row.get_by_test_id("member-role-select").select_option("viewer")
            page.get_by_test_id("member-demote-confirm").click()
            expect(row.get_by_test_id("member-role-select")).to_have_value("viewer", timeout=15000)
            expect(row.get_by_test_id("member-assign-count")).to_have_text("—", timeout=15000)
            view = admin.call("GET", f"/documents/{did}/assignments")
            assert all(item["assignees"][0]["name"] != "李四" for item in view["explicit"]), view["explicit"]
        g8()

        @step("G9 概述“编写分工”按文档：已分 / 总章数，未分配可去分配")
        def g9():
            page.goto(f"{WEB}/projects/{pid}")
            section = page.get_by_test_id("overview-assign-section")
            section.get_by_role("button", name="按文档").click()
            doc_row = section.locator(f'[data-testid=overview-assign-doc-row][data-document-id="{did}"]')
            expect(doc_row).to_contain_text("1 / 3 章", timeout=15000)
            expect(section.locator('[data-testid=overview-assign-doc-row]', has_text="测试计划")).to_contain_text("整篇：王测")
            doc_row.get_by_test_id("overview-assign-unassigned").click()
            expect(page.get_by_test_id("assign-panel")).to_be_visible(timeout=60000)
            expect(page.get_by_test_id("assign-filter-unassigned")).to_have_attribute("aria-pressed", "true")
            shot(page, "A07-06-概述去分配")
        g9()

        @step("G12 负责人操作失败要有提示（REVISION_CONFLICT）；负责人停用后行内“负责人已停用”与“需改派”筛选")
        def g12():
            open_doc(page, pid, did)
            page.get_by_test_id("assign-open").click()
            panel = page.get_by_test_id("assign-panel")
            expect(panel.get_by_test_id("assign-heading")).to_have_count(3)
            view = admin.call("GET", f"/documents/{did}/assignments")
            admin.call("PUT", f"/documents/{did}/assignment-mode", json={"mode": "open", "expectedAssignmentsRevision": view["assignmentsRevision"]})
            panel.get_by_test_id("assign-mode-toggle").click()
            page.get_by_test_id("assign-mode-confirm-ok").click()
            admin_name = admin.call("GET", "/auth/me")["user"]["name"]  # 已有 admin 账号的名字不一定是“系统管理员”
            expect(panel.get_by_test_id("assign-error")).to_contain_text(f"分工刚被 {admin_name} 调整过，已显示最新，请确认后再切换。", timeout=15000)
            view = admin.call("GET", f"/documents/{did}/assignments")
            if view["assignmentMode"] != "enforced":
                admin.call("PUT", f"/documents/{did}/assignment-mode", json={"mode": "enforced", "expectedAssignmentsRevision": view["assignmentsRevision"]})
            account = next(item for item in admin.call("GET", "/users", params={"account": f"wangce{RUN}"})["items"] if item["id"] == str(users["wangce"]))
            admin.call("POST", f"/users/{users['wangce']}/disable", json={"reason": "A07 验收", "expectedRevision": account["revision"]})
            try:
                open_doc(page, pid, did)
                page.get_by_test_id("assign-open").click()
                row = panel.get_by_test_id("assign-heading").filter(has_text="功能需求")
                expect(row.get_by_test_id("assign-needs-reassign")).to_have_text("负责人已停用", timeout=15000)
                panel.get_by_test_id("assign-filter-reassign").click()
                expect(panel.get_by_test_id("assign-heading")).to_have_count(1)
                shot(page, "A07-08-负责人已停用")
            finally:
                account = next(item for item in admin.call("GET", "/users", params={"account": f"wangce{RUN}"})["items"] if item["id"] == str(users["wangce"]))
                admin.call("POST", f"/users/{users['wangce']}/enable", json={"reason": "A07 验收", "expectedRevision": account["revision"]})
        g12()

        login(page2, f"liuyang{RUN}")
        liuyang = Api(f"liuyang{RUN}")

        @step("G10 3000 块的文档：编写人员改自己章节的末段，保存成功、不错位（27 M-1）")
        def g10():
            chapters = [(f"第{c}章", [f"第{c}章第{k}段内容。" for k in range(1, 100)]) for c in range(1, 31)]
            big = upload_doc(admin, pid, "大文档", make_word(chapters))
            heads = heading_ids(admin, big["id"])
            view = admin.call("GET", f"/documents/{big['id']}/assignments")
            admin.call("PUT", f"/documents/{big['id']}/assignments",
                       json={"objectIds": [heads["第30章"]], "assigneeIds": [str(users["liuyang"])], "expectedAssignmentsRevision": view["assignmentsRevision"]})
            assert len(admin.call("GET", f"/documents/{big['id']}/objects")["objects"]) >= 3000
            open_doc(page2, pid, big["id"])
            target = para(page2, "第30章第99段内容")
            expect(target).to_have_attribute("data-mine", "", timeout=60000)
            target.click()
            page2.keyboard.press("End")
            page2.keyboard.type("（刘洋）")
            page2.keyboard.press("Control+s")
            for _ in range(90):
                if "第30章第99段内容。（刘洋）" in admin.call("GET", f"/documents/{big['id']}")["markdownContent"]:
                    break
                page2.wait_for_timeout(500)
            else:
                raise AssertionError("大文档保存没有生效")
            expect(page2.get_by_test_id("save-rejected-banner")).to_have_count(0)
            wait_saved(page2, 60000)
        g10()

        @step("G11 没有分工的编写人员整篇只读并提示；分到章节后提示并可编辑（29 M-6）")
        def g11():
            page2.goto(f"{WEB}/projects/{pid}/documents/{did}")
            expect(page2.locator(".paper-content")).to_be_visible(timeout=60000)
            expect(page2.get_by_test_id("no-assign-notice")).to_contain_text("这篇文档还没有分配给你的部分，目前只读", timeout=30000)
            expect(page2.locator(".paper-content")).not_to_have_attribute("contenteditable", "true")
            heads = heading_ids(admin, did)
            view = admin.call("GET", f"/documents/{did}/assignments")
            admin.call("PUT", f"/documents/{did}/assignments",
                       json={"objectIds": [heads["附录"]], "assigneeIds": [str(users["liuyang"])], "expectedAssignmentsRevision": view["assignmentsRevision"]})
            expect(page2.get_by_test_id("assign-notice")).to_contain_text("新分配给你", timeout=90000)
            expect(page2.get_by_test_id("no-assign-notice")).to_have_count(0)
            expect(para(page2, "附录第一段")).to_have_attribute("data-mine", "")
            shot(page2, "A07-07-分到章节")
            assert liuyang.call("GET", f"/documents/{did}/assignments")["assignmentMode"] == "enforced"
        g11()

        @step("G13 只保存我负责的部分：负责人同时改了我负责的段 → 不覆盖，走冲突流程（28 复验 N-1，F5 conflict）")
        def g13():
            heads = heading_ids(admin, did)
            view = admin.call("GET", f"/documents/{did}/assignments")
            admin.call("PUT", f"/documents/{did}/assignments",
                       json={"objectIds": [heads["功能需求"]], "assigneeIds": [str(users["liuyang"])], "expectedAssignmentsRevision": view["assignmentsRevision"]})
            open_doc(page2, pid, did)
            func, appendix = para(page2, "功能需求第一段"), para(page2, "附录第一段")
            expect(func).to_have_attribute("data-mine", "", timeout=30000)
            # 先改负责人后改分工，刘洋的页面还不知道
            objects = admin.call("GET", f"/documents/{did}/objects")["objects"]
            record = admin.call("GET", f"/documents/{did}")
            body = "\n\n".join(item["markdown"] for item in objects) + "\n"
            changed = body.replace("附录第一段：术语表。", "附录第一段：术语表。（负责人改）")
            admin.call("PATCH", f"/documents/{did}/content", json={"expectedRevision": record["revision"], "markdownContent": changed,
                                                                 "allowMerge": True, "blockIds": [item["id"] for item in objects]})
            view = admin.call("GET", f"/documents/{did}/assignments")
            admin.call("PUT", f"/documents/{did}/assignments",
                       json={"objectIds": [heads["功能需求"]], "assigneeIds": [str(users["wangce"])], "expectedAssignmentsRevision": view["assignmentsRevision"]})
            func.click(); page2.keyboard.press("End"); page2.keyboard.type("X刘")
            expect(func).to_contain_text("X刘")
            appendix.click(); page2.keyboard.press("End"); page2.keyboard.type("Y刘")
            expect(appendix).to_contain_text("Y刘")
            page2.keyboard.press("Control+s")
            banner = page2.get_by_test_id("save-rejected-banner")
            expect(banner).to_be_visible(timeout=30000)
            banner.get_by_test_id("save-rejected-open").click()
            page2.get_by_test_id("save-rejected-save-mine").click()
            page2.wait_for_timeout(3000)
            expect(page2.get_by_text(re.compile("你负责的部分有 1 处在你编辑期间也被他人修改过，已按冲突处理"))).to_be_visible(timeout=30000)
            expect(page2.get_by_text("已保存你负责部分的修改")).to_have_count(0)
            saved = admin.call("GET", f"/documents/{did}")["markdownContent"]
            assert "（负责人改）" in saved, saved[-300:]                      # 负责人的并发修改没被覆盖
            assert "X刘" not in saved, saved[-300:]                           # 越权部分没保存
            expect(page2.get_by_test_id("conflict-mark").first).to_be_visible()  # 冲突卡（我的版本“Y刘”在冲突副本里）
            shot(page2, "A07-09-只保存我负责的部分-冲突")
        g13()

        browser.close()
    failed = [item for item in results if not item[1]]
    print(f"\n{len(results) - len(failed)}/{len(results)} 通过（{VIEW}）")
    for name, _ok, message in failed:
        print(f"  失败 {name}: {message}")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
