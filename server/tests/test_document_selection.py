from lxml import html
import markdown
import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from app.docx_export import edit_original_docx
from app.docx_parser import parse_docx

from app.document_selection import (
    SelectionError,
    decorate_selection_html,
    replace_selection,
    selection_blocks,
    validate_selection,
)


def _units(value):
    return len(value.encode("utf-16-le")) // 2


def _selection(source, block=0, start=0, end=None):
    item = selection_blocks(source)[block]
    end = len(item["text"]) if end is None else end
    return {"blockId": item["blockId"], "start": _units(item["text"][:start]),
            "end": _units(item["text"][:end]), "text": item["text"][start:end]}


def _root(source):
    return html.fragment_fromstring(markdown.markdown(source, extensions=["tables", "fenced_code"]), create_parent="div")


def _multi_selection(source, first=0, last=1, start=0, end=None):
    blocks = selection_blocks(source)[first:last + 1]
    end = len(blocks[-1]["text"]) if end is None else end
    return {"blockId": blocks[0]["blockId"], "endBlockId": blocks[-1]["blockId"],
            "start": _units(blocks[0]["text"][:start]), "end": _units(blocks[-1]["text"][:end]),
            "text": "\n\n".join([blocks[0]["text"][start:], *[block["text"] for block in blocks[1:-1]], blocks[-1]["text"][:end]])}


@pytest.mark.parametrize("replacement", ["合并后的正文。", "第一段\n\n第二段", "第一段\n\n第二段\n\n第三段"])
def test_continuous_paragraph_selection_keeps_exact_unselected_text_and_emphasis(replacement):
    source = "# 范围\n\n**前缀旧😀文**\n\n中间段\n\n*尾旧文后缀*\n\n相同原文"
    selected = _multi_selection(source, first=0, last=2, start=2, end=3)
    assert validate_selection(source, selected) == "旧😀文\n\n中间段\n\n尾旧文"
    updated = replace_selection(source, selected, replacement)
    nodes = _root(updated).xpath("./p")
    parts = replacement.split("\n\n")
    assert ["".join(node.itertext()) for node in nodes] == [
        "前缀" + parts[0] + ("后缀" if len(parts) == 1 else ""),
        *parts[1:-1], *([parts[-1] + "后缀"] if len(parts) > 1 else []), "相同原文"]
    assert nodes[0].xpath("./strong/text()")[0].startswith("前缀")
    assert nodes[-2].xpath("./em/text()") == ["后缀"]
    assert nodes[-1].text == "相同原文"


@pytest.mark.parametrize("gap", ["# 其他章节", "| A | B |\n| --- | --- |\n| x | y |", "![图](docx-image:x)", "<span>结构</span>", "- 列表项", "---", "```\n代码\n```"])
def test_multi_paragraph_range_cannot_swallow_unsupported_structure(gap):
    source = "首段\n\n" + gap + "\n\n尾段"
    blocks = selection_blocks(source)
    selected = {"blockId": blocks[0]["blockId"], "endBlockId": blocks[-1]["blockId"],
                "start": 0, "end": 2, "text": "首段\n\n尾段"}
    with pytest.raises(SelectionError, match="连续"):
        replace_selection(source, selected, "新正文")


def test_multi_paragraph_binding_rejects_stale_middle_unicode_and_reversed_ranges():
    source = "甲😀尾\n\n中间段\n\n乙😀尾"
    selected = _multi_selection(source, last=2, start=1, end=2)
    assert validate_selection(source, selected) == "😀尾\n\n中间段\n\n乙😀"
    for changed in (source.replace("中间段", "新段落"), source.replace("中间段", "# 中间段")):
        with pytest.raises(SelectionError):
            replace_selection(changed, selected, "替换")
    for extra in ({"end": 2}, {"start": 2}, {"endBlockId": selected["blockId"], "blockId": selected["endBlockId"]}, {"endBlockId": False}):
        with pytest.raises(SelectionError):
            replace_selection(source, selected | extra, "替换")


@pytest.mark.parametrize("replacement", ["合并正文", "第一段\n\n第二段", "第一段\n\n第二段\n\n第三段\n\n第四段"])
def test_cross_paragraph_word_export_keeps_unselected_native_runs_and_page_properties(tmp_path, replacement):
    source, output = tmp_path / "source.docx", tmp_path / "edited.docx"
    word = Document()
    word.add_heading("范围", level=1)
    word.sections[0].header.paragraphs[0].text = "原页眉"
    first = word.add_paragraph()
    first.paragraph_format.space_before = Pt(3)
    prefix = first.add_run("前缀")
    prefix.font.size, prefix.bold = Pt(19), True
    first.add_run("旧😀内容")
    word.add_paragraph("中间旧正文")
    last = word.add_paragraph()
    last.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    last.paragraph_format.space_after = Pt(9)
    last.add_run("尾旧内容")
    suffix = last.add_run("后缀")
    suffix.font.size, suffix.italic, suffix.font.color.rgb = Pt(17), True, RGBColor.from_string("224466")
    word.add_paragraph("其他正文")
    word.save(source)
    original = source.read_bytes()
    markdown = parse_docx(str(source)).markdown
    selected = _multi_selection(markdown, last=2, start=2, end=4)
    updated = replace_selection(markdown, selected, replacement)
    edit_original_docx(updated, source, output)
    result = Document(output)
    parts = replacement.split("\n\n")
    assert [paragraph.text for paragraph in result.paragraphs] == [
        "范围", "前缀" + parts[0] + ("后缀" if len(parts) == 1 else ""),
        *parts[1:-1], *([parts[-1] + "后缀"] if len(parts) > 1 else []), "其他正文"]
    prefix_run = next(run for run in result.paragraphs[1].runs if "前缀" in run.text)
    suffix_run = next(run for run in result.paragraphs[-2].runs if "后缀" in run.text)
    assert prefix_run.font.size == Pt(19) and prefix_run.bold is True
    assert suffix_run.font.size == Pt(17) and suffix_run.italic is True
    assert suffix_run.font.color.rgb == RGBColor.from_string("224466")
    assert result.paragraphs[1].paragraph_format.space_before == Pt(3)
    if len(parts) > 1:
        assert result.paragraphs[-2].alignment == WD_ALIGN_PARAGRAPH.RIGHT
        assert result.paragraphs[-2].paragraph_format.space_after == Pt(9)
    assert result.sections[0].header.paragraphs[0].text == "原页眉"
    assert source.read_bytes() == original


def test_repeated_unselected_prefix_and_suffix_keep_their_native_run_owners_when_merging(tmp_path):
    source, output = tmp_path / "source.docx", tmp_path / "merged.docx"
    word = Document()
    first = word.add_paragraph()
    first.add_run("repeat").font.size = Pt(19)
    first.add_run("chosen")
    last = word.add_paragraph()
    last.add_run("repeat")
    last.add_run("repeat").font.size = Pt(17)
    word.save(source)
    markdown = parse_docx(str(source)).markdown
    updated = replace_selection(markdown, _multi_selection(markdown, start=6, end=6), "repeat")
    edit_original_docx(updated, source, output)
    paragraph = Document(output).paragraphs[0]
    fonts = [run.font.size for run in paragraph.runs for _char in run.text]
    assert paragraph.text == "repeatrepeatrepeat" and fonts[:6] == [Pt(19)] * 6 and fonts[-6:] == [Pt(17)] * 6


def test_repeated_document_paragraphs_do_not_move_unselected_native_fonts(tmp_path):
    source, output = tmp_path / "repeated.docx", tmp_path / "edited.docx"
    word = Document()
    for size in (19, 17, 31):
        word.add_paragraph().add_run("repeat").font.size = Pt(size)
    word.save(source)
    markdown = parse_docx(str(source)).markdown
    edit_original_docx(replace_selection(markdown, _selection(markdown, block=1), "changed"), source, output)
    paragraphs = Document(output).paragraphs
    assert [paragraph.text for paragraph in paragraphs] == ["repeat", "changed", "repeat"]
    assert [paragraph.runs[0].font.size for paragraph in paragraphs] == [Pt(19), Pt(17), Pt(31)]


def test_repeated_paragraphs_and_words_use_the_bound_block_and_offset():
    source = "same word same\n\nsame word same\n\nsame word same"
    blocks = selection_blocks(source)
    assert len({item["blockId"] for item in blocks}) == 3
    descriptor = _selection(source, block=1, start=10, end=14)
    assert validate_selection(source, descriptor) == "same"
    assert replace_selection(source, descriptor, "changed") == "same word same\n\nsame word changed\n\nsame word same"


def test_utf16_offsets_do_not_split_astral_characters():
    source = "甲😀乙 same"
    descriptor = _selection(source, start=1, end=2)
    assert descriptor["start"] == 1 and descriptor["end"] == 3
    assert replace_selection(source, descriptor, "新字") == "甲新字乙 same"
    descriptor["end"] = 2
    with pytest.raises(SelectionError, match="Unicode"):
        replace_selection(source, descriptor, "x")


@pytest.mark.parametrize(("source", "start", "end", "expected"), [
    ("left **bold words** right", 6, 10, "left **bNEWwords** right"),
    ("left **bold** right", 2, 8, "leNEW**d** right"),
    ("left **bold** right", 7, 14, "left **boNEW**t"),
    ("left **bold** right", 0, 15, "NEW"),
    ("**bold** plain *italic*", 0, 17, "**NEW**"),
])
def test_partial_and_cross_run_edits_preserve_unselected_styles(source, start, end, expected):
    updated = replace_selection(source, _selection(source, start=start, end=end), "NEW")
    assert updated == expected


def test_whole_paragraph_keeps_start_style_and_other_document_blocks():
    source = "# Heading\n\n**first** and *second*\n\n| Name | Value |\n| --- | --- |\n| other | keep |\n\nTail paragraph"
    updated = replace_selection(source, _selection(source, end=16), "Replacement")
    root = _root(updated)
    assert root.xpath("./p[1]/strong/text()") == ["Replacement"]
    assert root.xpath("./h1/text()") == ["Heading"]
    assert root.xpath("./table/tbody/tr/td/text()") == ["other", "keep"]
    assert root.xpath("./p[last()]/text()") == ["Tail paragraph"]


def test_plain_ai_output_does_not_become_markup_or_html():
    source = "Prefix selected suffix"
    replacement = "**literal** <script>alert(1)</script> & [link](url)"
    updated = replace_selection(source, _selection(source, start=7, end=15), replacement)
    root = _root(updated)
    assert "".join(root.xpath("./p")[0].itertext()) == "Prefix " + replacement + " suffix"
    assert not root.xpath(".//strong|.//script|.//a")


def test_existing_escapes_entities_and_links_are_preserved():
    source = r"Before \*literal\* &amp; [keep label](https://example.invalid/x) after"
    text = selection_blocks(source)[0]["text"]
    start = text.index("label")
    updated = replace_selection(source, _selection(source, start=start, end=start + 5), "new")
    assert updated == r"Before \*literal\* &amp; [keep new](https://example.invalid/x) after"
    assert _root(updated).xpath(".//a/@href") == ["https://example.invalid/x"]


def test_single_list_item_and_generated_blank_lines_stay_in_that_item():
    source = "- first\n- **second** tail\n- third"
    updated = replace_selection(source, _selection(source, block=1), "one\n\ntwo")
    root = _root(updated)
    items = root.xpath("./ul/li")
    assert len(items) == 3
    assert "".join(items[0].itertext()) == "first"
    assert "".join(items[1].itertext()) == "one\n\ntwo"
    assert len(items[1].xpath(".//br")) == 2
    assert items[1].xpath("./strong")
    assert "".join(items[2].itertext()) == "third"
    decorated = html.fragment_fromstring(decorate_selection_html(updated), create_parent="div")
    assert len(decorated.xpath(".//*[@data-selection-id]")) == 3


def test_native_style_spans_keep_ids_and_text_offsets():
    source = "plain **bold**\n\nrepeat\n\nrepeat"
    styled = '<p style="font-size:12pt"><span style="color:red">plain </span><strong><span>bold</span></strong></p><p>repeat</p><p>repeat</p>'
    root = html.fragment_fromstring(decorate_selection_html(source, styled), create_parent="div")
    assert root.xpath("./p/@data-selection-id") == [item["blockId"] for item in selection_blocks(source)]
    assert root.xpath("./p[1]/@style") == ["font-size:12pt"]
    assert "".join(root.xpath("./p[1]")[0].itertext()) == "plain bold"


def test_native_preview_with_changed_text_is_not_misbound():
    decorated = decorate_selection_html("first\n\nsecond", '<p data-selection-id="forged">first</p><p>different</p>')
    assert "data-selection-id" not in decorated


def test_native_br_normalization_matches_markdown_textcontent():
    source = "first  \nsecond"
    root = html.fragment_fromstring(decorate_selection_html(source, "<p>first<br>second</p>"), create_parent="div")
    node = root.xpath("./p")[0]
    assert node.get("data-selection-id") == selection_blocks(source)[0]["blockId"]
    assert "".join(node.itertext()) == "first\nsecond"


def test_crlf_source_keeps_exact_source_offsets():
    source = "first  \r\nsecond\r\n\r\nTail"
    assert selection_blocks(source)[0]["text"] == "first\nsecond"
    assert replace_selection(source, _selection(source, start=6, end=12), "changed") == "first  \r\nchanged\r\n\r\nTail"


def test_unsupported_structures_do_not_receive_ids():
    source = "# title\n\nSafe paragraph\n\n| a | b |\n| --- | --- |\n| cell | value |\n\n![image](docx-image:x)\n\n```python\nprint('not a paragraph')\n```\n\n<span>raw html</span>\n\nAnother safe paragraph"
    root = html.fragment_fromstring(decorate_selection_html(source), create_parent="div")
    assert ["".join(node.itertext()) for node in root.xpath(".//*[@data-selection-id]")] == ["Safe paragraph", "Another safe paragraph"]
    assert not root.xpath(".//h1[@data-selection-id]|.//td[@data-selection-id]|.//pre[@data-selection-id]")


def test_a_changed_block_cannot_accept_an_old_descriptor():
    source = "old selected text"
    descriptor = _selection(source, start=4, end=12)
    with pytest.raises(SelectionError):
        replace_selection("new selected text", descriptor, "replacement")


@pytest.mark.parametrize("change", [
    {"start": True}, {"start": -1}, {"end": 100}, {"end": 0}, {"text": "forged"},
    {"blockId": "forged"}, {"unexpected": "field"},
])
def test_invalid_or_forged_descriptors_are_rejected(change):
    source = "selected text"
    descriptor = _selection(source)
    descriptor.update(change)
    with pytest.raises(SelectionError):
        replace_selection(source, descriptor, "new")


def test_inline_code_replacement_preserves_code_and_rejects_delimiter_injection():
    source = "Call `old_function()` now"
    descriptor = _selection(source, start=5, end=17)
    updated = replace_selection(source, descriptor, "new_function")
    assert _root(updated).xpath(".//code/text()") == ["new_function()"]
    with pytest.raises(SelectionError):
        replace_selection(source, descriptor, "`broken`")


def test_an_entity_that_cannot_be_proven_has_no_selection_id():
    source = "before &NotEqualTilde; after"
    assert selection_blocks(source) == []
    assert "data-selection-id" not in decorate_selection_html(source)
    with pytest.raises(SelectionError):
        replace_selection(source, {"blockId": "unproven", "start": 7, "end": 8, "text": "≂"}, "new")


def test_generated_control_characters_are_rejected():
    with pytest.raises(SelectionError, match="控制字符"):
        replace_selection("source", _selection("source"), "new\x00text")
