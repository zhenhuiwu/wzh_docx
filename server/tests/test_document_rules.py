"""The skill's chapter and table rules are enforced outside the model prompt."""

import pytest

from app.document_rules import DraftFormatError, section_rules, validate_draft, generated_body


@pytest.mark.parametrize("marker,label", [("```", "markdown"), ("```", "md"), ("~~~", "text"), ("```", "")])
def test_generated_body_unwraps_only_one_prose_envelope(marker, label):
    source = f"{marker}{label}\n正常正文。\n\n| 项目 |\n|---|\n| 数据 |\n{marker}"
    body = "正常正文。\n\n| 项目 |\n|---|\n| 数据 |"
    assert generated_body(source) == body
    for end in range(1, len(source) + 1):
        partial = generated_body(source[:end], complete=False)
        assert marker not in partial and not partial.startswith(label + "\n")
    assert validate_draft(body, {}) == body
    # Manual/API adoption still rejects fences instead of loosening the Word validator.
    with pytest.raises(DraftFormatError):
        validate_draft(source, {})


@pytest.mark.parametrize("source", ["```xml\n<p>正文</p>\n```", "<p>正文</p>", "正文\n```\n代码\n```", "```markdown\n正文", "```markdown\n<p>正文</p>\n```"])
def test_generated_body_does_not_bypass_html_inner_code_or_unclosed_fences(source):
    assert generated_body(source, complete=False) == "" or "<p>" not in generated_body(source, complete=False)
    with pytest.raises(DraftFormatError):
        generated_body(source)


def test_generated_whole_document_envelope_keeps_json_protocol():
    assert generated_body('```json\n{"edits":[]}\n```', whole_document=True) == '{"edits":[]}'
    with pytest.raises(DraftFormatError):
        generated_body('```json\n{"edits":[]}\n```')


@pytest.mark.parametrize('br', ['<br>', '<br/>', '<br />', '<BR>'])
def test_native_word_breaks_remain_valid_through_generation_and_adoption(br):
    body = f'| 要求 |\n| --- |\n| 第一行{br}第二行 |'
    assert generated_body(body) == validate_draft(body, {}) == body
    from app.documents import DocumentEdit, _validate_edits
    row = [None] * 16
    row[2], row[7], row[6] = '通用说明', '# 目标\n\n' + body, {}
    edited = body.replace('第二行', '更新行')
    assert _validate_edits(row, [DocumentEdit(headingIndex=0, content=edited)])[0].content == edited


@pytest.mark.parametrize('unsafe', ['<br onclick="run()">', '<br style="display:none">', '<script>run()</script>',
                                    '<img src="x" onerror="run()">', '<!--hidden-->', '```python\nrun()\n```'])
def test_preserved_markup_never_allows_attributes_scripts_comments_or_code(unsafe):
    with pytest.raises(DraftFormatError):
        validate_draft(unsafe, {}, existing=unsafe)


def test_word_literal_xml_names_are_not_treated_as_html():
    literal = r'接口参数为 \<request\>，保留原文。'
    assert validate_draft(generated_body(literal), {}) == literal


def test_imported_multiline_word_table_can_be_validated_and_exported(tmp_path):
    from docx import Document
    from app.docx_parser import parse_docx
    from app.docx_export import edit_original_docx
    word = Document()
    word.add_heading('目标', 1)
    table = word.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = '功能', '要求'
    table.cell(1, 0).text, table.cell(1, 1).text = '查询', '保留第一行\n超时30秒'
    original, exported = tmp_path / 'original.docx', tmp_path / 'edited.docx'
    word.save(original)
    source = parse_docx(str(original)).markdown
    body = source[source.index('|'):].strip()
    assert '<br>' in body
    result = validate_draft(generated_body(body.replace('30秒', '60秒')), {}, existing=body)
    edit_original_docx(source.replace(body, result), original, exported)
    saved = Document(exported)
    assert saved.tables[0].cell(1, 1).text == '保留第一行\n超时60秒'
    assert len(saved.tables[0].rows) == len(saved.tables[0].columns) == 2


def test_skill_rules_guard_word_body_and_tables():
    srs = section_rules(heading="3.2.1 自定义功能", document_name="软件需求规格说明")
    assert srs["skill"] == "word-fillter-438c-srs" and srs["chapter"]["id"] == "3.2"
    assert "content" not in srs["chapter"] and "generation_prompt" not in srs["chapter"]
    sdd = section_rules(heading="4.3.2 外部接口", template_name="软件设计说明", document_name="SRS")
    assert sdd["skill"] == "word-fillter-438c-sdd" and sdd["chapter"]["id"] == "4.3"
    assert not section_rules(heading="范围", document_name="通用说明")["skill"]
    refs = section_rules(heading="2 引用文档", document_name="软件需求规格说明")
    valid = "| 序号 | 名称 | 版本 | 发布日期 | 来源 |\n|---|---|---|---|---|\n|1|待补充|待补充|待补充|待补充|"
    assert validate_draft(valid, refs) == valid
    assert validate_draft("**要点**\n\n1. 已知内容\n2. 待补充", srs)
    invalid = [
        "## 新标题\n\n正文", "新标题\n====\n正文", "```xml\n正文\n```",
        '<p style="font-size:90px">正文</p>', "![图片](a.png)", "[链接](https://example.com)",
        "`代码`", "> 引用", '{"content":"正文"}', "- 第一项\n    - 嵌套项",
        valid.replace("| 来源 |", "| 新列 |"), valid + "\n| 2 | 待补充 |",
        valid.replace("|1|待补充|待补充|待补充|待补充|", "|1|待补充|待补充|待补充|待补充|多余字段|"),
    ]
    for content in invalid:
        with pytest.raises(DraftFormatError):
            validate_draft(content, refs)
    with pytest.raises(DraftFormatError):
        validate_draft("**加粗**", {}, selected_only=True)
    with pytest.raises(DraftFormatError):
        validate_draft("软件能力需求\n\n正文", srs, heading="3.2 软件能力需求")
    assert validate_draft("未完成的表格\n\n| 序号 |", refs, complete=False)
    assert validate_draft("普通改写文本。", {}, selected_only=True)
    # A type chosen in the request must remain enforced when the draft is adopted.
    from app.documents import DocumentEdit, _validate_edits
    row = [None] * 16
    row[2], row[7] = "通用说明", "# 2 引用文档\n\n正文"
    with pytest.raises(DraftFormatError):
        _validate_edits(row, [DocumentEdit(headingIndex=0, content="正文")], "按照软件需求规格说明编写")
    assert _validate_edits(row, [DocumentEdit(headingIndex=0, content=valid)], "按照软件需求规格说明编写")


def test_chapter_edit_keeps_original_and_visio_images_in_order():
    existing = "原文\n\n![原图](image_0.png)\n\n![流程图](docx-image:diagram-example)"
    edited = existing.replace("原文", "修改后的说明")
    assert validate_draft(edited, {}, preserve_images=True, existing=existing) == edited
    assert validate_draft("修改后的说明\n\n![原图](image_0.png)", {}, complete=False,
                          preserve_images=True, existing=existing)
    for invalid in ("只有正文", edited.replace("image_0.png", "https://example.com/image.png"),
                    edited.replace("![原图](image_0.png)\n\n", ""),
                    "![流程图](docx-image:diagram-example)\n\n![原图](image_0.png)"):
        with pytest.raises(DraftFormatError):
            validate_draft(invalid, {}, preserve_images=True, existing=existing)
    with pytest.raises(DraftFormatError):
        validate_draft("![流程图](docx-image:diagram-example)", {}, complete=False,
                       preserve_images=True, existing=existing)
