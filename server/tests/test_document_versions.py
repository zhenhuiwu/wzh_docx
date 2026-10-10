"""Durable native Word versions restore text, diagrams and exact accepted memory state."""

import base64
from io import BytesIO
import json
from uuid import uuid4
from zipfile import ZipFile

import pytest
from docx import Document
from fastapi.testclient import TestClient

import app.documents as documents
from app.config import clear_settings_cache
from app.db import connect, migrate
from app.document_diagrams import _diagram_path
from app.document_memory import clear_memory, load_memory, read_memory_history
from app.errors import AppError
from app.files import resolve_object, store_bytes
from app.identity import create_user
from app.main import create_app


def login(client, account="version-owner"):
    token = client.get("/api/v1/auth/csrf").json()["data"]["csrfToken"]
    response = client.post("/api/v1/auth/login", json={"account": account, "password": "version-test-password"},
                           headers={"X-CSRF-Token": token})
    assert response.status_code == 200, response.text
    return {"X-CSRF-Token": response.json()["data"]["csrfToken"]}


@pytest.fixture
def version_context(tmp_path, monkeypatch):
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path / 'versions.db'}")
    monkeypatch.setenv("FILE_DATA_DIR", str(tmp_path / "files"))
    clear_settings_cache()
    try:
        migrate()
        owner = create_user(account="version-owner", name="历史测试", password="version-test-password", roles=["system_admin"])
        create_user(account="version-other", name="其他用户", password="version-test-password", roles=["project_creator"])
        project = uuid4()
        with connect() as connection:
            connection.execute("INSERT INTO projects (id, name, owner_id) VALUES (%s, %s, %s)", (project, "中性历史项目", owner))
            connection.commit()
        word = Document()
        word.add_heading("范围", level=1)
        word.add_paragraph("原始正文。")
        data = BytesIO()
        word.save(data)
        stored = store_bytes(name="历史测试.docx", media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                             purpose="attachment", project_id=project, content=data.getvalue())
        with TestClient(create_app()) as client:
            headers = login(client)
            response = client.post(f"/api/v1/projects/{project}/documents",
                                   json={"name": "中性历史文档", "sourceType": "upload", "wordFileId": str(stored.id)}, headers=headers)
            assert response.status_code == 200, response.text
            document = response.json()["data"]
            yield client, headers, document, project, owner
    finally:
        clear_settings_cache()


def save(client, headers, document, revision, text, instruction=None):
    body = {"expectedRevision": revision, "markdownContent": "# 范围\n\n" + text}
    if instruction:
        body["memory"] = {"instruction": instruction, "scope": "section", "target": "范围"}
    result = client.patch(f"/api/v1/documents/{document}/content", json=body, headers=headers)
    assert result.status_code == 200, result.text
    return result.json()["data"]


def restore(client, headers, document, target, current):
    result = client.post(f"/api/v1/documents/{document}/versions/{target}/restore", json={"expectedRevision": current}, headers=headers)
    assert result.status_code == 200, result.text
    return result.json()["data"]


def stored_versions(document):
    with connect() as connection:
        return connection.execute("SELECT revision, memory_revisions FROM document_versions WHERE document_id=%s ORDER BY revision", (document,)).fetchall()


def test_create_and_each_edit_path_capture_durable_versions_with_pagination(version_context):
    client, headers, document, _, _ = version_context
    did = document["id"]
    first = client.get(f"/api/v1/documents/{did}/versions").json()["data"]
    assert first["items"][0]["kind"] == "created" and first["items"][0]["revision"] == 0
    save(client, headers, did, 0, "手工修改。")
    heading = client.post(f"/api/v1/documents/{did}/headings", json={"expectedRevision": 1, "afterIndex": 0, "placement": "sibling", "title": "计划"}, headers=headers)
    assert heading.status_code == 200, heading.text
    section = client.post(f"/api/v1/documents/{did}/sections/insert", json={"expectedRevision": 2, "headingIndex": 0,
                          "content": "AI 章节正文。", "mode": "replace", "instruction": "保持章节事实。"}, headers=headers)
    assert section.status_code == 200, section.text
    whole = client.post(f"/api/v1/documents/{did}/edits", json={"expectedRevision": 3, "edits": [{"headingIndex": 1, "content": "全文计划正文。"}],
                        "instruction": "统一全文术语。"}, headers=headers)
    assert whole.status_code == 200, whole.text
    pages = [client.get(f"/api/v1/documents/{did}/versions", params={"offset": offset, "limit": 2}).json()["data"] for offset in (0, 2, 4)]
    assert [item["revision"] for page in pages for item in page["items"]] == [4, 3, 2, 1, 0]
    assert [page["hasMore"] for page in pages] == [True, True, False]
    assert [row["kind"] for page in pages for row in page["items"]] == ["ai_document", "ai_section", "heading", "edit", "created"]
    original_response = client.get(f"/api/v1/documents/{did}/versions/0")
    assert original_response.status_code == 200, original_response.text
    original = original_response.json()["data"]
    assert original["markdownContent"] == document["markdownContent"] and "原始正文。" in original["html"]
    assert original["outline"] == document["outline"] and original["formatTemplate"] == document["formatTemplate"]
    memories = {row["revision"]: json.loads(row["memory_revisions"]) for row in stored_versions(did)}
    assert memories[3] == [3] and memories[4] == [3, 4]
    for query in ({"offset": -1}, {"limit": 0}, {"limit": 51}):
        assert client.get(f"/api/v1/documents/{did}/versions", params=query).status_code == 422
    assert client.get(f"/api/v1/documents/{did}/versions/999").status_code == 404
    login(client, "version-other")
    assert client.get(f"/api/v1/documents/{did}/versions").status_code == 404
    assert client.get(f"/api/v1/documents/{did}/versions/0").status_code == 404
    assert client.post(f"/api/v1/documents/{did}/versions/0/restore", json={"expectedRevision": 4}, headers=login(client, "version-other")).status_code == 404


def test_restore_reactivates_exact_historical_memory_and_never_rewinds_revision(version_context):
    client, headers, document, _, _ = version_context
    did = document["id"]
    save(client, headers, did, 0, "第一版正文。", "要求甲。")
    save(client, headers, did, 1, "第二版正文。", "要求乙。")
    undone = client.patch(f"/api/v1/documents/{did}/content", json={"expectedRevision": 2, "markdownContent": "# 范围\n\n第一版正文。",
                          "retractMemoryRevision": 2}, headers=headers)
    assert undone.status_code == 200, undone.text
    assert [item["revision"] for item in read_memory_history(did)] == [1]
    second = restore(client, headers, did, 2, 3)
    assert second["revision"] == 4 and "第二版正文。" in second["markdownContent"]
    assert {item["revision"] for item in read_memory_history(did)} == {1, 2}
    first = restore(client, headers, did, 1, 4)
    assert first["revision"] == 5 and "第一版正文。" in first["markdownContent"]
    assert [item["revision"] for item in read_memory_history(did)] == [1] and "要求乙。" not in load_memory(did)
    third = restore(client, headers, did, 4, 5)
    assert third["revision"] == 6 and "第二版正文。" in third["markdownContent"]
    assert {item["revision"] for item in read_memory_history(did)} == {1, 2}
    memories = {row["revision"]: json.loads(row["memory_revisions"]) for row in stored_versions(did)}
    assert memories == {0: [], 1: [1], 2: [1, 2], 3: [1], 4: [1, 2], 5: [1], 6: [1, 2]}
    with connect() as connection:
        clear_memory(connection, did)
        connection.commit()
    assert load_memory(did) == "" and read_memory_history(did) == []
    after_clear = restore(client, headers, did, 6, 6)
    assert after_clear["revision"] == 7 and {item["revision"] for item in read_memory_history(did)} == {1, 2}
    with connect() as connection:
        assert connection.execute("SELECT count(*) AS total FROM document_memory_events WHERE document_id=%s", (did,)).fetchone()["total"] == 2
    latest = client.get(f"/api/v1/documents/{did}/versions").json()["data"]["items"][0]
    assert latest["kind"] == "restore" and latest["restoredFromRevision"] == 6


def test_diagram_versions_restore_native_word_image_metadata_and_export(version_context):
    client, headers, document, _, _ = version_context
    did, gid = document["id"], uuid4()
    with connect() as connection:
        connection.execute("INSERT INTO document_diagrams(id,document_id,name,kind,source_revision,graph_json) VALUES(%s,%s,%s,%s,%s,%s)",
                           (gid, did, "测试图表", "flowchart", 0, "{}"))
        connection.commit()
    picture = _diagram_path(gid, "diagram.png")
    picture.parent.mkdir(parents=True)
    picture.write_bytes(base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII="))
    inserted = client.post(f"/api/v1/documents/{did}/diagrams/{gid}/insert", json={"expectedRevision": 0, "headingIndex": 0}, headers=headers)
    assert inserted.status_code == 200, inserted.text
    saved_response = client.get(f"/api/v1/documents/{did}/versions/1")
    assert saved_response.status_code == 200, saved_response.text
    saved = saved_response.json()["data"]
    assert saved["html"]
    # 终验 S-7：比较视图只取 markdown，不渲染 Word 版式
    light = client.get(f"/api/v1/documents/{did}/versions/1?html=false").json()["data"]
    assert light["html"] == "" and light["markdownContent"] == saved["markdownContent"]
    assert saved["kind"] == "diagram" and saved["formatTemplate"]["images"][0]["ref"] == f"docx-image:diagram-{gid}"
    plain = restore(client, headers, did, 0, 1)
    assert plain["revision"] == 2 and plain["formatTemplate"]["images"] == []
    returned = restore(client, headers, did, 1, 2)
    assert returned["revision"] == 3 and returned["formatTemplate"] == saved["formatTemplate"]
    export = client.post(f"/api/v1/documents/{did}/export", json={"expectedRevision": 3}, headers=headers)
    assert export.status_code == 200, export.text
    with connect() as connection:
        key = connection.execute("SELECT storage_key FROM files WHERE id=%s", (export.json()["data"]["fileId"],)).fetchone()[0]
    from app.config import get_settings
    with ZipFile(resolve_object(get_settings().file_data_dir, key)) as exported:
        assert any(name.startswith("word/media/") for name in exported.namelist())


def test_failed_or_conflicting_restore_leaves_versions_and_memory_unchanged(version_context, monkeypatch):
    client, headers, document, _, _ = version_context
    did = document["id"]
    save(client, headers, did, 0, "已保存的新正文。", "必须保留的要求。")
    before = [(row["revision"], row["memory_revisions"]) for row in stored_versions(did)], load_memory(did)
    conflict = client.post(f"/api/v1/documents/{did}/versions/0/restore", json={"expectedRevision": 0}, headers=headers)
    assert conflict.status_code == 409
    assert ([(row["revision"], row["memory_revisions"]) for row in stored_versions(did)], load_memory(did)) == before

    def blocked(*args, **kwargs):
        raise AppError(422, "AI_FORMAT_INVALID", "原稿无法安全恢复。")

    with monkeypatch.context() as patch:
        patch.setattr(documents, "_check_word", blocked)
        failed = client.post(f"/api/v1/documents/{did}/versions/0/restore", json={"expectedRevision": 1}, headers=headers)
        assert failed.status_code == 422
    assert ([(row["revision"], row["memory_revisions"]) for row in stored_versions(did)], load_memory(did)) == before

    def concurrent(*args, **kwargs):
        with connect() as writer:
            writer.execute("UPDATE documents SET revision=revision+1,markdown_content=%s WHERE id=%s", ("# 范围\n\n其他窗口保存。", did))
            writer.commit()

    with monkeypatch.context() as patch:
        patch.setattr(documents, "_check_word", concurrent)
        raced = client.post(f"/api/v1/documents/{did}/versions/0/restore", json={"expectedRevision": 1}, headers=headers)
        assert raced.status_code == 409
    assert client.get(f"/api/v1/documents/{did}").json()["data"]["markdownContent"] == "# 范围\n\n其他窗口保存。"
    assert ([(row["revision"], row["memory_revisions"]) for row in stored_versions(did)], load_memory(did)) == before


def test_existing_document_lazily_keeps_only_real_available_versions(version_context):
    client, _, _, project, owner = version_context
    did = uuid4()
    with connect() as connection:
        connection.execute("INSERT INTO documents(id,project_id,name,owner_id,revision,markdown_content) VALUES(%s,%s,%s,%s,%s,%s)",
                           (did, project, "旧项目历史", owner, 7, "# 范围\n\n现有正文。"))
        connection.commit()
    for _ in range(2):
        history = client.get(f"/api/v1/documents/{did}/versions").json()["data"]
        assert [item["revision"] for item in history["items"]] == [7] and history["items"][0]["kind"] == "baseline"
    assert client.get(f"/api/v1/documents/{did}/versions/6").status_code == 404
    with connect() as connection:
        connection.execute("DELETE FROM documents WHERE id=%s", (did,))
        connection.commit()
    assert stored_versions(did) == []


def test_restore_summary_failure_rolls_back_new_version_text_and_memory(version_context, monkeypatch):
    client, headers, document, _, _ = version_context
    did = document["id"]
    save(client, headers, did, 0, "已采纳的正文。", "已采纳的要求。")
    before_versions = [(row["revision"], row["memory_revisions"]) for row in stored_versions(did)]
    before_memory = load_memory(did)

    def fail(*args):
        raise RuntimeError("isolated restore summary failure")

    with monkeypatch.context() as patch:
        patch.setattr("app.document_memory._refresh_summary", fail)
        failed = client.post(f"/api/v1/documents/{did}/versions/0/restore", json={"expectedRevision": 1}, headers=headers)
        assert failed.status_code == 500, failed.text
    current = client.get(f"/api/v1/documents/{did}").json()["data"]
    assert current["revision"] == 1 and "已采纳的正文。" in current["markdownContent"]
    assert [(row["revision"], row["memory_revisions"]) for row in stored_versions(did)] == before_versions
    assert load_memory(did) == before_memory and [item["revision"] for item in read_memory_history(did)] == [1]


def test_explicit_restore_of_current_or_same_text_version_creates_a_new_revision(version_context):
    client, headers, document, _, _ = version_context
    did = document["id"]
    original = document["markdownContent"]
    same = restore(client, headers, did, 0, 0)
    assert same["revision"] == 1 and same["markdownContent"] == original
    assert same["formatTemplate"] == document["formatTemplate"]
    second = restore(client, headers, did, 0, 1)
    assert second["revision"] == 2 and second["markdownContent"] == original
    current = restore(client, headers, did, 2, 2)
    assert current["revision"] == 3 and current["markdownContent"] == original
    snapshots = stored_versions(did)
    assert [row["revision"] for row in snapshots] == [0, 1, 2, 3]
    assert all(json.loads(row["memory_revisions"]) == [] for row in snapshots)


def test_first_new_save_after_memory_clear_records_only_newly_accepted_requirements(version_context):
    client, headers, document, _, _ = version_context
    did = document["id"]
    save(client, headers, did, 0, "旧要求正文。", "清除前的要求。")
    with connect() as connection:
        clear_memory(connection, did)
        connection.commit()
    assert load_memory(did) == "" and read_memory_history(did) == []
    # Inspecting current history cannot mutate the saved acceptance snapshot.
    client.get(f"/api/v1/documents/{did}/versions")
    save(client, headers, did, 1, "清除后的手工正文。")
    memories = {row["revision"]: json.loads(row["memory_revisions"]) for row in stored_versions(did)}
    assert memories[1] == [1] and memories[2] == []
    save(client, headers, did, 2, "新的要求正文。", "清除后新采纳的要求。")
    memories = {row["revision"]: json.loads(row["memory_revisions"]) for row in stored_versions(did)}
    assert memories[3] == [3]
    assert [item["revision"] for item in read_memory_history(did)] == [3]
    returned = restore(client, headers, did, 1, 3)
    assert returned["revision"] == 4
    assert [item["revision"] for item in read_memory_history(did)] == [1]
    assert "清除前的要求。" in load_memory(did) and "清除后新采纳的要求。" not in load_memory(did)
