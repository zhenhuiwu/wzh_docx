"""Spec 30 §4.3: Word heading paragraphs kept as body text are detected and locked."""

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.docx_parser import parse_docx
from app.docx_preview import render_document_html, word_heading_like

SRS = next((Path(__file__).resolve().parents[2] / "438C-2021全套模板").glob("*[[]SRS[]]*.docx"))


def test_custom_style_based_on_heading_is_heading_like():
    word = Document(SRS)
    paragraph = next(item for item in word.paragraphs if item.text.strip() == "标识")
    assert paragraph.style.name.startswith("样式 标题 2") and word_heading_like(paragraph) == 2
    body = next(item for item in word.paragraphs if item.text.startswith("本条应描述本文档所适用软件的完整标识"))
    assert word_heading_like(body) is None


def test_outline_level_and_chinese_style_name(tmp_path):
    word = Document()
    outlined = word.add_paragraph("大纲级别段落")
    p_pr = outlined._p.get_or_add_pPr()
    level = OxmlElement("w:outlineLvl")
    level.set(qn("w:val"), "2")
    p_pr.append(level)
    style = word.styles.add_style("标题 3 自定义", WD_STYLE_TYPE.PARAGRAPH)
    named = word.add_paragraph("中文样式名", style=style)
    plain = word.add_paragraph("1.1 正文里写的编号不算标题")
    assert word_heading_like(outlined) == 3 and word_heading_like(named) == 3 and word_heading_like(plain) is None


def test_editing_html_marks_word_heading_paragraphs():
    markdown = parse_docx(str(SRS)).markdown
    html = render_document_html(SRS, markdown)
    assert 'data-word-heading="2"' in html
    marked = html.split('data-word-heading="2"', 1)[1][:400]
    assert "标识" in marked
