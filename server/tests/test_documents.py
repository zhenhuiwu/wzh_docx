"""Whole-document CAS saves retain the original native Word package."""

from io import BytesIO
from uuid import uuid4
from zipfile import ZipFile

from docx import Document
from docx.shared import Pt
from fastapi.testclient import TestClient
import pytest

from app.db import connect, migrate
from app.files import store_bytes
from app.identity import create_user
from app.main import create_app


@pytest.fixture
def document_context():
    migrate()
    owner = create_user(account="document-owner", name="文档测试", password="document-test-password", roles=["system_admin"])
    project, other = uuid4(), uuid4()
    with connect() as connection:
        for pid in (project, other):
            connection.execute("INSERT INTO projects(id,name,owner_id) VALUES(%s,%s,%s)", (pid, str(pid), owner))
        connection.commit()
    word = Document()
    word.sections[0].header.paragraphs[0].text = "固定页眉"
    word.sections[0].footer.paragraphs[0].text = "固定页脚"
    word.add_heading("范围", level=1)
    paragraph = word.add_paragraph("原始正文。")
    paragraph.paragraph_format.space_after = Pt(13)
    paragraph.runs[0].font.name = "Microsoft YaHei"
    paragraph.runs[0].font.size = Pt(12)
    table = word.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "项目", "结论"
    table.cell(1, 0).text, table.cell(1, 1).text = "固定测试", "通过"
    buffer = BytesIO()
    word.save(buffer)
    source = buffer.getvalue()
    stored = store_bytes(name="中性原稿.docx", media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                         purpose="attachment", project_id=project, content=source)
    with TestClient(create_app()) as client:
        token = client.get("/api/v1/auth/csrf").json()["data"]["csrfToken"]
        login = client.post("/api/v1/auth/login", json={"account": "document-owner", "password": "document-test-password"},
                            headers={"X-CSRF-Token": token})
        assert login.status_code == 200, login.text
        headers = {"X-CSRF-Token": login.json()["data"]["csrfToken"]}
        yield client, headers, project, other, stored, source


def create_document(context, name="中性文档"):
    client, headers, project, _, stored, _ = context
    response = client.post(f"/api/v1/projects/{project}/documents",
                           json={"name": name, "sourceType": "upload", "wordFileId": str(stored.id)}, headers=headers)
    assert response.status_code == 200, response.text
    return response.json()["data"]


def test_current_save_and_download_keep_native_word_properties(document_context):
    client, headers, _, _, _, source = document_context
    document = create_document(document_context)
    did = document["id"]
    assert document["revision"] == 0 and document["outline"][0]["text"] == "范围"
    markdown = document["markdownContent"].replace("原始正文。", "已保存正文。")
    saved = client.patch(f"/api/v1/documents/{did}/content", json={"expectedRevision": 0, "markdownContent": markdown}, headers=headers)
    assert saved.status_code == 200, saved.text
    assert saved.json()["data"]["revision"] == 1
    assert client.get(f"/api/v1/documents/{did}").json()["data"]["markdownContent"] == markdown
    preview = client.get(f"/api/v1/documents/{did}/preview")
    assert preview.status_code == 200 and "已保存正文。" in preview.json()["data"]["html"]
    exported = client.post(f"/api/v1/documents/{did}/export", json={"expectedRevision": 1}, headers=headers)
    assert exported.status_code == 200, exported.text
    download = client.get(f"/api/v1/files/{exported.json()['data']['fileId']}/content")
    assert download.status_code == 200, download.text
    edited = Document(BytesIO(download.content))
    assert [p.text for p in edited.paragraphs] == ["范围", "已保存正文。"]
    assert edited.paragraphs[1].paragraph_format.space_after == Pt(13)
    assert edited.paragraphs[1].runs[0].font.name == "Microsoft YaHei"
    assert edited.paragraphs[1].runs[0].font.size == Pt(12)
    assert [[cell.text for cell in row.cells] for row in edited.tables[0].rows] == [["项目", "结论"], ["固定测试", "通过"]]
    with ZipFile(BytesIO(source)) as original, ZipFile(BytesIO(download.content)) as result:
        for name in original.namelist():
            if name != "word/document.xml":
                assert result.read(name) == original.read(name), name


def test_stale_save_and_export_cannot_replace_the_current_document(document_context):
    client, headers, *_ = document_context
    document = create_document(document_context)
    did = document["id"]
    first = document["markdownContent"].replace("原始正文。", "第一窗口正文。")
    saved = client.patch(f"/api/v1/documents/{did}/content", json={"expectedRevision": 0, "markdownContent": first}, headers=headers)
    assert saved.status_code == 200, saved.text
    stale = client.patch(f"/api/v1/documents/{did}/content", json={"expectedRevision": 0,
                         "markdownContent": document["markdownContent"].replace("原始正文。", "过期窗口正文。")}, headers=headers)
    assert stale.status_code == 409 and stale.json()["error"]["code"] == "REVISION_CONFLICT"
    assert client.get(f"/api/v1/documents/{did}").json()["data"]["markdownContent"] == first
    assert client.post(f"/api/v1/documents/{did}/export", json={"expectedRevision": 0}, headers=headers).status_code == 409
    history = client.get(f"/api/v1/documents/{did}/versions").json()["data"]["items"]
    assert [version["revision"] for version in history] == [1, 0]
    assert "过期窗口正文。" not in client.get(f"/api/v1/documents/{did}/versions/1").json()["data"]["markdownContent"]


def test_current_document_creation_uses_names_and_project_scoped_sources(document_context):
    client, headers, project, other, stored, _ = document_context
    created = create_document(document_context)
    duplicate = client.post(f"/api/v1/projects/{project}/documents",
                            json={"name": created["name"], "sourceType": "upload", "wordFileId": str(stored.id)}, headers=headers)
    assert duplicate.status_code == 409 and duplicate.json()["error"]["code"] == "DUPLICATE_NAME"
    unrelated = client.post(f"/api/v1/projects/{other}/documents",
                            json={"name": "其他项目", "sourceType": "upload", "wordFileId": str(stored.id)}, headers=headers)
    assert unrelated.status_code == 404
    old_contract = client.post(f"/api/v1/projects/{project}/documents", json={"name": "旧接口文档",
                               "purpose": "测试", "templateVersionId": str(uuid4()), "ownerId": str(uuid4())}, headers=headers)
    assert old_contract.status_code == 422
    template = client.post("/api/v1/templates", json={"name": "当前模板", "category": "说明"}, headers=headers)
    assert template.status_code == 200, template.text
    template_document = client.post(f"/api/v1/templates/{template.json()['data']['id']}/documents",
                                    json={"name": "范围说明", "wordFileId": str(stored.id)}, headers=headers)
    assert template_document.status_code == 200, template_document.text
    structure = template_document.json()["data"]["importStructure"]
    structure["reviewed"] = True
    reviewed = client.patch(f"/api/v1/template-documents/{template_document.json()['data']['id']}/structure",
                            json={"expectedRevision": 0, "importStructure": structure}, headers=headers)
    assert reviewed.status_code == 200, reviewed.text
    from_template = client.post(f"/api/v1/projects/{project}/documents", json={"name": "模板文档", "sourceType": "template",
                                "templateDocumentId": template_document.json()["data"]["id"]}, headers=headers)
    assert from_template.status_code == 200, from_template.text
    assert from_template.json()["data"]["markdownContent"] == created["markdownContent"]
    # 终验 S-3：详情带模板组名，工作台模板标签显示它
    detail = client.get(f"/api/v1/documents/{from_template.json()['data']['id']}").json()["data"]
    assert detail["templateName"] == "当前模板" and detail["templateDocumentName"] == "范围说明"
