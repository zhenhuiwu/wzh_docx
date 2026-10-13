"""The SQLite setup must support the browser's main workflow."""

from io import BytesIO
from typing import ClassVar

from docx import Document
from fastapi.testclient import TestClient
from langchain_core.language_models.fake_chat_models import FakeMessagesListChatModel
from langchain_core.messages import AIMessage, ToolMessage
from langchain_core.messages import AIMessageChunk
from langchain_core.outputs import ChatGenerationChunk
import json

from app.config import clear_settings_cache
from app.db import connect, migrate
from app.identity import create_user
from app.main import create_app
from app.document_agent import _skill_files
from app.document_selection import selection_blocks


def test_sqlite_login_project_document_export(tmp_path, monkeypatch):
    skill_files = _skill_files()
    for name in ("word-fillter-438c-srs", "word-fillter-438c-sdd"):
        prefix = f"/skills/gjb438c/{name}"
        assert "name: " + name in skill_files[f"{prefix}/SKILL.md"]["content"]
        assert skill_files[f"{prefix}/references/template.md"]["content"]
        config = json.loads(skill_files[f"{prefix}/templates/config.json"]["content"])
        assert isinstance(config["content"], dict) and config["project"]["name"]["content"] == ""
        rules = json.loads(skill_files[f"{prefix}/templates/project.json"]["content"])
        assert "structure" in rules
        assert "四川大学" not in json.dumps(rules, ensure_ascii=False)
        assert all(not section.get("content") and not section.get("subsections") for section in rules["structure"].values())
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path / 'app.db'}")
    monkeypatch.setenv("FILE_DATA_DIR", str(tmp_path / "files"))
    monkeypatch.setenv("AI_BASE_URL", "")
    monkeypatch.setenv("AI_MODEL", "")
    monkeypatch.setenv("AI_API_KEY", "")
    clear_settings_cache()
    try:
        assert migrate() == [
            "100_sqlite_init.sql",
            "101_sqlite_runtime.sql",
            "102_sqlite_project_client_contact.sql",
            "103_sqlite_document_memory.sql",
            "104_test_management.sql",
            "105_document_diagrams.sql",
            "106_document_memory_target.sql",
            "107_document_memory_retractions.sql",
            "108_document_versions.sql",
            "109_writing_skills.sql",
            "110_project_member_owner_backfill.sql",
            "111_document_objects.sql",
            "112_document_changes.sql",
            "113_document_conflicts.sql",
            "114_document_objects_deleted_after.sql",
            "115_document_conflict_mine_snapshots.sql",
            "116_project_members_v2.sql",
            "117_template_versions.sql",
            "118_document_assignments.sql",
            "120_tasks_templates.sql",
        ]
        assert migrate() == []
        admin_id = create_user(
            account="admin", name="管理员", password="correct-horse-1", roles=["system_admin"]
        )
        with TestClient(create_app()) as client:
            challenge = client.get("/api/v1/auth/csrf").json()["data"]["csrfToken"]
            login = client.post(
                "/api/v1/auth/login",
                json={"account": "admin", "password": "correct-horse-1"},
                headers={"X-CSRF-Token": challenge},
            )
            assert login.status_code == 200
            csrf = login.json()["data"]["csrfToken"]
            assert client.get("/api/v1/auth/me").status_code == 200
            for index, name in enumerate(("验收项目甲", "验收项目乙")):
                result = client.post(
                    "/api/v1/projects",
                    json={
                        "name": name,
                        "ownerId": str(admin_id),
                        "status": "not_started",
                        "group": "工程一组",
                        "department": "研发部",
                        "organization": "甲方单位测试",
                        "clientContact": f"联系人{index}",
                    },
                    headers={"X-CSRF-Token": csrf, "Idempotency-Key": f"project-{index}"},
                )
                assert result.status_code == 201, result.text
                assert result.json()["data"]["clientContact"] == f"联系人{index}"
                project_id = result.json()["data"]["id"]
            page = client.get("/api/v1/projects?page=2&pageSize=1")
            assert page.status_code == 200
            assert page.json()["data"]["total"] == 2
            assert len(page.json()["data"]["items"]) == 1
            first_page = client.get("/api/v1/projects?page=1&pageSize=1").json()["data"]
            assert first_page["page"] == 1 and first_page["pageSize"] == 1
            assert first_page["items"][0]["id"] != page.json()["data"]["items"][0]["id"]
            assert client.get("/api/v1/projects?page=1&pageSize=20").json()["data"]["total"] == 2
            assert client.get("/api/v1/projects?page=0").status_code == 422
            assert client.get("/api/v1/projects?pageSize=51").status_code == 422
            assert client.get("/api/v1/projects?name=验收项目甲").json()["data"]["total"] == 1
            assert client.get("/api/v1/templates").status_code == 200
            assert client.get("/api/v1/workbench").status_code == 200

            word = Document()
            word.add_paragraph("SQLite 文档")
            data = BytesIO()
            word.save(data)
            uploaded = client.post(
                "/api/v1/files",
                files={"file": ("sample.docx", data.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                data={"purpose": "attachment", "projectId": project_id},
                headers={"X-CSRF-Token": csrf, "Idempotency-Key": "sample-docx"},
            )
            assert uploaded.status_code == 201, uploaded.text
            created = client.post(
                f"/api/v1/projects/{project_id}/documents",
                json={"name": "验收文档", "sourceType": "upload", "wordFileId": uploaded.json()["data"]["id"]},
                headers={"X-CSRF-Token": csrf},
            )
            assert created.status_code == 200, created.text
            assert "SQLite 文档" in created.json()["data"]["markdownContent"]
            assert "等待解析" not in created.json()["data"]["markdownContent"]
            document_id = created.json()["data"]["id"]
            with connect() as connection:
                storage_key = connection.execute(
                    "SELECT storage_key FROM files WHERE id = %s", [uploaded.json()["data"]["id"]]
                ).fetchone()[0]
            preview_dir = tmp_path / "files" / "previews"
            preview_dir.mkdir(parents=True)
            (preview_dir / f"{storage_key}.pdf").write_bytes(b"%PDF-1.4\n%%EOF")
            original = client.get(f"/api/v1/documents/{document_id}/original-preview")
            assert client.get(f"/api/v1/documents/{document_id}/word-preview?scope=edited&expectedRevision=1").status_code == 409
            assert client.get(f"/api/v1/documents/{document_id}/word-preview?scope=edited").status_code == 409
            edited_pdf = preview_dir / f"{document_id}.0.format4.pdf"
            edited_pdf.write_bytes(b"%PDF-1.4\n%%EOF")
            layout = client.get(f"/api/v1/documents/{document_id}/word-preview?scope=edited&expectedRevision=0")
            assert layout.status_code == 200 and layout.headers['content-type'] == 'application/pdf'
            assert 'no-cache' in layout.headers['cache-control']
            assert client.get(f"/api/v1/documents/{document_id}/word-preview?scope=invalid").status_code == 422
            assert original.status_code == 200 and original.headers["content-type"] == "application/pdf"
            root = client.post(
                f"/api/v1/documents/{document_id}/headings",
                json={"expectedRevision": 0, "title": "实施方案"},
                headers={"X-CSRF-Token": csrf},
            )
            assert root.status_code == 200, root.text
            assert root.json()["data"]["headingIndex"] == 0
            child = client.post(
                f"/api/v1/documents/{document_id}/headings",
                json={"expectedRevision": 1, "afterIndex": 0, "placement": "child", "title": "范围"},
                headers={"X-CSRF-Token": csrf},
            )
            assert child.status_code == 200, child.text
            assert "## 范围" in child.json()["data"]["markdownContent"]
            sibling = client.post(
                f"/api/v1/documents/{document_id}/headings",
                json={"expectedRevision": 2, "afterIndex": 1, "placement": "sibling", "title": "进度"},
                headers={"X-CSRF-Token": csrf},
            )
            assert sibling.status_code == 200, sibling.text
            assert sibling.json()["data"]["headingIndex"] == 2
            unavailable = client.post(
                f"/api/v1/documents/{document_id}/sections/draft",
                json={"expectedRevision": 3, "headingIndex": 1, "instruction": "写范围"},
                headers={"X-CSRF-Token": csrf},
            )
            assert unavailable.status_code == 503
            monkeypatch.setenv("AI_BASE_URL", "http://127.0.0.1:11434/v1")
            monkeypatch.setenv("AI_MODEL", "test-model")
            clear_settings_cache()
            class FakeAgentModel(FakeMessagesListChatModel):
                observed_messages: ClassVar[list] = []

                def bind_tools(self, tools, **kwargs):
                    return self

                def _generate(self, messages, **kwargs):
                    self.observed_messages.extend(messages)
                    return super()._generate(messages, **kwargs)

                async def _astream(self, messages, **kwargs):
                    response = self._generate(messages).generations[0].message
                    if response.tool_calls:
                        yield ChatGenerationChunk(message=AIMessageChunk(content="", tool_call_chunks=[
                            {"name": call["name"], "args": json.dumps(call["args"]), "id": call["id"], "index": index}
                            for index, call in enumerate(response.tool_calls)
                        ]))
                    else:
                        for character in response.content:
                            yield ChatGenerationChunk(message=AIMessageChunk(content=character))

            model = FakeAgentModel(responses=[
                AIMessage(content="", tool_calls=[
                    {"name": "read_project_overview", "args": {}, "id": "overview-1"},
                ]),
                AIMessage(content="", tool_calls=[
                    {"name": "prepare_document_edit", "args": {"objective": "编写范围正文", "sections": [1]}, "id": "prepare-section"},
                ]),
                AIMessage(content="范围正文。"),
            ])
            monkeypatch.setattr("app.document_agent.ChatOpenAI", lambda **kwargs: model)
            drafted = client.post(
                f"/api/v1/documents/{document_id}/sections/draft",
                json={"expectedRevision": 3, "headingIndex": 1, "instruction": "写范围"},
                headers={"X-CSRF-Token": csrf},
            )
            assert drafted.status_code == 200, drafted.text
            assert drafted.json()["data"]["content"] == "范围正文。"
            assert any(isinstance(message, ToolMessage) and "验收项目乙" in message.content
                       for message in model.observed_messages)
            assert any("本次按普通文档正文编写，不启用或自动推断438C技能" in str(message.content)
                       for message in model.observed_messages)
            assert any(isinstance(message, ToolMessage) and '"ready": true' in message.content.lower()
                       for message in model.observed_messages)
            assert any("文档名称：验收文档" in str(message.content) for message in model.observed_messages)
            streamed = client.post(
                f"/api/v1/documents/{document_id}/sections/draft?stream=true",
                json={"expectedRevision": 3, "headingIndex": 1},
                headers={"X-CSRF-Token": csrf},
            )
            events = [json.loads(line) for line in streamed.text.splitlines()]
            assert streamed.status_code == 200, streamed.text
            assert events[0]["type"] == "start" and events[-1]["type"] == "done"
            assert events[-1]["content"] == "范围正文。", streamed.text
            drafts = [event for event in events if event["type"] == "draft"]
            assert drafts and len(drafts[0]["content"]) < len(events[-1]["content"])
            assert events[-1]["html"] == "<p>范围正文。</p>"
            assert client.get(f"/api/v1/documents/{document_id}").json()["data"]["revision"] == 3

            async def interrupted(**kwargs):
                yield "未完成的正文"
                raise RuntimeError("secret provider details")
            with monkeypatch.context() as stream_patch:
                stream_patch.setattr("app.documents.stream_section_text", interrupted)
                failed = client.post(
                    f"/api/v1/documents/{document_id}/sections/draft?stream=true",
                    json={"expectedRevision": 3, "headingIndex": 1},
                    headers={"X-CSRF-Token": csrf},
                )
                failure_events = [json.loads(line) for line in failed.text.splitlines()]
                assert failure_events[-1]["type"] == "error"
                assert all(event["type"] != "done" for event in failure_events)
                assert "secret" not in failed.text
            inserted = client.post(
                f"/api/v1/documents/{document_id}/sections/insert",
                json={"expectedRevision": 3, "headingIndex": 1, "content": drafted.json()["data"]["content"], "instruction": "术语统一使用文档工作台。"},
                headers={"X-CSRF-Token": csrf},
            )
            assert inserted.status_code == 200, inserted.text
            assert "## 范围\n\n范围正文。\n\n## 进度" in inserted.json()["data"]["markdownContent"]
            assert "术语统一使用文档工作台。" in client.get(f"/api/v1/documents/{document_id}/memory").json()["data"]["summary"]
            styled_export = client.post(
                f"/api/v1/documents/{document_id}/export",
                json={"expectedRevision": 4},
                headers={"X-CSRF-Token": csrf},
            )
            assert styled_export.status_code == 200, styled_export.text
            styled_file = client.get(f"/api/v1/files/{styled_export.json()['data']['fileId']}/content")
            styled_doc = Document(BytesIO(styled_file.content))
            assert [(p.text, p.style.name) for p in styled_doc.paragraphs if p.text in {"实施方案", "范围"}] == [
                ("实施方案", "Heading 1"), ("范围", "Heading 2")
            ]
            conflict = client.post(
                f"/api/v1/documents/{document_id}/sections/insert",
                json={"expectedRevision": 3, "headingIndex": 1, "content": "旧草稿"},
                headers={"X-CSRF-Token": csrf},
            )
            assert conflict.status_code == 409
            saved = client.patch(
                f"/api/v1/documents/{document_id}/content",
                json={"expectedRevision": 4, "markdownContent": "# 已更新\n\n正文"},
                headers={"X-CSRF-Token": csrf},
            )
            assert saved.status_code == 200, saved.text
            exported = client.post(
                f"/api/v1/documents/{document_id}/export",
                json={"expectedRevision": saved.json()["data"]["revision"]},
                headers={"X-CSRF-Token": csrf},
            )
            assert exported.status_code == 200, exported.text
            download = client.get(f"/api/v1/files/{exported.json()['data']['fileId']}/content")
            assert download.status_code == 200 and download.content.startswith(b"PK")

            # A replacement stops before child/sibling headings; selection drafts are read-only.
            section_source = "# 主标题\n\n旧正文\n\n## 下级标题\n\n保留子节\n\n# 其他标题\n\n保留其他节"
            prepared = client.patch(
                f"/api/v1/documents/{document_id}/content",
                json={"expectedRevision": saved.json()["data"]["revision"], "markdownContent": section_source},
                headers={"X-CSRF-Token": csrf},
            ).json()["data"]
            calls = []
            def fake_draft(**kwargs):
                calls.append(kwargs)
                return "清晰的正文。"
            monkeypatch.setattr("app.documents.draft_section_text", fake_draft)
            selected_block = next(block for block in selection_blocks(section_source) if block["text"] == "旧正文")
            selected = client.post(
                f"/api/v1/documents/{document_id}/sections/draft",
                json={"expectedRevision": prepared["revision"], "scope": "selection", "selectedText": "旧正文",
                      "selection": {"blockId": selected_block["blockId"], "start": 0, "end": 3, "text": "旧正文"}},
                headers={"X-CSRF-Token": csrf},
            )
            assert selected.status_code == 200, selected.text
            assert calls[-1]["existing"] == "旧正文" and calls[-1]["selected_only"] is True
            assert calls[-1]["document_name"] == "验收文档"
            default_draft = client.post(
                f"/api/v1/documents/{document_id}/sections/draft",
                json={"expectedRevision": prepared["revision"], "headingIndex": 0},
                headers={"X-CSRF-Token": csrf},
            )
            assert default_draft.status_code == 200
            assert client.get(f"/api/v1/documents/{document_id}").json()["data"]["markdownContent"] == section_source
            # The same format guard covers JSON, streaming, and adoption (including stopped drafts).
            async def bad_stream(**kwargs):
                yield "额外标题\n====\n\n正文"
            with monkeypatch.context() as format_patch:
                format_patch.setattr("app.documents.draft_section_text", lambda **kwargs: "额外标题\n====\n\n正文")
                format_patch.setattr("app.documents.stream_section_text", bad_stream)
                for suffix in ("", "?stream=true"):
                    bad = client.post(
                        f"/api/v1/documents/{document_id}/sections/draft{suffix}",
                        json={"expectedRevision": prepared["revision"], "headingIndex": 0},
                        headers={"X-CSRF-Token": csrf},
                    )
                    if suffix:
                        rejected = [json.loads(line) for line in bad.text.splitlines()]
                        assert rejected[-1]["type"] == "error" and not any(event["type"] == "done" for event in rejected)
                    else:
                        assert bad.status_code == 502 and bad.json()["error"]["code"] == "AI_FORMAT_INVALID"
                adoption = client.post(
                    f"/api/v1/documents/{document_id}/sections/insert",
                    json={"expectedRevision": prepared["revision"], "headingIndex": 0,
                          "content": "额外标题\n====\n\n正文", "mode": "replace"},
                    headers={"X-CSRF-Token": csrf},
                )
                assert adoption.status_code == 422
            unchanged = client.get(f"/api/v1/documents/{document_id}").json()["data"]
            assert unchanged["revision"] == prepared["revision"] and unchanged["markdownContent"] == section_source
            replaced = client.post(
                f"/api/v1/documents/{document_id}/sections/insert",
                json={"expectedRevision": prepared["revision"], "headingIndex": 0, "content": "新正文", "mode": "replace"},
                headers={"X-CSRF-Token": csrf},
            )
            assert replaced.status_code == 200, replaced.text
            assert replaced.json()["data"]["markdownContent"] == section_source.replace("旧正文", "新正文")
            assert client.post(
                f"/api/v1/documents/{document_id}/sections/draft",
                json={"expectedRevision": replaced.json()["data"]["revision"]},
                headers={"X-CSRF-Token": csrf},
            ).status_code == 422
            # Whole-document edits use separate module agents, retain the outline and commit atomically.
            plan = {"edits": [{"headingIndex": 0, "content": "文档工作台的新正文。"}, {"headingIndex": 2, "content": "文档工作台的其他节。"}]}
            whole_models = [FakeAgentModel(responses=[
                AIMessage(content="", tool_calls=[{"name": "prepare_document_edit",
                    "args": {"objective": "统一术语，保留子节。", "sections": [index]}, "id": f"whole-prepare-{index}"}]),
                AIMessage(content=json.dumps({"replacements": replacements}, ensure_ascii=False)),
            ]) for index, replacements in enumerate([
                [{"before": "新正文", "after": plan["edits"][0]["content"]}], [],
                [{"before": "保留其他节", "after": plan["edits"][1]["content"]}],
            ])]
            module_models = iter(whole_models)
            from app.document_agent import draft_section_text as real_draft
            with monkeypatch.context() as whole_patch:
                whole_patch.setattr("app.document_agent.ChatOpenAI", lambda **kwargs: next(module_models))
                whole_patch.setattr("app.documents.draft_section_text", real_draft)
                whole_draft = client.post(
                    f"/api/v1/documents/{document_id}/sections/draft?stream=true",
                    json={"expectedRevision": replaced.json()["data"]["revision"], "scope": "document", "instruction": "统一术语，保留子节。"},
                    headers={"X-CSRF-Token": csrf},
                )
                whole_events = [json.loads(line) for line in whole_draft.text.splitlines()]
                assert whole_events[-1]["type"] == "done", whole_draft.text
                assert whole_events[-1]["edits"] == plan["edits"]
                assert len(whole_models) == 3 and all(model.observed_messages for model in whole_models)
                assert any("术语统一使用文档工作台。" in str(message.content)
                           for model in whole_models for message in model.observed_messages)
                assert "统一术语，保留子节。" not in client.get(f"/api/v1/documents/{document_id}/memory").json()["data"]["summary"]
                # A module cannot return whole-document JSON to edit another module.
                outside_model = FakeAgentModel(responses=[
                    AIMessage(content="", tool_calls=[{"name": "prepare_document_edit",
                        "args": {"objective": "只修改第一节", "sections": [0]}, "id": "prepare-bounds"}]),
                    AIMessage(content=json.dumps({"edits": [plan["edits"][1]]}, ensure_ascii=False)),
                    AIMessage(content=json.dumps({"edits": [plan["edits"][1]]}, ensure_ascii=False)),
                ])
                whole_patch.setattr("app.document_agent.ChatOpenAI", lambda **kwargs: outside_model)
                outside = client.post(f"/api/v1/documents/{document_id}/sections/draft",
                    json={"expectedRevision": replaced.json()["data"]["revision"], "scope": "document", "instruction": "只修改第一节"},
                    headers={"X-CSRF-Token": csrf})
                assert outside.status_code == 502 and "replacements" in outside.json()["error"]["message"]
                assert client.get(f"/api/v1/documents/{document_id}").json()["data"]["revision"] == replaced.json()["data"]["revision"]
            whole_applied = client.post(
                f"/api/v1/documents/{document_id}/edits",
                json={"expectedRevision": replaced.json()["data"]["revision"], **plan, "instruction": "统一术语，保留子节。"},
                headers={"X-CSRF-Token": csrf},
            )
            assert whole_applied.status_code == 200, whole_applied.text
            assert whole_applied.json()["data"]["markdownContent"] == section_source.replace("旧正文", "文档工作台的新正文。").replace("保留其他节", "文档工作台的其他节。")
            assert "统一术语，保留子节。" in client.get(f"/api/v1/documents/{document_id}/memory").json()["data"]["summary"]
            duplicate = client.post(f"/api/v1/documents/{document_id}/edits", json={"expectedRevision": whole_applied.json()["data"]["revision"], "edits": [plan["edits"][0], plan["edits"][0]]}, headers={"X-CSRF-Token": csrf})
            assert duplicate.status_code == 422
            stale = client.post(f"/api/v1/documents/{document_id}/edits", json={"expectedRevision": replaced.json()["data"]["revision"], **plan}, headers={"X-CSRF-Token": csrf})
            assert stale.status_code == 409
            assert client.get(f"/api/v1/documents/{document_id}").json()["data"]["revision"] == whole_applied.json()["data"]["revision"]
            # A concurrent save during the Word preflight must win without losing its content or adding stale memory.
            def concurrent_save(connection, row, content):
                with connect() as writer:
                    writer.execute("UPDATE documents SET revision=revision+1 WHERE id=%s", (document_id,))
                    writer.commit()
            with monkeypatch.context() as race_patch:
                race_patch.setattr("app.documents._check_word", concurrent_save)
                raced = client.post(f"/api/v1/documents/{document_id}/edits", json={"expectedRevision": whole_applied.json()["data"]["revision"], **plan, "instruction": "过期的要求不能写入记忆"}, headers={"X-CSRF-Token": csrf})
                assert raced.status_code == 409, raced.text
            assert "过期的要求不能写入记忆" not in client.get(f"/api/v1/documents/{document_id}/memory").json()["data"]["summary"]
            current_revision = client.get(f"/api/v1/documents/{document_id}").json()["data"]["revision"]
            long_body = "已知事实。" * 4100
            long_section = client.patch(
                f"/api/v1/documents/{document_id}/content",
                json={"expectedRevision": current_revision, "markdownContent": "# 长章节\n\n" + long_body},
                headers={"X-CSRF-Token": csrf},
            ).json()["data"]
            for suffix in ("", "?stream=true"):
                too_long = client.post(
                    f"/api/v1/documents/{document_id}/sections/draft{suffix}",
                    json={"expectedRevision": long_section["revision"], "headingIndex": 0},
                    headers={"X-CSRF-Token": csrf},
                )
                assert too_long.status_code == 422
            assert client.get(f"/api/v1/documents/{document_id}").json()["data"]["markdownContent"].endswith(long_body)
    finally:
        clear_settings_cache()
