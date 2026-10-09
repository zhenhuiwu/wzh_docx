"""Verify plan/document linkage against the real SQLite schema and HTTP routes."""

from io import BytesIO
from uuid import uuid4

from docx import Document
from fastapi.testclient import TestClient
import pytest

from app.config import clear_settings_cache
from app.db import connect, migrate
from app.files import store_bytes
from app.identity import create_user
from app.main import create_app


@pytest.fixture
def plan_context(tmp_path, monkeypatch):
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path / 'plans.db'}")
    monkeypatch.setenv("FILE_DATA_DIR", str(tmp_path / "files"))
    clear_settings_cache()
    try:
        migrate()
        owner = create_user(account="plan-owner", name="中性计划测试", password="plan-test-password", roles=["system_admin"])
        projects = [uuid4(), uuid4()]
        with connect() as connection:
            for project in projects:
                connection.execute("INSERT INTO projects(id,name,owner_id) VALUES(%s,%s,%s)",
                                   (project, str(project), owner))
            connection.commit()
        word = Document()
        word.add_heading("范围", level=1)
        word.add_paragraph("中性计划测试正文。")
        buffer = BytesIO()
        word.save(buffer)
        files = [store_bytes(name="中性计划.docx", media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                             purpose="attachment", project_id=project, content=buffer.getvalue()) for project in projects]
        with TestClient(create_app()) as client:
            csrf = client.get("/api/v1/auth/csrf").json()["data"]["csrfToken"]
            signed = client.post("/api/v1/auth/login", json={"account": "plan-owner", "password": "plan-test-password"},
                                 headers={"X-CSRF-Token": csrf})
            assert signed.status_code == 200, signed.text
            headers = {"X-CSRF-Token": signed.json()["data"]["csrfToken"]}
            yield client, headers, projects, files
    finally:
        clear_settings_cache()


def _plan(client, headers, project, name, purpose):
    response = client.post(f"/api/v1/projects/{project}/document-plans", json={"name": name, "purpose": purpose}, headers=headers)
    assert response.status_code == 201, response.text
    return response.json()["data"]


def _document(client, headers, project, file, name):
    response = client.post(f"/api/v1/projects/{project}/documents",
                           json={"name": name, "sourceType": "upload", "wordFileId": str(file.id)}, headers=headers)
    assert response.status_code == 200, response.text
    return response.json()["data"]


def test_existing_plan_links_after_document_creation_without_a_purpose_field(plan_context):
    client, headers, projects, files = plan_context
    project = projects[0]
    plan = _plan(client, headers, project, "验证报告", "归档")
    assert plan["documentId"] is None
    document = _document(client, headers, project, files[0], "验证报告")
    response = client.get(f"/api/v1/projects/{project}/document-plans")
    assert response.status_code == 200, response.text
    assert response.json()["data"]["items"] == [{**plan, "documentId": document["id"]}]
    with connect() as connection:
        stored = connection.execute("SELECT name,purpose FROM document_plans WHERE id=%s", (plan["id"],)).fetchone()
        assert stored["name"] == "验证报告" and stored["purpose"] == "归档"
        assert connection.execute("SELECT purpose FROM documents WHERE id=%s", (document["id"],)).fetchone()[0] == ""


def test_document_created_first_is_linked_when_the_plan_is_created(plan_context):
    client, headers, projects, files = plan_context
    document = _document(client, headers, projects[0], files[0], "先建文档")
    plan = _plan(client, headers, projects[0], "先建文档", "功能验证")
    assert plan["documentId"] == document["id"]


def test_matching_names_are_scoped_to_the_project_and_pagination_keeps_one_row_per_plan(plan_context):
    client, headers, projects, files = plan_context
    other_document = _document(client, headers, projects[1], files[1], "同名报告")
    current_plan = _plan(client, headers, projects[0], "同名报告", "归档")
    _plan(client, headers, projects[0], "待编写报告", "测试")
    own_document = _document(client, headers, projects[0], files[0], "同名报告")
    pages = [client.get(f"/api/v1/projects/{projects[0]}/document-plans", params={"page": page, "pageSize": 1}).json()["data"]
             for page in (1, 2, 3)]
    items = [item for page in pages for item in page["items"]]
    assert [page["total"] for page in pages] == [2, 2, 2]
    assert len(items) == 2 and len({item["id"] for item in items}) == 2
    linked = next(item for item in items if item["id"] == current_plan["id"])
    assert linked["documentId"] == own_document["id"] and linked["documentId"] != other_document["id"]
    assert next(item for item in items if item["name"] == "待编写报告")["documentId"] is None
