"""Neutral Word typography and editable-HTML contract regressions."""

from base64 import b64decode
from io import BytesIO
import json

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from lxml import etree
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from lxml import html
import pytest

from app.docx_export import edit_original_docx, _html_blocks, _inline_chunks
from app.docx_parser import parse_docx
from app.docx_preview import render_document_html


def _root(source, markdown=None, images=None):
    return html.fragment_fromstring(render_document_html(source, markdown if markdown is not None else parse_docx(str(source)).markdown, images), create_parent="div")


def _css(node):
    return dict(item.split(":", 1) for item in node.get("style", "").split(";") if ":" in item)


def _span(root, text):
    return next(node for node in root.xpath(".//span") if node.text == text)


def test_preview_supplies_native_paper_size_and_margins(tmp_path):
    source = tmp_path / 'paper.docx'
    doc = Document()
    doc.add_paragraph('原稿纸面')
    section = doc.sections[-1]
    section.page_width, section.page_height = Pt(600), Pt(800)
    section.left_margin, section.right_margin = Pt(70), Pt(80)
    section.top_margin, section.bottom_margin = Pt(60), Pt(90)
    doc.save(source)
    root = _root(source)
    assert json.loads(root[0].get('data-word-page')) == {
        'width':'450pt','minHeight':'650pt','paddingTop':'60pt','paddingBottom':'90pt',
        'paddingLeft':'70pt','paddingRight':'80pt'}
    assert root.text_content() == '原稿纸面'


def test_preview_character_indents_line_units_and_minimum_line_height(tmp_path):
    source = tmp_path / 'character-indent.docx'
    doc = Document()
    paragraph = doc.add_paragraph('正文排版检查')
    paragraph.runs[0].font.size = Pt(14)
    properties = paragraph._p.get_or_add_pPr()
    indent = OxmlElement('w:ind')
    for key, value in {'firstLine':'120','firstLineChars':'200','leftChars':'100'}.items():
        indent.set(qn('w:' + key), value)
    properties.append(indent)
    spacing = properties.get_or_add_spacing()
    for key, value in {'line':'240','lineRule':'atLeast','before':'20','beforeLines':'50'}.items():
        spacing.set(qn('w:' + key), value)
    doc.save(source)
    css = _css(_root(source)[0])
    assert css['text-indent'] == '2em' and css['margin-left'] == '1em'
    assert css['line-height'] == 'max(12pt, 1em)'
    assert css['font-size'] == '14pt'
    assert css['margin-top'] == 'calc(0.5 * max(12pt, 1em))'


def test_preview_mixed_east_asian_and_latin_fonts_in_same_run(tmp_path):
    source = tmp_path / 'mixed-fonts.docx'
    doc = Document()
    paragraph = doc.add_paragraph('中文ABC123中文')
    fonts = paragraph.runs[0]._r.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn('w:ascii'), 'Times New Roman')
    fonts.set(qn('w:eastAsia'), '宋体')
    doc.save(source)
    root = _root(source)
    assert root.text_content() == '中文ABC123中文'
    assert _css(_span(root, 'ABC123'))['font-family'] == '"Times New Roman"'
    assert all(_css(node)['font-family'] == '"宋体"' for node in root.xpath('.//span') if node.text == '中文')


def test_theme_font_wins_in_same_properties_but_direct_font_overrides_inherited_theme(tmp_path):
    from app.docx_preview import _theme
    source = tmp_path / 'font-theme.docx'
    doc = Document()
    first = doc.add_paragraph('主题字体')
    fonts = first.runs[0]._r.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn('w:eastAsia'), 'Arial')
    fonts.set(qn('w:eastAsiaTheme'), 'majorEastAsia')
    second = doc.add_paragraph('直接字体')
    second.runs[0]._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '黑体')
    doc.save(source)
    root = _root(source)
    assert _css(_span(root,'主题字体'))['font-family'] == json.dumps(_theme(doc)[1]['majorea'],ensure_ascii=False)
    assert _css(_span(root,'直接字体'))['font-family'] == '"黑体"'


def test_theme_east_asian_font_uses_document_theme_language(tmp_path):
    from app.docx_preview import _theme
    doc = Document()
    language = doc.settings.element.xpath('./w:themeFontLang')[0]
    language.set(qn('w:eastAsia'), 'ja-JP')
    doc.add_heading('主题语言验收', 1)
    source = tmp_path / 'theme-language.docx'
    doc.save(source)
    assert _theme(doc)[1]['majorea'] == 'ＭＳ ゴシック'
    assert _css(_span(_root(source), '主题语言验收'))['font-family'] == '"ＭＳ ゴシック"'


def test_preview_resolves_paragraph_and_character_style_inheritance(tmp_path):
    source = tmp_path / "inherited.docx"
    doc = Document()
    base = doc.styles.add_style("Inherited body", WD_STYLE_TYPE.PARAGRAPH)
    base.base_style = doc.styles["Normal"]
    base.font.name, base.font.size = "Arial", Pt(13)
    base.element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
    base.font.color.rgb, base.font.italic = RGBColor(18, 52, 86), True
    base.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    base.paragraph_format.left_indent, base.paragraph_format.first_line_indent = Pt(14), Pt(7)
    base.paragraph_format.space_after, base.paragraph_format.line_spacing = Pt(6), 1.5
    child = doc.styles.add_style("Inherited child", WD_STYLE_TYPE.PARAGRAPH)
    child.base_style = base
    paragraph = doc.add_paragraph(style=child)
    paragraph.add_run("继承字体")
    emphasis = doc.styles.add_style("Inherited emphasis", WD_STYLE_TYPE.CHARACTER)
    emphasis.font.bold, emphasis.font.size = True, Pt(15)
    run = paragraph.add_run("字符样式")
    run.style = emphasis
    overridden = paragraph.add_run("局部覆盖")
    overridden.italic, overridden.font.color.rgb = False, RGBColor(0, 0, 0)
    doc.sections[0].header.paragraphs[0].text = "页眉不得进入编辑正文"
    doc.sections[0].footer.paragraphs[0].text = "页脚不得进入编辑正文"
    doc.save(source)
    original_bytes = source.read_bytes()
    root = _root(source)
    assert "页眉" not in root.text_content() and "页脚" not in root.text_content()
    assert _css(root[0]).items() >= {
        "text-align": "center", "margin-left": "14pt", "text-indent": "7pt",
        "margin-top": "0pt", "margin-bottom": "6pt", "line-height": "1.5",
    }.items()
    css = _css(_span(root, "继承字体"))
    assert (css["font-family"], css["font-size"], css["color"], css["font-style"]) == ('"Arial"', "13pt", "#123456", "italic")
    assert _css(_span(root, "字符样式"))["font-weight"] == "700"
    assert _css(_span(root, "字符样式"))["font-size"] == "15pt"
    assert _css(_span(root, "局部覆盖"))["font-style"] == "normal"
    assert _css(_span(root, "局部覆盖"))["color"] == "#000000"
    assert source.read_bytes() == original_bytes


def test_preview_preserves_mixed_runs_and_only_markdown_emphasis_is_semantic(tmp_path):
    source = tmp_path / "mixed.docx"
    doc = Document()
    heading = doc.add_heading("样式对应", 1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.runs[0].font.size = Pt(20)
    paragraph = doc.add_paragraph()
    native = paragraph.add_run("原有重点")
    native.bold, native.font.size, native.font.color.rgb = True, Pt(14), RGBColor(160, 0, 0)
    ordinary = paragraph.add_run("普通正文")
    ordinary.bold, ordinary.font.size, ordinary.font.color.rgb = False, Pt(12), RGBColor(0, 0, 0)
    doc.save(source)
    original = _root(source)
    assert not original.xpath(".//strong|.//em")
    assert _css(_span(original, "原有重点"))["font-weight"] == "700"
    assert _css(_span(original, "普通正文"))["font-weight"] == "400"
    target = parse_docx(str(source)).markdown.replace("原有重点普通正文", "原有重点**更新重点**普通正文")
    root = _root(source, target)
    assert root.xpath("./h1")[0].text_content() == "样式对应"
    assert _css(root[0])["text-align"] == "center"
    assert _css(_span(root, "样式对应"))["font-size"] == "20pt"
    assert [node.text_content() for node in root.xpath(".//strong")] == ["更新重点"]
    assert root[1].text_content() == "原有重点更新重点普通正文"
    assert _css(_span(root, "普通正文"))["font-size"] == "12pt"
    assert _css(_span(root, "普通正文"))["color"] == "#000000"
    assert _css(_span(root, "普通正文"))["font-weight"] == "400"


def test_preview_projects_native_table_cell_typography_and_shading(tmp_path):
    source = tmp_path / "table.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=2, style="Table Grid")
    for cell, text in zip((cell for row in table.rows for cell in row.cells), ("项目", "内容", "状态", "原值")):
        cell.text = text
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        cell.paragraphs[0].runs[0].font.size = Pt(9)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "DDEEFF")
    table.cell(0, 0)._tc.get_or_add_tcPr().append(shading)
    doc.save(source)
    target = parse_docx(str(source)).markdown.replace("原值", "新值")
    root = _root(source, target)
    assert len(root.xpath(".//tr")) == 2
    assert root.xpath(".//td")[-1].text_content() == "新值"
    assert _css(root[0])["border-collapse"] == "collapse"
    assert _css(root.xpath(".//th")[0])["background-color"] == "#DDEEFF"
    assert _css(root.xpath(".//th")[0])["text-align"] == "right"
    assert _css(root.xpath(".//th")[0])["border-top"] == "0.5pt solid #000000"
    assert _css(_span(root, "新值"))["font-size"] == "9pt"


def test_preview_resolves_theme_color_and_first_row_table_style(tmp_path):
    source = tmp_path / "theme-table.docx"
    doc = Document()
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "FFFFFF")  # Cached fallback must not override theme.
    color.set(qn("w:themeColor"), "accent1")
    doc.styles["Normal"].element.get_or_add_rPr().append(color)
    paragraph = doc.add_paragraph()
    paragraph.add_run("主题颜色")
    paragraph.add_run("局部颜色").font.color.rgb = RGBColor(160, 0, 0)
    style = doc.styles.add_style("QA conditional table", WD_STYLE_TYPE.TABLE)
    style.base_style = doc.styles["Table Grid"]
    conditional = OxmlElement("w:tblStylePr")
    conditional.set(qn("w:type"), "firstRow")
    rpr = OxmlElement("w:rPr")
    rpr.append(OxmlElement("w:b"))
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "22")
    rpr.append(size)
    conditional.append(rpr)
    tcpr, shade = OxmlElement("w:tcPr"), OxmlElement("w:shd")
    shade.set(qn("w:fill"), "DDEEFF")
    tcpr.append(shade)
    conditional.append(tcpr)
    style.element.append(conditional)
    table = doc.add_table(rows=2, cols=1, style=style)
    table.cell(0, 0).text, table.cell(1, 0).text = "表头", "普通单元格"
    theme = next(relation.target_part.blob for relation in doc.part.rels.values() if relation.reltype.endswith("/theme"))
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    accent = etree.fromstring(theme).xpath(".//a:clrScheme/a:accent1/a:srgbClr/@val", namespaces=ns)[0]
    doc.save(source)
    root = _root(source)
    assert _css(_span(root, "主题颜色"))["color"] == "#" + accent
    assert _css(_span(root, "局部颜色"))["color"] == "#A00000"
    assert _css(root.xpath(".//th")[0])["background-color"] == "#DDEEFF"
    assert _css(_span(root, "表头"))["font-weight"] == "700"
    assert _css(_span(root, "表头"))["font-size"] == "11pt"
    assert "font-weight" not in _css(_span(root, "普通单元格"))


def test_preview_lists_keep_independent_levels_and_legacy_numbered_p_text(tmp_path):
    source = tmp_path / "lists.docx"
    doc = Document()
    paragraph = doc.add_paragraph("原编号正文")
    numbering = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    numbering.get_or_add_numId().val, numbering.get_or_add_ilvl().val = 5, 0
    doc.add_paragraph("普通正文")
    doc.save(source)
    target = parse_docx(str(source)).markdown + "\n\n1. 主项\n    - 子项\n        1. **孙项**"
    root = _root(source, target)
    assert root[0].tag == "p" and root[0].text_content() == "原编号正文"
    assert root[0].get("data-word-marker") == "1."
    assert root[1].get("data-word-marker") is None
    assert [node.text_content().strip() for node in root.xpath(".//li[not(./ul|./ol)]")] == ["孙项"]
    items = root.xpath(".//li")
    assert len(items) == 3 and len(root.xpath(".//li/ul|.//li/ol")) == 2
    assert [_css(node)["margin-left"] for node in items] == ["36pt", "18pt", "18pt"]
    assert [node.text_content() for node in root.xpath(".//strong")] == ["孙项"]


def test_preview_strips_raw_html_attributes_and_unsafe_link_schemes(tmp_path):
    source = tmp_path / "safe.docx"
    doc = Document()
    doc.add_paragraph("安全链接")
    doc.save(source)
    target = '<p onclick="alert(1)"><span style="background:url(https://invalid)">安全</span><a href="javascript:alert(1)" onmouseover="alert(2)">链接</a></p>'
    root = _root(source, target)
    assert root.text_content() == "安全链接"
    assert not root.xpath(".//*[@onclick or @onmouseover]")
    assert not root.xpath(".//a/@href")
    assert "invalid" not in html.tostring(root, encoding="unicode")


def test_preview_resolves_only_owned_images_and_preserves_occurrences(tmp_path):
    source = tmp_path / "images.docx"
    doc = Document()
    png = b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jP1kAAAAASUVORK5CYII=")
    doc.add_picture(BytesIO(png))
    doc.add_picture(BytesIO(png))
    doc.save(source)
    root = _root(source)
    images = root.xpath(".//img")
    assert len(images) == 2
    assert all(image.get("src").startswith("data:image/png;base64,") for image in images)
    assert all(image.get("data-word-image-ref") == "image_0.png" for image in images)


def test_markdown_hard_break_exports_one_word_line_break(tmp_path):
    source, output = tmp_path / "break.docx", tmp_path / "export.docx"
    doc = Document()
    doc.add_paragraph("第一行\n第二行")
    doc.save(source)
    target = "第一行  \n第二行"
    block = _html_blocks(target)[0]
    assert "".join(text for text, _, _ in _inline_chunks(block)) == "第一行\n第二行"
    edit_original_docx(target, source, output)
    assert Document(output).paragraphs[0].text == "第一行\n第二行"
    assert len(Document(output).paragraphs[0]._p.xpath(".//w:br")) == 1
    root = _root(source, target)
    assert len(root.xpath(".//br")) == 1
    assert _css(root.xpath(".//br/following-sibling::span")[0])["white-space"] == "normal"
    # Old imports stored native Word manual breaks as Markdown soft breaks.
    # Keep that text semantic while making its actual Word break visible.
    imported = _root(source)
    assert not imported.xpath(".//br")
    assert "\n" in imported.text_content()
    assert _css(next(span for span in imported.xpath(".//span") if "\n" in (span.text or "")))["white-space"] == "pre-wrap"


def test_preview_propagates_complex_edit_protection_without_touching_source(tmp_path):
    source = tmp_path / "field.docx"
    doc = Document()
    paragraph = doc.add_paragraph("原正文")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "DATE")
    paragraph._p.append(field)
    doc.save(source)
    before = source.read_bytes()
    with pytest.raises(ValueError, match="复杂 Word"):
        render_document_html(source, "新正文")
    assert source.read_bytes() == before


def _long_doc(tmp_path, text):
    source = tmp_path / "long.docx"
    doc = Document()
    doc.add_paragraph("第1段：原始内容 1。")
    paragraph = doc.add_paragraph()
    paragraph.add_run(text[: len(text) // 2]).bold = True
    paragraph.add_run(text[len(text) // 2:])
    doc.add_paragraph("第3段：原始内容 3。")
    doc.save(source)
    return source


@pytest.mark.parametrize("chars,budget", [(20_000, 0.5), (990_000, 5.0)])
def test_preview_long_paragraph_is_not_quadratic(tmp_path, chars, budget):
    """测试工程师 N7：长段落预览曾按整段 SequenceMatcher（二次方），2 万字 46.8s、99 万字拖住整个服务。
    现在先去公共前后缀、中间长时线性对齐：2 万字 ≤ 0.5s，99 万字几秒内（含解析 docx 本身）。"""
    import time
    import random
    rnd = random.Random(chars)
    text = "".join(rnd.choice("长文字甲乙丙丁戊己庚辛子丑寅卯辰巳午未申酉戌亥") for _ in range(chars))
    source = _long_doc(tmp_path, text)
    # 正文与原稿不同：中间插了字、两处改字（最坏：前后缀去不掉多少，中间很长）
    edited = "第1段：原始内容 1。\n\n" + "改" + text[:chars // 3] + "插入的字" + text[chars // 3:-5] + "尾改\n\n第3段：原始内容 3。"
    for markdown in (parse_docx(str(source)).markdown, edited, "第1段：原始内容 1。\n\n" + "长" * chars + "\n\n第3段：原始内容 3。"):
        started = time.perf_counter()
        root = html.fragment_fromstring(render_document_html(source, markdown, None), create_parent="div")
        elapsed = time.perf_counter() - started
        assert elapsed <= budget, f"{chars} 字预览 {elapsed:.2f}s"
        assert "第3段：原始内容 3。" in root.text_content()


def test_preview_alignment_keeps_run_styles_on_long_paragraph(tmp_path):
    """线性对齐仍把原稿的加粗映射到对应文字上（前半段加粗、后半段不加粗）。"""
    import random
    rnd = random.Random(3)
    text = "".join(rnd.choice("长文字甲乙丙丁戊己庚辛子丑寅卯辰巳午未申酉戌亥") for _ in range(6000))
    source = _long_doc(tmp_path, text)
    markdown = parse_docx(str(source)).markdown.replace(text[3490:3500], text[3490:3500] + "新增", 1)
    root = _root(source, markdown)
    bold = "".join(node.text or "" for node in root.xpath(".//span") if "font-weight:700" in node.get("style", ""))
    plain = "".join(node.text or "" for node in root.xpath(".//span") if "font-weight:700" not in node.get("style", ""))
    assert bold == text[:3000], (len(bold), bold[-20:])
    assert text[3000:3500] + "新增" + text[3500:] in plain
