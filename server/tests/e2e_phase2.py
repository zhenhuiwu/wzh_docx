"""阶段2 浏览器验收（S8）：1920×1080，非管理员账号。

环境变量（测试 L5 / 复测 N5）：E2E_API、E2E_WEB **必须设置**（不设直接报错，避免打到别人的服务）；
E2E_SHOTS（验收截图目录，默认系统临时目录/e2e_phase2/shots，不再默认写协作目录）、
E2E_OUT（失败截图与下载，默认系统临时目录/e2e_phase2）；
DATABASE_URL / FILE_DATA_DIR 必须指向被测服务的同一个库（脚本直接建账号）。

运行（先用 /tmp/wzh_phase2_run/live/start.sh 起服务）：
    E2E_API=http://127.0.0.1:8041/api/v1 E2E_WEB=http://127.0.0.1:5211 E2E_SHOTS=/workspace/wzh_collab/screens/phase2 \
    DATABASE_URL=sqlite:////tmp/wzh_phase2_run/live/app.db FILE_DATA_DIR=/tmp/wzh_phase2_run/live/files \
        uv run python tests/e2e_phase2.py
覆盖：成员可见项目与文档列表、非成员 404、438C 模板文档编辑保存与导出（失败提示在导出按钮旁）、“文档标识号”行可保存、
保存中/重试中/保存失败与“撤销这一段的修改”、离开确认、两账号不同段落合并、他人更新（本地无修改自动刷新 / 有修改提示条）、
同段冲突（正文与目录标记、定位、并排对比、处理后仍可进入“已处理”）、自动保存冲突提示条、负责人成员冲突入口、
离线恢复与草稿恢复条、多标签提示、复制标签页（tabId 不同、A 的未保存草稿不被 B 删除）、只读成员与版本记录。
截图写入 /workspace/wzh_collab/screens/phase2/。
"""

from __future__ import annotations

import json
import os
import re
import sys
import tempfile
import time
import uuid
from pathlib import Path

import httpx
from playwright.sync_api import Page, expect, sync_playwright

ROOT = Path(__file__).resolve().parents[2]
# 测试 L5 / 复测 N5：地址必须显式给出；输出默认写系统临时目录。
API = os.environ.get("E2E_API", "")
WEB = os.environ.get("E2E_WEB", "")
OUT = Path(os.environ.get("E2E_OUT", str(Path(tempfile.gettempdir()) / "e2e_phase2")))
SHOTS = Path(os.environ.get("E2E_SHOTS", str(OUT / "shots")))
PASSWORD = "E2e-phase2-pass-1"
RUN = uuid.uuid4().hex[:6]
results: list[tuple[str, bool, str]] = []
open_pages: list = []



def open_mode(api, doc: dict) -> dict:
    """A07 之前写的场景假定编写人员能改整篇：新建文档（默认按分工编辑）改为不限制分工（open）。"""
    view = api.call("GET", f"/documents/{doc['id']}/assignments")
    if view["assignmentMode"] != "open":
        api.call("PUT", f"/documents/{doc['id']}/assignment-mode",
                 json={"mode": "open", "expectedAssignmentsRevision": view["assignmentsRevision"]})
    return doc

def step(name: str):
    def wrap(fn):
        def run(*args, **kwargs):
            if os.environ.get("E2E_ONLY") and os.environ["E2E_ONLY"] not in name:  # 调试用：只跑名字含该词的步骤
                print(f"SKIP {name}", flush=True)
                return
            try:
                fn(*args, **kwargs)
                results.append((name, True, ""))
                print(f"PASS {name}", flush=True)
            except Exception as error:  # noqa: BLE001 - 汇总后统一失败
                results.append((name, False, f"{type(error).__name__}: {error}"))
                print(f"FAIL {name}: {error}", flush=True)
                for index, page in enumerate(open_pages):
                    if not page.is_closed():
                        OUT.mkdir(parents=True, exist_ok=True)
                        page.screenshot(path=str(OUT / f"fail-{len(results)}-{index}.png"))
                # 审查员建议 9：失败后继续跑后面的步骤，最后汇总（E2E_STOP_ON_FAIL=1 时仍中止）
                if os.environ.get("E2E_STOP_ON_FAIL") == "1":
                    raise
        return run
    return wrap


class Api:
    def __init__(self, account: str):
        self.client = httpx.Client(base_url=API, timeout=120)
        csrf = self.client.get("/auth/csrf").json()["data"]["csrfToken"]
        response = self.client.post("/auth/login", json={"account": account, "password": PASSWORD}, headers={"X-CSRF-Token": csrf})
        response.raise_for_status()
        self.csrf = response.json()["data"]["csrfToken"]

    def call(self, method: str, path: str, **kwargs):
        headers = {"X-CSRF-Token": self.csrf, **kwargs.pop("headers", {})}
        response = self.client.request(method, path, headers=headers, **kwargs)
        if response.status_code >= 400:
            raise AssertionError(f"{method} {path} → {response.status_code} {response.text[:300]}")
        return response.json()["data"] if response.content else None


def seed() -> dict:
    sys.path.insert(0, str(ROOT / "server"))
    from app.db import connect
    from app.identity import create_user
    users = {}
    with connect() as connection:
        users["admin"] = create_user(account=f"e2e-admin-{RUN}", name="验收管理员", password=PASSWORD, roles=["system_admin"], connection=connection)
        for key, name in (("zhang", "张工"), ("li", "李工"), ("wang", "王工"), ("zhao", "赵工")):
            # 与 visual 的新建用户一致：globalRoles 默认 []，只靠项目角色。
            users[key] = create_user(account=f"e2e-{key}-{RUN}", name=name, password=PASSWORD, roles=[], allow_no_roles=True, connection=connection)
        connection.commit()
    admin = Api(f"e2e-admin-{RUN}")
    project = admin.call("POST", "/projects", json={"name": f"阶段2验收项目-{RUN}", "ownerId": str(users["zhang"]), "status": "active"},
                         headers={"Idempotency-Key": f"e2e-project-{RUN}"})
    for key, role in (("li", "editor"), ("wang", "viewer")):
        revision = admin.call("GET", f"/projects/{project['id']}")["membersRevision"]  # A05 R5：成员名单独立版本号
        admin.call("POST", f"/projects/{project['id']}/members:batch",  # PUT 只改已有成员（A05 20 I-3），加入走批量接口
                   json={"userIds": [str(users[key])], "role": role, "expectedMembersRevision": revision})
    zhang = Api(f"e2e-zhang-{RUN}")
    source = next((ROOT / "438C-2021全套模板").glob("*[[]SRS[]]*.docx"))
    upload = zhang.client.post("/files", files={"file": (source.name, source.read_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                               data={"purpose": "attachment", "projectId": project["id"]},
                               headers={"X-CSRF-Token": zhang.csrf, "Idempotency-Key": f"e2e-upload-{RUN}"})
    assert upload.status_code == 201, upload.text
    document = open_mode(zhang, zhang.call("POST", f"/projects/{project['id']}/documents",
                          json={"name": "软件需求规格说明", "sourceType": "upload", "wordFileId": upload.json()["data"]["id"]}))
    # 本机（Linux）没有 Microsoft Word：带目录域的 438C 模板导出需要 Word 刷新目录，会返回 503 WORD_UNAVAILABLE。
    # 另建一份不带目录的 Word 走界面导出下载；438C 文档的导出管线（edit_original_docx）在脚本内直接校验。
    from docx import Document as WordDocument
    from io import BytesIO
    plain = WordDocument()
    plain.add_heading("概述", 1)
    plain.add_paragraph("本文档用于验证界面导出下载的完整链路。")
    buffer = BytesIO()
    plain.save(buffer)
    upload2 = zhang.client.post("/files", files={"file": ("导出验收.docx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                                data={"purpose": "attachment", "projectId": project["id"]},
                                headers={"X-CSRF-Token": zhang.csrf, "Idempotency-Key": f"e2e-upload2-{RUN}"})
    assert upload2.status_code == 201, upload2.text
    plain_doc = open_mode(zhang, zhang.call("POST", f"/projects/{project['id']}/documents",
                           json={"name": "导出验收", "sourceType": "upload", "wordFileId": upload2.json()["data"]["id"]}))
    return {"users": users, "project": project, "document": document, "plain": plain_doc, "zhang": zhang, "source": source}


def login(page: Page, account: str, landing: str | None = None):
    page.goto(f"{WEB}/login")
    page.fill("#staff-id", account)
    page.fill("#password", PASSWORD)
    page.locator("form button[type=submit]").click()
    # 30 §2 / 验收 P1-4：管理员、项目创建者与项目负责人落在「项目」，只有编写人员落在「我的编写」
    page.wait_for_url(re.compile(rf"/({landing})(\?|$)" if landing else r"/(projects|my-writing)"), timeout=30000)


def shot(page: Page, name: str):
    SHOTS.mkdir(parents=True, exist_ok=True)
    page.screenshot(path=str(SHOTS / f"{name}.png"))


def open_doc(page: Page, data: dict):
    page.goto(f"{WEB}/projects/{data['project']['id']}/documents/{data['document']['id']}")
    expect(page.locator(".paper-content")).to_be_visible(timeout=60000)
    expect(page.get_by_test_id("save-status")).to_contain_text(re.compile("保存"), timeout=30000)


def paragraph(page: Page, text: str):
    return page.locator(".paper-content p", has_text=text).first


def type_after(page: Page, text: str, addition: str):
    # 合并/撤销后编辑区会按服务器正文重绘，重绘期间不可编辑
    expect(page.locator(".paper-content")).to_have_attribute("contenteditable", "true", timeout=30000)
    target = paragraph(page, text)
    target.scroll_into_view_if_needed()
    target.click()
    page.keyboard.press("End")
    page.keyboard.type(addition)


def save_now(page: Page):
    page.keyboard.press("Control+s")


def wait_saved(page: Page, timeout: int = 30000):
    expect(page.get_by_test_id("save-status")).to_contain_text(re.compile("刚刚保存|已保存"), timeout=timeout)


def pick_paragraphs(markdown: str, count: int) -> list[str]:
    picked = []
    for block in markdown.split("\n\n"):
        block = block.strip()
        if not block or block[0] in "#|!-*>0123456789[" or "\n" in block or len(block) < 12 or any(c in block for c in "*_`[]<>\\"):
            continue
        key = block[:10]
        if any(key in item or item[:10] in block for item in picked):
            continue
        picked.append(block[:10])
        if len(picked) == count:
            break
    assert len(picked) == count, picked
    return picked


def require_targets() -> None:
    missing = [name for name, value in (("E2E_API", API), ("E2E_WEB", WEB), ("DATABASE_URL", os.environ.get("DATABASE_URL", ""))) if not value]
    if missing:
        raise SystemExit(f"请设置环境变量 {', '.join(missing)}（被测服务的接口、前端地址和同一个数据库），脚本不再使用默认值。")


def main() -> int:
    require_targets()
    data = seed()
    markdown = data["zhang"].call("GET", f"/documents/{data['document']['id']}")["markdownContent"]
    p_export, p_a, p_b, p_conflict, p_offline, p_member, p_upd, p_upd2, p_dup_a, p_dup_b, p_block, p_del = pick_paragraphs(markdown, 12)
    doc_id = data["document"]["id"]

    def server_markdown() -> str:
        return data["zhang"].call("GET", f"/documents/{doc_id}")["markdownContent"]

    def api_append(paragraph_key: str, addition: str):
        """张工通过接口改一段（模拟另一个人的保存）。"""
        record = data["zhang"].call("GET", f"/documents/{doc_id}")
        blocks = record["markdownContent"].split("\n\n")
        index = next(i for i, block in enumerate(blocks) if block.strip().startswith(paragraph_key))
        blocks[index] = blocks[index].rstrip() + addition
        data["zhang"].call("PATCH", f"/documents/{doc_id}/content",
                           json={"expectedRevision": record["revision"], "markdownContent": "\n\n".join(blocks),
                                 "allowMerge": True, "clientSaveKey": str(uuid.uuid4())})
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch()
        ctx_a = browser.new_context(viewport={"width": 1920, "height": 1080}, accept_downloads=True, locale="zh-CN")
        ctx_b = browser.new_context(viewport={"width": 1920, "height": 1080}, accept_downloads=True, locale="zh-CN")
        a, b = ctx_a.new_page(), ctx_b.new_page()
        open_pages.extend([a, b])
        login(a, f"e2e-zhang-{RUN}", landing="projects")
        login(b, f"e2e-li-{RUN}", landing="my-writing")

        @step("成员可见项目与文档列表（P0-4 / N-S5）")
        def members_visible():
            # 编写人员登录落在「我的编写」（只列分给自己的部分），项目可见性到「项目」列表核对
            b.goto(f"{WEB}/projects")
            expect(b.get_by_text(data["project"]["name"]).first).to_be_visible(timeout=30000)
            b.goto(f"{WEB}/projects/{data['project']['id']}/documents")
            expect(b.get_by_text("软件需求规格说明").first).to_be_visible(timeout=30000)
            expect(b.get_by_text("项目不存在或你无权访问")).to_have_count(0)
            shot(b, "01-成员-文档列表")
            viewer = browser.new_context(viewport={"width": 1920, "height": 1080}, locale="zh-CN").new_page()
            login(viewer, f"e2e-wang-{RUN}", landing="projects")
            viewer.goto(f"{WEB}/projects/{data['project']['id']}/documents")
            expect(viewer.get_by_text("软件需求规格说明").first).to_be_visible(timeout=30000)
            expect(viewer.get_by_role("button", name="新建文档")).to_have_count(0)
            shot(viewer, "02-只读成员-文档列表")
            expect(viewer.get_by_role("link", name=re.compile("^查看：")).first).to_be_visible()
            viewer.goto(f"{WEB}/projects/{data['project']['id']}/documents/{doc_id}")
            expect(viewer.locator(".paper-content")).to_be_visible(timeout=60000)
            expect(viewer.get_by_test_id("save-status")).to_contain_text("只读", timeout=15000)
            expect(viewer.get_by_role("button", name="添加标题")).to_be_disabled()
            shot(viewer, "14-只读成员-工作台")
            viewer.context.close()
            outsider = browser.new_context(viewport={"width": 1920, "height": 1080}, locale="zh-CN").new_page()
            login(outsider, f"e2e-zhao-{RUN}", landing="projects")
            outsider.goto(f"{WEB}/projects/{data['project']['id']}/documents/{doc_id}")
            expect(outsider.get_by_test_id("load-error")).to_contain_text("无权访问", timeout=30000)
            shot(outsider, "22-非成员-404")
            outsider.context.close()
        members_visible()

        @step("438C 模板文档编辑保存并导出")
        def edit_and_export():
            open_doc(a, data)
            type_after(a, p_export, "（张工编辑保存）")
            save_now(a)
            wait_saved(a)
            shot(a, "03-编辑保存")
            server = data["zhang"].call("GET", f"/documents/{data['document']['id']}")["markdownContent"]
            assert "（张工编辑保存）" in server
            # 438C 文档：导出管线写回原 Word（目录刷新需 Windows Word，本机跳过）
            from tempfile import TemporaryDirectory
            from docx import Document as WordDocument
            from app.docx_export import edit_original_docx
            with TemporaryDirectory() as directory:
                output = Path(directory) / "out.docx"
                edit_original_docx(server, data["source"], output)
                assert any("（张工编辑保存）" in p.text for p in WordDocument(str(output)).paragraphs)
            a.get_by_test_id("export-word").click()
            # 界面验收 S9：失败说明在导出按钮旁（危险色），文案说明是服务器缺 Word
            expect(a.get_by_test_id("export-error")).to_contain_text(re.compile("导出失败：服务器未安装 Microsoft Word"), timeout=60000)
            shot(a, "03a-导出失败-按钮旁")
            a.get_by_test_id("export-error").get_by_role("button", name="知道了").click()
            # 不带目录的文档：界面导出下载
            a.goto(f"{WEB}/projects/{data['project']['id']}/documents/{data['plain']['id']}")
            expect(a.locator(".paper-content")).to_be_visible(timeout=60000)
            type_after(a, "完整链路", "（导出前编辑）")
            save_now(a)
            wait_saved(a)
            with a.expect_download(timeout=120000) as download:
                a.get_by_test_id("export-word").click()
            path = download.value.path()
            assert any("（导出前编辑）" in p.text for p in WordDocument(str(path)).paragraphs)
            shot(a, "03b-导出下载")
        edit_and_export()

        @step("“文档标识号”行（含书签）可修改保存（验收 M1）")
        def identifier_line():
            open_doc(a, data)
            type_after(a, "文档标识号", "（E2E-001）")
            save_now(a)
            wait_saved(a)
            assert "（E2E-001）" in server_markdown()
            expect(a.get_by_test_id("save-blocked")).to_have_count(0)
        identifier_line()

        @step("保存中、保存失败说明条与“撤销这一段的修改”、离开确认")
        def save_states():
            open_doc(a, data)
            # 保存中：先扣住保存请求，截图后放行
            held = []
            a.route("**/documents/*/content", lambda route: held.append(route) if route.request.method == "PATCH" else route.continue_())
            type_after(a, p_block, "（慢保存）")
            save_now(a)
            expect(a.get_by_test_id("save-status")).to_contain_text("保存中", timeout=10000)
            shot(a, "01-保存中")
            for route in held:
                route.continue_()
            a.unroute("**/documents/*/content")
            wait_saved(a)
            # 保存失败：模拟服务器拒绝这一段（Word 受保护内容），说明条给出位置与“撤销这一段的修改”
            def reject(route):
                if route.request.method != "PATCH":
                    return route.continue_()
                # 与服务器一样，blockText 取提交正文里出问题的那一块
                sent = json.loads(route.request.post_data or "{}").get("markdownContent", "")
                block_text = next((b for b in sent.split("\n\n") if "（会被拒）" in b), "")
                route.fulfill(status=422, content_type="application/json", body=json.dumps({"error": {
                    "code": "WORD_BLOCK_PROTECTED", "message": "该段含 Word 域或目录，不能在这里修改。",
                    "details": {"location": {"headingPath": ["范围"], "ordinal": 1, "kind": "paragraph"},
                                "reason": "该段含 Word 域或目录，不能在这里修改。", "blockText": block_text}}}))
            a.route("**/documents/*/content", reject)
            type_after(a, p_block, "（会被拒）")
            save_now(a)
            banner = a.get_by_test_id("save-blocked")
            expect(banner).to_be_visible(timeout=15000)
            expect(banner.get_by_test_id("revert-block")).to_be_visible()
            shot(a, "12-保存失败-撤销这一段")
            a.unroute("**/documents/*/content")
            banner.get_by_test_id("revert-block").click()
            expect(banner).to_have_count(0, timeout=20000)
            wait_saved(a)
            assert "（会被拒）" not in server_markdown()
            # 离开确认：保存一直失败时离开，页面内对话框
            a.route("**/documents/*/content", lambda route: route.abort() if route.request.method == "PATCH" else route.continue_())
            type_after(a, p_block, "（离开前未保存）")
            a.get_by_role("button", name="返回文档列表").click()
            expect(a.get_by_role("alertdialog")).to_be_visible(timeout=20000)
            shot(a, "13-离开确认")
            a.get_by_role("button", name="留在本页").click()
            a.unroute("**/documents/*/content")
            a.evaluate("window.dispatchEvent(new Event('online'))")
            wait_saved(a, 60000)
            assert "（离开前未保存）" in server_markdown()
        save_states()

        @step("两账号改不同段落自动合并")
        def merge_two_accounts():
            open_doc(a, data)
            open_doc(b, data)  # B 打开时的版本
            type_after(a, p_a, "（张工改A段）")
            save_now(a)
            wait_saved(a)
            type_after(b, p_b, "（李工改B段）")
            save_now(b)
            wait_saved(b)
            expect(b.get_by_text(re.compile("张工 更新了 \\d+ 处内容，已合并到你的文档"))).to_be_visible(timeout=15000)
            expect(paragraph(b, "（张工改A段）")).to_be_visible(timeout=15000)
            shot(b, "04-两人合并-提示")
            server = data["zhang"].call("GET", f"/documents/{data['document']['id']}")["markdownContent"]
            assert "（张工改A段）" in server and "（李工改B段）" in server
        merge_two_accounts()

        @step("他人更新：本地无修改自动刷新，有修改时提示条（验收 M6 / 审查 S3）")
        def others_updated():
            open_doc(b, data)
            api_append(p_upd, "（张工接口更新）")
            b.evaluate("document.dispatchEvent(new Event('visibilitychange'))")
            expect(paragraph(b, "（张工接口更新）")).to_be_visible(timeout=20000)
            expect(b.get_by_text("已显示 张工 的最新修改")).to_be_visible(timeout=10000)
            expect(b.locator(".paper-content .blk-updated").first).to_be_visible(timeout=5000)
            shot(b, "17-他人更新-本地无修改-自动刷新")
            type_after(b, p_upd2, "（李工未保存）")
            api_append(p_upd, "（张工再次更新）")
            b.evaluate("document.dispatchEvent(new Event('visibilitychange'))")
            bar = b.get_by_test_id("others-updated")
            expect(bar).to_be_visible(timeout=10000)
            expect(bar).to_contain_text("张工 更新了")
            shot(b, "16-他人更新-本地有修改")
            try:  # 自动保存（停止输入 3 秒）合并后提示条会自己收起，点不到也算通过
                bar.get_by_role("button", name=re.compile("立即刷新|保存后刷新")).click(timeout=2000)
            except Exception:  # noqa: BLE001
                pass
            expect(bar).to_have_count(0, timeout=20000)
            wait_saved(b)
            expect(paragraph(b, "（张工再次更新）")).to_be_visible(timeout=15000)
            server = server_markdown()
            assert "（李工未保存）" in server and "（张工再次更新）" in server
        others_updated()

        @step("他人更新时正在打字：一个字都不丢（终验 F-M1）")
        def typing_while_others_update():
            open_doc(b, data)
            # 1) 正在输入时到来的更新：推迟自动刷新（显示提示条），已打的字照常保存
            target = paragraph(b, p_upd2)
            target.scroll_into_view_if_needed()
            target.click()
            b.keyboard.press("End")
            b.keyboard.type("（打字甲", delay=60)
            api_append(p_upd, "（张工第三次）")
            b.evaluate("document.dispatchEvent(new Event('visibilitychange'))")
            b.keyboard.type("乙丙丁戊己）", delay=150)
            wait_saved(b, 40000)
            expect(paragraph(b, "（打字甲乙丙丁戊己）")).to_be_visible(timeout=15000)
            expect(paragraph(b, "（张工第三次）")).to_be_visible(timeout=40000)
            server = server_markdown()
            assert "（打字甲乙丙丁戊己）" in server and "（张工第三次）" in server, server[-400:]
            # 2) 本地干净、自动刷新刚开始时马上输入：放弃这次刷新或刷新后再输入，字都在
            wait_saved(b)
            b.wait_for_timeout(2500)  # 超过“最近 2 秒在输入”的窗口，让自动刷新真的开始
            api_append(p_upd, "（张工第四次）")
            b.evaluate("document.dispatchEvent(new Event('visibilitychange'))")
            b.wait_for_timeout(300)
            target = paragraph(b, p_upd2)
            target.click()
            b.keyboard.press("End")
            b.keyboard.type("（紧接着输入）", delay=80)
            wait_saved(b, 40000)
            expect(paragraph(b, "（紧接着输入）")).to_be_visible(timeout=15000)
            server = server_markdown()
            assert "（紧接着输入）" in server and "（张工第四次）" in server, server[-400:]
            # 3) 自动保存合并了他人修改、载入重绘期间继续打字：有“正在合并…”提示（aria-busy），
            #    编辑区不锁、不重绘，字留在编辑区；下次保存按原基准再合并，停手后再显示他人修改
            # 审查员 M4：F-M1 后合并重绘先在后台取版式、不锁编辑区，载入提示条 350ms 后才出现，快的重绘看不到任何提示。
            # 把后台取版式（/preview）拖慢 4 秒，确认重绘已在途（请求已发出）再继续打字——真正走到“重绘期间输入”的路径
            b.evaluate("""() => { const original = window.__origFetch || window.fetch; window.__origFetch = original; window.__previewDelayed = 0;
                window.fetch = (input, init) => { const url = String(input instanceof Request ? input.url : input);
                    if (!url.includes('/preview')) return original(input, init);
                    window.__previewDelayed += 1; return new Promise(resolve => setTimeout(resolve, 4000)).then(() => original(input, init)); }; }""")
            api_append(p_upd, "（张工第五次）")
            target = paragraph(b, p_upd2)
            target.click()
            b.keyboard.press("End")
            b.keyboard.type("（锁前）", delay=40)
            # 停止输入 3 秒触发自动保存，服务器合并张工的修改后重绘编辑区；重绘一开始就继续打字
            b.wait_for_function("window.__previewDelayed > 0", timeout=30000, polling=20)
            b.keyboard.type("（锁定中输入）", delay=60)
            b.evaluate("() => { window.fetch = window.__origFetch; }")
            wait_saved(b, 40000)
            b.wait_for_timeout(3500)
            wait_saved(b, 40000)
            server = server_markdown()
            def around(word: str) -> str:
                at = server.find(word)
                return server[max(0, at - 120):at + 160] if at >= 0 else f"缺少 {word}"
            assert "（锁前）" in server and "（张工第五次）" in server, (around("（锁前）"), around("（张工第五次）"), around(p_upd2[:8]))
            assert "锁定中输入" in server, around("（锁前）")
            expect(b.locator(".paper-content")).to_contain_text("（张工第五次）", timeout=20000)
            expect(b.locator(".paper-content")).to_contain_text("（锁定中输入）")
            shot(b, "28-他人更新时打字不丢字")
        typing_while_others_update()

        @step("重复段落（两个“无。”）：重绘时写回不落进另一个同文段落（终审必改）")
        def duplicate_paragraph_carry():
            def append_probe(tag: str):
                record = data["zhang"].call("GET", f"/documents/{doc_id}")
                tail = f"\n\n探针{tag}甲段落。\n\n无。\n\n探针{tag}乙段落。\n\n无。\n\n探针{tag}丙段落。\n"
                data["zhang"].call("PATCH", f"/documents/{doc_id}/content", json={
                    "expectedRevision": record["revision"], "markdownContent": record["markdownContent"].rstrip("\n") + tail,
                    "allowMerge": True, "clientSaveKey": str(uuid.uuid4())})

            def insert_before(tag: str):
                record = data["zhang"].call("GET", f"/documents/{doc_id}")
                md = record["markdownContent"].replace(f"探针{tag}甲段落。", f"张工在{tag}甲前插入的一段。\n\n探针{tag}甲段落。")
                data["zhang"].call("PATCH", f"/documents/{doc_id}/content", json={
                    "expectedRevision": record["revision"], "markdownContent": md, "allowMerge": True, "clientSaveKey": str(uuid.uuid4())})

            # 预览请求在页面里改慢（mode=delay）或先失败一次再慢（mode=fail-then-delay，走锁定兜底路径）
            slow_preview = """(mode) => { const original = window.__origFetch || window.fetch; window.__origFetch = original; let first = true;
                window.fetch = (input, init) => {
                    const url = String(input instanceof Request ? input.url : input);
                    if (!url.includes('/preview')) return original(input, init);
                    if (mode === 'fail-then-delay' && first) { first = false; return Promise.reject(new TypeError('probe')); }
                    window.__previewDelayed = (window.__previewDelayed || 0) + 1;
                    return new Promise(resolve => setTimeout(resolve, 4000)).then(() => original(input, init)); }; }"""
            ctx = browser.new_context(viewport={"width": 1920, "height": 1080}, locale="zh-CN")
            page = ctx.new_page()
            login(page, f"e2e-li-{RUN}", landing="my-writing")

            def second_none(tag: str):
                return paragraph(page, f"探针{tag}乙段落").locator("xpath=following-sibling::p[1]")

            # 1) 不锁的写回路径（审查员探针）：重绘载入期间在乙后的“无。”里打字
            append_probe("A")
            open_doc(page, data)
            type_after(page, "探针A丙段落", "（李工改丙A）")
            insert_before("A")
            page.evaluate(slow_preview, "delay")
            # 审查员 M4：确认重绘真的在进行（后台取版式的请求已发出、被拖慢 4 秒），再在乙后的“无。”里打字——走到重绘期间写回的路径
            # （方案 C 后不锁的后台重绘不显示载入提示条，不能再用 editor-refreshing 判断）
            page.wait_for_function("(window.__previewDelayed || 0) > 0", timeout=30000, polling=20)
            target = second_none("A")
            target.click()
            page.keyboard.press("End")
            page.keyboard.type("（写在乙A后的无）", delay=40)
            page.wait_for_timeout(6000)
            wait_saved(page, 40000)
            page.wait_for_timeout(4000)
            wait_saved(page, 40000)
            page.evaluate("() => { window.fetch = window.__origFetch; }")
            server = server_markdown()
            seg = server[server.find("张工在A甲前"):]
            assert "探针A乙段落。\n\n无。（写在乙A后的无）" in seg, repr(seg[:300])
            assert "探针A甲段落。\n\n无。\n\n" in seg, repr(seg[:300])
            expect(page.locator(".paper-content")).to_contain_text("张工在A甲前插入的一段。", timeout=20000)
            # 2) 锁定重绘的兜底路径：后台取版式失败一次 → 锁定重绘；光标在乙后的“无。”，锁定期间敲的键解锁后写回
            append_probe("B")
            page.reload()
            expect(page.locator(".paper-content")).to_be_visible(timeout=60000)
            wait_saved(page)
            type_after(page, "探针B丙段落", "（李工改丙B）")
            target = second_none("B")
            target.click()
            page.keyboard.press("End")
            insert_before("B")
            page.evaluate(slow_preview, "fail-then-delay")
            expect(page.locator(".paper-content")).to_have_attribute("aria-busy", "true", timeout=20000)
            expect(page.locator(".paper-content.refreshing")).to_have_count(1, timeout=20000)
            page.keyboard.type("LOCKB", delay=40)
            expect(page.locator(".paper-content.refreshing")).to_have_count(0, timeout=30000)
            page.wait_for_timeout(4000)
            wait_saved(page, 40000)
            page.evaluate("() => { window.fetch = window.__origFetch; }")
            server = server_markdown()
            seg = server[server.find("张工在B甲前"):]
            assert "探针B乙段落。\n\n无。LOCKB" in seg, repr(seg[:300])
            assert "探针B甲段落。\n\n无。\n\n" in seg, repr(seg[:300])
            shot(page, "29-重复段落写回不串段")
            ctx.close()
        duplicate_paragraph_carry()

        @step("同段冲突保留草稿，面板可处理")
        def same_paragraph_conflict():
            open_doc(a, data)
            open_doc(b, data)
            type_after(a, p_conflict, "（张工的写法）")
            save_now(a)
            wait_saved(a)
            type_after(b, p_conflict, "（李工的写法）")
            save_now(b)
            panel = b.locator(".conflict-panel")
            expect(panel).to_be_visible(timeout=20000)
            expect(panel.locator(".ccard")).to_have_count(1, timeout=20000)
            expect(panel.locator(".ccard-side.mine")).to_contain_text("（李工的写法）")
            expect(panel.locator(".ccard-side.theirs")).to_contain_text("（张工的写法）")
            expect(b.get_by_test_id("conflict-toggle")).to_contain_text("冲突 1")
            # 方案 C：冲突部分不写入正文——编辑区里这一段冻结（红框）、仍显示我的版本，正文是对方的版本
            expect(b.locator(".paper-content [data-frozen]", has_text="（李工的写法）")).to_have_count(1, timeout=15000)
            wait_saved(b)
            server = data["zhang"].call("GET", f"/documents/{data['document']['id']}")["markdownContent"]
            assert "（张工的写法）" in server and "（李工的写法）" not in server, "冲突段不该写入正文"
            # 正文冲突标记与目录红点（验收 M5）
            mark = b.get_by_test_id("conflict-mark")
            expect(mark).to_have_count(1, timeout=15000)
            expect(mark).to_contain_text("与 张工 的修改冲突")
            expect(b.locator(".paper-content .blk-conflict")).to_have_count(1)
            expect(b.locator(".outline-dot.conflict")).to_have_count(1)
            panel.get_by_test_id("conflict-locate").click()
            shot(b, "05-同段冲突-面板与正文标记")
            # 卡片逐词差异（验收 M4）
            expect(panel.locator(".ccard-side.mine ins")).to_contain_text("（李工的写法）")
            # 并排对比弹窗
            panel.get_by_test_id("conflict-compare").click()
            dialog = b.get_by_test_id("compare-dialog")
            expect(dialog).to_be_visible(timeout=10000)
            expect(b.get_by_test_id("compare-merged")).to_contain_text("（李工的写法）")
            shot(b, "08-并排对比")
            dialog.get_by_role("button", name="取消").click()
            expect(dialog).to_have_count(0)
            # 处理中：只有被点的按钮显示“正在处理…”（审查 S6）
            held = []
            b.route("**/conflicts/*/resolve", lambda route: held.append(route))
            panel.get_by_role("button", name=re.compile("保留我的")).click()
            expect(panel.get_by_role("button", name="正在处理…")).to_have_count(1, timeout=5000)
            shot(b, "11-冲突处理中")
            # 终验回归 I-4：处理结果返回后大文档重绘较慢，期间接着在别的段落打字，字不丢
            other = paragraph(b, p_offline)
            other.scroll_into_view_if_needed()
            other.click()
            b.keyboard.press("End")
            for route in held:
                route.continue_()
            b.unroute("**/conflicts/*/resolve")
            b.keyboard.type("（处理冲突后马上打字）", delay=80)
            expect(panel.locator(".ccard")).to_have_count(0, timeout=20000)
            b.wait_for_timeout(3500)
            wait_saved(b, 40000)
            server = data["zhang"].call("GET", f"/documents/{data['document']['id']}")["markdownContent"]
            assert "（处理冲突后马上打字）" in server, "处理冲突后重绘期间打的字丢了"
            # 全部处理后仍能进入“已处理”（验收 M2）
            toggle = b.get_by_test_id("conflict-toggle")
            expect(toggle).to_have_text("冲突", timeout=20000)
            expect(b.get_by_test_id("conflict-mark")).to_have_count(0, timeout=15000)
            expect(panel.locator(".conflict-handled[open]")).to_be_visible(timeout=10000)
            server = data["zhang"].call("GET", f"/documents/{data['document']['id']}")["markdownContent"]
            assert "（李工的写法）" in server, "保留我的版本后正文应为李工的写法"
            expect(paragraph(b, "（李工的写法）")).to_be_visible(timeout=15000)
            shot(b, "06-同段冲突-处理后-已处理")
        same_paragraph_conflict()

        @step("自动保存冲突提示条与负责人成员冲突入口（审查 S2 / 验收 M3）")
        def member_conflict():
            open_doc(b, data)
            type_after(b, p_member, "（李工自动保存）")
            api_append(p_member, "（张工抢先）")
            bar = b.get_by_test_id("auto-conflict")
            expect(bar).to_be_visible(timeout=30000)  # 停止输入 3 秒后自动保存，遇冲突只提示
            # 方案 C 文案（14 §3.4）：冲突段冻结，可以继续写，续写存为冲突副本
            expect(bar).to_contain_text("这一段和 张工 的修改冲突。")
            expect(bar).to_contain_text("不会丢失")
            expect(b.locator(".conflict-panel")).to_have_count(0)
            shot(b, "06-自动保存冲突-提示条")
            open_doc(a, data)
            toggle = a.get_by_test_id("conflict-toggle")
            expect(toggle).to_have_text("成员冲突 1", timeout=20000)
            toggle.click()
            panel = a.locator(".conflict-panel")
            expect(panel.locator(".ccard")).to_have_count(1, timeout=20000)
            expect(panel.get_by_role("tab", name=re.compile("成员冲突"))).to_have_attribute("aria-selected", "true")
            expect(panel).to_contain_text("提交人：李工")
            expect(panel.get_by_role("button", name="下载提交人的副本")).to_be_visible()
            shot(a, "18-负责人-成员冲突")
            # 李工处理掉，后续步骤不受影响
            b.get_by_test_id("auto-conflict").get_by_role("button", name="查看冲突").click()
            b.locator(".conflict-panel").get_by_role("button", name=re.compile("使用对方的版本")).click()
            b.locator(".conflict-panel").locator(".ccard-confirm").get_by_role("button", name="使用对方的版本").click()
            expect(b.locator(".conflict-panel .ccard")).to_have_count(0, timeout=20000)
        member_conflict()

        @step("离线编辑本机保留，联网后自动保存；关闭后重开可恢复")
        def offline_recovery():
            open_doc(b, data)
            ctx_b.set_offline(True)
            type_after(b, p_offline, "（离线写的一句）")
            save_now(b)
            expect(b.get_by_test_id("save-status")).to_contain_text(re.compile("离线|重试|未保存"), timeout=30000)
            shot(b, "07-离线-本机保留")
            ctx_b.set_offline(False)
            b.evaluate("window.dispatchEvent(new Event('online'))")
            wait_saved(b, 60000)
            server = data["zhang"].call("GET", f"/documents/{data['document']['id']}")["markdownContent"]
            assert "（离线写的一句）" in server
            # 关闭页面前断网：草稿留在本机，重新打开后“恢复并保存”
            ctx_b.set_offline(True)
            type_after(b, p_offline, "（断网关页）")
            b.wait_for_timeout(500)
            b.close()
            ctx_b.set_offline(False)
            page = ctx_b.new_page()
            open_doc(page, data)
            bar = page.get_by_test_id("draft-recovery")
            expect(bar).to_be_visible(timeout=20000)
            shot(page, "08-重开-草稿恢复条")
            bar.get_by_role("button", name="恢复并保存").click()
            expect(bar).to_have_count(0, timeout=30000)
            wait_saved(page)
            server = data["zhang"].call("GET", f"/documents/{data['document']['id']}")["markdownContent"]
            assert "（断网关页）" in server
            shot(page, "09-草稿恢复后")
        offline_recovery()

        @step("会话过期→重新登录→回到文档：继续编辑前出现恢复条（终验回归 H1）")
        def session_expired_recovery():
            ctx = browser.new_context(viewport={"width": 1920, "height": 1080}, locale="zh-CN")
            page = ctx.new_page()
            login(page, f"e2e-li-{RUN}", landing="my-writing")
            open_doc(page, data)
            type_after(page, p_offline, "（会话过期前未保存）")
            page.wait_for_timeout(800)
            ctx.clear_cookies()  # 会话失效
            save_now(page)
            page.wait_for_url(re.compile(r"/login"), timeout=30000)
            assert page.evaluate("Object.values(localStorage).some(v => v.includes('会话过期前未保存'))"), "草稿不在本机"
            page.fill("#staff-id", f"e2e-li-{RUN}")
            page.fill("#password", PASSWORD)
            page.locator("form button[type=submit]").click()
            page.wait_for_url(re.compile(r"/documents/"), timeout=30000)  # SPA 回到文档，不整页刷新
            expect(page.locator(".paper-content")).to_be_visible(timeout=60000)
            # 恢复条出现前编辑区不可输入；立刻打字也不会覆盖本标签的旧草稿
            page.locator(".paper-content").click()
            page.keyboard.type("（回来立刻打字）")
            bar = page.get_by_test_id("draft-recovery")
            expect(bar).to_be_visible(timeout=20000)
            expect(page.locator(".paper-content")).to_have_attribute("contenteditable", "false")
            shot(page, "09b-会话过期重新登录-草稿恢复条")
            bar.get_by_role("button", name="恢复并保存").click()
            expect(bar).to_have_count(0, timeout=30000)
            wait_saved(page)
            server = data["zhang"].call("GET", f"/documents/{data['document']['id']}")["markdownContent"]
            assert "（会话过期前未保存）" in server, "恢复后服务器没有过期前的内容"
            # 之后继续编辑：两段内容都在，刷新后不再有恢复条
            type_after(page, p_offline, "（回来后新输入）")
            save_now(page)
            wait_saved(page)
            server = data["zhang"].call("GET", f"/documents/{data['document']['id']}")["markdownContent"]
            assert "（会话过期前未保存）" in server and "（回来后新输入）" in server
            page.reload()
            expect(page.locator(".paper-content")).to_be_visible(timeout=60000)
            page.wait_for_timeout(2000)
            expect(bar).to_have_count(0)
            ctx.close()
        session_expired_recovery()

        @step("草稿边角：本标签旧草稿改存；恢复的一份没存上前不删草稿、不恢复下一份（最后一审建议 3）")
        def draft_corner_cases():
            record = data["zhang"].call("GET", f"/documents/{doc_id}")
            tail = "\n\n草稿探针甲段落。\n\n草稿探针乙段落。\n\n草稿探针丙段落。\n"
            data["zhang"].call("PATCH", f"/documents/{doc_id}/content", json={
                "expectedRevision": record["revision"], "markdownContent": record["markdownContent"].rstrip("\n") + tail,
                "allowMerge": True, "clientSaveKey": str(uuid.uuid4())})
            fail_saves = lambda route: route.abort() if route.request.method == "PATCH" else route.continue_()  # noqa: E731
            ctx = browser.new_context(viewport={"width": 1920, "height": 1080}, locale="zh-CN")
            page = ctx.new_page()
            open_pages.append(page)
            login(page, f"e2e-li-{RUN}", landing="my-writing")
            # 1) 本标签（page）写下草稿 O：保存一直失败
            open_doc(page, data)
            page.route("**/documents/*/content", fail_saves)
            type_after(page, "草稿探针乙段落", "（O本标签）")
            draft_has = "text => Object.entries(localStorage).some(([k, v]) => k.startsWith('doc-workbench.draft.v2:') && v.includes(text))"
            page.wait_for_function(draft_has, arg="（O本标签）", timeout=20000)
            # 2) 另一个标签写下草稿 X 后关闭
            other = ctx.new_page()
            open_doc(other, data)
            other.route("**/documents/*/content", fail_saves)
            type_after(other, "草稿探针甲段落", "（X别的标签）")
            other.wait_for_function(draft_has, arg="（X别的标签）", timeout=20000)
            # 关页时的 keepalive 保存不受 route/断网模拟拦截、会落库；这里让它失败，模拟“没来得及存上就关了”
            other.evaluate("""() => { const original = window.fetch;
                window.fetch = (url, init) => init && init.keepalive ? Promise.reject(new TypeError('offline')) : original(url, init); }""")
            other.close()
            page.wait_for_timeout(1000)
            # 3) 本标签离开文档再回来（同一页面），两份草稿都列出
            page.get_by_role("button", name="返回文档列表").click()
            expect(page.get_by_role("alertdialog")).to_be_visible(timeout=20000)
            page.get_by_test_id("leave-confirm").click()
            expect(page.locator(".paper-content")).to_have_count(0, timeout=30000)
            page.go_back()
            expect(page.locator(".paper-content")).to_be_visible(timeout=60000)
            bar = page.get_by_test_id("draft-recovery")
            expect(bar).to_be_visible(timeout=20000)
            pick = bar.locator("select")
            expect(pick.locator("option")).to_have_count(2, timeout=10000)
            stored = lambda: page.evaluate("JSON.stringify(Object.fromEntries(Object.entries(localStorage).filter(([k]) => k.startsWith('doc-workbench.draft.v2:'))))")  # noqa: E731
            # 选中 X（看差异里有 X 的那份）
            bar.get_by_role("button", name="查看差异").click()
            if "（X别的标签）" not in page.locator(".draft-recovery").inner_text():
                pick.select_option(index=1 - int(pick.input_value()))
            expect(page.locator(".draft-recovery")).to_contain_text("（X别的标签）", timeout=10000)
            # 4) 恢复 X（保存失败，进入重试）：本标签的旧草稿 O 改存到新键；X 没存上之前草稿不删、
            #    恢复条停在“恢复中…”（不能切换/恢复下一份/舍弃），编辑区保持锁定
            bar.get_by_role("button", name="恢复并保存").click()
            page.wait_for_timeout(3000)  # 经过至少一次自动重试
            # 恢复条仍是 X 这份（差异里是 X 的修改），没有换成下一份 O
            expect(page.locator(".draft-diff")).to_contain_text("（X别的标签）", timeout=1000)
            drafts = stored()
            assert "-kept-" in drafts and "（O本标签）" in drafts, "本标签的旧草稿 O 没有改存"
            assert "（X别的标签）" not in server_markdown()
            expect(page.locator(".wb-notice.info")).to_contain_text("还没存上服务器", timeout=20000)
            expect(bar.get_by_role("button", name="恢复中…")).to_be_disabled()
            expect(pick).to_be_disabled()
            expect(bar.get_by_role("button", name="舍弃")).to_be_disabled()
            expect(page.locator(".paper-content")).to_have_attribute("contenteditable", "false")
            # 5) 网络恢复：X 存上服务器后才显示在编辑区、删掉 X 草稿，恢复条换成 O
            page.unroute("**/documents/*/content")
            page.evaluate("window.dispatchEvent(new Event('online'))")
            expect(page.locator(".paper-content")).to_contain_text("（X别的标签）", timeout=60000)
            assert "（X别的标签）" in server_markdown()
            expect(pick).to_have_count(0, timeout=20000)  # 只剩 O
            expect(bar.get_by_role("button", name="恢复并保存")).to_be_enabled()
            assert "（X别的标签）" not in stored() and "（O本标签）" in stored()
            # 6) 再恢复 O：两份都在服务器上
            bar.get_by_role("button", name="恢复并保存").click()
            expect(bar).to_have_count(0, timeout=30000)
            wait_saved(page, 30000)
            server = server_markdown()
            assert "（X别的标签）" in server and "（O本标签）" in server, server[-300:]
            shot(page, "30-两份草稿依次恢复")
            ctx.close()
        draft_corner_cases()

        @step("冲突未处理时拒绝恢复下一份草稿并提示，前一份的冲突内容不丢（测试工程师 F-H4）")
        def draft_refused_while_conflict():
            from e2e_selfconf import make_doc
            ctx = browser.new_context(viewport={"width": 1920, "height": 1080}, locale="zh-CN")
            page = ctx.new_page()
            open_pages.append(page)
            login(page, f"e2e-li-{RUN}", landing="my-writing")
            did = make_doc(data, f"草稿冲突-{RUN}", 0)
            url = f"{WEB}/projects/{data['project']['id']}/documents/{did}"
            # 先写出一份真实草稿当模板（断网输入），取得账号与格式
            page.goto(url)
            expect(page.locator(".paper-content")).to_be_visible(timeout=60000)
            wait_saved(page, 60000)
            ctx.set_offline(True)
            type_after(page, "第4段", "（模板）")
            key = "doc-workbench.draft.v2:"
            page.wait_for_function(f"() => Object.keys(localStorage).some(k => k.startsWith('{key}') && k.includes('{did}'))", timeout=20000)
            template = page.evaluate(f"""() => {{ const k = Object.keys(localStorage).find(k => k.startsWith('{key}') && k.includes('{did}'));
                const v = JSON.parse(localStorage.getItem(k)); return [k, v]; }}""")
            # 离开时的 keepalive 保存不受断网模拟拦截：让它失败，模板内容不上服务器
            page.evaluate("""() => { const original = window.fetch;
                window.fetch = (url, init) => init && init.keepalive ? Promise.reject(new TypeError('offline')) : original(url, init); }""")
            ctx.set_offline(False)
            page.goto(f"{WEB}/projects")
            page.evaluate("localStorage.clear()")
            account = template[0][len(key):].split(":")[0]
            record = data["zhang"].call("GET", f"/documents/{did}")
            base_md, rev = record["markdownContent"], record["revision"]
            # 两份草稿都改第 2 段；之后张工也改第 2 段
            data["zhang"].call("PATCH", f"/documents/{did}/content", json={"expectedRevision": rev, "allowMerge": True, "clientSaveKey": str(uuid.uuid4()),
                               "markdownContent": base_md.replace("第2段：原始内容 2。", "第2段：原始内容 2（对方改）。")})
            now = int(time.time() * 1000)
            drafts = []
            for n, (tag, age) in enumerate((("【A稿】", 60000), ("【B稿】", 30000))):
                draft = dict(template[1], docId=did, tabId=f"e2e-tab-{n}-{RUN}", docRevision=rev, mergeBase=rev, savedAt=now - age,
                             markdown=base_md.replace("第2段：原始内容 2。", f"第2段：原始内容 2。{tag}"), pendingSave=None)
                draft.pop("baseMarkdown", None)
                drafts.append(draft)
            page.evaluate("""([account, did, drafts]) => { for (const d of drafts) localStorage.setItem(`doc-workbench.draft.v2:${account}:${did}:${d.tabId}`, JSON.stringify(d));
                localStorage.setItem(`doc-workbench.draft.v2.index:${account}:${did}`, JSON.stringify(drafts.map(d => ({ tabId: d.tabId, savedAt: d.savedAt })))); }""",
                          [account, did, drafts])
            page.goto(url)
            bar = page.get_by_test_id("draft-recovery")
            expect(bar).to_be_visible(timeout=60000)
            bar.get_by_role("button", name="查看差异").click()
            pick = bar.locator("select")
            if "【A稿】" not in page.locator(".draft-recovery").inner_text():
                pick.select_option(index=1 - int(pick.input_value()))
            expect(page.locator(".draft-recovery")).to_contain_text("【A稿】", timeout=10000)
            bar.get_by_role("button", name="恢复并保存").click()
            expect(page.locator(".wb-notice.info")).to_contain_text("冲突", timeout=30000)
            page.wait_for_timeout(3000)
            assert "还没存上服务器" not in page.locator("body").inner_text(), "冲突时仍显示“正在自动重试”"
            # 恢复条换成 B；冲突未处理时恢复 B 被拒绝并提示
            expect(page.locator(".draft-recovery")).to_contain_text("【B稿】", timeout=20000)
            # 设计师 T-S1：拒绝原因常驻在恢复条里，“恢复并保存”置灰并带提示（不再点了才弹短提示）
            restore = bar.get_by_role("button", name="恢复并保存")
            expect(restore).to_be_disabled(timeout=20000)
            expect(restore).to_have_attribute("title", "请先处理未完成的冲突")
            expect(page.locator(".draft-recovery")).to_contain_text("冲突")
            page.wait_for_timeout(2000)
            from app.db import connect
            with connect() as connection:
                rows = connection.execute("SELECT status, payload_json FROM document_conflicts WHERE document_id=%s", (did,)).fetchall()
            mines = [(row["status"], json.dumps(json.loads(row["payload_json"]).get("conflicts"), ensure_ascii=False)) for row in rows]
            assert len(mines) == 1 and mines[0][0] == "open" and "【A稿】" in mines[0][1], mines
            server = data["zhang"].call("GET", f"/documents/{did}")["markdownContent"]
            assert "（对方改）" in server and "【B稿】" not in server, server[:300]
            assert page.evaluate("JSON.stringify(localStorage)").count("【B稿】") >= 1, "B 草稿被删"
            shot(page, "31-冲突未处理时拒绝恢复下一份草稿")
            ctx.close()
        draft_refused_while_conflict()

        @step("载入锁定（检查本机草稿）期间点正文打的字，解锁后写回原处（测试工程师 F-L1）")
        def typing_during_load_lock():
            from e2e_selfconf import make_doc
            ctx = browser.new_context(viewport={"width": 1920, "height": 1080}, locale="zh-CN")
            page = ctx.new_page()
            open_pages.append(page)
            login(page, f"e2e-li-{RUN}", landing="my-writing")
            hits = 0
            for n in range(3):
                did = make_doc(data, f"载入锁定-{n}-{RUN}", 0)
                page.goto(f"{WEB}/projects/{data['project']['id']}/documents/{did}")
                target = page.locator(".paper-content p", has_text="第2段")
                try:
                    page.wait_for_function("""() => { const e = document.querySelector('.paper-content');
                        return e && e.getAttribute('contenteditable') !== 'true' && [...e.querySelectorAll('p')].some(p => p.textContent.includes('第2段')); }""",
                                           polling="raf", timeout=60000)
                except AssertionError:
                    pass
                locked = page.locator(".paper-content").get_attribute("contenteditable") != "true"
                box = target.bounding_box()
                page.mouse.click(box["x"] + box["width"] - 2, box["y"] + box["height"] / 2)
                page.keyboard.type(f"LOCKx{n}y", delay=40)  # 逐键输入（中文输入法的 Process 键另有提示）
                page.wait_for_timeout(1500)
                wait_saved(page, 30000)
                server = data["zhang"].call("GET", f"/documents/{did}")["markdownContent"]
                ok = f"LOCKx{n}y" in server and f"LOCKx{n}y" in page.locator(".paper-content").inner_text()
                notice = page.locator(".wb-notice").all_inner_texts()
                assert ok or any("没能写入" in text for text in notice), (locked, server[:200], notice)
                hits += locked
                assert ok, (locked, server[:200], notice)
            assert hits >= 1, "三次都没赶上载入锁定"
            ctx.close()
        typing_during_load_lock()

        @step("保存后马上继续打字：不与自己冲突、后打的字不丢（终验 F-H2）；打开时有载入提示（F-L1）")
        def save_then_keep_typing():
            from e2e_selfconf import MARK, make_doc
            ctx = browser.new_context(viewport={"width": 1920, "height": 1080}, locale="zh-CN")
            page = ctx.new_page()
            open_pages.append(page)
            # 终验 F-L1：打开文档后检查本机草稿的约 0.25 秒内正文不可编辑，要有载入提示。
            # 设计师 T-S2 后载入提示延迟 350ms 才出现，真实的 250ms 锁定看不到提示：与 e2e_scheme_c 的 F-L1 一样把这一个 250ms 等待拉长到 2.5s
            page.add_init_script("""(() => { const raw = window.setTimeout; window.setTimeout = (fn, ms, ...rest) => raw(fn, ms === 250 ? 2500 : ms, ...rest); })()""")
            page.add_init_script("""window.__notices = []; new MutationObserver(() => {
                const node = document.querySelector('[data-testid=editor-refreshing]');
                if (node && !window.__notices.includes(node.textContent)) window.__notices.push(node.textContent);
            }).observe(document, { childList: true, subtree: true, characterData: true });""")
            login(page, f"e2e-li-{RUN}", landing="my-writing")
            for big, gap in ((0, 0), (1508, 500)):
                fh_doc = make_doc(data, f"F-H2-{big}-{RUN}", big)
                page.goto(f"{WEB}/projects/{data['project']['id']}/documents/{fh_doc}")
                expect(page.locator(".paper-content")).to_be_visible(timeout=60000)
                expect(page.get_by_test_id("save-status")).to_contain_text("保存", timeout=60000)
                type_after(page, "第3段", MARK[:5])
                notices = page.evaluate("window.__notices")
                assert any("正在检查本机草稿" in text for text in notices), notices
                record = data["zhang"].call("GET", f"/documents/{fh_doc}")
                data["zhang"].call("PATCH", f"/documents/{fh_doc}/content", json={
                    "expectedRevision": record["revision"], "allowMerge": True, "clientSaveKey": str(uuid.uuid4()),
                    "markdownContent": record["markdownContent"].replace("第7段：原始内容 7。", "第7段：原始内容 7（对方改）。")})
                page.keyboard.press("Control+s")
                page.wait_for_timeout(gap)
                page.keyboard.type(MARK[5:], delay=60)
                page.wait_for_timeout(5000)
                wait_saved(page, 60000)
                server = data["zhang"].call("GET", f"/documents/{fh_doc}")["markdownContent"]
                assert MARK in server and "（对方改）" in server, (big, server[:400])
                expect(page.locator(".paper-content")).to_contain_text(MARK, timeout=20000)
                expect(page.locator(".paper-content")).to_contain_text("（对方改）", timeout=30000)
                assert page.locator("text=已单独保存").count() == 0, "与自己冲突"
                assert data["zhang"].call("GET", f"/documents/{fh_doc}/conflicts/summary").get("allOpen") in (0, None)
            shot(page, "31-保存后马上继续打字")
            ctx.close()
        save_then_keep_typing()

        @step("同一账号多标签提示（非阻断，可关闭）")
        def multi_tab():
            open_doc(a, data)
            second = ctx_a.new_page()
            open_pages.append(second)
            open_doc(second, data)
            expect(second.get_by_test_id("multi-tab")).to_be_visible(timeout=15000)
            expect(a.get_by_test_id("multi-tab")).to_be_visible(timeout=15000)
            shot(second, "10-多标签提示")
            second.get_by_test_id("multi-tab").get_by_role("button", name="知道了").click()
            expect(second.get_by_test_id("multi-tab")).to_have_count(0)
            second.close()
        multi_tab()

        @step("复制标签页：tabId 不同，A 的未保存草稿不被 B 删除（审查 M1）")
        def duplicate_tab():
            page = ctx_a.new_page()
            open_pages.append(page)
            open_doc(page, data)
            with ctx_a.expect_page() as info:
                page.evaluate("window.open(location.href)")
            dup = info.value
            open_pages.append(dup)
            expect(dup.locator(".paper-content")).to_be_visible(timeout=60000)
            expect(dup.get_by_test_id("save-status")).to_contain_text(re.compile("保存"), timeout=30000)
            tab_a = page.evaluate("sessionStorage.getItem('doc-workbench.tab')")
            tab_b = dup.evaluate("sessionStorage.getItem('doc-workbench.tab')")
            assert tab_a and tab_b and tab_a != tab_b, (tab_a, tab_b)
            expect(dup.get_by_test_id("multi-tab")).to_be_visible(timeout=15000)
            page.route("**/documents/*/content", lambda route: route.abort() if route.request.method == "PATCH" else route.continue_())
            type_after(page, p_dup_a, "（A未保存）")
            save_now(page)
            expect(page.get_by_test_id("save-status")).to_contain_text(re.compile("重试|离线"), timeout=30000)
            shot(page, "03-重试中")
            type_after(dup, p_dup_b, "（B已保存）")
            save_now(dup)
            wait_saved(dup)
            dup.wait_for_timeout(1000)
            suffix = f":{doc_id}:{tab_a}"
            draft = page.evaluate("""suffix => { for (let i = 0; i < localStorage.length; i++) { const k = localStorage.key(i);
                if (k.startsWith('doc-workbench.draft.v2:') && k.endsWith(suffix)) return localStorage.getItem(k); } return null; }""", suffix)
            assert draft and "（A未保存）" in draft, "B 保存后 A 的本机草稿必须还在"
            dup.close()
            page.unroute("**/documents/*/content")
            page.evaluate("window.dispatchEvent(new Event('online'))")
            wait_saved(page, 60000)
            server = server_markdown()
            assert "（A未保存）" in server and "（B已保存）" in server
            page.get_by_role("button", name="版本记录").click()
            expect(page.locator(".history-panel")).to_be_visible(timeout=15000)
            expect(page.locator(".history-panel")).to_contain_text("张工")
            shot(page, "20-版本记录")
            page.close()
        duplicate_tab()

        @step("对方删段的冲突：标记不落在标题上，恢复回原位置（复审 M3）")
        def deleted_conflict():
            li = Api(f"e2e-li-{RUN}")
            record = data["zhang"].call("GET", f"/documents/{doc_id}")
            base, blocks = record["revision"], record["markdownContent"].split("\n\n")
            index = next(i for i, block in enumerate(blocks) if block.strip().startswith(p_del))
            victim, before = blocks[index], blocks[index - 1]
            data["zhang"].call("PATCH", f"/documents/{doc_id}/content", json={
                "expectedRevision": base, "markdownContent": "\n\n".join(blocks[:index] + blocks[index + 1:]),
                "allowMerge": True, "clientSaveKey": str(uuid.uuid4())})
            mine = blocks[:]
            mine[index] = victim.rstrip() + "（李工改了被删的段）"
            response = li.client.patch(f"/documents/{doc_id}/content", headers={"X-CSRF-Token": li.csrf}, json={
                "expectedRevision": base, "markdownContent": "\n\n".join(mine), "allowMerge": True, "clientSaveKey": str(uuid.uuid4())})
            assert response.status_code == 409 and response.json()["error"]["details"]["conflicts"][0]["reason"] == "DELETED", response.text[:300]
            b = ctx_b.new_page()  # 前面离线步骤关掉了最初的 b
            open_pages.append(b)
            open_doc(b, data)
            mark = b.get_by_test_id("conflict-mark")
            expect(mark).to_have_count(1, timeout=20000)
            expect(mark).to_contain_text("该段已被 张工 删除")
            expect(b.locator(".paper-content").locator("h1.blk-conflict, h2.blk-conflict, h3.blk-conflict, h4.blk-conflict, h5.blk-conflict, h6.blk-conflict")).to_have_count(0)
            expect(b.get_by_test_id("conflict-gap")).to_have_count(1)
            mark.click()
            panel = b.locator(".conflict-panel")
            expect(panel.locator(".ccard")).to_have_count(1, timeout=20000)
            shot(b, "26-对方删段-大致位置标记")
            panel.get_by_role("button", name="恢复为新段落").click()
            expect(panel.locator(".ccard")).to_have_count(0, timeout=20000)
            restored = server_markdown().split("\n\n")
            at = next(i for i, block in enumerate(restored) if "（李工改了被删的段）" in block)
            assert restored[at - 1].strip() == before.strip(), (restored[at - 1][:40], before[:40])
            expect(paragraph(b, "（李工改了被删的段）")).to_be_visible(timeout=15000)
        deleted_conflict()

        browser.close()
    failed = [item for item in results if not item[1]]
    print(f"\n{len(results) - len(failed)}/{len(results)} passed")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
