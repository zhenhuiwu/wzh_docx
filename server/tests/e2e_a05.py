"""A05 成员管理浏览器验收（16 §2.12 / §2.13，1920×1080）。

环境变量：E2E_API、E2E_WEB、DATABASE_URL、FILE_DATA_DIR **必须设置**（脚本直接在被测库里建账号）；
E2E_SHOTS（截图目录，默认系统临时目录/e2e_a05/shots）、E2E_OUT（失败截图）。

运行（先用 /tmp/wzh_a05_run/live/start.sh 起服务，8051/5221）：
    E2E_API=http://127.0.0.1:8051/api/v1 E2E_WEB=http://127.0.0.1:5221 E2E_SHOTS=/tmp/wzh_a05_run/shots \
    DATABASE_URL=sqlite:////tmp/wzh_a05_run/live/app.db FILE_DATA_DIR=/tmp/wzh_a05_run/live/files \
        uv run python tests/e2e_a05.py
覆盖：V1 无成员导航；V2 添加编写人员（默认编写人员）；V3 停用账号不在候选里；V4 普通账号进入所分配项目并能编写、没有管理入口、
未加入的项目 404；V5 降为只读后正在编辑的人保存即只读、可下载；V6 最后负责人（前端禁用 + 后端 409）；V7 停用的负责人不计数；
V8 并发 409 提示并刷新；V9 自身降级确认；V10 负责人移除自己跳回列表；V11 归档只读（成员区与文档）；V14 主负责人改指；
V15 成员操作不让项目资料 409；V16 缺版本号 422；V17 reviewer；V18 成员不能自己退出；V19–V22 用户管理停用唯一负责人
（含 Q-2：已归档项目也挡，列表标“（已归档）”）；V23 Q-1 归档项目事实/测试任务/知识库关联写入 409、读照常。
admin/admin123 登录、不改密（用户决定）。
"""

from __future__ import annotations

import os
import re
import sys
import tempfile
import uuid
from io import BytesIO
from pathlib import Path

import httpx
from playwright.sync_api import Page, expect, sync_playwright

ROOT = Path(__file__).resolve().parents[2]
API = os.environ.get("E2E_API", "")
WEB = os.environ.get("E2E_WEB", "")
OUT = Path(os.environ.get("E2E_OUT", str(Path(tempfile.gettempdir()) / "e2e_a05")))
SHOTS = Path(os.environ.get("E2E_SHOTS", str(OUT / "shots")))
PASSWORD = "E2e-a05-pass-1"
ADMIN, ADMIN_PASSWORD = "admin", "admin123"
RUN = uuid.uuid4().hex[:6]
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
results: list[tuple[str, bool, str]] = []
open_pages: list[Page] = []



def open_mode(api, doc: dict) -> dict:
    """A07 之前写的场景假定编写人员能改整篇：新建文档（默认按分工编辑）改为不限制分工（open）。"""
    view = api.call("GET", f"/documents/{doc['id']}/assignments")
    if view["assignmentMode"] != "open":
        api.call("PUT", f"/documents/{doc['id']}/assignment-mode",
                 json={"mode": "open", "expectedAssignmentsRevision": view["assignmentsRevision"]})
    return doc

def step(name: str):
    def wrap(fn):
        def run(*args, **kwargs):
            if os.environ.get("E2E_ONLY") and os.environ["E2E_ONLY"] not in name:
                print(f"SKIP {name}", flush=True)
                return
            try:
                fn(*args, **kwargs)
                results.append((name, True, ""))
                print(f"PASS {name}", flush=True)
            except Exception as error:  # noqa: BLE001 - 汇总后统一失败
                results.append((name, False, f"{type(error).__name__}: {error}"))
                print(f"FAIL {name}: {error}", flush=True)
                OUT.mkdir(parents=True, exist_ok=True)
                for index, page in enumerate(open_pages):
                    if not page.is_closed():
                        page.screenshot(path=str(OUT / f"fail-{len(results)}-{index}.png"))
                if os.environ.get("E2E_STOP_ON_FAIL") == "1":
                    raise
        return run
    return wrap


class Api:
    def __init__(self, account: str, password: str = PASSWORD):
        self.client = httpx.Client(base_url=API, timeout=120)
        csrf = self.client.get("/auth/csrf").json()["data"]["csrfToken"]
        response = self.client.post("/auth/login", json={"account": account, "password": password}, headers={"X-CSRF-Token": csrf})
        response.raise_for_status()
        self.csrf = response.json()["data"]["csrfToken"]

    def raw(self, method: str, path: str, **kwargs) -> httpx.Response:
        headers = {"X-CSRF-Token": self.csrf, **kwargs.pop("headers", {})}
        return self.client.request(method, path, headers=headers, **kwargs)

    def call(self, method: str, path: str, **kwargs):
        response = self.raw(method, path, **kwargs)
        if response.status_code >= 400:
            raise AssertionError(f"{method} {path} → {response.status_code} {response.text[:300]}")
        return response.json()["data"] if response.content else None

    def members_revision(self, project: str) -> int:
        return self.call("GET", f"/projects/{project}/members")["membersRevision"]

    def set_role(self, project: str, user, role: str):
        """已是成员 → PUT 改角色；不是成员 → members:batch 加入（PUT 只改已有成员，A05 20 I-3）。"""
        snapshot = self.call("GET", f"/projects/{project}/members")
        if any(item["userId"] == str(user) for item in snapshot["items"]):
            return self.call("PUT", f"/projects/{project}/members/{user}", json={"role": role, "expectedMembersRevision": snapshot["membersRevision"]})
        return self.call("POST", f"/projects/{project}/members:batch",
                         json={"userIds": [str(user)], "role": role, "expectedMembersRevision": snapshot["membersRevision"]})

    def project(self, name: str, owner) -> dict:
        return self.call("POST", "/projects", json={"name": name, "ownerId": str(owner), "status": "active"},
                         headers={"Idempotency-Key": f"a05-{uuid.uuid4().hex}"})

    def disable(self, user) -> httpx.Response:
        return self.raw("POST", f"/users/{user}/disable", json={"expectedRevision": _user_revision(user), "reason": "A05 验收停用"})


def _user_revision(user) -> int:
    from app.db import connect
    with connect() as connection:
        return int(connection.execute("select revision from users where id = %s", (user,)).fetchone()[0])


PEOPLE = (("lisi", "李四", "软件一部"), ("wangce", "王测", "测试部"), ("zhao", "赵停用", "软件二部"), ("sunmin", "孙敏", "软件一部"),
          ("zhoutao", "周涛", "测试部"), ("owner2", "钱负责", "软件一部"), ("zhang2", "张二", "软件一部"), ("owner3", "吴三", "软件二部"),
          ("solo", "郑唯一", "软件二部"), ("helper", "冯帮手", "软件一部"), ("other", "陈其他", "软件二部"))


def make_document(admin: "Api", project_id: str) -> dict:
    from docx import Document as WordDocument
    word = WordDocument()
    word.add_heading("概述", 1)
    word.add_paragraph("本文档用于验证成员权限调整后的编写与只读。")
    word.add_paragraph("第二段正文，编写人员会在这里补充内容。")
    buffer = BytesIO()
    word.save(buffer)
    upload = admin.client.post("/files", files={"file": ("成员验收.docx", buffer.getvalue(), DOCX)},
                               data={"purpose": "attachment", "projectId": project_id},
                               headers={"X-CSRF-Token": admin.csrf, "Idempotency-Key": f"a05-upload-{uuid.uuid4().hex}"})
    assert upload.status_code == 201, upload.text
    return open_mode(admin, admin.call("POST", f"/projects/{project_id}/documents",
                      json={"name": "软件需求规格说明", "sourceType": "upload", "wordFileId": upload.json()["data"]["id"]}))


def seed() -> dict:
    sys.path.insert(0, str(ROOT / "server"))
    from app.db import connect
    from app.identity import create_user
    users = {}
    with connect() as connection:
        row = connection.execute("select id from users where account_key = %s", (ADMIN,)).fetchone()
        users["admin"] = row[0] if row else create_user(account=ADMIN, name="系统管理员", password=ADMIN_PASSWORD, roles=["system_admin"], connection=connection)
        for key, name, department in PEOPLE:
            users[key] = create_user(account=f"{key}{RUN}", name=name, password=PASSWORD, roles=[], department=department,
                                     allow_no_roles=True, connection=connection)
        connection.commit()
    admin = Api(ADMIN, ADMIN_PASSWORD)
    # 管理员显示名以 /auth/me 为准，不写死“系统管理员”（20：现场库里管理员可能改过名）
    admin_name = admin.call("GET", "/auth/me")["user"]["name"]
    project = admin.project(f"综合保障信息系统（二期）{RUN}", users["admin"])
    other = admin.project(f"未加入的项目{RUN}", users["admin"])
    document = make_document(admin, project["id"])
    # 背景成员：王测（编写）、周涛（只读）
    admin.set_role(project["id"], users["wangce"], "editor")
    admin.set_role(project["id"], users["zhoutao"], "viewer")
    assert admin.disable(users["zhao"]).status_code == 200
    return {"users": users, "project": project, "other": other, "document": document, "admin": admin, "admin_name": admin_name}


def login(page: Page, account: str, password: str = PASSWORD):
    page.context.clear_cookies()
    page.goto(f"{WEB}/login")
    page.fill("#staff-id", account)
    page.fill("#password", password)
    page.locator("form button[type=submit]").click()
    page.wait_for_url(re.compile(r"/(projects|my-writing)"), timeout=30000)  # 编写人员落地页为“我的编写”


def shot(page: Page, name: str):
    SHOTS.mkdir(parents=True, exist_ok=True)
    page.screenshot(path=str(SHOTS / f"{name}.png"))


def row(page: Page, user) -> object:
    return page.locator(f'[data-testid=member-row][data-user-id="{user}"]')


def open_overview(page: Page, project: str):
    page.goto(f"{WEB}/projects/{project}")
    expect(page.get_by_test_id("members-section")).to_be_visible(timeout=30000)
    expect(page.get_by_test_id("members-count")).to_be_visible(timeout=30000)


def require_targets() -> None:
    missing = [name for name, value in (("E2E_API", API), ("E2E_WEB", WEB), ("DATABASE_URL", os.environ.get("DATABASE_URL", "")),
                                        ("FILE_DATA_DIR", os.environ.get("FILE_DATA_DIR", ""))) if not value]
    if missing:
        raise SystemExit(f"请设置环境变量 {', '.join(missing)}（被测服务的接口、前端地址和同一个数据库），脚本不使用默认值。")


def main() -> int:  # noqa: C901 - 验收脚本按用例顺序展开
    require_targets()
    data = seed()
    users, admin, admin_name = data["users"], data["admin"], data["admin_name"]
    pid, doc = data["project"]["id"], data["document"]["id"]

    with sync_playwright() as playwright:
        browser = playwright.chromium.launch()
        ctx_a = browser.new_context(viewport={"width": 1920, "height": 1080}, accept_downloads=True)
        ctx_b = browser.new_context(viewport={"width": 1920, "height": 1080}, accept_downloads=True)
        page = ctx_a.new_page()
        page2 = ctx_b.new_page()
        open_pages.extend([page, page2])
        login(page, ADMIN, ADMIN_PASSWORD)

        @step("V1 项目内只有概述/文档编写，没有成员导航；概述里有成员与权限")
        def v1():
            open_overview(page, pid)
            expect(page.get_by_role("navigation", name="项目内导航").get_by_role("link")).to_have_text(["项目概述", "文档编写"])
            expect(page.locator('aside[aria-label="主导航"] >> text=成员')).to_have_count(0)
            expect(page.get_by_test_id("members-count")).to_have_text("3")
            expect(page.get_by_test_id("frame-owners")).to_contain_text(admin_name)
        v1()

        @step("V2 添加编写人员（默认编写人员，多选，弹窗截图）")
        def v2():
            page.get_by_test_id("members-add").click()
            dialog = page.get_by_test_id("member-add-dialog")
            expect(dialog).to_be_visible()
            expect(dialog.get_by_test_id("member-candidate-search")).to_be_focused()
            dialog.get_by_test_id("member-candidate-search").fill(RUN)
            expect(dialog.locator(f'[data-testid=member-candidate][data-user-id="{users["lisi"]}"]')).to_be_visible(timeout=10000)
            expect(dialog.locator(f'[data-testid=member-candidate][data-user-id="{users["admin"]}"]')).to_have_count(0, timeout=10000)  # 防抖后的结果
            # 已在项目中的显示为灰色、不可勾选
            expect(dialog.locator(f'[data-testid=member-candidate][data-user-id="{users["wangce"]}"]')).to_contain_text("已是编写人员")
            expect(dialog.locator(f'[data-testid=member-candidate][data-user-id="{users["wangce"]}"] input')).to_be_disabled()
            dialog.locator(f'[data-testid=member-candidate][data-user-id="{users["lisi"]}"]').click()
            dialog.locator(f'[data-testid=member-candidate][data-user-id="{users["sunmin"]}"]').click()
            expect(dialog.get_by_test_id("member-add-role-editor")).to_be_checked()
            expect(dialog.get_by_test_id("member-add-submit")).to_have_text("添加 2 人")
            shot(page, "A05-02-添加成员弹窗")
            dialog.get_by_test_id("member-add-submit").click()
            expect(dialog).to_be_hidden(timeout=10000)
            expect(row(page, users["lisi"]).get_by_test_id("member-role-select")).to_have_value("editor")
            expect(row(page, users["sunmin"]).get_by_test_id("member-role-select")).to_have_value("editor")
            expect(page.get_by_test_id("members-count")).to_have_text("5")
        v2()

        @step("V3 停用账号不在候选里")
        def v3():
            page.get_by_test_id("members-add").click()
            dialog = page.get_by_test_id("member-add-dialog")
            dialog.get_by_test_id("member-candidate-search").fill(f"zhao{RUN}")
            expect(dialog).to_contain_text("没有找到账号", timeout=10000)
            expect(dialog.locator("[data-testid=member-candidate]")).to_have_count(0)
            dialog.get_by_test_id("member-add-submit").click()
            expect(dialog.get_by_test_id("member-add-error")).to_have_text("请至少选择一个账号。")
            page.keyboard.press("Escape")
            expect(dialog).to_be_hidden()
        v3()

        @step("V6 最后负责人：前端禁用降级和移除，后端 409 LAST_OWNER")
        def v6():
            admin_row = row(page, users["admin"])
            expect(admin_row.locator("[data-testid=member-role-select] option[value=viewer]")).to_be_disabled()
            expect(admin_row.get_by_test_id("member-remove")).to_be_disabled()
            expect(admin_row).to_contain_text("唯一启用的负责人")
            response = admin.raw("PUT", f"/projects/{pid}/members/{users['admin']}", json={"role": "viewer", "expectedMembersRevision": admin.members_revision(pid)})
            assert response.status_code == 409 and response.json()["error"]["code"] == "LAST_OWNER", response.text
            shot(page, "A05-01-成员与权限-管理视图")
        v6()

        @step("V4 普通账号进入所分配项目并能编写；概述里没有管理入口；未加入的项目 404")
        def v4():
            login(page2, f"lisi{RUN}")
            expect(page2.locator(f'[data-testid=project-my-role][data-project-id="{pid}"]')).to_have_text("编写人员")
            expect(page2.locator(f'[data-testid=project-my-role][data-project-id="{data["other"]["id"]}"]')).to_have_count(0)
            page2.goto(f"{WEB}/projects/{pid}/documents/{doc}")
            expect(page2.locator(".paper-content")).to_have_attribute("contenteditable", "true", timeout=60000)
            expect(page2.get_by_test_id("save-status")).not_to_contain_text("只读")
            open_overview(page2, pid)
            expect(page2.get_by_test_id("members-add")).to_have_count(0)
            expect(page2.get_by_test_id("member-remove")).to_have_count(0)
            expect(page2.get_by_test_id("member-role-select")).to_have_count(0)
            expect(page2.get_by_test_id("members-note")).to_have_text(f"如需调整权限，请联系负责人 {admin_name}。")
            shot(page2, "A05-03-成员与权限-只读视图（编写人员）")
            page2.goto(f"{WEB}/projects/{data['other']['id']}")
            expect(page2.get_by_text("项目不存在或你无权访问。")).to_be_visible(timeout=30000)
        v4()

        @step("V5 降为只读后，正在编辑的人保存即变只读，可下载未保存的修改")
        def v5():
            page2.goto(f"{WEB}/projects/{pid}/documents/{doc}")
            expect(page2.locator(".paper-content")).to_have_attribute("contenteditable", "true", timeout=60000)
            expect(page2.get_by_test_id("save-status")).to_contain_text(re.compile("保存"), timeout=30000)
            open_overview(page, pid)
            select = row(page, users["lisi"]).get_by_test_id("member-role-select")
            select.select_option("viewer")
            expect(row(page, users["lisi"]).get_by_test_id("member-role-status")).to_have_text(re.compile("已保存|正在保存…"))
            expect(select).to_have_value("viewer")
            target = page2.locator(".paper-content p", has_text="第二段正文").first
            target.click()
            page2.keyboard.press("End")
            page2.keyboard.type("补充一句。")
            page2.keyboard.press("Control+s")
            expect(page2.get_by_text("你的权限已调整为只读成员")).to_be_visible(timeout=30000)
            expect(page2.get_by_role("button", name="下载我的修改").first).to_be_visible()
            expect(page2.get_by_test_id("save-status")).to_contain_text("只读")
        v5()

        @step("V18 编写人员/只读成员不能自己退出（界面无移除，DELETE 403）")
        def v18():
            open_overview(page2, pid)
            expect(page2.get_by_test_id("member-remove")).to_have_count(0)
            lisi = Api(f"lisi{RUN}")
            response = lisi.raw("DELETE", f"/projects/{pid}/members/{users['lisi']}?expectedMembersRevision={admin.members_revision(pid)}")
            assert response.status_code == 403, response.text
        v18()

        @step("V7 停用的负责人不计数")
        def v7():
            admin.set_role(pid, users["owner2"], "owner")
            open_overview(page, pid)
            expect(row(page, users["admin"]).get_by_test_id("member-remove")).to_be_enabled()
            assert admin.disable(users["owner2"]).status_code == 200
            open_overview(page, pid)
            expect(row(page, users["owner2"]).get_by_test_id("member-disabled")).to_be_visible()
            expect(row(page, users["owner2"]).get_by_test_id("member-role-select")).to_be_disabled()
            expect(row(page, users["admin"]).get_by_test_id("member-remove")).to_be_disabled()
            expect(row(page, users["owner2"]).get_by_test_id("member-remove")).to_be_enabled()
        v7()

        @step("V8 并发：别人刚调整过名单 → 409 提示、刷新名单、下拉停在服务端值")
        def v8():
            open_overview(page, pid)
            other_admin = Api(ADMIN, ADMIN_PASSWORD)
            other_admin.call("DELETE", f"/projects/{pid}/members/{users['wangce']}?expectedMembersRevision={other_admin.members_revision(pid)}")
            select = row(page, users["sunmin"]).get_by_test_id("member-role-select")
            select.select_option("viewer")
            expect(page.get_by_test_id("members-notice")).to_have_text(re.compile(r"成员名单刚被 .+ 于 .+ 调整过"), timeout=10000)
            expect(row(page, users["wangce"])).to_have_count(0)
            expect(select).to_have_value("editor")
        v8()

        @step("V15 成员操作不让“编辑项目”表单 409")
        def v15():
            open_overview(page, pid)
            page.get_by_role("button", name="编辑项目").click()
            page.locator(".edit-form textarea").fill(f"保障资源管理、维修计划与器材库存模块的需求与设计文档。{RUN}")
            admin.set_role(pid, users["sunmin"], "viewer")
            page.get_by_role("button", name="保存项目").click()
            expect(page.locator(".edit-form")).to_have_count(0, timeout=10000)
            expect(page.locator(".description")).to_contain_text(RUN)
        v15()

        @step("V16/V17 接口：缺版本号 422；reviewer 不能分配，旧数据显示“审核人员（只读）”")
        def v16():
            response = admin.raw("DELETE", f"/projects/{pid}/members/{users['sunmin']}")
            assert response.status_code == 422, response.text
            response = admin.raw("PUT", f"/projects/{pid}/members/{users['helper']}", json={"role": "reviewer", "expectedMembersRevision": admin.members_revision(pid)})
            assert response.status_code == 422, response.text
            from app.db import connect
            with connect() as connection:
                connection.execute("insert into project_members(project_id, user_id, role) values (%s, %s, 'reviewer')", (pid, users["helper"]))
                connection.commit()
            open_overview(page, pid)
            expect(row(page, users["helper"]).get_by_test_id("member-role-tag")).to_have_text("审核人员（只读）")
            expect(row(page, users["helper"]).get_by_test_id("member-role-select")).to_have_count(0)
        v16()

        @step("V9 自身降级要确认；确认后区块变只读")
        def v9():
            admin.set_role(pid, users["zhang2"], "owner")
            login(page2, f"zhang2{RUN}")
            open_overview(page2, pid)
            row(page2, users["zhang2"]).get_by_test_id("member-role-select").select_option("editor")
            dialog = page2.get_by_test_id("member-self-dialog")
            expect(dialog).to_be_visible()
            expect(dialog).to_contain_text("把自己改为编写人员？")
            dialog.get_by_test_id("member-self-confirm").click()
            expect(page2.get_by_test_id("members-add")).to_have_count(0, timeout=10000)
            expect(page2.get_by_test_id("members-note")).to_have_text(re.compile("如需调整权限，请联系负责人"))
            expect(page2.get_by_test_id("members-info")).to_have_text("你的项目权限已调整为编写人员。")
        v9()

        @step("V10 负责人移除自己 → 跳到项目列表并提示，列表里没有该项目（移除确认截图）")
        def v10():
            admin.set_role(pid, users["owner3"], "owner")
            login(page2, f"owner3{RUN}")
            open_overview(page2, pid)
            row(page2, users["owner3"]).get_by_test_id("member-remove").click()
            dialog = page2.get_by_test_id("member-remove-dialog")
            expect(dialog).to_contain_text("退出该项目？")
            dialog.get_by_test_id("member-remove-confirm").click()
            page2.wait_for_url(re.compile(r"/projects$"), timeout=10000)
            expect(page2.get_by_text(re.compile("你已退出 "))).to_be_visible(timeout=10000)
            expect(page2.locator(f'[data-testid=project-my-role][data-project-id="{pid}"]')).to_have_count(0)
            # 移除他人的确认弹窗截图
            open_overview(page, pid)
            row(page, users["sunmin"]).get_by_test_id("member-remove").click()
            remove = page.get_by_test_id("member-remove-dialog")
            expect(remove).to_contain_text(f"将 孙敏（sunmin{RUN}）移出项目？")
            shot(page, "A05-04-移除确认")
            remove.get_by_test_id("member-remove-confirm").click()
            expect(remove).to_be_hidden(timeout=10000)
            expect(row(page, users["sunmin"])).to_have_count(0)
        v10()

        @step("V14 主负责人改指：主负责人移除自己后，页头与列表显示新负责人")
        def v14():
            p3 = admin.project(f"主负责人改指{RUN}", users["admin"])
            admin.set_role(p3["id"], users["other"], "owner")
            open_overview(page, p3["id"])
            expect(page.get_by_test_id("frame-owners")).to_have_text(f"负责人：{admin_name}、陈其他")
            row(page, users["admin"]).get_by_test_id("member-remove").click()
            page.get_by_test_id("member-remove-dialog").get_by_test_id("member-remove-confirm").click()
            # 系统管理员不是成员也能管理：留在页面；主负责人改为陈其他
            expect(page.get_by_test_id("members-info")).to_contain_text("主负责人改为 陈其他", timeout=10000)
            expect(page.get_by_test_id("frame-owners")).to_have_text("负责人：陈其他")
            # 改指提示 4 秒后自动收起（21 S-5）
            expect(page.get_by_test_id("members-info")).to_have_count(0, timeout=8000)
            login(page2, f"other{RUN}")
            expect(page2.locator(f'[data-testid=project-my-role][data-project-id="{p3["id"]}"]')).to_have_text("负责人")
            page2.goto(f"{WEB}/projects/{p3['id']}")
            expect(page2.get_by_test_id("frame-owners")).to_have_text("负责人：陈其他", timeout=30000)
        v14()

        @step("V19–V22 用户管理：停用唯一启用负责人被拒；归档项目在界面上恢复、补负责人、再归档后停用成功（21 M-1）")
        def v19():
            p2 = admin.project(f"唯一负责人项目{RUN}", users["solo"])
            p3b = admin.project(f"另有负责人项目{RUN}", users["solo"])
            admin.set_role(p3b["id"], users["helper"], "owner")
            # Q-2：已归档项目里的唯一负责人也挡停用，列表里标“（已归档）”
            p2a = admin.project(f"已归档唯一负责人项目{RUN}", users["solo"])
            admin.call("POST", f"/projects/{p2a['id']}/archive",
                       json={"expectedRevision": admin.call("GET", f"/projects/{p2a['id']}")["revision"], "reason": "A05 Q-2 验收"})
            page.goto(f"{WEB}/users")
            page.locator(".query input").nth(1).fill(f"solo{RUN}")
            page.locator(".query button[type=submit]").click()
            page.get_by_role("button", name="停用").first.click()
            page.locator("#ask-reason").fill("离职")
            page.locator("dialog[open] button[type=submit]").click()
            blocked = page.get_by_test_id("user-disable-blocked")
            expect(blocked).to_be_visible(timeout=10000)
            expect(blocked).to_contain_text(re.compile("是以下 2 个项目唯一的启用负责人"))
            expect(blocked.locator(f'[data-testid=user-disable-project][data-project-id="{p2["id"]}"]')).to_be_visible()
            expect(blocked.locator(f'[data-testid=user-disable-project][data-project-id="{p3b["id"]}"]')).to_have_count(0)
            archived_row = blocked.locator(f'[data-testid=user-disable-project][data-project-id="{p2a["id"]}"]')
            expect(archived_row).to_contain_text(f"已归档唯一负责人项目{RUN}（已归档）")
            expect(blocked.locator(f'[data-testid=user-disable-project][data-project-id="{p2["id"]}"]')).not_to_contain_text("（已归档）")
            expect(blocked).to_contain_text("标“（已归档）”的项目需先恢复")
            expect(archived_row.get_by_test_id("user-disable-project-restore")).to_have_text("打开并恢复")
            expect(page.locator("dialog[open] h2")).to_have_text(f"不能停用 郑唯一（solo{RUN}）")
            shot(page, "A05-06-用户管理-不能停用唯一负责人")
            # V22：账号仍是启用
            assert _user_status(users["solo"]) == "active"
            # 项目链接在新标签页打开到 #members（新标签页校验滚动到成员区）
            with ctx_a.expect_page() as popup_info:
                blocked.locator(f'[data-testid=user-disable-project][data-project-id="{p2["id"]}"] a').first.click()
            popup = popup_info.value
            expect(popup.get_by_test_id("members-section")).to_be_in_viewport(timeout=30000)
            popup.close()
            admin.set_role(p2["id"], users["helper"], "owner")
            # 归档项目成员只读：接口 409 PROJECT_ARCHIVED
            response = admin.raw("POST", f"/projects/{p2a['id']}/members:batch",
                                 json={"userIds": [str(users["helper"])], "role": "owner", "expectedMembersRevision": admin.members_revision(p2a["id"])})
            assert response.status_code == 409 and response.json()["error"]["code"] == "PROJECT_ARCHIVED", response.text
            # 21 M-1：全程走界面——“打开并恢复”→ 恢复确认 → 添加负责人 → 页头［归档项目］→ 回到用户管理重新检查
            with ctx_a.expect_page() as popup_info:
                archived_row.get_by_test_id("user-disable-project-restore").click()
            popup = popup_info.value
            restore = popup.get_by_test_id("project-archive-dialog")
            expect(restore).to_be_visible(timeout=30000)
            expect(restore).to_contain_text(f"恢复 已归档唯一负责人项目{RUN}？")
            restore.get_by_test_id("project-archive-confirm").click()
            expect(restore.get_by_test_id("project-archive-error")).to_have_text("请填写恢复原因。")
            restore.get_by_test_id("project-archive-reason").fill("补负责人")
            shot(popup, "A05-08-恢复项目确认")
            restore.get_by_test_id("project-archive-confirm").click()
            expect(restore).to_be_hidden(timeout=10000)
            expect(popup.get_by_test_id("frame-archived")).to_have_count(0)
            expect(popup.get_by_test_id("members-add")).to_be_visible(timeout=10000)
            popup.get_by_test_id("members-add").click()
            add = popup.get_by_test_id("member-add-dialog")
            add.get_by_test_id("member-candidate-search").fill(f"helper{RUN}")
            add.locator(f'[data-testid=member-candidate][data-user-id="{users["helper"]}"]').click()
            add.get_by_test_id("member-add-role-owner").check()
            add.get_by_test_id("member-add-submit").click()
            expect(add).to_be_hidden(timeout=10000)
            expect(row(popup, users["helper"]).get_by_test_id("member-role-select")).to_have_value("owner")
            popup.get_by_test_id("project-archive").click()
            archive = popup.get_by_test_id("project-archive-dialog")
            archive.get_by_test_id("project-archive-reason").fill("补完再归档")
            archive.get_by_test_id("project-archive-confirm").click()
            expect(archive).to_be_hidden(timeout=10000)
            expect(popup.get_by_test_id("frame-archived")).to_be_visible()
            expect(popup.get_by_test_id("project-restore")).to_be_visible()
            expect(popup.get_by_test_id("members-restore")).to_be_visible()
            shot(popup, "A05-09-补负责人后重新归档")
            popup.close()
            page.get_by_test_id("user-disable-recheck").click()
            expect(page.locator("dialog[open]")).to_have_count(0, timeout=10000)
            expect(page.get_by_text("账号已停用。")).to_be_visible()
            assert _user_status(users["solo"]) == "disabled"
            open_overview(page, p2["id"])
            expect(row(page, users["solo"]).get_by_test_id("member-disabled")).to_be_visible()
        v19()

        @step("V11 归档项目：成员区只读、文档只读（文案按枚举），成员接口 409 PROJECT_ARCHIVED")
        def v11():
            admin.set_role(pid, users["lisi"], "editor")
            revision = admin.call("GET", f"/projects/{pid}")["revision"]
            admin.call("POST", f"/projects/{pid}/archive", json={"expectedRevision": revision, "reason": "A05 归档验收"})
            open_overview(page, pid)
            expect(page.get_by_test_id("members-notice")).to_contain_text("项目已归档，成员和权限不能调整。恢复项目后才能修改。")
            # 21 S-1/S-3：页头“已归档”，概述“已归档（原：进行中）”；说明行不再说“可以添加成员”；管理员有［恢复项目］
            expect(page.get_by_test_id("frame-archived")).to_have_text("已归档")
            expect(page.get_by_test_id("project-status")).to_have_text("已归档（原：进行中）")
            expect(page.get_by_test_id("members-note")).not_to_contain_text("可以添加成员")
            expect(page.get_by_test_id("project-restore")).to_be_visible()
            expect(page.get_by_test_id("member-role-select")).to_have_count(0)
            expect(page.get_by_test_id("members-add")).to_have_count(0)
            expect(page.get_by_test_id("member-remove")).to_have_count(0)
            shot(page, "A05-05-成员与权限-归档只读")
            response = admin.raw("PUT", f"/projects/{pid}/members/{users['lisi']}", json={"role": "viewer", "expectedMembersRevision": admin.members_revision(pid)})
            assert response.status_code == 409 and response.json()["error"]["code"] == "PROJECT_ARCHIVED", response.text
            login(page2, f"lisi{RUN}")
            page2.goto(f"{WEB}/projects/{pid}/documents/{doc}")
            expect(page2.locator(".paper-content")).to_be_visible(timeout=60000)
            expect(page2.get_by_test_id("save-status")).to_contain_text("只读", timeout=30000)
            expect(page2.get_by_test_id("save-status")).to_have_attribute("title", "项目已归档，可以查看和导出，不能修改正文。")
            expect(page2.get_by_test_id("archived-notice")).to_have_text("项目已归档，可以查看和导出，不能修改正文。")  # 20 L-4
            shot(page2, "A05-07-归档项目文档只读")
            # 编写人员没有恢复权限：没有按钮，提示去找系统管理员
            open_overview(page2, pid)
            expect(page2.get_by_test_id("project-restore")).to_have_count(0)
            expect(page2.get_by_test_id("members-notice")).to_have_text("项目已归档，成员和权限不能调整。如需恢复项目，请联系系统管理员。")
            page2.goto(f"{WEB}/projects/{pid}/documents")
            expect(page2.get_by_test_id("frame-archived-notice")).to_contain_text("项目已归档，文档只能查看和导出。")
            # 项目列表：默认不含归档项目，状态选“已归档”能找到，筛选写进地址（21 M-1、S-9）
            page.goto(f"{WEB}/projects?name={RUN}")
            expect(page.get_by_role("link", name=f"打开项目：综合保障信息系统（二期）{RUN}")).to_have_count(0, timeout=30000)
            page.locator(".project-query select").select_option("archived")
            page.locator(".project-query button[type=submit]").click()
            expect(page.get_by_role("link", name=f"打开项目：综合保障信息系统（二期）{RUN}")).to_be_visible(timeout=30000)
            expect(page).to_have_url(re.compile(r"status=archived"))
            expect(page.get_by_test_id("project-archived").first).to_have_text("已归档")
            shot(page, "A05-10-项目列表-已归档筛选")
            page.reload()
            expect(page.locator(".project-query select")).to_have_value("archived")
            expect(page.get_by_role("link", name=f"打开项目：综合保障信息系统（二期）{RUN}")).to_be_visible(timeout=30000)
        v11()

        @step("V23 Q-1 归档项目完全只读：事实、测试任务、知识库关联写入 409 PROJECT_ARCHIVED，读照常（资料页当前未开放路由，只验接口）")
        def v23():
            revision = admin.call("GET", f"/projects/{pid}")["revision"]
            for method, path, body in (
                ("POST", f"/projects/{pid}/facts", {"key": "a05_q1", "label": "归档后事实", "value": "x", "decisionReason": "验收"}),
                ("PUT", f"/projects/{pid}/knowledge-bindings", {"expectedProjectRevision": revision, "baseIds": []}),
                ("POST", "/test-tasks", {"name": "归档后任务", "projectId": pid, "planFileId": str(uuid.uuid4())}),
            ):
                response = admin.raw(method, path, json=body)
                assert response.status_code == 409 and response.json()["error"]["code"] == "PROJECT_ARCHIVED", (path, response.text)
            assert admin.raw("GET", f"/projects/{pid}/facts").status_code == 200
            assert admin.raw("GET", f"/projects/{pid}/knowledge-bindings").status_code == 200
        v23()

        @step("V24 PUT 只改已有成员：对非成员 404 MEMBER_NOT_FOUND，不会把人加回来（20 I-3）")
        def v24():
            p4 = admin.project(f"PUT 不加人{RUN}", users["admin"])
            response = admin.raw("PUT", f"/projects/{p4['id']}/members/{users['lisi']}",
                                 json={"role": "editor", "expectedMembersRevision": admin.members_revision(p4["id"])})
            assert response.status_code == 404 and response.json()["error"]["code"] == "MEMBER_NOT_FOUND", response.text
            assert all(item["userId"] != str(users["lisi"]) for item in admin.call("GET", f"/projects/{p4['id']}/members")["items"])
        v24()

        @step("V25 没有启用的负责人：页头、概述、列表显示“暂无启用的负责人”，不显示停用账号的姓名（19 S-1）")
        def v25():
            p5 = admin.project(f"无启用负责人{RUN}", users["owner3"])
            from app.db import connect
            with connect() as connection:  # 存量数据：主负责人账号在 A05 之前就已停用
                connection.execute("update users set status = 'disabled' where id = %s", (users["owner3"],))
                connection.commit()
            open_overview(page, p5["id"])
            expect(page.get_by_test_id("frame-no-owner")).to_have_text("暂无启用的负责人")
            expect(page.get_by_test_id("frame-owners")).to_have_text("负责人：—")
            expect(page.get_by_test_id("project-no-owner")).to_be_visible()
            expect(page.get_by_test_id("members-no-owner")).to_contain_text("请把一位启用的成员设为负责人")
            shot(page, "A05-11-暂无启用的负责人")
            page.goto(f"{WEB}/projects?name={RUN}")
            expect(page.locator(f'[data-testid=project-no-owner][data-project-id="{p5["id"]}"]')).to_be_visible(timeout=30000)
        v25()

        @step("V26 被移出的人停在概述页和工作台：切回窗口即更新；工作台提示条有［返回项目列表］、空闲时不提下载（20 I-6、L-2，21 S-4）")
        def v26():
            p6 = admin.project(f"移出即更新{RUN}", users["admin"])
            admin.set_role(p6["id"], users["lisi"], "editor")
            login(page2, f"lisi{RUN}")
            open_overview(page2, p6["id"])
            snapshot = admin.call("GET", f"/projects/{p6['id']}/members")
            admin.call("DELETE", f"/projects/{p6['id']}/members/{users['lisi']}?expectedMembersRevision={snapshot['membersRevision']}")
            page2.evaluate("window.dispatchEvent(new Event('focus'))")
            expect(page2.get_by_text("项目不存在或你无权访问。")).to_be_visible(timeout=15000)
            # 工作台：空闲时（没有未保存内容）被移出，下一次轮询发现
            doc6 = make_document(admin, p6["id"])
            admin.set_role(p6["id"], users["lisi"], "editor")
            page2.goto(f"{WEB}/projects/{p6['id']}/documents/{doc6['id']}")
            expect(page2.locator(".paper-content")).to_have_attribute("contenteditable", "true", timeout=60000)
            snapshot = admin.call("GET", f"/projects/{p6['id']}/members")
            admin.call("DELETE", f"/projects/{p6['id']}/members/{users['lisi']}?expectedMembersRevision={snapshot['membersRevision']}")
            page2.evaluate("document.dispatchEvent(new Event('visibilitychange'))")  # 可见状态下等同立即轮询
            banner = page2.get_by_test_id("save-blocked")
            expect(banner).to_be_visible(timeout=30000)
            expect(banner).to_contain_text("项目不存在或你无权访问。")
            expect(banner).not_to_contain_text("无法保存")
            expect(banner).not_to_contain_text("可以下载")
            expect(banner.get_by_role("button", name="下载我的修改")).to_have_count(0)
            expect(banner.get_by_test_id("gone-back")).to_have_text("返回项目列表")
            shot(page2, "A05-12-工作台-被移出（空闲）")
            banner.get_by_test_id("gone-back").click()
            page2.wait_for_url(re.compile(r"/projects$"), timeout=10000)
        v26()

        browser.close()
    failed = [item for item in results if not item[1]]
    print(f"\n{len(results) - len(failed)}/{len(results)} passed")
    for name, _, message in failed:
        print(f"FAILED {name}: {message}")
    return 1 if failed else 0


def _user_status(user) -> str:
    from app.db import connect
    with connect() as connection:
        return str(connection.execute("select status from users where id = %s", (user,)).fetchone()[0])


if __name__ == "__main__":
    sys.exit(main())
