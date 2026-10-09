"""Stored test observations supplement the current Word without losing history."""

from io import BytesIO
import json
import struct
from uuid import uuid4
from zipfile import ZipFile
import zlib

from docx import Document
from docx.shared import Cm
from fastapi.testclient import TestClient

from app.config import clear_settings_cache
from app.db import connect, migrate
from app.files import store_bytes
from app.identity import create_user
from app.main import create_app
from app.test_executor import run_directory


def test_stored_results_insert_into_current_word_with_evidence_versions_and_guards(tmp_path, monkeypatch):
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path / 'results.db'}")
    monkeypatch.setenv("FILE_DATA_DIR", str(tmp_path / "files"))
    monkeypatch.setenv("AI_API_KEY", "fixture-api-secret")
    clear_settings_cache()
    try:
        migrate()
        owner = create_user(account="result-owner", name="文档验收", password="result-password-1", roles=["system_admin"])
        projects = [uuid4(), uuid4()]
        with connect() as connection:
            for index, project in enumerate(projects):
                connection.execute("INSERT INTO projects(id,name,owner_id) VALUES(%s,%s,%s)", (project, f"验收项目{index}", owner))
            connection.commit()
        source = Document()
        source.sections[0].left_margin = Cm(3)
        source.sections[0].header.paragraphs[0].text = "原页眉"
        source.add_heading("测试结果", 1)
        source.add_paragraph("原正文。")
        table = source.add_table(rows=1, cols=2)
        table.style = "Light Shading Accent 1"
        table.cell(0, 0).text, table.cell(0, 1).text = "原表", "必须保留"
        source.add_heading("后续章节", 1)
        source.add_paragraph("后续原文。")
        data = BytesIO()
        source.save(data)
        media = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        stored = store_bytes(name="测试文档.docx", media_type=media, purpose="attachment", project_id=projects[0], content=data.getvalue())
        other_file = store_bytes(name="其他文档.docx", media_type=media, purpose="attachment", project_id=projects[1], content=data.getvalue())

        def chunk(kind, value):
            return struct.pack(">I", len(value)) + kind + value + struct.pack(">I", zlib.crc32(kind + value))

        png = (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 1920, 1080, 8, 2, 0, 0, 0))
               + chunk(b"IDAT", zlib.compress((b"\0" + b"\x80\x80\x80" * 1920) * 1080)) + chunk(b"IEND", b""))
        run_id, case_ids = str(uuid4()), [str(uuid4()) for _ in range(3)]
        cases = [
            {"caseId": case_ids[0], "title": "项目查询", "kind": "functional", "status": "passed", "expected": "无匹配项目", "planReference": "计划第1节", "sourceReference": "ProjectList.vue", "steps": [
                {"action": "assert_text", "description": "检查空态", "status": "passed", "actual": "已出现无匹配项目", "screenshot": "full.png", "reportScreenshot": "old-detail.png", "focusRect": {"x": 480, "y": 270, "width": 960, "height": 270}},
            ]},
            {"caseId": case_ids[1], "title": "函数测试", "kind": "unit", "status": "failed", "expected": "返回3", "error": "实际返回2 fixture-password", "log": "FAILED test_add\nassert 2 == 3\nfixture-api-secret\n<script>literal log</script>", "steps": []},
            {"caseId": case_ids[2], "title": "并发测试", "kind": "concurrency", "status": "passed", "expected": "无失败请求", "metrics": {"concurrency": 5, "requested": 20, "completed": 20, "failed": 0, "averageMs": 10, "p95Ms": 15.65, "maxLatencyMs": 1000}, "steps": []},
        ]
        with TestClient(create_app()) as client:
            csrf = client.get("/api/v1/auth/csrf").json()["data"]["csrfToken"]
            authenticated = client.post("/api/v1/auth/login", json={"account": "result-owner", "password": "result-password-1"}, headers={"X-CSRF-Token": csrf})
            headers = {"X-CSRF-Token": authenticated.json()["data"]["csrfToken"]}

            def create_doc(project, file, name):
                response = client.post(f"/api/v1/projects/{project}/documents", json={"name": name, "sourceType": "upload", "wordFileId": str(file.id)}, headers=headers)
                assert response.status_code == 200, response.text
                return response.json()["data"]

            document = create_doc(projects[0], stored, "测试结果文档")
            other = create_doc(projects[1], other_file, "其他项目文档")
            did = document["id"]
            edited = client.patch(f"/api/v1/documents/{did}/content", json={"expectedRevision": 0, "markdownContent": document["markdownContent"].replace("原正文。", "用户已编辑的正文。")}, headers=headers)
            assert edited.status_code == 200, edited.text
            created = client.post("/api/v1/test-tasks", json={"name": "真实结果验收夹具", "projectId": str(projects[0]), "planFileId": str(stored.id), "testPassword": "fixture-password"}, headers=headers)
            assert created.status_code == 201, created.text
            task = created.json()["data"]
            snapshot = {**task, "testPassword": "fixture-password"}
            with connect() as connection:
                connection.execute("INSERT INTO test_task_runs(id,task_id,status,snapshot_json,result_json,ended_at) VALUES(%s,%s,'failed',%s,%s,CURRENT_TIMESTAMP)",
                                   (run_id, task["id"], json.dumps(snapshot), json.dumps({"status": "failed", "cases": cases})))
                connection.commit()
            folder = run_directory(run_id)
            (folder / "full.png").write_bytes(png)
            (folder / "old-detail.png").write_bytes(b"must never replace original evidence")
            url = f"/api/v1/test-tasks/{task['id']}/documents/insert"
            body = {"runId": run_id, "documentId": did, "expectedRevision": 1, "headingIndex": 0, "caseIds": case_ids}
            assert client.post(url, json={**body, "expectedRevision": 0}, headers=headers).status_code == 409
            assert client.post(url, json={**body, "documentId": other["id"]}, headers=headers).status_code == 422
            assert client.post(url, json={**body, "caseIds": [str(uuid4())]}, headers=headers).status_code == 422
            assert client.post(url, json={**body, "caseIds": [case_ids[0], case_ids[0]]}, headers=headers).status_code == 422
            assert client.post(url, json={**body, "headingIndex": 99}, headers=headers).status_code == 422
            # Missing passed-case evidence must not partially save the unit/performance results.
            (folder / "full.png").unlink()
            assert client.post(url, json=body, headers=headers).status_code == 422
            assert client.get(f"/api/v1/documents/{did}").json()["data"]["revision"] == 1
            (folder / "full.png").write_bytes(png)
            inserted = client.post(url, json=body, headers=headers)
            assert inserted.status_code == 200, inserted.text
            assert inserted.json()["data"] == {"documentId": did, "projectId": str(projects[0]), "documentName": "测试结果文档", "revision": 2, "insertedCount": 3, "alreadyInsertedCount": 0}
            current = client.get(f"/api/v1/documents/{did}").json()["data"]
            assert current["originalDocxFileId"] == str(stored.id)
            assert "用户已编辑的正文。" in current["markdownContent"]
            assert current["markdownContent"].index("实测记录") < current["markdownContent"].index("# 后续章节")
            assert "实际结果：失败" in current["markdownContent"] and "assert 2 == 3" in current["markdownContent"]
            assert "fixture-password" not in json.dumps(current) and "fixture-api-secret" not in json.dumps(current)
            assert "&lt;script&gt;literal log&lt;/script&gt;" in current["markdownContent"]
            assert len(current["formatTemplate"]["images"]) == 1 and len(current["formatTemplate"]["testResults"]) == 3
            assert current["outline"][0]["children"][0]["text"].endswith("实测记录")
            duplicate = client.post(url, json=body, headers=headers)
            assert duplicate.status_code == 200 and duplicate.json()["data"]["revision"] == 2
            assert duplicate.json()["data"]["insertedCount"] == 0 and duplicate.json()["data"]["alreadyInsertedCount"] == 3
            versions = client.get(f"/api/v1/documents/{did}/versions").json()["data"]["items"]
            assert versions[0]["kind"] == "test_results" and len(versions) == 3
            exported = client.post(f"/api/v1/documents/{did}/export", json={"expectedRevision": 2}, headers=headers)
            assert exported.status_code == 200, exported.text
            content = client.get(f"/api/v1/files/{exported.json()['data']['fileId']}/content").content
            word = Document(BytesIO(content))
            assert word.sections[0].left_margin == source.sections[0].left_margin
            assert word.sections[0].header.paragraphs[0].text == "原页眉"
            assert word.tables[0].cell(0, 1).text == "必须保留" and word.tables[1].style == word.tables[0].style
            assert word.tables[1].cell(6, 1).text == "15.65"
            assert len(word.inline_shapes) == 1
            picture = word.inline_shapes[0]._inline.xpath("ancestor::w:p")[0]
            actual, caption = list(picture.itersiblings(preceding=True))[:2]
            assert caption.xpath('./w:pPr/w:keepNext') and actual.xpath('./w:pPr/w:keepNext')
            assert not picture.xpath('./w:pPr/w:keepNext')
            following = next(paragraph for paragraph in word.paragraphs if paragraph.text == "类型：单元测试")
            assert not following._p.xpath('./w:pPr/w:keepNext')
            crop = word.inline_shapes[0]._inline.graphic.graphicData.pic.blipFill.find("{http://schemas.openxmlformats.org/drawingml/2006/main}srcRect")
            assert dict(crop.attrib) == {"l": "25000", "t": "25000", "r": "25000", "b": "50000"}
            with ZipFile(BytesIO(data.getvalue())) as before, ZipFile(BytesIO(content)) as after:
                assert before.read("word/styles.xml") == after.read("word/styles.xml")
                assert [after.read(name) for name in after.namelist() if name.startswith("word/media/")] == [png]
            # Restoring the prior version removes provenance/images and permits a new supplementation.
            restored = client.post(f"/api/v1/documents/{did}/versions/1/restore", json={"expectedRevision": 2}, headers=headers)
            assert restored.status_code == 200, restored.text
            plain = client.get(f"/api/v1/documents/{did}").json()["data"]
            assert plain["revision"] == 3 and not plain["formatTemplate"].get("testResults")
            again = client.post(url, json={**body, "expectedRevision": 3, "headingIndex": None, "caseIds": [case_ids[0]]}, headers=headers)
            assert again.status_code == 200 and again.json()["data"]["insertedCount"] == 1
            reinserted = client.get(f"/api/v1/documents/{did}").json()["data"]
            assert reinserted["markdownContent"].index("实测记录") > reinserted["markdownContent"].index("# 后续章节")
            with connect() as connection:
                connection.execute("UPDATE test_task_runs SET status='running' WHERE id=%s", (run_id,))
                connection.commit()
            assert client.post(url, json={**body, "expectedRevision": 4}, headers=headers).status_code == 409
            assert client.get(f"/api/v1/documents/{did}").json()["data"]["revision"] == 4
    finally:
        clear_settings_cache()
