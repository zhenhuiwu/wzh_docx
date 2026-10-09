"""Edited Word export keeps the uploaded document's layout."""

from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from copy import deepcopy
from base64 import b64decode
from io import BytesIO
import pytest

from app.docx_export import edit_original_docx
from app.docx_parser import parse_docx, _table_to_markdown
from app.document_rules import SKILL_ROOT


def test_export_edits_original_docx_in_place(tmp_path):
    source = tmp_path / "source.docx"
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(3), Cm(2)
    section.header.paragraphs[0].text = "保留页眉"
    section.footer.paragraphs[0].text = "保留页脚"
    doc.add_heading("原有标题", level=1)
    paragraph = doc.add_paragraph()
    paragraph.add_run("保留").bold = True
    paragraph.add_run("原有正文")
    table = doc.add_table(rows=2, cols=2, style="Table Grid")
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "内容"
    table.cell(1, 0).text = "状态"
    table.cell(1, 1).text = "旧值"
    doc.save(source)

    original_markdown = parse_docx(str(source)).markdown
    unchanged = tmp_path / "unchanged.docx"
    edit_original_docx(original_markdown, source, unchanged)
    assert unchanged.read_bytes() == source.read_bytes()

    edited = tmp_path / "edited.docx"
    edit_original_docx(
        original_markdown.replace("旧值", "新值") + "\n\n## 新章节\n\n新增正文",
        source,
        edited,
    )
    result = Document(edited)
    original = Document(source)
    assert (result.sections[0].page_width, result.sections[0].page_height) == (
        original.sections[0].page_width, original.sections[0].page_height
    )
    assert (result.sections[0].top_margin, result.sections[0].bottom_margin) == (
        original.sections[0].top_margin, original.sections[0].bottom_margin
    )
    assert result.sections[0].header.paragraphs[0].text == "保留页眉"
    assert result.sections[0].footer.paragraphs[0].text == "保留页脚"
    assert result.paragraphs[1].runs[0].bold is True
    assert result.tables[0].style.name == "Table Grid"
    assert result.tables[0].cell(1, 1).text == "新值"
    assert [(p.text, p.style.name) for p in result.paragraphs[-2:]] == [
        ("新章节", "Heading 2"), ("新增正文", "Normal")
    ]


def test_only_test_screenshots_keep_their_caption_and_result_on_the_same_page(tmp_path):
    png = b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII=")
    source, output = tmp_path / "source.docx", tmp_path / "results.docx"
    word = Document()
    word.add_paragraph("步骤 0：原稿图片")
    word.add_paragraph("通过；原图保留")
    word.add_picture(BytesIO(png))
    word.save(source)
    content = parse_docx(str(source)).markdown + "\n\n步骤 1：实际操作\n\n通过；已验证\n\n![操作截图](docx-image:test-fixture)\n\n后续普通正文\n\n步骤 2：图表说明\n\n通过；图表不绑定\n\n![图表](docx-image:diagram-fixture)"
    edit_original_docx(content, source, output, image_sources={"docx-image:test-fixture": {"data": png, "focusRect": None}, "docx-image:diagram-fixture": png})
    result = Document(output)
    assert len(result.inline_shapes) == 3
    for index, shape in enumerate(result.inline_shapes):
        picture = shape._inline.xpath("ancestor::w:p")[0]
        preceding = list(picture.itersiblings(preceding=True))[:2]
        assert len(preceding) == 2
        assert all(bool(paragraph.xpath('./w:pPr/w:keepNext')) is (index == 1) for paragraph in preceding)
        assert not picture.xpath('./w:pPr/w:keepNext')
    following = next(paragraph for paragraph in result.paragraphs if paragraph.text == "后续普通正文")
    assert not following._p.xpath('./w:pPr/w:keepNext')


def test_ai_content_inherits_local_body_and_table_prototypes(tmp_path):
    source = tmp_path / "source.docx"
    doc = Document()
    cover = doc.add_paragraph("封面")
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.runs[0].font.size = Pt(28)
    doc.add_heading("正文", level=1)
    body = doc.add_paragraph("原有正文")
    body.alignment = WD_ALIGN_PARAGRAPH.LEFT
    body.paragraph_format.first_line_indent = Cm(0.8)
    body.runs[0].font.size = Pt(12)
    for text in ("第一项", "第二项"):
        doc.add_paragraph(text, style="List Number")
    table = doc.add_table(rows=3, cols=2, style="Table Grid")
    table.cell(0, 0).merge(table.cell(0, 1)).text = "合并表头"
    table.cell(1, 0).text, table.cell(1, 1).text = "名称", "内容"
    table.cell(2, 0).text, table.cell(2, 1).text = "原项目", "原内容"
    for cell in table.rows[2].cells:
        cell.paragraphs[0].runs[0].font.size = Pt(10)
        cell.paragraphs[0].paragraph_format.space_after = Pt(2)
    header_xml = table.rows[0]._tr.xml
    doc.save(source)
    original = parse_docx(str(source)).markdown
    target = original.replace("原有正文", "原有正文\n\n新增**重点**及普通文字。")
    target = target.replace("1. 第二项", "1. 第二项\n1. **新增项**")
    target = target.rstrip() + "\n| 新项目 | 新内容 |\n"
    output = tmp_path / "output.docx"
    edit_original_docx(target, source, output)
    result = Document(output)
    added = next(p for p in result.paragraphs if p.text.startswith("新增重点"))
    assert added.alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert added.paragraph_format.first_line_indent == Document(source).paragraphs[2].paragraph_format.first_line_indent
    assert all(run.font.size == Pt(12) for run in added.runs)
    assert next(run for run in added.runs if run.text == "重点").bold is True
    added_list = next(p for p in result.paragraphs if p.text == "新增项")
    assert added_list.style.name == "List Number" and added_list.runs[0].bold
    assert result.tables[0].rows[0]._tr.xml == header_xml
    assert result.tables[0].style.name == "Table Grid" and len(result.tables[0].rows) == 4
    cell = result.tables[0].cell(3, 1)
    assert cell.text == "新内容" and cell.paragraphs[0].runs[0].font.size == Pt(10)
    assert cell.paragraphs[0].paragraph_format.space_after == Pt(2)
    # Shrinking data rows must also retain the merged header.
    edit_original_docx(original.replace("| 原项目 | 原内容 |", ""), source, output)
    shrunk = Document(output)
    assert len(shrunk.tables[0].rows) == 2 and shrunk.tables[0].rows[0]._tr.xml == header_xml


def test_bundled_438c_templates_keep_cover_toc_and_grow_reference_table(tmp_path):
    for source in SKILL_ROOT.glob("*/documents/*.docx"):
        doc = Document(source)
        original = parse_docx(str(source)).markdown
        output = tmp_path / "template.docx"
        edit_original_docx(original, source, output)
        assert output.read_bytes() == source.read_bytes()
        reference_index = next(index for index, table in enumerate(doc.tables)
            if [cell.text.strip() for cell in table.rows[0].cells] == ["序号", "名称", "版本", "发布日期", "来源"])
        reference = doc.tables[reference_index]
        table_md = _table_to_markdown(reference).strip()
        edited = original.replace(table_md, table_md + "\n| 3 | 测试规范 | V1 | 待补充 | 测试组 |")
        assert edited != original
        edit_original_docx(edited, source, output)
        result = Document(output)
        assert len(result.tables[reference_index].rows) == len(reference.rows) + 1
        assert result.tables[reference_index].rows[0]._tr.xml == reference.rows[0]._tr.xml
        assert result.tables[0]._tbl.xml == doc.tables[0]._tbl.xml
        assert [p._p.xml for p in result.paragraphs if p._p.xpath(".//w:fldChar")] == [
            p._p.xml for p in doc.paragraphs if p._p.xpath(".//w:fldChar")]


def test_export_preserves_section_boundaries_and_rejects_protected_deletions(tmp_path):
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_heading("第一节", level=1)
    boundary = doc.add_paragraph("分节前的正文")
    boundary._p.get_or_add_pPr().append(deepcopy(doc.sections[0]._sectPr))
    doc.add_heading("第二节", level=1)
    field = doc.add_paragraph("保留域")
    simple = OxmlElement("w:fldSimple")
    simple.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}instr", "DATE")
    field._p.append(simple)
    table = doc.add_table(rows=2, cols=2, style="Table Grid")
    table.cell(0, 0).merge(table.cell(0, 1)).text = "合并表头"
    table.cell(1, 0).text, table.cell(1, 1).text = "项目", "原值"
    doc.add_heading("第三节", level=1)
    doc.add_paragraph("尾部正文")
    doc.save(source)
    original = parse_docx(str(source)).markdown
    output = tmp_path / "output.docx"
    edit_original_docx(original.replace("分节前的正文", "分节前的正文\n\n新增正文"), source, output)
    result = Document(output)
    assert len(result.sections) == len(Document(source).sections)
    assert result.tables[0]._tbl.xml == Document(source).tables[0]._tbl.xml
    assert next(p for p in result.paragraphs if p.text == "保留域")._p.xml == field._p.xml
    output.unlink()
    for target in (
        original.replace("保留域", "修改域"),
        original.replace("保留域\n\n", ""),
        original.replace(_table_to_markdown(table).strip(), ""),
        original.replace("分节前的正文\n\n", ""),
    ):
        with pytest.raises(ValueError):
            edit_original_docx(target, source, output)
        assert not output.exists()


def test_export_repeated_images_keep_their_occurrences(tmp_path):
    source, output = tmp_path / "source.docx", tmp_path / "output.docx"
    png = b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jP1kAAAAASUVORK5CYII=")
    doc = Document()
    doc.add_heading("图片检查", level=1)
    doc.add_paragraph("原有正文")
    pictures = doc.add_paragraph()
    pictures.add_run().add_picture(BytesIO(png))
    pictures.add_run().add_picture(BytesIO(png))
    doc.add_picture(BytesIO(png))
    doc.save(source)
    parsed = parse_docx(str(source))
    assert len(parsed.images) == 1
    assert parsed.markdown.count("![图片 1](image_0.png)") == 3
    edit_original_docx(parsed.markdown, source, output)
    assert output.read_bytes() == source.read_bytes()
    edit_original_docx(parsed.markdown.replace("原有正文", "修改后的正文"), source, output)
    original, result = Document(source), Document(output)
    assert len(result.inline_shapes) == 3
    assert [p._p.xml for p in result.paragraphs if p._p.xpath(".//w:drawing")] == [
        p._p.xml for p in original.paragraphs if p._p.xpath(".//w:drawing")]
