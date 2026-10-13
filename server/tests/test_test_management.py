"""Plan/test history is durable and cannot confuse blocked work with passes."""

import json
import sys
from io import BytesIO
from types import SimpleNamespace

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.config import clear_settings_cache
from app.db import connect, migrate
from app.identity import create_user
from app.main import create_app
from app.test_management import _interrupt_result, _source_path, _validate_cases, recover_interrupted_tasks, source_files
from app.test_executor import execute_run as real_execute_run, initial_run_result


def test_plan_scope_api_revision_and_truthful_history(tmp_path, monkeypatch):
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path / 'tests.db'}")
    monkeypatch.setenv("FILE_DATA_DIR", str(tmp_path / "files"))
    monkeypatch.setenv("AI_API_KEY", "")
    clear_settings_cache()
    try:
        assert "104_test_management.sql" in migrate()
        owner = create_user(account="test-admin", name="测试", password="test-password-1", roles=["system_admin"])
        source = tmp_path / "source"
        source.mkdir()
        (source / "feature.py").write_text("def add(a, b): return a + b\n", encoding="utf-8")
        (source / "index.html").write_text('<label>Name<input></label>', encoding="utf-8")
        (source / ".env").write_text("API_KEY=do-not-read", encoding="utf-8")
        (source / "secrets.json").write_text('{}', encoding="utf-8")
        (source / "node_modules").mkdir()
        (source / "node_modules" / "ignore.js").write_text("ignored", encoding="utf-8")
        for directory in (".uv-cache", ".pytest_cache", ".qa-demo"):
            (source / directory).mkdir()
            (source / directory / "ignore.py").write_text("ignored", encoding="utf-8")
        source_task = {"backendRoot": str(source)}
        assert source_files(source_task, "backend") == ["feature.py", "index.html"]
        shared = {"backendRoot": str(source), "frontendRoot": str(source)}
        unit = {"cases": [{"title": "同一物理源码", "kind": "unit", "planReference": "计划", "sourceReference": "backend feature.py add", "expected": "返回3", "steps": [{"action": "unit", "runner": "pytest", "root": "backend", "testFile": "test_add.py", "testCode": "from feature import add\ndef test_add():\n assert add(1,2)==3"}]}]}
        assert not _validate_cases(unit, shared, [("frontend", "feature.py")])[0]["needsInput"]
        for path in ("../outside.py", ".env", "secrets.json", "node_modules/ignore.js", ".uv-cache/ignore.py", ".pytest_cache/ignore.py", ".qa-demo/ignore.py"):
            with pytest.raises(ValueError):
                _source_path(source_task, "backend", path)

        with TestClient(create_app()) as client:
            csrf = client.get("/api/v1/auth/csrf").json()["data"]["csrfToken"]
            login = client.post("/api/v1/auth/login", json={"account": "test-admin", "password": "test-password-1"}, headers={"X-CSRF-Token": csrf})
            assert login.status_code == 200
            headers = {"X-CSRF-Token": login.json()["data"]["csrfToken"]}
            project = client.post("/api/v1/projects", json={"name": "测试任务验收", "ownerId": str(owner), "status": "active"}, headers={**headers, "Idempotency-Key": "test-management-project"})
            assert project.status_code == 201, project.text
            project_id = project.json()["data"]["id"]
            word, content = Document(), BytesIO()
            word.add_heading("测试计划", 1)
            word.add_paragraph("对加法函数进行单元测试，1+2应当为3；缺少接口环境的用例应当阻塞。")
            word.save(content)
            upload = client.post("/api/v1/files", files={"file": ("测试计划.docx", content.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                                 data={"purpose": "attachment", "projectId": project_id}, headers={**headers, "Idempotency-Key": "test-plan-word"})
            assert upload.status_code == 201, upload.text
            create = client.post("/api/v1/test-tasks", json={"name": "加法验收", "projectId": project_id, "planFileId": upload.json()["data"]["id"], "backendRoot": str(source),
                                "testUsername": "tester", "testPassword": "private-test-password"}, headers=headers)
            assert create.status_code == 201, create.text
            task = create.json()["data"]
            assert "testPassword" not in task and task["planName"] == "测试计划.docx"
            assert task["executionMode"] == "agent"
            task_id = task["id"]
            assert client.post(f"/api/v1/test-tasks/{task_id}/stop", headers=headers).status_code == 409
            with connect() as connection:
                connection.execute("UPDATE test_tasks SET status='analyzing' WHERE id=%s", (task_id,))
                connection.commit()
            analyzing_stop = client.post(f"/api/v1/test-tasks/{task_id}/stop", headers=headers)
            assert analyzing_stop.status_code == 409 and "分析" in analyzing_stop.json()["error"]["message"]
            with connect() as connection:
                connection.execute("UPDATE test_tasks SET status='draft' WHERE id=%s", (task_id,))
                connection.commit()
            invalid_mode = client.post("/api/v1/test-tasks", json={"name": "无效模式", "projectId": project_id,
                "planFileId": upload.json()["data"]["id"], "executionMode": "shell"}, headers=headers)
            assert invalid_mode.status_code == 422
            with connect() as connection:
                stored = json.loads(connection.execute("SELECT config_json FROM test_tasks WHERE id=%s", (task_id,)).fetchone()[0])
                assert stored["executionMode"] == "agent"
                del stored["executionMode"]  # Existing tasks keep their original fixed-step behavior.
                connection.execute("UPDATE test_tasks SET config_json=%s WHERE id=%s", (json.dumps(stored), task_id))
                connection.commit()
            assert client.get(f"/api/v1/test-tasks/{task_id}").json()["data"]["executionMode"] == "steps"
            assert client.get("/api/v1/test-tasks?page=1&pageSize=1&name=加法").json()["data"]["total"] == 1
            assert client.get("/api/v1/test-tasks?page=2&pageSize=1").json()["data"]["items"] == []
            assert client.get("/api/v1/test-tasks?pageSize=51").status_code == 422

            raw = {"cases": [{"title": "缺少断言", "kind": "functional", "planReference": "计划第1章", "expected": "页面显示正常", "steps": [{"action": "open", "url": "http://localhost:5173"}]}]}
            cases = _validate_cases(raw, {**task, "frontendUrl": "http://localhost:5173"})
            assert "独立断言" in cases[0]["needsInput"]
            cases[0]["needsInput"] = ""  # Admission is covered separately; this test exercises runner outcomes.
            monkeypatch.setattr("app.test_management.analyze_task", lambda _task: cases)
            analyzed = client.post(f"/api/v1/test-tasks/{task_id}/analyze", headers=headers)
            assert analyzed.status_code == 200, analyzed.text
            analyzed_task = analyzed.json()["data"]
            assert analyzed_task["revision"] == 1 and analyzed_task["status"] == "ready"
            assert client.post(f"/api/v1/test-tasks/{task_id}/run", json={"expectedRevision": 0}, headers=headers).status_code == 409
            case_id = cases[0]["id"]
            changed = client.patch(f"/api/v1/test-tasks/{task_id}/cases/{case_id}", json={"enabled": False, "expectedRevision": 1}, headers=headers)
            assert changed.status_code == 200 and changed.json()["data"]["revision"] == 2
            assert client.post(f"/api/v1/test-tasks/{task_id}/run", json={"expectedRevision": 2}, headers=headers).status_code == 422
            client.patch(f"/api/v1/test-tasks/{task_id}/cases/{case_id}", json={"enabled": True, "expectedRevision": 2}, headers=headers)
            fake_executor = SimpleNamespace(initial_run_result=initial_run_result, execute_run=lambda _task, selected, _run_id, _cancelled: {"status": "blocked", "cases": [{"caseId": selected[0]["id"], "title": selected[0]["title"], "status": "blocked", "steps": [], "error": "private-test-password"}]})
            monkeypatch.setitem(sys.modules, "app.test_executor", fake_executor)
            executed = client.post(f"/api/v1/test-tasks/{task_id}/run", json={"expectedRevision": 3}, headers=headers)
            assert executed.status_code == 202, executed.text
            result = client.get(f"/api/v1/test-tasks/{task_id}").json()["data"]
            assert result["status"] == "blocked" and result["runs"][0]["status"] == "blocked"
            assert "private-test-password" not in json.dumps(result)
            assert result["runs"][0]["result"]["cases"][0]["error"] == "[已隐藏]"
            expected_changed = client.patch(f"/api/v1/test-tasks/{task_id}/cases/{case_id}", json={"expected": "新的预期结果", "expectedRevision": 3}, headers=headers)
            assert expected_changed.status_code == 200
            changed_case = expected_changed.json()["data"]["cases"][0]
            assert changed_case["expected"] == "新的预期结果" and "断言尚未同步" in changed_case["needsInput"]
            assert changed_case["steps"] == cases[0]["steps"]
            fake_executor.execute_run = real_execute_run
            rejected = client.post(f"/api/v1/test-tasks/{task_id}/run", json={"expectedRevision": 4}, headers=headers)
            assert rejected.status_code == 422 and "断言尚未同步" in rejected.json()["error"]["message"]
            assert client.get(f"/api/v1/test-tasks/{task_id}").json()["data"]["runs"][0]["id"] == result["runs"][0]["id"]
            assert client.post(f"/api/v1/test-tasks/{task_id}/analyze", headers=headers).json()["data"]["revision"] == 5

            def crash_with_saved_evidence(_task, selected, run_id, _cancelled):
                from app.test_management import save_run_progress
                save_run_progress(run_id, {"status": "running", "cases": [{"caseId": selected[0]["id"], "title": "中途失败", "kind": "functional", "status": "running", "steps": [{"action": "open", "status": "passed", "screenshot": "retained.png"}]}]})
                raise RuntimeError("private-test-password")
            fake_executor.execute_run = crash_with_saved_evidence
            client.post(f"/api/v1/test-tasks/{task_id}/run", json={"expectedRevision": 5}, headers=headers)
            failed_run = client.get(f"/api/v1/test-tasks/{task_id}").json()["data"]["runs"][0]
            assert failed_run["status"] == "failed" and failed_run["result"]["cases"][0]["status"] == "failed"
            assert failed_run["result"]["cases"][0]["steps"][0]["screenshot"] == "retained.png"
            assert "private-test-password" not in json.dumps(failed_run)

            fake_executor.execute_run = lambda *_: {"status": "completed", "cases": []}
            client.post(f"/api/v1/test-tasks/{task_id}/run", json={"expectedRevision": 5}, headers=headers)
            assert client.get(f"/api/v1/test-tasks/{task_id}").json()["data"]["runs"][0]["status"] == "failed"

            def recovered_worker(_task, selected, _run_id, _cancelled):
                recover_interrupted_tasks()
                return {"status": "completed", "cases": [{"caseId": selected[0]["id"], "status": "passed"}]}
            fake_executor.execute_run = recovered_worker
            client.post(f"/api/v1/test-tasks/{task_id}/run", json={"expectedRevision": 5}, headers=headers)
            interrupted = client.get(f"/api/v1/test-tasks/{task_id}").json()["data"]["runs"][0]
            assert interrupted["status"] == "failed" and "重启" in interrupted["result"]["error"]

            def late_stop(_task, selected, run_id, _cancelled):
                with connect() as connection:
                    connection.execute("UPDATE test_task_runs SET stop_requested=1 WHERE id=%s", (run_id,))
                    connection.commit()
                return {"status": "completed", "cases": [{"caseId": selected[0]["id"], "status": "passed"}]}
            fake_executor.execute_run = late_stop
            client.post(f"/api/v1/test-tasks/{task_id}/run", json={"expectedRevision": 5}, headers=headers)
            assert client.get(f"/api/v1/test-tasks/{task_id}").json()["data"]["runs"][0]["status"] == "stopped"
            with connect() as connection:
                run = connection.execute("SELECT snapshot_json FROM test_task_runs WHERE task_id=%s", (task_id,)).fetchone()
                assert json.loads(run[0])["testPassword"] == "private-test-password"
                connection.execute("UPDATE test_tasks SET status='running' WHERE id=%s", (task_id,))
                connection.execute("UPDATE test_task_runs SET status='running' WHERE id=(SELECT id FROM test_task_runs WHERE task_id=%s ORDER BY started_at DESC,rowid DESC LIMIT 1)", (task_id,))
                connection.commit()
            recover_interrupted_tasks()
            recovered = client.get(f"/api/v1/test-tasks/{task_id}").json()["data"]
            assert recovered["status"] == "failed" and recovered["runs"][0]["status"] == "failed"
            assert "重启" in recovered["runs"][0]["result"]["error"]
    finally:
        clear_settings_cache()


def test_edit_configuration_and_run_admission_preserve_history(tmp_path, monkeypatch):
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path / 'editing.db'}")
    monkeypatch.setenv("FILE_DATA_DIR", str(tmp_path / "files"))
    monkeypatch.setenv("AI_API_KEY", "")
    clear_settings_cache()
    try:
        migrate()
        owner = create_user(account="edit-admin", name="配置验收", password="edit-password-1", roles=["system_admin"])
        source = tmp_path / "source"
        source.mkdir()
        word, content = Document(), BytesIO()
        word.add_paragraph("查询项目名称。未提供后端源码时不得执行单元测试。")
        word.save(content)
        with TestClient(create_app()) as client:
            csrf = client.get("/api/v1/auth/csrf").json()["data"]["csrfToken"]
            login = client.post("/api/v1/auth/login", json={"account": "edit-admin", "password": "edit-password-1"}, headers={"X-CSRF-Token": csrf})
            assert login.status_code == 200
            headers = {"X-CSRF-Token": login.json()["data"]["csrfToken"]}
            project = client.post("/api/v1/projects", json={"name": "配置与恢复验收", "ownerId": str(owner), "status": "active"}, headers={**headers, "Idempotency-Key": "edit-project"})
            assert project.status_code == 201
            project_id = project.json()["data"]["id"]

            def upload(name):
                response = client.post("/api/v1/files", files={"file": (name, content.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                                       data={"purpose": "attachment", "projectId": project_id}, headers={**headers, "Idempotency-Key": "edit-" + name.encode("utf-8").hex()})
                assert response.status_code == 201, response.text
                return response.json()["data"]["id"]

            plan, replacement, report = upload("原计划.docx"), upload("新计划.docx"), upload("原报告.docx")
            payload = {"name": "原任务", "projectId": project_id, "planFileId": plan, "reportFileId": report,
                       "frontendRoot": str(source), "frontendUrl": "http://localhost:5173", "backendUrl": "http://localhost:8001",
                       "testUsername": "tester", "testPassword": "retained-private-password", "dataScope": "允许只读查询"}
            created = client.post("/api/v1/test-tasks", json=payload, headers=headers)
            assert created.status_code == 201, created.text
            task = created.json()["data"]
            url = f"/api/v1/test-tasks/{task['id']}"
            cases = _validate_cases({"cases": [
                {"title": "只读查询", "kind": "functional", "planReference": "查询项目", "expected": "地址为项目列表",
                 "steps": [{"action": "open", "url": "/projects"}, {"action": "assert_url", "url": "/projects"}]},
                {"title": "缺少源码", "kind": "unit", "planReference": "函数测试", "expected": "函数返回符合预期", "steps": [], "needsInput": "需要后端源码目录"}
            ]}, task)
            interrupted = _interrupt_result(initial_run_result(cases), "服务重启中断了测试。")
            assert interrupted["status"] == "failed"
            assert all(case["status"] == "blocked" and "未执行" in case["error"] for case in interrupted["cases"])
            monkeypatch.setattr("app.test_management.analyze_task", lambda _task: cases)
            assert client.post(url + "/analyze", headers=headers).json()["data"]["revision"] == 1
            selected_ids = []

            def successful(_task, selected, _run_id, _cancelled):
                selected_ids.extend(case["id"] for case in selected)
                return {"status": "completed", "cases": [{"caseId": case["id"], "title": case["title"], "kind": case["kind"], "status": "passed", "steps": []} for case in selected]}

            monkeypatch.setitem(sys.modules, "app.test_executor", SimpleNamespace(initial_run_result=initial_run_result, execute_run=successful))
            started = client.post(url + "/run", json={"expectedRevision": 1}, headers=headers)
            assert started.status_code == 202, started.text
            assert selected_ids == [cases[0]["id"]]
            initial = started.json()["data"]["runs"][0]["result"]
            assert [case["status"] for case in initial["cases"]] == ["queued"]
            completed = client.get(url).json()["data"]
            assert completed["status"] == "completed" and completed["runs"][0]["status"] == "completed"
            with connect() as connection:
                snapshot = connection.execute("SELECT snapshot_json FROM test_task_runs WHERE task_id=%s", (task["id"],)).fetchone()[0]
            rejected = client.post(url + "/run", json={"expectedRevision": 1, "caseIds": [cases[1]["id"]]}, headers=headers)
            assert rejected.status_code == 422 and rejected.json()["error"]["code"] == "TEST_INPUT_MISSING"
            assert "需要后端源码目录" in rejected.json()["error"]["message"]
            assert len(client.get(url).json()["data"]["runs"]) == 1
            disabled = client.patch(url + f"/cases/{cases[1]['id']}", json={"enabled": False, "expectedRevision": 1}, headers=headers)
            assert disabled.status_code == 200 and disabled.json()["data"]["revision"] == 2
            assert client.post(url + "/run", json={"expectedRevision": 2, "caseIds": [cases[1]["id"]]}, headers=headers).status_code == 422
            before = client.get(url).json()["data"]
            unchanged = client.patch(url, json={"expectedRevision": 2, "name": " 原任务 ", "frontendRoot": str(source), "frontendUrl": "http://localhost:5173/", "planFileId": plan, "reportFileId": report}, headers=headers)
            assert unchanged.status_code == 200, unchanged.text
            assert unchanged.json()["data"]["revision"] == 2 and unchanged.json()["data"]["cases"] == before["cases"]
            renamed = client.patch(url, json={"expectedRevision": 2, "name": "改名后的任务"}, headers=headers)
            assert renamed.status_code == 200 and renamed.json()["data"]["revision"] == 3
            assert renamed.json()["data"]["cases"] == before["cases"] and renamed.json()["data"]["status"] == before["status"]
            assert client.patch(url, json={"expectedRevision": 2, "dataScope": "新范围"}, headers=headers).status_code == 409
            for invalid in ({"projectId": project_id}, {"frontendRoot": str(tmp_path / "missing")}, {"frontendUrl": "ftp://localhost"}, {"frontendUrl": None}, {"planFileId": None}, {"testPassword": None}):
                assert client.patch(url, json={"expectedRevision": 3, **invalid}, headers=headers).status_code == 422
            for status in ("running", "analyzing"):
                with connect() as connection:
                    connection.execute("UPDATE test_tasks SET status=%s WHERE id=%s", (status, task["id"]))
                    connection.commit()
                assert client.patch(url, json={"expectedRevision": 3, "dataScope": "新范围"}, headers=headers).status_code == 409
            with connect() as connection:
                connection.execute("UPDATE test_tasks SET status='ready' WHERE id=%s", (task["id"],))
                connection.commit()
            changed = client.patch(url, json={"expectedRevision": 3, "dataScope": "仅在隔离浏览器只读查询"}, headers=headers)
            assert changed.status_code == 200, changed.text
            changed_task = changed.json()["data"]
            assert changed_task["revision"] == 4 and changed_task["status"] == "new" and changed_task["cases"] == []
            assert "testPassword" not in changed_task and "retained-private-password" not in json.dumps(changed_task)
            with connect() as connection:
                config = json.loads(connection.execute("SELECT config_json FROM test_tasks WHERE id=%s", (task["id"],)).fetchone()[0])
                assert config["testPassword"] == "retained-private-password"
                assert connection.execute("SELECT snapshot_json FROM test_task_runs WHERE task_id=%s", (task["id"],)).fetchone()[0] == snapshot
            replaced = client.patch(url, json={"expectedRevision": 4, "planFileId": replacement, "reportFileId": None, "testPassword": ""}, headers=headers)
            assert replaced.status_code == 200, replaced.text
            assert replaced.json()["data"]["revision"] == 5 and replaced.json()["data"]["planName"] == "新计划.docx"
            assert replaced.json()["data"]["reportFileId"] is None
            with connect() as connection:
                config = json.loads(connection.execute("SELECT config_json FROM test_tasks WHERE id=%s", (task["id"],)).fetchone()[0])
                assert config["testPassword"] == ""
                assert connection.execute("SELECT snapshot_json FROM test_task_runs WHERE task_id=%s", (task["id"],)).fetchone()[0] == snapshot
            cleared_again = client.patch(url, json={"expectedRevision": 5, "reportFileId": "", "testPassword": ""}, headers=headers)
            assert cleared_again.status_code == 200 and cleared_again.json()["data"]["revision"] == 5
            monkeypatch.setattr("app.test_management.analyze_task", lambda _task: [cases[1]])
            assert client.post(url + "/analyze", headers=headers).json()["data"]["revision"] == 6
            assert client.post(url + "/run", json={"expectedRevision": 6}, headers=headers).status_code == 422
            assert len(client.get(url).json()["data"]["runs"]) == 1
    finally:
        clear_settings_cache()
