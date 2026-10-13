"""A05 成员管理（17 §2 / 03 决策）：成员名单独立版本号、最后负责人、停用账号、归档只读、审计、候选人、可见性。

这些用例全部针对 A05 新契约，在 0022d04（A05 之前）上会失败：那里没有 membersRevision、members:batch、
member-candidates、member-events、LAST_OWNER/PROJECT_ARCHIVED/USER_DISABLED 等。
"""

from __future__ import annotations

import random
import threading
from io import BytesIO
from uuid import uuid4

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.db import connect, migrate
from app.files import store_bytes
from app.identity import create_user

PASSWORD = "a05-member-pw"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _login(client, account):
    client.cookies.clear()
    token = client.get("/api/v1/auth/csrf").json()["data"]["csrfToken"]
    response = client.post("/api/v1/auth/login", json={"account": account, "password": PASSWORD}, headers={"X-CSRF-Token": token})
    assert response.status_code == 200, response.text
    return {"X-CSRF-Token": response.json()["data"]["csrfToken"]}


class Ctx:
    def __init__(self, client, users, project):
        self.client, self.users, self.project = client, users, project
        self.headers = {}

    def as_(self, name):
        self.headers = _login(self.client, f"a05-{name}")
        return self

    def get(self, path, **kw):
        return self.client.get(f"/api/v1{path}", **kw)

    def req(self, method, path, body=None):
        return self.client.request(method, f"/api/v1{path}", json=body, headers=self.headers)

    def members(self):
        response = self.get(f"/projects/{self.project}/members")
        assert response.status_code == 200, response.text
        return response.json()["data"]

    def rev(self):
        with connect() as connection:
            return int(connection.execute("select members_revision from projects where id = %s", (self.project,)).fetchone()[0])

    def project_rev(self):
        with connect() as connection:
            return int(connection.execute("select revision from projects where id = %s", (self.project,)).fetchone()[0])

    def put(self, who, role, revision=None):
        body = {"role": role, "expectedMembersRevision": self.rev() if revision is None else revision}
        return self.req("PUT", f"/projects/{self.project}/members/{self.users[who]}", body)

    def delete(self, who, revision=None):
        revision = self.rev() if revision is None else revision
        return self.req("DELETE", f"/projects/{self.project}/members/{self.users[who]}?expectedMembersRevision={revision}")

    def batch(self, names, role="editor", revision=None):
        body = {"userIds": [str(self.users[n]) for n in names], "role": role,
                "expectedMembersRevision": self.rev() if revision is None else revision}
        return self.req("POST", f"/projects/{self.project}/members:batch", body)

    def events(self):
        with connect() as connection:
            keys = ("action", "old_role", "new_role", "members_revision")
            return [{key: row[i] for i, key in enumerate(keys)} for row in connection.execute(
                "select action, old_role, new_role, members_revision from project_member_events where project_id = %s order by rowid",
                (self.project,)).fetchall()]

    def disable(self, who):
        self.as_("admin")
        return self.req("POST", f"/users/{self.users[who]}/disable", {"expectedRevision": _user_revision(self.users[who]), "reason": "A05 测试停用"})


def _user_revision(user_id):
    with connect() as connection:
        return int(connection.execute("select revision from users where id = %s", (user_id,)).fetchone()[0])


@pytest.fixture
def ctx():
    migrate()
    users = {}
    with connect() as connection:
        for name, display, roles in (("admin", "管理员甲", ["system_admin"]), ("alice", "张工", ["project_creator"]),
                                     ("bob", "李工", ["project_creator"]), ("carol", "王工", []), ("dave", "赵工", []),
                                     ("outsider", "外人", ["project_creator"]), ("auditor", "审计员", ["auditor"])):
            users[name] = create_user(account=f"a05-{name}", name=display, password=PASSWORD, roles=roles,
                                      department="测试部", allow_no_roles=True, connection=connection)
        connection.commit()
    with TestClient(__import__("app.main", fromlist=["create_app"]).create_app()) as client:
        context = Ctx(client, users, None)
        context.as_("alice")
        created = client.post("/api/v1/projects", headers={**context.headers, "Idempotency-Key": str(uuid4())},
                              json={"name": "成员项目", "ownerId": str(users["alice"]), "status": "active"})
        assert created.status_code == 201, created.text
        context.project = created.json()["data"]["id"]
        yield context


def _active_owner_count(project):
    with connect() as connection:
        return int(connection.execute(
            """select count(*) from project_members m join users u on u.id = m.user_id
               where m.project_id = %s and m.role = 'owner' and u.status = 'active'""", (project,)).fetchone()[0])


# ---------- 基本契约 ----------

def test_new_project_starts_with_members_revision_and_owner_audit(ctx):
    data = ctx.members()
    assert data["membersRevision"] == 1 and data["activeOwnerCount"] == 1 and data["canManageMembers"] is True
    assert [(item["name"], item["role"], item["isMe"]) for item in data["items"]] == [("张工", "owner", True)]
    # A07 钩子：本期分工计数恒为 0
    assert data["items"][0]["addedAt"] and data["items"][0]["assignments"] == {"sections": 0, "blocks": 0, "documents": 0}
    assert ctx.events() == [{"action": "add", "old_role": None, "new_role": "owner", "members_revision": 1}]
    detail = ctx.get(f"/projects/{ctx.project}").json()["data"]
    assert detail["membersRevision"] == 1
    assert [owner["name"] for owner in detail["owners"]] == ["张工"] and detail["owners"][0]["isPrimary"] is True


def test_add_change_remove_bump_members_revision_only(ctx):
    project_revision = ctx.project_rev()
    added = ctx.batch(["bob", "carol"], "editor")
    assert added.status_code == 200, added.text
    body = added.json()["data"]
    assert body["membersRevision"] == 2 and len(body["added"]) == 2 and body["failed"] == []
    changed = ctx.put("carol", "viewer")
    assert changed.status_code == 200, changed.text
    assert changed.json()["data"]["item"]["role"] == "viewer" and changed.json()["data"]["membersRevision"] == 3
    assert changed.json()["data"]["assignmentsCleared"] == {"sections": 0, "blocks": 0, "documents": 0}
    removed = ctx.delete("carol")
    assert removed.status_code == 200, removed.text
    assert removed.json()["data"]["membersRevision"] == 4 and removed.json()["data"]["removedSelf"] is False
    # R5：成员操作不改项目 revision，项目资料编辑不会因此 409
    assert ctx.project_rev() == project_revision
    patched = ctx.req("PATCH", f"/projects/{ctx.project}", {"expectedRevision": project_revision, "description": "改简介"})
    assert patched.status_code == 200, patched.text
    assert ctx.rev() == 4
    assert [event["action"] for event in ctx.events()] == ["add", "add", "add", "change_role", "remove"]
    # 移除非成员：幂等 204，不加版本、不写审计
    again = ctx.delete("carol")
    assert again.status_code == 204 and ctx.rev() == 4 and len(ctx.events()) == 5
    # 同角色：200，不加版本、不写审计
    same = ctx.put("bob", "editor")
    assert same.status_code == 200 and ctx.rev() == 4 and len(ctx.events()) == 5


def test_stale_members_revision_is_409_with_last_actor(ctx):
    assert ctx.batch(["bob"]).status_code == 200
    stale = ctx.put("bob", "viewer", revision=1)
    assert stale.status_code == 409
    error = stale.json()["error"]
    assert error["code"] == "REVISION_CONFLICT" and error["details"]["currentRevision"] == 2
    assert error["details"]["lastActor"]["name"] == "张工"
    assert ctx.delete("bob", revision=1).status_code == 409
    assert ctx.batch(["carol"], revision=1).status_code == 409
    assert ctx.rev() == 2


def test_missing_revision_and_mismatched_user_are_422(ctx):
    path = f"/projects/{ctx.project}/members/{ctx.users['bob']}"
    missing = ctx.req("PUT", path, {"role": "editor"})
    assert missing.status_code == 422
    assert {item["field"] for item in missing.json()["error"]["fieldErrors"]} == {"expectedMembersRevision"}
    assert ctx.req("DELETE", path).status_code == 422
    assert ctx.req("DELETE", path + "?expectedMembersRevision=abc").status_code == 422
    assert ctx.req("POST", f"/projects/{ctx.project}/members:batch",
                   {"userIds": [str(ctx.users["bob"])], "role": "editor"}).status_code == 422
    mismatch = ctx.req("PUT", path, {"role": "editor", "expectedMembersRevision": 1, "userId": str(ctx.users["carol"])})
    assert mismatch.status_code == 422
    unknown = ctx.req("PUT", f"/projects/{ctx.project}/members/{uuid4()}", {"role": "editor", "expectedMembersRevision": 1})
    assert unknown.status_code == 404
    empty = ctx.req("POST", f"/projects/{ctx.project}/members:batch", {"userIds": [], "role": "editor", "expectedMembersRevision": 1})
    assert empty.status_code == 422
    assert ctx.rev() == 1


@pytest.mark.parametrize("role", ["reviewer", "admin", ""])
def test_reviewer_and_unknown_roles_cannot_be_assigned(ctx, role):
    response = ctx.put("bob", role)
    assert response.status_code == 422
    assert response.json()["error"]["fieldErrors"][0]["field"] == "role"
    assert ctx.batch(["bob"], role).status_code == 422
    assert ctx.rev() == 1


def test_existing_reviewer_rows_are_kept_and_can_be_changed(ctx):
    with connect() as connection:
        connection.execute("insert into project_members(project_id, user_id, role) values (%s, %s, 'reviewer')",
                           (ctx.project, ctx.users["bob"]))
        connection.commit()
    roles = {item["name"]: item["role"] for item in ctx.members()["items"]}
    assert roles == {"张工": "owner", "李工": "reviewer"}
    assert ctx.put("bob", "viewer").status_code == 200


# ---------- R1/R2 最后负责人 ----------

def test_last_owner_cannot_be_demoted_or_removed(ctx):
    for response in (ctx.put("alice", "editor"), ctx.delete("alice")):
        assert response.status_code == 409
        assert response.json()["error"]["code"] == "LAST_OWNER"
    assert ctx.rev() == 1 and _active_owner_count(ctx.project) == 1
    assert ctx.batch(["bob"], "owner").status_code == 200
    demoted = ctx.put("alice", "editor")
    assert demoted.status_code == 200, demoted.text
    # R4：原主负责人不再是负责人 → 主负责人改为剩下的负责人，写 owner_reassigned
    assert demoted.json()["data"]["ownerChanged"]["to"]["name"] == "李工"
    assert demoted.json()["data"]["ownerId"] == str(ctx.users["bob"])
    assert [event["action"] for event in ctx.events()][-1] == "owner_reassigned"


def test_disabled_owner_does_not_count(ctx):
    assert ctx.batch(["bob"], "owner").status_code == 200
    assert ctx.disable("bob").status_code == 200  # 张工仍是启用负责人，可以停用李工
    ctx.as_("alice")
    data = ctx.members()
    assert data["activeOwnerCount"] == 1
    bob = next(item for item in data["items"] if item["name"] == "李工")
    assert bob["userStatus"] == "disabled"
    blocked = ctx.put("alice", "viewer")
    assert blocked.status_code == 409 and blocked.json()["error"]["code"] == "LAST_OWNER"
    # 停用账号可以降级/移除，不能升级
    assert ctx.put("bob", "editor").status_code == 200
    upgraded = ctx.put("bob", "owner")
    assert upgraded.status_code == 422 and upgraded.json()["error"]["code"] == "USER_DISABLED"
    assert ctx.delete("bob").status_code == 200


def test_disabling_sole_owner_is_409_and_lists_projects(ctx):
    blocked = ctx.disable("alice")
    assert blocked.status_code == 409, blocked.text
    error = blocked.json()["error"]
    assert error["code"] == "LAST_OWNER"
    assert [item["name"] for item in error["details"]["projects"]] == ["成员项目"] and error["details"]["total"] == 1
    with connect() as connection:
        assert connection.execute("select status from users where id = %s", (ctx.users["alice"],)).fetchone()[0] == "active"
    ctx.as_("alice")
    assert ctx.batch(["bob"], "owner").status_code == 200
    assert ctx.disable("alice").status_code == 200


def test_disabling_sole_owner_of_archived_project_is_still_blocked(ctx):
    """Q-2（2026-10-11 用户授权决定）：归档项目里的唯一启用负责人照样挡停用；列表里带 archived 标记供界面标“（已归档）”。"""
    _archive(ctx)
    blocked = ctx.disable("alice")
    assert blocked.status_code == 409, blocked.text
    error = blocked.json()["error"]
    assert error["code"] == "LAST_OWNER" and error["details"]["total"] == 1
    assert error["details"]["projects"] == [{"id": ctx.project, "name": "成员项目", "archived": True}]
    with connect() as connection:
        assert connection.execute("select status from users where id = %s", (ctx.users["alice"],)).fetchone()[0] == "active"
    # 归档项目成员只读：要先恢复才能加第二负责人，之后才能停用
    ctx.as_("alice")
    assert ctx.batch(["bob"], "owner").json()["error"]["code"] == "PROJECT_ARCHIVED"
    ctx.as_("admin")
    assert ctx.req("POST", f"/projects/{ctx.project}/unarchive", {"expectedRevision": ctx.project_rev(), "reason": "恢复"}).status_code == 200
    ctx.as_("alice")
    assert ctx.batch(["bob"], "owner").status_code == 200
    assert ctx.disable("alice").status_code == 200


def test_disabled_accounts_cannot_be_added(ctx):
    assert ctx.disable("dave").status_code == 200
    ctx.as_("alice")
    single = ctx.put("dave", "viewer")  # PUT 不再加人（20 I-3）：非成员一律 404 MEMBER_NOT_FOUND
    assert single.status_code == 404 and single.json()["error"]["code"] == "MEMBER_NOT_FOUND"
    batch = ctx.batch(["dave", "carol"])
    assert batch.status_code == 200
    assert [(item["code"]) for item in batch.json()["data"]["failed"]] == ["USER_DISABLED"]
    assert batch.json()["data"]["added"] == [str(ctx.users["carol"])]


def test_batch_does_not_change_existing_members(ctx):
    assert ctx.batch(["bob"], "owner").status_code == 200
    result = ctx.batch(["bob", "carol"], "viewer").json()["data"]
    assert [item["code"] for item in result["failed"]] == ["ALREADY_MEMBER"]
    assert {item["name"]: item["role"] for item in result["items"]}["李工"] == "owner"


# ---------- 可见性与权限 ----------

def test_editor_viewer_cannot_manage_or_leave(ctx):
    assert ctx.batch(["bob"], "editor").status_code == 200
    assert ctx.batch(["carol"], "viewer").status_code == 200
    for who in ("bob", "carol"):
        ctx.as_(who)
        data = ctx.members()
        assert data["canManageMembers"] is False and len(data["items"]) == 3
        assert ctx.delete(who).status_code == 403  # U-6：不能自己退出
        assert ctx.put(who, "owner").status_code == 403
        denied = ctx.put("dave", "viewer")
        assert denied.status_code == 403 and denied.json()["error"]["details"]["ownerName"] == "张工"
        assert ctx.get(f"/projects/{ctx.project}/member-candidates").status_code == 403
        assert ctx.get(f"/projects/{ctx.project}/member-events").status_code == 403
    assert ctx.rev() == 3


def test_non_members_get_404_everywhere(ctx):
    ctx.as_("outsider")
    p = ctx.project
    responses = [ctx.get(f"/projects/{p}/members"), ctx.get(f"/projects/{p}/member-candidates"),
                 ctx.get(f"/projects/{p}/member-events"), ctx.put("bob", "editor", revision=1),
                 ctx.delete("bob", revision=1), ctx.batch(["bob"], revision=1)]
    assert [response.status_code for response in responses] == [404] * 6
    missing = ctx.get(f"/projects/{uuid4()}/members")
    assert missing.status_code == 404


def test_auditor_reads_members_and_events_but_cannot_write(ctx):
    ctx.as_("auditor")
    data = ctx.members()
    assert data["canManageMembers"] is False and data["projectRole"] == "auditor"
    assert ctx.get(f"/projects/{ctx.project}/member-events").status_code == 200
    assert ctx.get(f"/projects/{ctx.project}/member-candidates").status_code == 403
    assert ctx.put("bob", "editor", revision=1).status_code == 403
    listed = ctx.get("/projects").json()["data"]["items"]
    assert [(item["name"], item["myRole"]) for item in listed] == [("成员项目", None)]


def test_admin_manages_without_being_member(ctx):
    ctx.as_("admin")
    data = ctx.members()
    assert data["canManageMembers"] is True and data["projectRole"] is None
    assert ctx.batch(["bob"], "owner").status_code == 200
    assert ctx.put("alice", "viewer").status_code == 200
    events = ctx.get(f"/projects/{ctx.project}/member-events").json()["data"]["items"]
    assert events[0]["actor"]["name"] == "管理员甲"


def test_project_list_my_role_and_owners(ctx):
    ctx.batch(["bob"], "editor")
    ctx.as_("bob")
    items = ctx.get("/projects").json()["data"]["items"]
    assert [(item["myRole"], [owner["name"] for owner in item["owners"]]) for item in items] == [("editor", ["张工"])]
    assert ctx.get("/projects?myRole=owner").json()["data"]["total"] == 0
    assert ctx.get("/projects?myRole=editor").json()["data"]["total"] == 1
    assert ctx.get("/projects?myRole=reviewer").status_code == 422
    ctx.as_("alice")
    assert ctx.get("/projects?myRole=owner").json()["data"]["total"] == 1


def test_candidates_search_escapes_wildcards_and_skips_disabled(ctx):
    with connect() as connection:
        for account, name in (("a05-pct", "百分%号"), ("a05-und", "下划_线"), ("a05-plain", "普通人")):
            create_user(account=account, name=name, password=PASSWORD, roles=[], allow_no_roles=True, connection=connection)
        connection.commit()
    assert ctx.disable("dave").status_code == 200
    ctx.as_("alice")
    base = f"/projects/{ctx.project}/member-candidates"
    assert [item["name"] for item in ctx.get(f"{base}?q=%").json()["data"]["items"]] == ["百分%号"]
    assert [item["name"] for item in ctx.get(f"{base}?q=_").json()["data"]["items"]] == ["下划_线"]
    everyone = ctx.get(f"{base}?pageSize=50").json()["data"]
    names = {item["name"]: item for item in everyone["items"]}
    assert "赵工" not in names and names["张工"]["isMember"] is True and names["张工"]["currentRole"] == "owner"
    assert set(names["李工"]) == {"id", "name", "account", "department", "isMember", "currentRole"}
    assert ctx.get(f"{base}?pageSize=1000").status_code == 422
    assert ctx.get(f"{base}?page=0").status_code == 422
    paged = ctx.get(f"{base}?pageSize=1&page=2").json()["data"]
    assert len(paged["items"]) == 1 and paged["total"] == everyone["total"]


# ---------- R3 归档只读 ----------

def _document(ctx):
    word = Document()
    word.add_heading("范围", level=1)
    word.add_paragraph("原始正文。")
    buffer = BytesIO()
    word.save(buffer)
    stored = store_bytes(name="成员.docx", media_type=DOCX, purpose="attachment", project_id=ctx.project, content=buffer.getvalue())
    ctx.as_("admin")
    created = ctx.req("POST", f"/projects/{ctx.project}/documents", {"name": "成员文档", "sourceType": "upload", "wordFileId": str(stored.id)})
    assert created.status_code == 200, created.text
    return created.json()["data"]


def _archive(ctx):
    ctx.as_("admin")
    revision = ctx.project_rev()
    archived = ctx.req("POST", f"/projects/{ctx.project}/archive", {"expectedRevision": revision, "reason": "A05 归档测试"})
    assert archived.status_code == 200, archived.text


def test_archived_project_is_read_only_for_members_and_documents(ctx):
    document = _document(ctx)
    ctx.as_("alice")
    assert ctx.batch(["bob"], "editor").status_code == 200
    _archive(ctx)
    for who in ("alice", "admin"):
        ctx.as_(who)
        for response in (ctx.put("bob", "viewer"), ctx.delete("bob"), ctx.batch(["carol"])):
            assert response.status_code == 409, response.text
            assert response.json()["error"]["code"] == "PROJECT_ARCHIVED"
    ctx.as_("alice")
    data = ctx.members()
    assert data["readonlyReason"] == "archived" and data["canManageMembers"] is False
    permissions = ctx.get(f"/projects/{ctx.project}").json()["data"]["permissions"]
    assert permissions["canEdit"] is False and permissions["canManageMembers"] is False and permissions["readonlyReason"] == "archived"
    ctx.as_("bob")
    detail = ctx.get(f"/documents/{document['id']}")
    assert detail.status_code == 200
    assert detail.json()["data"]["permissions"]["canEdit"] is False
    assert detail.json()["data"]["permissions"]["readonlyReason"] == "archived"
    saved = ctx.req("PATCH", f"/documents/{document['id']}/content", {"expectedRevision": 0, "markdownContent": "# 改\n"})
    assert saved.status_code == 409 and saved.json()["error"]["code"] == "PROJECT_ARCHIVED"
    assert saved.json()["error"]["details"]["readonlyReason"] == "archived"
    exported = ctx.req("POST", f"/documents/{document['id']}/export", {"expectedRevision": 0})
    assert exported.status_code == 200, exported.text
    objects = ctx.get(f"/documents/{document['id']}/objects")
    assert objects.status_code == 200 and objects.json()["data"]["permissions"]["readonlyReason"] == "archived"
    listing = ctx.get(f"/projects/{ctx.project}/documents").json()["data"]
    assert listing["permissions"]["canEdit"] is False
    patched = ctx.as_("alice").req("PATCH", f"/projects/{ctx.project}", {"expectedRevision": ctx.project_rev(), "description": "x"})
    assert patched.status_code == 409 and patched.json()["error"]["code"] == "PROJECT_ARCHIVED"
    assert ctx.rev() == 2


def _code(response):
    return response.status_code, (response.json().get("error") or {}).get("code")


def test_archived_project_blocks_facts_test_tasks_and_bindings(ctx):
    """Q-1（2026-10-11 用户授权决定）：归档项目完全只读。事实、测试任务、知识库绑定的写入一律 409 PROJECT_ARCHIVED，
    读照常；权限检查先于归档检查（无权限仍 403、非成员仍 404）。"""
    word, buffer = Document(), BytesIO()
    word.add_heading("测试计划", level=1)
    word.add_paragraph("计划正文。")
    word.save(buffer)
    plan = store_bytes(name="计划.docx", media_type=DOCX, purpose="attachment", project_id=ctx.project, content=buffer.getvalue())
    ctx.as_("alice")
    assert ctx.batch(["bob"], "editor").status_code == 200
    ctx.as_("admin")
    fact = ctx.req("POST", f"/projects/{ctx.project}/facts", {"key": "software_version", "label": "软件版本", "value": "V1.0",
                                                              "decisionReason": "归档前录入"})
    assert fact.status_code == 201, fact.text
    fact = fact.json()["data"]
    task = ctx.req("POST", "/test-tasks", {"name": "归档前任务", "projectId": ctx.project, "planFileId": str(plan.id)})
    assert task.status_code == 201, task.text
    task = task.json()["data"]
    bindings_before = ctx.get(f"/projects/{ctx.project}/knowledge-bindings").json()["data"]
    _archive(ctx)
    project_revision = ctx.project_rev()

    ctx.as_("admin")
    blocked = {
        "create fact": ctx.req("POST", f"/projects/{ctx.project}/facts", {"key": "k2", "label": "新事实", "value": "x", "decisionReason": "r"}),
        "revise fact": ctx.req("PATCH", f"/facts/{fact['id']}", {"expectedRevision": fact["revision"], "value": "V2.0", "decisionReason": "r"}),
        "confirm fact": ctx.req("POST", f"/facts/{fact['id']}/confirm", {"expectedRevision": fact["revision"], "reason": "确认"}),
        "bindings": ctx.req("PUT", f"/projects/{ctx.project}/knowledge-bindings", {"expectedProjectRevision": project_revision, "baseIds": []}),
        "create task": ctx.req("POST", "/test-tasks", {"name": "归档后任务", "projectId": ctx.project, "planFileId": str(plan.id)}),
        "patch task": ctx.req("PATCH", f"/test-tasks/{task['id']}", {"expectedRevision": task["revision"], "name": "改名"}),
        "analyze": ctx.req("POST", f"/test-tasks/{task['id']}/analyze"),
        "patch case": ctx.req("PATCH", f"/test-tasks/{task['id']}/cases/c1", {"enabled": False}),
        "run": ctx.req("POST", f"/test-tasks/{task['id']}/run", {"expectedRevision": task["revision"]}),
    }
    for name, response in blocked.items():
        assert response.status_code == 409, (name, response.text)
        error = response.json()["error"]
        assert error["code"] == "PROJECT_ARCHIVED" and error["details"]["readonlyReason"] == "archived", (name, error)

    # 权限检查先于归档：编写人员改绑定仍 403；非成员仍 404（不泄露项目存在）
    ctx.as_("bob")
    assert _code(ctx.req("PUT", f"/projects/{ctx.project}/knowledge-bindings",
                         {"expectedProjectRevision": project_revision, "baseIds": []})) == (403, "FORBIDDEN")
    ctx.as_("outsider")
    assert ctx.req("POST", f"/projects/{ctx.project}/facts", {"key": "k3", "label": "x", "value": "x", "decisionReason": "r"}).status_code == 404

    # 读照常，且没有被写入
    ctx.as_("admin")
    facts = ctx.get(f"/projects/{ctx.project}/facts")
    assert facts.status_code == 200 and [item["key"] for item in facts.json()["data"]["items"]] == ["software_version"]
    assert facts.json()["data"]["items"][0]["revision"] == fact["revision"]
    assert ctx.get(f"/projects/{ctx.project}/knowledge-bindings").json()["data"]["items"] == bindings_before["items"]
    detail = ctx.get(f"/test-tasks/{task['id']}")
    assert detail.status_code == 200 and detail.json()["data"]["name"] == "归档前任务"
    assert detail.json()["data"]["revision"] == task["revision"]
    listed = ctx.get(f"/test-tasks?projectId={ctx.project}").json()["data"]
    assert [item["name"] for item in listed["items"]] == ["归档前任务"]
    assert ctx.project_rev() == project_revision

    # 恢复后可写
    restored = ctx.req("POST", f"/projects/{ctx.project}/unarchive", {"expectedRevision": project_revision, "reason": "恢复"})
    assert restored.status_code == 200, restored.text
    renamed = ctx.req("PATCH", f"/test-tasks/{task['id']}", {"expectedRevision": task["revision"], "name": "恢复后改名"})
    assert renamed.status_code == 200, renamed.text
    revised = ctx.req("PATCH", f"/facts/{fact['id']}", {"expectedRevision": fact["revision"], "value": "V2.0", "decisionReason": "恢复后"})
    assert revised.status_code == 200, revised.text


def test_objects_poll_reports_role_change(ctx):
    document = _document(ctx)
    ctx.as_("alice")
    ctx.batch(["bob"], "editor")
    ctx.as_("bob")
    assert ctx.get(f"/documents/{document['id']}/objects").json()["data"]["permissions"]["canEdit"] is True
    ctx.as_("alice")
    assert ctx.put("bob", "viewer").status_code == 200
    ctx.as_("bob")
    assert ctx.get(f"/documents/{document['id']}/objects").json()["data"]["permissions"]["canEdit"] is False
    ctx.as_("alice")
    assert ctx.delete("bob").status_code == 200
    ctx.as_("bob")
    assert ctx.get(f"/documents/{document['id']}/objects").status_code == 404


# ---------- 并发与不变量 ----------

def test_two_owners_demoting_each_other_concurrently_keep_one_owner(ctx):
    assert ctx.batch(["bob"], "owner").status_code == 200
    revision = ctx.rev()
    app = ctx.client.app
    barrier = threading.Barrier(2, timeout=30)
    results = {}

    def demote(actor, target):
        with TestClient(app) as client:
            headers = _login(client, f"a05-{actor}")
            barrier.wait()
            # 各自读到的版本都是最新的（如前端刚刷新过），只有事务内的 R1 检查能挡住第二个
            response = client.put(f"/api/v1/projects/{ctx.project}/members/{ctx.users[target]}",
                                  json={"role": "editor", "expectedMembersRevision": revision}, headers=headers)
            results[actor] = (response.status_code, response.json().get("error", {}).get("code"))

    threads = [threading.Thread(target=demote, args=pair) for pair in (("alice", "bob"), ("bob", "alice"))]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join()
    codes = sorted(results.values())
    assert codes[0][0] == 200 and codes[1][0] in (403, 409), results
    assert _active_owner_count(ctx.project) == 1


def test_disable_and_demote_race_keeps_an_active_owner(ctx):
    assert ctx.batch(["bob"], "owner").status_code == 200
    revision = ctx.rev()
    app = ctx.client.app
    barrier = threading.Barrier(2, timeout=30)
    results = {}

    def disable_bob():
        with TestClient(app) as client:
            headers = _login(client, "a05-admin")
            user_revision = _user_revision(ctx.users["bob"])
            barrier.wait()
            response = client.post(f"/api/v1/users/{ctx.users['bob']}/disable",
                                   json={"expectedRevision": user_revision, "reason": "并发停用"}, headers=headers)
            results["disable"] = response.status_code

    def demote_alice():
        with TestClient(app) as client:
            headers = _login(client, "a05-alice")
            barrier.wait()
            response = client.put(f"/api/v1/projects/{ctx.project}/members/{ctx.users['alice']}",
                                  json={"role": "editor", "expectedMembersRevision": revision}, headers=headers)
            results["demote"] = response.status_code

    threads = [threading.Thread(target=disable_bob), threading.Thread(target=demote_alice)]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join()
    assert sorted(results.values()) == [200, 409], results
    assert _active_owner_count(ctx.project) == 1


def test_disable_sole_owner_and_add_second_owner_race_is_consistent(ctx):
    """17 §2.7：“停用唯一启用 owner”与“添加第二位 owner”并发，两种顺序都得到一致结果。"""
    revision = ctx.rev()
    app = ctx.client.app
    barrier = threading.Barrier(2, timeout=30)
    results = {}

    def disable_alice():
        with TestClient(app) as client:
            headers = _login(client, "a05-admin")
            user_revision = _user_revision(ctx.users["alice"])
            barrier.wait()
            results["disable"] = client.post(f"/api/v1/users/{ctx.users['alice']}/disable",
                                             json={"expectedRevision": user_revision, "reason": "并发停用"}, headers=headers).status_code

    def add_bob():
        with TestClient(app) as client:
            headers = _login(client, "a05-admin")
            barrier.wait()
            results["add"] = client.post(f"/api/v1/projects/{ctx.project}/members:batch",
                                         json={"userIds": [str(ctx.users["bob"])], "role": "owner",
                                               "expectedMembersRevision": revision}, headers=headers).status_code

    threads = [threading.Thread(target=disable_alice), threading.Thread(target=add_bob)]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join()
    assert results["add"] == 200, results
    assert results["disable"] in (200, 409), results
    # 先加 bob 再停用：停用成功，只剩 bob 一位启用负责人；先停用：alice 是唯一负责人被 409 挡下，之后加 bob → 两位
    assert _active_owner_count(ctx.project) == (1 if results["disable"] == 200 else 2)
    with connect() as connection:
        alice = connection.execute("select status from users where id = %s", (ctx.users["alice"],)).fetchone()[0]
        owner_id = connection.execute("select owner_id from projects where id = %s", (ctx.project,)).fetchone()[0]
    assert (alice == "disabled") == (results["disable"] == 200)
    # 主负责人：alice 被停用后仍可能留在 owner_id 上（停用不改指，见 18 偏差），但启用负责人只剩 bob 时必须是 bob 或 alice
    assert str(owner_id) in {str(ctx.users["alice"]), str(ctx.users["bob"])}


def test_random_member_operations_keep_invariants(ctx):
    """17 §2.7 fuzz：添加、改角色、移除、停用、启用随机 1000 步；每步后启用负责人 ≥1（或本来就没有且这一步没让它减少），
    且读接口的有效主负责人是启用中的负责人（S-8：按 Q-3 + 读时计算的口径）。"""
    rng = random.Random(116)
    ctx.as_("admin")
    names = ["alice", "bob", "carol", "dave"]
    successes = 0
    for _ in range(1000):
        who = rng.choice(names)
        op = rng.choice(["put", "put", "delete", "batch", "disable", "enable"])
        before, owners_before = ctx.rev(), _active_owner_count(ctx.project)
        if op == "put":
            response = ctx.put(who, rng.choice(["owner", "editor", "viewer"]))
        elif op == "delete":
            response = ctx.delete(who)
        elif op == "batch":
            response = ctx.batch(rng.sample(names, 2), rng.choice(["owner", "editor", "viewer"]))
        elif op == "disable":
            response = ctx.req("POST", f"/users/{ctx.users[who]}/disable", {"expectedRevision": _user_revision(ctx.users[who]), "reason": "fuzz"})
        else:
            response = ctx.req("POST", f"/users/{ctx.users[who]}/enable", {"expectedRevision": _user_revision(ctx.users[who]), "reason": "fuzz"})
        assert response.status_code in (200, 204, 404, 409, 422), response.text
        if response.status_code in (404, 409, 422):
            assert response.json()["error"]["code"] in ("LAST_OWNER", "USER_DISABLED", "MEMBER_NOT_FOUND"), response.text
            assert ctx.rev() == before
        if response.status_code == 404:  # PUT 只改已有成员（20 I-3）
            assert op == "put" and f"a05-{who}" not in {item["account"] for item in ctx.members()["items"]}
        changed = op in ("put", "delete", "batch") and response.status_code == 200 and ctx.rev() != before
        if changed:
            assert ctx.rev() == before + 1
            successes += 1
        owners_after = _active_owner_count(ctx.project)
        assert owners_after >= 1 or owners_after >= owners_before
        # 有效主负责人（19 审查 S-2/S-4，读时计算，每一步之后都查，不只在成员变更之后）：
        # 有启用负责人时，详情 ownerId 就是 owners[0]、是启用中的负责人；owners[] 只含启用账号；没有启用负责人时为 null
        detail = ctx.get(f"/projects/{ctx.project}").json()["data"]
        assert {owner["status"] for owner in detail["owners"]} <= {"active"}
        if owners_after:
            assert detail["owners"][0]["id"] == detail["ownerId"] and detail["owners"][0]["isPrimary"]
            with connect() as connection:
                role = connection.execute(
                    """select m.role, u.status from project_members m join users u on u.id = m.user_id
                       where m.project_id = %s and m.user_id = %s""", (ctx.project, detail["ownerId"])).fetchone()
            assert role is not None and role[0] == "owner" and role[1] == "active"
        else:
            assert detail["ownerId"] is None and detail["owners"] == []
    assert ctx.rev() == 1 + successes
    assert {event["members_revision"] for event in ctx.events()} == set(range(1, ctx.rev() + 1))
