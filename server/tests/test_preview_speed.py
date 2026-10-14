"""整篇预览提速：缓存不改变结果，原稿变化时缓存失效。"""
import os

from docx import Document

from app.docx_parser import parse_docx, parse_docx_cached
from app.text_memo import TextMemo


def _word(path, text):
    word = Document()
    word.add_heading("范围", 1)
    word.add_paragraph(text)
    word.save(path)


def test_parse_docx_cached_matches_and_invalidates(tmp_path):
    path = tmp_path / "a.docx"
    _word(path, "原始正文。")
    first = parse_docx_cached(str(path))
    assert first.markdown == parse_docx(str(path)).markdown
    assert parse_docx_cached(str(path)) is first
    _word(path, "改过的正文，长度也不同。")
    stat = path.stat()
    os.utime(path, ns=(stat.st_atime_ns, stat.st_mtime_ns + 1_000_000))
    assert "改过的正文" in parse_docx_cached(str(path)).markdown


def test_text_memo_returns_same_result_and_bounds_size():
    calls = []

    def upper(value):
        calls.append(value)
        return value.upper()
    memo = TextMemo(upper, max_items=2)
    assert memo("a") == "A" and memo("a") == "A" and calls == ["a"]
    memo("b")
    memo("c")
    assert memo("a") == "A" and calls == ["a", "b", "c", "a"]  # 超过条数上限被淘汰后重新计算
