"""Current template groups accept usable, authorized Word files."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from uuid import UUID

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.db import connect, migrate
from app.files import resolve_object, store_bytes
from app.identity import create_user

PASSWORD = "correct-horse-1"
WORD_MEDIA = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@pytest.fixture()
def client(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setenv("FILE_DATA_DIR", str(tmp_path / "files"))
    from app.config import clear_settings_cache
    from app.main import create_app

    clear_settings_cache()
    migrate()
    with TestClient(create_app()) as test_client:
        yield test_client
    clear_settings_cache()


def _login(client: TestClient, account: str, roles=None) -> str:
    create_user(account=account, name=account, password=PASSWORD, roles=roles or ["system_admin"])
    csrf = client.get("/api/v1/auth/csrf").json()["data"]["csrfToken"]
    response = client.post("/api/v1/auth/login", json={"account": account, "password": PASSWORD}, headers={"X-CSRF-Token": csrf})
    assert response.status_code == 200, response.text
    return response.json()["data"]["csrfToken"]


def _group(client: TestClient, csrf: str, name="规范模板") -> dict:
    response = client.post("/api/v1/templates", json={"name": name, "category": "软件文档"}, headers={"X-CSRF-Token": csrf})
    assert response.status_code == 200, response.text
    return response.json()["data"]


def _word() -> bytes:
    document = Document()
    document.add_heading("概述", 1)
    document.add_paragraph("模板正文")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _import(client, csrf, group, file_id, name="需求说明"):
    return client.post(f"/api/v1/templates/{group['id']}/documents", json={"name": name, "wordFileId": str(file_id)}, headers={"X-CSRF-Token": csrf})


def test_template_manager_can_upload_and_import_word(client: TestClient) -> None:
    csrf = _login(client, "template-owner", ["template_manager"])
    group = _group(client, csrf)
    upload = client.post("/api/v1/files", files={"file": ("模板.docx", _word(), WORD_MEDIA)}, data={"purpose": "template"}, headers={"X-CSRF-Token": csrf, "Idempotency-Key": "template-word"})
    assert upload.status_code == 201, upload.text
    imported = _import(client, csrf, group, upload.json()["data"]["id"])
    assert imported.status_code == 200, imported.text
    assert imported.json()["data"]["formatTemplate"]
    assert client.get(f"/api/v1/templates/{group['id']}").json()["data"]["documentCount"] == 1
    assert _import(client, csrf, group, upload.json()["data"]["id"]).status_code == 409


@pytest.mark.parametrize("file_type,content", [("text/plain", b"text"), (WORD_MEDIA, b"invalid docx")])
def test_invalid_word_does_not_create_successful_template(client, file_type, content) -> None:
    csrf = _login(client, "template-admin")
    group = _group(client, csrf)
    stored = store_bytes(name="bad.docx", media_type=file_type, purpose="template", content=content)
    failed = _import(client, csrf, group, stored.id)
    assert failed.status_code == 422, failed.text
    assert client.get(f"/api/v1/templates/{group['id']}/documents").json()["data"]["items"] == []


def test_missing_word_object_is_not_imported(client) -> None:
    from app.config import get_settings
    csrf = _login(client, "template-admin")
    group = _group(client, csrf)
    stored = store_bytes(name="missing.docx", media_type=WORD_MEDIA, purpose="template", content=_word())
    resolve_object(get_settings().file_data_dir, stored.storage_key).unlink()
    assert _import(client, csrf, group, stored.id).status_code == 404


def test_template_import_cannot_copy_another_projects_private_file(client) -> None:
    csrf = _login(client, "template-admin")
    group = _group(client, csrf)
    owner = create_user(account="project-owner", name="负责人", password=PASSWORD, roles=["project_creator"])
    project_id = UUID("11111111-1111-4111-8111-111111111111")
    with connect() as connection:
        with connection.transaction():
            connection.execute("INSERT INTO projects (id, name, owner_id) VALUES (%s, %s, %s)", (project_id, "私有项目", owner))
    stored = store_bytes(name="private.docx", media_type=WORD_MEDIA, purpose="attachment", content=_word(), project_id=project_id)
    other = TestClient(client.app)
    other_csrf = _login(other, "limited-template-manager", ["template_manager"])
    assert _import(other, other_csrf, group, stored.id).status_code == 404
    assert client.get(f"/api/v1/templates/{group['id']}/documents").json()["data"]["items"] == []


def test_template_paging_literal_filter_and_rename_conflicts(client) -> None:
    csrf = _login(client, "template-admin")
    groups = [_group(client, csrf, name) for name in ["模板_%", "模板一", "模板二"]]
    first = client.get("/api/v1/templates", params={"page": 1, "pageSize": 2}).json()["data"]
    second = client.get("/api/v1/templates", params={"page": 2, "pageSize": 2}).json()["data"]
    assert first["total"] == second["total"] == 3
    assert len(first["items"]) == 2 and len(second["items"]) == 1
    assert not {item["id"] for item in first["items"]} & {item["id"] for item in second["items"]}
    literal = client.get("/api/v1/templates", params={"name": "_%"}).json()["data"]
    assert literal["total"] == 1 and literal["items"][0]["name"] == "模板_%"
    rename = client.patch(f"/api/v1/templates/{groups[0]['id']}", json={"expectedRevision": 0, "name": groups[1]["name"]}, headers={"X-CSRF-Token": csrf})
    assert rename.status_code == 409, rename.text
    for params in [{"page": "bad"}, {"page": 0}, {"pageSize": 0}, {"pageSize": -1}, {"pageSize": 101}]:
        assert client.get("/api/v1/templates", params=params).status_code == 422
