"""Bound selections are atomic and checked against the native Word before generation."""
from unittest.mock import Mock
from io import BytesIO
import json
import pytest
from lxml import html
from docx import Document
from docx.oxml import OxmlElement
from test_document_versions import version_context, save, login  # noqa: F401 - pytest fixtures
import app.documents as documents
from app.config import get_settings
from app.db import connect
from app.files import resolve_object, store_bytes
from app.document_selection import selection_blocks


def descriptor(client, did, index=0, text=None):
    response = client.get(f'/api/v1/documents/{did}/preview')
    assert response.status_code == 200, response.text
    node = html.fragment_fromstring(response.json()['data']['html'], create_parent='div').xpath('//*[@data-selection-id]')[index]
    value = ''.join(node.itertext())
    text = value if text is None else text
    start = len(value[:value.index(text)].encode('utf-16-le')) // 2
    return dict(blockId=node.get('data-selection-id'), start=start, end=start+len(text.encode('utf-16-le'))//2, text=text)


def test_repeated_paragraph_adoption_snapshots_memory_and_checks_revision(version_context):
    client, headers, doc, _, _ = version_context
    did = doc['id']
    save(client, headers, did, 0, '相同🙂订单。\n\n相同🙂订单。')
    selection = descriptor(client, did, 1, '🙂订单')
    body = dict(expectedRevision=1, selection=selection, content='归档记录', instruction='仅改选中位置。')
    result = client.post(f'/api/v1/documents/{did}/selection', headers=headers, json=body)
    assert result.status_code == 200, result.text
    assert result.json()['data']['markdownContent'] == '# 范围\n\n相同🙂订单。\n\n相同归档记录。'
    assert result.json()['data']['revision'] == 2
    assert client.get(f'/api/v1/documents/{did}/versions/1').json()['data']['markdownContent'].count('相同🙂订单。') == 2
    assert client.get(f'/api/v1/documents/{did}/memory').json()['data']['items'][0]['instruction'] == '仅改选中位置。'
    assert client.post(f'/api/v1/documents/{did}/selection', headers=headers, json=body).status_code == 409


def test_stale_anchor_and_other_owner_cannot_modify(version_context):
    client, headers, doc, _, _ = version_context
    did = doc['id']
    selection = descriptor(client, did)
    save(client, headers, did, 0, '已经修改的正文。')
    body = dict(expectedRevision=1, selection=selection, content='不得写入')
    assert client.post(f'/api/v1/documents/{did}/selection', json=body, headers=headers).status_code == 422
    assert client.post(f'/api/v1/documents/{did}/selection', json=body, headers=login(client, 'version-other')).status_code == 404
    with connect() as conn:
        assert conn.execute('SELECT revision FROM documents WHERE id=%s', (did,)).fetchone()[0] == 1


def test_native_protected_selection_is_rejected_before_model_call(version_context, monkeypatch):
    client, headers, doc, _, _ = version_context
    did = doc['id']
    selection = descriptor(client, did)
    settings = get_settings().model_copy(update={'ai_base_url':'http://localhost/v1','ai_model':'offline','ai_api_key':''})
    monkeypatch.setattr(documents, 'get_settings', lambda: settings)
    draft = Mock(return_value='替换后的正文。')
    monkeypatch.setattr(documents, 'draft_section_text', draft)
    body = dict(expectedRevision=0, headingIndex=0, selectedText=selection['text'], selection=selection, scope='selection')
    result = client.post(f'/api/v1/documents/{did}/sections/draft', json=body, headers=headers)
    assert result.status_code == 200, result.text
    assert draft.call_count == 1
    with connect() as conn:
        source = conn.execute('SELECT storage_key FROM files WHERE id=%s', (doc['originalDocxFileId'],)).fetchone()[0]
    path = resolve_object(settings.file_data_dir, source)
    word = Document(path)
    field = OxmlElement('w:fldSimple'); field.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}instr', ' DATE ')
    word.paragraphs[1]._p.append(field); word.save(path)
    result = client.post(f'/api/v1/documents/{did}/sections/draft', json=body, headers=headers)
    assert result.status_code == 422, result.text
    assert draft.call_count == 1
    assert client.get(f'/api/v1/documents/{did}').json()['data']['revision'] == 0


def test_memory_failure_rolls_back_selected_text_and_new_version(version_context, monkeypatch):
    client, headers, doc, _, _ = version_context
    did = doc['id']
    selection = descriptor(client, did)
    def fail(*args, **kwargs):
        raise RuntimeError('forced memory storage failure')
    monkeypatch.setattr(documents, 'record_memory', fail)
    result = client.post(f'/api/v1/documents/{did}/selection', headers=headers,
                         json=dict(expectedRevision=0, selection=selection, content='不会留下的改动', instruction='回滚测试'))
    assert result.status_code == 500, result.text
    assert client.get(f'/api/v1/documents/{did}').json()['data']['revision'] == 0
    assert [item['revision'] for item in client.get(f'/api/v1/documents/{did}/versions').json()['data']['items']] == [0]


@pytest.mark.parametrize('streaming', [False, True])
@pytest.mark.parametrize('content', ['合并后的正文。', '第一段新正文。\n\n第二段新正文。\n\n第三段新正文。'])
def test_multiline_selection_generates_adopts_and_exports_without_changing_other_content(version_context, monkeypatch, streaming, content):
    client, headers, _, project, _ = version_context
    word = Document()
    word.add_heading('范围', level=1)
    word.add_paragraph('前缀🙂旧第一段。')
    word.add_paragraph('旧第二段后缀。')
    word.add_heading('其他章节', level=1)
    word.add_paragraph('不得修改的正文。')
    data = BytesIO(); word.save(data)
    stored = store_bytes(name='跨段验收.docx', media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                         purpose='attachment', project_id=project, content=data.getvalue())
    response = client.post(f'/api/v1/projects/{project}/documents',
                           json={'name':'跨段验收','sourceType':'upload','wordFileId':str(stored.id)}, headers=headers)
    assert response.status_code == 200, response.text
    doc = response.json()['data']; did = doc['id']
    blocks = selection_blocks(doc['markdownContent'])
    selection = dict(blockId=blocks[0]['blockId'], endBlockId=blocks[1]['blockId'], start=2, end=4,
                     text='🙂旧第一段。\n\n旧第二段')
    settings = get_settings().model_copy(update={'ai_base_url':'http://localhost/v1','ai_model':'offline','ai_api_key':''})
    monkeypatch.setattr(documents, 'get_settings', lambda: settings)
    draft = Mock(return_value=content)
    monkeypatch.setattr(documents, 'draft_section_text', draft)
    async def generate(**kwargs):
        draft(**kwargs)
        yield content
    monkeypatch.setattr(documents, 'stream_section_text', generate)
    body = dict(expectedRevision=0, selectedText=selection['text'], selection=selection, scope='selection')
    result = client.post(f'/api/v1/documents/{did}/sections/draft', params={'stream': streaming}, json=body, headers=headers)
    assert result.status_code == 200, result.text
    generated = json.loads(result.text.splitlines()[-1]) if streaming else result.json()['data']
    if streaming:
        assert generated['type'] == 'done', result.text
    assert generated['content'] == content and draft.call_count == 1
    assert draft.call_args.kwargs['existing'] == selection['text'] and draft.call_args.kwargs['selected_only'] is True
    assert client.get(f'/api/v1/documents/{did}').json()['data']['revision'] == 0
    adoption = dict(expectedRevision=0, selection=selection, content=content, instruction='只改写所选的两个段落。')
    written = client.post(f'/api/v1/documents/{did}/selection', json=adoption, headers=headers)
    assert written.status_code == 200, written.text
    assert written.json()['data']['revision'] == 1
    assert written.json()['data']['markdownContent'] == '# 范围\n\n前缀' + content + '后缀。\n\n# 其他章节\n\n不得修改的正文。'
    assert client.get(f'/api/v1/documents/{did}/versions/0').json()['data']['markdownContent'] == doc['markdownContent']
    assert client.get(f'/api/v1/documents/{did}/memory').json()['data']['items'][0]['instruction'] == adoption['instruction']
    assert client.post(f'/api/v1/documents/{did}/selection', json=adoption, headers=headers).status_code == 409
    exported = client.post(f'/api/v1/documents/{did}/export', json={'expectedRevision':1}, headers=headers)
    assert exported.status_code == 200, exported.text
    with connect() as connection:
        key = connection.execute('SELECT storage_key FROM files WHERE id=%s', (exported.json()['data']['fileId'],)).fetchone()[0]
    assert Document(resolve_object(settings.file_data_dir, key)).paragraphs[-1].text == '不得修改的正文。'
