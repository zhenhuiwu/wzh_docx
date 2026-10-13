"""AI fills the uploaded table and Word keeps its native layout properties."""
import json

import pytest
from docx import Document
from docx.shared import Cm
from langchain_core.messages import AIMessage

from app.document_rules import DraftTableError, validate_draft
from app.documents import _section_update, _validate_edits, DocumentEdit
from app.docx_export import edit_original_docx
from app.docx_parser import parse_docx
from tests.test_document_agent import ObservedModel, call, context, offline_context  # noqa: F401
from tests.test_document_edit_output import run


TABLE = '| 项目 | 内容 |\n| --- | --- |\n| 名称 |  |\n| 用途 |  |'
FILLED = TABLE.replace('| 名称 |  |', '| 名称 | 测试软件 |').replace('| 用途 |  |', '| 用途 | 文档编写 |')


@pytest.mark.parametrize('invalid', [
    '名称：测试软件。用途：文档编写。',
    TABLE + '\n\n' + FILLED,
    '| 项目 | 内容 |\n| --- | --- |\n| 名称 | 测试软件 |',
    '| 项目 | 内容 | 说明 |\n| --- | --- | --- |\n| 名称 | 测试软件 | 新列 |\n| 用途 | 文档编写 | 新列 |',
])
def test_table_rebuild_is_rejected_at_generation_and_adoption(invalid):
    with pytest.raises(DraftTableError):
        validate_draft(invalid, {}, existing=TABLE)
    row = [None] * 16
    row[2], row[6], row[7] = '通用文档', {}, '# 目标\n\n' + TABLE
    with pytest.raises(DraftTableError):
        _validate_edits(row, [DocumentEdit(headingIndex=0, content=invalid)])


def test_generic_table_cannot_be_appended_again():
    row = [None] * 16
    row[2], row[6], row[7] = '通用文档', {}, '# 目标\n\n' + TABLE
    with pytest.raises(ValueError, match='重复'):
        _section_update(row, 0, FILLED, 'append')
    assert _section_update(row, 0, FILLED, 'replace').count('| 项目 | 内容 |') == 1


@pytest.mark.parametrize('streaming', [False, True])
@pytest.mark.parametrize('whole', [False, True])
def test_duplicate_table_is_corrected_to_original_cells(monkeypatch, streaming, whole):
    bad = TABLE + '\n\n' + FILLED
    output = lambda body: json.dumps({'replacements': [{'before': TABLE, 'after': body}]}) if whole else body
    model = ObservedModel(responses=[call('prepare_document_edit', objective='填写原表。'),
                                    AIMessage(content=output(bad)), AIMessage(content=output(FILLED))])
    values = context(model, monkeypatch, existing=TABLE, full_markdown='# 目标\n\n' + TABLE,
                     whole_document=whole, sequential_document=whole, edit_in_place=whole)
    result = run(values, streaming)
    assert (json.loads(result)['edits'][0]['content'] if whole else result) == FILLED
    assert len(model.observed) == 3


def test_fill_preserves_native_table_when_surrounding_paragraphs_expand(tmp_path, monkeypatch):
    import app.docx_export as exporter
    doc = Document()
    doc.add_heading('目标', 1)
    table = doc.add_table(rows=3, cols=2, style='Table Grid')
    table.autofit = False
    table.columns[0].width, table.columns[1].width = Cm(3), Cm(10)
    for row, values in zip(table.rows, [('项目', '内容'), ('名称', ''), ('用途', '')]):
        for cell, text in zip(row.cells, values):
            cell.text = text
    doc.add_paragraph('旧说明。')
    source, destination = tmp_path/'original.docx', tmp_path/'filled.docx'
    doc.save(source)
    old = parse_docx(str(source)).markdown
    body = old[old.index('|'):].split('\n\n')[0]
    filled = body.replace('| 名称 |  |', '| 名称 | 测试软件 |').replace('| 用途 |  |', '| 用途 | 文档编写 |')
    edited = old.replace(body, '新增说明。\n\n'+filled).replace('旧说明。', '修改说明。\n\n补充说明。')
    original_insert = exporter._insert_block
    def insert(doc, target, *args, **kwargs):
        assert target.tag != 'table', 'existing table was recreated'
        return original_insert(doc, target, *args, **kwargs)
    monkeypatch.setattr(exporter, '_insert_block', insert)
    edit_original_docx(edited, source, destination)
    saved = Document(destination)
    assert len(saved.tables) == 1
    assert saved.tables[0]._tbl.tblPr.xml == table._tbl.tblPr.xml
    assert saved.tables[0]._tbl.tblGrid.xml == table._tbl.tblGrid.xml
    assert [[c.text for c in r.cells] for r in saved.tables[0].rows] == [
        ['项目', '内容'], ['名称', '测试软件'], ['用途', '文档编写']]
