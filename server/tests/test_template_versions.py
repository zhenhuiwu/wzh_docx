"""A06 模板版本快照与完整编辑（17 §3，16 §3.8）：在 3a1f5f3（A06 之前）上全部失败——那里没有版本表、草稿、发布、
恢复、变基、上传新 Word、对比和样式影响接口。用真实 438C SRS 模板走一遍 A6-1…A6-14 的后端部分。"""

from __future__ import annotations

import copy
import json
import sqlite3
import threading
from io import BytesIO
from pathlib import Path
from uuid import uuid4

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app import template_merge as tm
from app.db import connect, migrate
from app.files import store_bytes
from app.identity import create_user
from app.template_import import parse_template
from app.template_word import new_block_id

PASSWORD = "a06-template-pw"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
ROOT = Path(__file__).resolve().parents[2]
RACE_DOCUMENTS = 100  # 17 §3.11：发布与新建文档并发 100 次
SRS = next((ROOT / "438C-2021全套模板").glob("*[[]SRS[]]*.docx"))


def _login(client, account):
    client.cookies.clear()
    token = client.get("/api/v1/auth/csrf").json()["data"]["csrfToken"]
    response = client.post("/api/v1/auth/login", json={"account": account, "password": PASSWORD}, headers={"X-CSRF-Token": token})
    assert response.status_code == 200, response.text
    return {"X-CSRF-Token": response.json()["data"]["csrfToken"]}


class Api:
    def __init__(self, client):
        self.client = client
        self.headers: dict = {}

    def as_(self, who):
        self.headers = _login(self.client, f"a06-{who}")
        return self

    def get(self, path, **kw):
        return self.client.get(f"/api/v1{path}", **kw)

    def send(self, method, path, body=None):
        return self.client.request(method, f"/api/v1{path}", json=body, headers=self.headers)


@pytest.fixture()
def api():
    migrate()
    from app.main import create_app
    with TestClient(create_app()) as client:
        tag = uuid4().hex[:6]
        users = {}
        for who, roles in {"admin": ["system_admin"], "maint": ["template_manager"], "creator": ["project_creator"],
                           "auditor": ["auditor"], "plain": []}.items():
            users[who] = create_user(account=f"a06-{who}-{tag}", name=f"A06{who}", password=PASSWORD, roles=roles, allow_no_roles=True)
        context = Api(client)
        context.tag = tag
        context.users = users
        project = uuid4()
        with connect() as connection:
            connection.execute("INSERT INTO projects (id,name,owner_id) VALUES (%s,%s,%s)", (project, f"A06 项目 {tag}", users["admin"]))
            connection.execute("INSERT INTO project_members (project_id,user_id,role) VALUES (%s,%s,'owner')", (project, users["admin"]))
            connection.commit()
        context.project = project
        original_as = context.as_
        context.as_ = lambda who: original_as(f"{who}-{tag}")
        yield context


def _template(api: Api, content: bytes | None = None) -> tuple[dict, str]:
    api.as_("admin")
    content = content or SRS.read_bytes()
    file = store_bytes(name="SRS.docx", media_type=DOCX, purpose="template", content=content)
    group = api.send("POST", "/templates", {"name": f"438C {uuid4().hex[:6]}", "category": "438C"}).json()["data"]
    structure = api.send("POST", "/template-documents/parse", {"wordFileId": str(file.id)}).json()["data"]
    structure["reviewed"] = True
    created = api.send("POST", f"/templates/{group['id']}/documents", {"name": "软件需求规格说明", "wordFileId": str(file.id), "importStructure": structure})
    assert created.status_code == 200, created.text
    return created.json()["data"], file.sha256


def _new_doc(api: Api, td_id: str, name: str) -> dict:
    response = api.send("POST", f"/projects/{api.project}/documents", {"name": name, "sourceType": "template", "templateDocumentId": td_id})
    assert response.status_code == 200, response.text
    return response.json()["data"]


def _base(api: Api, td_id: str) -> dict:
    response = api.get(f"/template-documents/{td_id}/draft")
    assert response.status_code == 200, response.text
    data = response.json()["data"]
    return data["draft"] or data["base"]


def _top(blocks):
    return [block for block in blocks if not block.get("tableId")]


def _chapter(blocks, heading):
    top = _top(blocks)
    start = top.index(heading)
    end = next((i for i in range(start + 1, len(top)) if top[i]["kind"] == "heading" and top[i]["headingLevel"] <= heading["headingLevel"]), len(top))
    return top[start:end]


def _with_children(blocks, order):
    children = {}
    for block in blocks:
        if block.get("tableId"):
            children.setdefault(block["tableId"], []).append(block)
    return [item for block in order for item in [block] + children.get(block["id"], [])]


def _a61_edits(structure: dict) -> tuple[dict, dict]:
    """16 §3.8 A6-1：改标题文字和级别、改一段正文、新增一章、删除一章、移动一章、改单元格、改指导、改正文字号。"""
    draft = copy.deepcopy(structure)
    blocks = draft["blocks"]
    top = _top(blocks)
    heads = [block for block in top if block["kind"] == "heading"]
    expect = {}
    body = next(block for block in top if block["kind"] == "body" and not block["lock"]["text"])
    body["text"] += "（已修改）"
    expect["body"] = (body["id"], body["text"])
    heading = next(block for block in heads[2:] if not block["lock"]["text"] and block["headingLevel"] < 4)
    heading["text"] += "（改）"
    heading["headingLevel"] += 1
    expect["heading"] = (heading["id"], heading["text"], heading["headingLevel"])
    guidance = next(block for block in top if block["kind"] == "guidance")
    guidance["text"] += "（指导修改）"
    expect["guidance"] = (guidance["id"], guidance["text"])
    table = next(block for block in top if block.get("isTable") and block["tableEditable"] in {"cells", "partial"})
    locked = {tuple(cell) for cell in table["lockedCells"]}
    r, c = next((r, c) for r in range(len(table["cells"]) - 1, -1, -1) for c in range(len(table["cells"][r])) if (r, c) not in locked)
    table["cells"][r][c] = (table["cells"][r][c] + "改").strip()
    expect["cell"] = (table["id"], copy.deepcopy(table["cells"]))
    # 删除一章：找一章全部可删（不含表格）的末级章节
    removable = next(h for h in heads[3:] if h is not heading and all(not b.get("isTable") and not b["lock"]["delete"] for b in _chapter(blocks, h))
                     and body not in _chapter(blocks, h) and guidance not in _chapter(blocks, h))
    removed = {block["id"] for block in _chapter(blocks, removable)}
    expect["removed"] = removed
    blocks[:] = [block for block in blocks if block["id"] not in removed]
    # 新增一章（标题 + 正文），放在第一章末尾
    level = heads[0]["headingLevel"]
    first = _chapter(blocks, heads[0])
    new_heading = {"id": new_block_id(), "path": None, "kind": "heading", "text": "新增章节", "sourceText": "", "headingLevel": level + 1,
                   "headingId": None, "tableId": None, "protected": False, "isTable": False}
    new_body = {"id": new_block_id(), "path": None, "kind": "body", "text": "新增章节的正文。", "sourceText": "", "headingLevel": None,
                "headingId": new_heading["id"], "tableId": None, "protected": False, "isTable": False}
    top = _top(blocks)
    position = top.index(first[-1]) + 1
    top[position:position] = [new_heading, new_body]
    expect["added"] = (new_heading["id"], new_body["id"])
    # 移动一章：把倒数第二个一级章节移到最后
    firsts = [block for block in top if block["kind"] == "heading" and block["headingLevel"] == level]
    a, b = top.index(firsts[-2]), top.index(firsts[-1])
    top = top[:a] + top[b:] + top[a:b]
    expect["order"] = [block["id"] for block in top]
    draft["blocks"] = _with_children(blocks, top)
    for block in draft["blocks"]:
        if block["id"] == new_body["id"]:
            pass
    styles = copy.deepcopy(structure.get("_styles") or {})
    return draft, expect | {"styles": styles}


def test_partial_unique_indexes_and_canonical_hash(api):
    td, _ = _template(api)
    with connect() as connection:
        version = connection.execute("SELECT * FROM template_document_versions WHERE template_document_id=%s", (td["id"],)).fetchone()
        assert version["status"] == "published" and version["version"] == 1 and version["edited_file_id"] is None
        with pytest.raises(sqlite3.IntegrityError):
            connection.execute("""INSERT INTO template_document_versions (id, template_document_id, version, status, source_file_id, source_sha256,
                                  structure_json, format_template_json, structure_sha256) VALUES (%s,%s,NULL,'published',%s,'x','{}','{}','x')""",
                               (uuid4(), td["id"], version["source_file_id"]))
        connection.rollback()
        connection.execute("""INSERT INTO template_document_versions (id, template_document_id, version, status, source_file_id, source_sha256,
                              structure_json, format_template_json, structure_sha256) VALUES (%s,%s,NULL,'draft',%s,'x','{}','{}','x')""",
                           (uuid4(), td["id"], version["source_file_id"]))
        with pytest.raises(sqlite3.IntegrityError):
            connection.execute("""INSERT INTO template_document_versions (id, template_document_id, version, status, source_file_id, source_sha256,
                                  structure_json, format_template_json, structure_sha256) VALUES (%s,%s,NULL,'draft',%s,'x','{}','{}','x')""",
                               (uuid4(), td["id"], version["source_file_id"]))
        connection.rollback()
    structure = parse_template(SRS.read_bytes())
    shuffled = json.loads(json.dumps(structure))
    shuffled["revision"] = 99
    shuffled["baseFileId"] = "x"
    shuffled = {key: shuffled[key] for key in reversed(list(shuffled))}
    assert tm.canonical_hash(structure) == tm.canonical_hash(shuffled)
    shuffled["blocks"][0]["text"] += "x"
    assert tm.canonical_hash(structure) != tm.canonical_hash(shuffled)


def test_a06_edit_publish_new_doc_old_doc_diff_restore(api):
    td, source_sha = _template(api)
    td_id = td["id"]
    assert td["currentVersion"]["seq"] == 1 and td["hasDraft"] is False
    old_doc = _new_doc(api, td_id, "老文档")
    detail = api.get(f"/documents/{old_doc['id']}").json()["data"]
    assert detail["template"]["seq"] == 1 and detail["template"]["legacySnapshot"] is False
    with connect() as connection:
        old_snapshot = connection.execute("SELECT format_template FROM documents WHERE id=%s", (old_doc["id"],)).fetchone()[0]
    base = _base(api, td_id)
    assert base["draftRevision"] == 0 and base["parentVersion"]["seq"] == 1
    tables = [block for block in base["structure"]["blocks"] if block.get("isTable")]
    assert tables and all(block["tableEditable"] in {"cells", "partial", "none"} for block in tables)
    draft, expect = _a61_edits(base["structure"])
    styles = copy.deepcopy(base["format"]["styles"])
    styles["正文"]["size"] = 12
    saved = api.send("PUT", f"/template-documents/{td_id}/draft", {"structure": draft, "format": {"styles": styles}, "expectedDraftRevision": 0})
    assert saved.status_code == 200, saved.text
    assert saved.json()["data"]["draftRevision"] == 1
    # 并发：后保存的一方 409，带最后修改人
    stale = api.send("PUT", f"/template-documents/{td_id}/draft", {"structure": draft, "format": {"styles": styles}, "expectedDraftRevision": 0})
    assert stale.status_code == 409 and stale.json()["error"]["code"] == "REVISION_CONFLICT"
    assert stale.json()["error"]["details"]["lastActor"]["name"] == "A06admin"
    # 样式影响预览（K-10）
    impact = api.send("POST", f"/template-documents/{td_id}/draft:format-impact", {"changes": [{"style": "正文", "prop": "size", "value": 12}]})
    assert impact.status_code == 200, impact.text
    item = impact.json()["data"]["items"][0]
    assert item["affectedParagraphs"] > 0 and item["directFormatParagraphs"] <= item["affectedParagraphs"] and item["willClear"] == item["directFormatParagraphs"]
    summary = api.send("POST", f"/template-documents/{td_id}/draft:format-impact", {}).json()["data"]["items"]
    assert [(entry["style"], entry["prop"], entry["after"]) for entry in summary] == [("正文", "size", 12)]
    # 发布
    assert api.send("POST", f"/template-documents/{td_id}/versions", {"draftRevision": 1, "note": ""}).status_code == 422
    assert api.send("POST", f"/template-documents/{td_id}/versions", {"draftRevision": 1, "note": "x" * 201}).status_code == 422
    published = api.send("POST", f"/template-documents/{td_id}/versions", {"draftRevision": 1, "note": "A6-1 八处修改"})
    assert published.status_code == 201, published.text
    v2 = published.json()["data"]["version"]
    assert v2["seq"] == 2 and v2["editedSha256"] and v2["sourceSha256"] == source_sha
    versions = api.get(f"/template-documents/{td_id}/versions").json()["data"]
    assert [item["seq"] for item in versions["items"]] == [2, 1] and versions["current"]["seq"] == 2 and versions["draft"] is None
    assert versions["original"]["sha256"] == source_sha and versions["usage"]["total"] == 1
    v1 = versions["items"][1]
    original = api.get(f"/template-documents/{td_id}/versions/{v1['versionId']}/file?kind=source")
    assert original.status_code == 200 and original.content == SRS.read_bytes() and original.headers["etag"] == f'"{source_sha}"'
    edited = api.get(f"/template-documents/{td_id}/versions/{v2['versionId']}/file?kind=edited")
    assert edited.status_code == 200
    version_structure = api.get(f"/template-documents/{td_id}/versions/{v2['versionId']}").json()["data"]["structure"]
    by_id = {block["id"]: block for block in version_structure["blocks"]}
    assert by_id[expect["body"][0]]["text"] == expect["body"][1]
    assert (by_id[expect["heading"][0]]["text"], by_id[expect["heading"][0]]["headingLevel"]) == expect["heading"][1:]
    assert by_id[expect["guidance"][0]]["text"] == expect["guidance"][1]
    assert by_id[expect["cell"][0]]["cells"] == expect["cell"][1]
    assert not expect["removed"] & set(by_id)
    assert all(by_id[block_id]["path"] for block_id in expect["added"])
    assert [block["id"] for block in _top(version_structure["blocks"]) if not block.get("virtual")] == expect["order"]
    from app.template_word import read_styles
    assert read_styles(Document(BytesIO(edited.content)))["正文"]["size"] == 12
    # 对比（A6-3）
    diff = api.get(f"/template-documents/{td_id}/versions/{v1['versionId']}/diff/{v2['versionId']}").json()["data"]
    assert diff["summary"]["format"] == 1 and diff["summary"]["headings"] >= 3 and diff["summary"]["guidance"] >= 1 and diff["summary"]["tables"] >= 1
    changes = {item["blockId"]: item["change"] for item in diff["blocks"]}
    assert changes[expect["added"][0]] == "added" and changes[expect["body"][0]] == "modified"
    assert all(changes[block_id] == "removed" for block_id in expect["removed"] if block_id in changes)
    # 新文档用第 2 版；老文档拷贝不变（A6-2）
    new_doc = _new_doc(api, td_id, "新文档")
    assert "## 新增章节" in new_doc["markdownContent"] and "（已修改）" in new_doc["markdownContent"]
    assert new_doc["formatTemplate"]["importStructure"]["revision"] == 2
    assert api.get(f"/documents/{new_doc['id']}").json()["data"]["template"]["seq"] == 2
    old_detail = api.get(f"/documents/{old_doc['id']}").json()["data"]["template"]
    assert old_detail["seq"] == 1 and old_detail["latestSeq"] == 2
    with connect() as connection:
        assert connection.execute("SELECT format_template FROM documents WHERE id=%s", (old_doc["id"],)).fetchone()[0] == old_snapshot
    # 无草稿 / 无变化
    assert api.send("POST", f"/template-documents/{td_id}/versions", {"draftRevision": 1, "note": "再来"}).status_code == 409
    same = _base(api, td_id)
    assert api.send("PUT", f"/template-documents/{td_id}/draft", {"structure": same["structure"], "format": same["format"], "expectedDraftRevision": 0}).status_code == 200
    unchanged = api.send("POST", f"/template-documents/{td_id}/versions", {"draftRevision": 1, "note": "无变化"})
    assert unchanged.status_code == 409 and unchanged.json()["error"]["code"] == "TEMPLATE_UNCHANGED"
    # 恢复第 1 版（A6-4）
    restored = api.send("POST", f"/template-documents/{td_id}/draft:restore", {"versionId": v1["versionId"], "expectedDraftRevision": 1})
    assert restored.status_code == 200, restored.text
    restored_data = restored.json()["data"]
    assert restored_data["restoredFrom"]["seq"] == 1 and restored_data["notePrefill"] == "恢复自第 1 版"
    v3 = api.send("POST", f"/template-documents/{td_id}/versions", {"draftRevision": restored_data["draftRevision"], "note": restored_data["notePrefill"]})
    assert v3.status_code == 201, v3.text
    assert v3.json()["data"]["version"]["restoredFrom"]["seq"] == 1
    v3_structure = api.get(f"/template-documents/{td_id}/versions/{v3.json()['data']['version']['versionId']}").json()["data"]["structure"]
    v1_structure = api.get(f"/template-documents/{td_id}/versions/{v1['versionId']}").json()["data"]["structure"]
    def view(structure):
        # 被删章节里的编写指导恢复为“虚拟指导”（不进 Word，只在结构里），所以连同虚拟块一起比较
        return [(block["kind"], block["text"], block["headingLevel"]) for block in _top(structure["blocks"])]
    assert view(v3_structure) == view(v1_structure)
    assert api.get(f"/template-documents/{td_id}/versions/{v2['versionId']}/file?kind=edited").status_code == 200


def _edit_text(structure, block_id, suffix):
    draft = copy.deepcopy(structure)
    block = next(item for item in draft["blocks"] if item["id"] == block_id)
    block["text"] += suffix
    return draft


def test_a613_stale_publish_and_rebase(api):
    td, _ = _template(api)
    td_id = td["id"]
    base = _base(api, td_id)
    editable = [block["id"] for block in _top(base["structure"]["blocks"]) if block["kind"] in {"body", "heading"} and not block["lock"]["text"]]
    x, y, z = editable[1], editable[4], editable[7]
    v1 = base["parentVersion"]["versionId"]
    # A：改 x、z 并发布第 2 版
    draft = _edit_text(_edit_text(base["structure"], x, "（A）"), z, "（A）")
    assert api.send("PUT", f"/template-documents/{td_id}/draft", {"structure": draft, "expectedDraftRevision": 0}).status_code == 200
    assert api.send("POST", f"/template-documents/{td_id}/versions", {"draftRevision": 1, "note": "A 的修改"}).status_code == 201
    # B：仍基于第 1 版，改 x、y（草稿已被 A 发布，带上 parentVersionId 重新建草稿）
    mine = _edit_text(_edit_text(base["structure"], x, "（B）"), y, "（B）")
    saved = api.send("PUT", f"/template-documents/{td_id}/draft", {"structure": mine, "expectedDraftRevision": 1, "parentVersionId": v1})
    assert saved.status_code == 200, saved.text
    stale = api.send("POST", f"/template-documents/{td_id}/versions", {"draftRevision": 1, "note": "B 的修改"})
    assert stale.status_code == 409 and stale.json()["error"]["code"] == "TEMPLATE_DRAFT_STALE"
    assert stale.json()["error"]["details"]["currentVersion"]["seq"] == 2
    dry = api.send("POST", f"/template-documents/{td_id}/draft:rebase", {"expectedDraftRevision": 1, "dryRun": True})
    assert dry.status_code == 200, dry.text
    data = dry.json()["data"]
    assert [(item["blockId"], item["field"]) for item in data["conflicts"]] == [(x, "text")]
    assert sum(data["theirsChanges"]["summary"].values()) == 2
    refused = api.send("POST", f"/template-documents/{td_id}/draft:rebase", {"expectedDraftRevision": 1})
    assert refused.status_code == 409 and refused.json()["error"]["code"] == "TEMPLATE_REBASE_CONFLICT"
    merged = api.send("POST", f"/template-documents/{td_id}/draft:rebase",
                      {"expectedDraftRevision": 1, "resolutions": [{"blockId": x, "field": "text", "take": "mine"}]})
    assert merged.status_code == 200, merged.text
    assert merged.json()["data"]["parentVersion"]["seq"] == 2
    v3 = api.send("POST", f"/template-documents/{td_id}/versions", {"draftRevision": 2, "note": "合并后发布"})
    assert v3.status_code == 201, v3.text
    structure = api.get(f"/template-documents/{td_id}/versions/{v3.json()['data']['version']['versionId']}").json()["data"]["structure"]
    texts = {block["id"]: block["text"] for block in structure["blocks"]}
    assert texts[x].endswith("（B）") and texts[y].endswith("（B）") and texts[z].endswith("（A）")


def test_upload_new_word_alignment_and_unplaced_guidance(api):
    td, _ = _template(api)
    td_id = td["id"]
    base = _base(api, td_id)
    same = store_bytes(name="SRS.docx", media_type=DOCX, purpose="template", content=SRS.read_bytes())
    unchanged = api.send("POST", f"/template-documents/{td_id}/draft:upload", {"wordFileId": str(same.id), "expectedDraftRevision": 0})
    assert unchanged.status_code == 409 and unchanged.json()["error"]["code"] == "TEMPLATE_UNCHANGED"
    document = Document(SRS)
    guidance = next(block for block in _top(base["structure"]["blocks"]) if block["kind"] == "guidance")
    edited = 0
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == guidance["sourceText"].strip():
            paragraph._p.getparent().remove(paragraph._p)   # 指导所在段落被删 → 指导“未落位”
            edited += 1
            break
    body = next(block for block in _top(base["structure"]["blocks"]) if block["kind"] == "body" and not block["lock"]["text"])
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == body["sourceText"].strip() and paragraph.runs:
            paragraph.runs[0].text = paragraph.runs[0].text + "（Word 中修改）"
            edited += 1
            break
    assert edited == 2
    output = BytesIO()
    document.save(output)
    file = store_bytes(name="SRS-改.docx", media_type=DOCX, purpose="template", content=output.getvalue())
    uploaded = api.send("POST", f"/template-documents/{td_id}/draft:upload", {"wordFileId": str(file.id), "expectedDraftRevision": 0})
    assert uploaded.status_code == 200, uploaded.text
    data = uploaded.json()["data"]
    assert data["uploaded"] is True and data["alignment"]["matched"] > 100
    assert [item["blockId"] for item in data["alignment"]["unmatched"]] == [guidance["id"]] and data["unplacedCount"] == 1
    assert next(block for block in data["structure"]["blocks"] if block["id"] == body["id"])["text"].find("（Word 中修改）") >= 0
    blocked = api.send("POST", f"/template-documents/{td_id}/versions", {"draftRevision": data["draftRevision"], "note": "上传"})
    assert blocked.status_code == 422 and blocked.json()["error"]["details"]["items"][0]["blockId"] == guidance["id"]
    # 指定位置：放到第一个标题之后
    structure = data["structure"]
    blocks = structure["blocks"]
    item = next(block for block in blocks if block["id"] == guidance["id"])
    blocks.remove(item)
    item.pop("unplaced")
    first = next(i for i, block in enumerate(blocks) if block["kind"] == "heading")
    blocks.insert(first + 1, item)
    placed = api.send("PUT", f"/template-documents/{td_id}/draft", {"structure": structure, "format": data["format"], "expectedDraftRevision": data["draftRevision"]})
    assert placed.status_code == 200, placed.text
    published = api.send("POST", f"/template-documents/{td_id}/versions", {"draftRevision": placed.json()["data"]["draftRevision"], "note": "上传新 Word"})
    assert published.status_code == 201, published.text
    version = published.json()["data"]["version"]
    assert version["sourceSha256"] == file.sha256 and version["editedSha256"] is None
    doc = _new_doc(api, td_id, "上传后新建")
    assert "（Word 中修改）" in doc["markdownContent"]


def test_negative_locks_limits_and_permissions(api):
    td, _ = _template(api)
    td_id = td["id"]
    base = _base(api, td_id)
    url = f"/template-documents/{td_id}/draft"
    locked = next((block for block in _top(base["structure"]["blocks"]) if block["lock"]["text"] and block["kind"] == "body" and not block.get("isTable")), None)
    if locked is not None:
        response = api.send("PUT", url, {"structure": _edit_text(base["structure"], locked["id"], "改"), "expectedDraftRevision": 0})
        assert response.status_code == 422 and response.json()["error"]["code"] == "TEMPLATE_EDIT_UNSUPPORTED"
        assert response.json()["error"]["details"]["items"][0]["blockId"] == locked["id"]
    undeletable = next(block for block in _top(base["structure"]["blocks"]) if block["lock"]["delete"] and not block.get("isTable"))
    draft = copy.deepcopy(base["structure"])
    draft["blocks"] = [block for block in draft["blocks"] if block["id"] != undeletable["id"]]
    response = api.send("PUT", url, {"structure": draft, "expectedDraftRevision": 0})
    assert response.status_code == 422 and response.json()["error"]["details"]["items"][0] == {"blockId": undeletable["id"], "message": undeletable["lock"]["delete"]}
    table = next(block for block in _top(base["structure"]["blocks"]) if block.get("isTable"))
    draft = copy.deepcopy(base["structure"])
    draft["blocks"] = [block for block in draft["blocks"] if block["id"] != table["id"] and block.get("tableId") != table["id"]]
    assert api.send("PUT", url, {"structure": draft, "expectedDraftRevision": 0}).status_code == 422
    partial = next((block for block in _top(base["structure"]["blocks"]) if block.get("isTable") and block["lockedCells"]), None)
    if partial is not None:
        draft = copy.deepcopy(base["structure"])
        target = next(block for block in draft["blocks"] if block["id"] == partial["id"])
        r, c = partial["lockedCells"][0]
        target["cells"][r][c] += "改"
        response = api.send("PUT", url, {"structure": draft, "expectedDraftRevision": 0})
        assert response.status_code == 422
        problem = response.json()["error"]["details"]["items"][0]
        assert problem["cell"] == [r, c] and problem["message"] in {reason["message"] for reason in partial["tableEditReasons"]}
    big = copy.deepcopy(base["structure"])
    big["blocks"] += [{"id": new_block_id(), "path": None, "kind": "body", "text": "x", "headingLevel": None} for _ in range(4001)]
    assert api.send("PUT", url, {"structure": big, "expectedDraftRevision": 0}).status_code == 413
    forged = copy.deepcopy(base["structure"])
    forged["blocks"][3]["path"] = forged["blocks"][5]["path"]
    assert api.send("PUT", url, {"structure": forged, "expectedDraftRevision": 0}).status_code == 422
    bad_style = api.send("PUT", url, {"structure": base["structure"], "format": {"styles": {"封面标题": {"size": 12}}}, "expectedDraftRevision": 0})
    assert bad_style.status_code == 422
    # 权限（17 §3.9）
    api.as_("creator")
    assert api.get(f"/template-documents/{td_id}/versions").status_code == 200
    assert api.send("PUT", url, {"structure": base["structure"], "expectedDraftRevision": 0}).status_code == 403
    assert api.send("POST", f"/template-documents/{td_id}/versions", {"draftRevision": 1, "note": "x"}).status_code == 403
    api.as_("auditor")
    assert api.get(f"/template-documents/{td_id}/versions").status_code == 403
    assert api.get(f"/template-documents/{td_id}/draft").status_code == 403
    api.as_("maint")
    assert api.send("PUT", url, {"structure": _edit_text(base["structure"], next(b["id"] for b in _top(base["structure"]["blocks"]) if b["kind"] == "guidance"), "改"),
                                 "expectedDraftRevision": 0}).status_code == 200
    assert api.send("POST", f"/template-documents/{td_id}/versions", {"draftRevision": 1, "note": "维护者发布"}).status_code == 201
    # 审计员看文档详情里的模板版本（只读）
    api.as_("admin")
    doc = _new_doc(api, td_id, "审计查看")
    api.as_("auditor")
    assert api.get(f"/documents/{doc['id']}").json()["data"]["template"]["seq"] == 2


def test_concurrent_publish_exactly_one(api):
    td, _ = _template(api)
    td_id = td["id"]
    base = _base(api, td_id)
    guidance = next(b["id"] for b in _top(base["structure"]["blocks"]) if b["kind"] == "guidance")
    assert api.send("PUT", f"/template-documents/{td_id}/draft", {"structure": _edit_text(base["structure"], guidance, "并发"), "expectedDraftRevision": 0}).status_code == 200
    results = []
    headers = dict(api.headers)
    cookies = dict(api.client.cookies)

    def go():
        from app.main import create_app
        with TestClient(create_app()) as other:
            other.cookies.update(cookies)
            results.append(other.post(f"/api/v1/template-documents/{td_id}/versions", json={"draftRevision": 1, "note": "并发发布"}, headers=headers).status_code)
    threads = [threading.Thread(target=go) for _ in range(2)]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join()
    assert sorted(results) == [201, 409]
    with connect() as connection:
        assert connection.execute("SELECT count(*) FROM template_document_versions WHERE template_document_id=%s AND status='published'", (td_id,)).fetchone()[0] == 1


def test_backfill_binds_matching_documents_and_marks_legacy(api, tmp_path):
    from app.migrate_template_versions import backfill
    api.as_("admin")
    content = SRS.read_bytes()
    file = store_bytes(name="SRS.docx", media_type=DOCX, purpose="template", content=content)
    structure = parse_template(content)
    structure["reviewed"] = True
    structure["revision"] = 3
    group = api.send("POST", "/templates", {"name": f"旧模板 {uuid4().hex[:6]}", "category": "438C"}).json()["data"]
    td_id, unreviewed_id = uuid4(), uuid4()
    same, legacy = copy.deepcopy(structure), copy.deepcopy(structure)
    same["revision"] = 1
    next(block for block in legacy["blocks"] if block["kind"] == "guidance")["text"] += "（旧）"
    with connect() as connection:
        connection.execute("INSERT INTO template_documents (id, template_id, name, word_file_id, format_template) VALUES (%s,%s,%s,%s,%s)",
                           (td_id, group["id"], "已审核", file.id, json.dumps({"importStructure": structure})))
        connection.execute("INSERT INTO template_documents (id, template_id, name, word_file_id, format_template) VALUES (%s,%s,%s,%s,%s)",
                           (unreviewed_id, group["id"], "未审核", file.id, json.dumps({})))
        for name, snapshot in (("同结构", same), ("旧结构", legacy), ("无结构", None)):
            connection.execute("""INSERT INTO documents (id, project_id, name, source_type, template_document_id, format_template, markdown_content,
                                  outline, status, owner_id, revision) VALUES (%s,%s,%s,'template',%s,%s,'','[]','draft',%s,0)""",
                               (uuid4(), api.project, name, td_id, json.dumps({"importStructure": snapshot} if snapshot else {}), api.users["admin"]))
        connection.commit()
    report = backfill()
    assert report["published"] >= 1 and report["drafts"] >= 1 and report["bound"] == 1 and report["legacy"] == 2
    assert report["backup"] and Path(report["backup"]).exists()
    again = backfill()
    assert again["templateDocuments"] == 0 and again["bound"] == 0
    with connect() as connection:
        rows = connection.execute("SELECT name, template_version_id FROM documents WHERE template_document_id=%s ORDER BY name", (td_id,)).fetchall()
        assert {row["name"]: row["template_version_id"] is not None for row in rows} == {"同结构": True, "旧结构": False, "无结构": False}
        assert connection.execute("SELECT status FROM template_document_versions WHERE template_document_id=%s", (unreviewed_id,)).fetchone()[0] == "draft"
    usage = api.get(f"/template-documents/{td_id}/versions").json()["data"]["usage"]
    assert sorted((item["documentName"], item["legacy"], item["seq"]) for item in usage["visible"]) == [
        ("同结构", False, 1), ("无结构", True, None), ("旧结构", True, None)]
    legacy_doc = next(item for item in usage["visible"] if item["documentName"] == "旧结构")
    detail = api.get(f"/documents/{legacy_doc['documentId']}").json()["data"]["template"]
    assert detail["legacySnapshot"] is True and detail["seq"] is None
    # 未审核的模板：核对确认后成为第 1 版；之后 PATCH /structure 返回 TEMPLATE_VERSIONED
    pending = api.get(f"/template-documents/{unreviewed_id}").json()["data"]["importStructure"]
    pending["reviewed"] = True
    confirmed = api.send("PATCH", f"/template-documents/{unreviewed_id}/structure", {"expectedRevision": 0, "importStructure": pending})
    assert confirmed.status_code == 200, confirmed.text
    assert confirmed.json()["data"]["currentVersion"]["seq"] == 1


def test_find_or_store_upload_not_docx_and_publish_racing_new_documents(api):
    """17 §3.11：find_or_store 命中/逐字节比对/未命中；上传非 docx 得到 422；发布与新建文档并发，每篇文档都绑定存在的版本。"""
    from app.template_versions import find_or_store

    td, _ = _template(api)
    td_id = td["id"]
    content = SRS.read_bytes()
    with connect() as connection:
        first, digest = find_or_store(connection, content, "SRS.docx")
        again, _ = find_or_store(connection, content, "SRS.docx")
        other, _ = find_or_store(connection, content + b"\0", "SRS-2.docx")
        connection.commit()
    assert again == first and other != first and len(digest) == 64
    junk = store_bytes(name="不是Word.docx", media_type=DOCX, purpose="template", content=b"not a docx at all")
    bad = api.send("POST", f"/template-documents/{td_id}/draft:upload", {"wordFileId": str(junk.id), "expectedDraftRevision": 0})
    assert bad.status_code == 422, bad.text

    base = _base(api, td_id)
    guidance = next(b["id"] for b in _top(base["structure"]["blocks"]) if b["kind"] == "guidance")
    created: list = []
    errors: list = []
    headers = dict(api.headers)
    cookies = dict(api.client.cookies)

    def create_many():
        from app.main import create_app
        with TestClient(create_app()) as other_client:
            other_client.cookies.update(cookies)
            for index in range(RACE_DOCUMENTS):
                response = other_client.post(f"/api/v1/projects/{api.project}/documents", headers=headers,
                                             json={"name": f"并发 {index}", "sourceType": "template", "templateDocumentId": td_id})
                (created if response.status_code == 200 else errors).append(response.json()["data"]["id"] if response.status_code == 200 else response.text)

    worker = threading.Thread(target=create_many)
    worker.start()
    revision = 0
    for round_ in range(3):
        structure = _edit_text(_base(api, td_id)["structure"], guidance, f"并发第 {round_} 次")
        saved = api.send("PUT", f"/template-documents/{td_id}/draft", {"structure": structure, "expectedDraftRevision": revision})
        assert saved.status_code == 200, saved.text
        published = api.send("POST", f"/template-documents/{td_id}/versions", {"draftRevision": saved.json()["data"]["draftRevision"], "note": f"并发 {round_}"})
        assert published.status_code == 201, published.text
    worker.join()
    assert not errors and len(created) == RACE_DOCUMENTS
    with connect() as connection:
        rows = connection.execute(f"""SELECT d.template_version_id, v.status FROM documents d LEFT JOIN template_document_versions v
                                       ON v.id = d.template_version_id WHERE d.id IN ({','.join(['%s'] * len(created))})""", created).fetchall()
    assert len(rows) == RACE_DOCUMENTS and all(row["template_version_id"] and row["status"] in {"published", "superseded"} for row in rows)

