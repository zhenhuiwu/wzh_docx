"""代码审查 M8：保存路径（_check_word → edit_original_docx → _edit_paragraph）里长段落改动后，未改动的字要保留原样式。
7800032 的贪心对齐在中间部分超过 2000 字时错配：1 万字插 10 处，保留率最差 0.559；要求接近 1.0。"""

import random
import re
import time

import pytest
from docx import Document
from lxml import html as LH

from app.docx_export import _edit_paragraph
from app.fast_diff import bounded_opcodes

PHRASES = ["系统应当支持", "用户在提交申请后", "管理员可以查看", "数据需要定期备份", "，", "。", "并且", "需求规格说明书",
           "接口返回的结果", "在规定时间内完成", "相关人员", "审批流程", "文档", "保存", "格式", "第三方服务", "日志记录",
           # 插入标记里的字（【】、改、数字）在原文里也出现——贪心对齐会被它们拉偏（审查员复现用的真实语料即如此）
           "【注1】", "修改", "第2条", "改进", "【说明】", "3", "0"]


def _text(rng: random.Random, length: int) -> str:
    out = ""
    while len(out) < length:
        out += rng.choice(PHRASES)
    return out[:length]


def _keep_ratio(length: int, edits: int, seed: int) -> tuple[float, float]:
    rng = random.Random(seed)
    src = _text(rng, length)
    document = Document()
    paragraph = document.add_paragraph()
    bold: list[bool] = []
    pos = k = 0
    while pos < length:
        n = rng.randint(10, 60)
        piece = src[pos:pos + n]
        paragraph.add_run(piece).bold = k % 2 == 1
        bold += [k % 2 == 1] * len(piece)
        pos += n
        k += 1
    target = list(src)
    for number, q in enumerate(sorted((rng.randrange(length + 1) for _ in range(edits)), reverse=True)):
        target[q:q] = list(f"【改{number}】")
    tgt = "".join(target)
    element = LH.fragment_fromstring(f"<p>{tgt}</p>")
    started = time.perf_counter()
    _edit_paragraph(paragraph, element)
    elapsed = time.perf_counter() - started
    assert paragraph.text == tgt
    got = [bool(run.bold) for run in paragraph.runs for _ in run.text]
    marks = {x for m in re.finditer(r"【改\d+】", tgt) for x in range(m.start(), m.end())}
    keep = total = i = 0
    for j in range(len(tgt)):
        if j in marks:
            continue
        total += 1
        keep += got[j] == bold[i]
        i += 1
    return keep / total, elapsed


@pytest.mark.parametrize("length,edits", [(3000, 2), (3000, 10), (10000, 2), (10000, 10), (30000, 10)])
def test_unchanged_text_keeps_run_style_in_long_paragraph(length, edits):
    results = [_keep_ratio(length, edits, seed) for seed in range(3)]
    worst = min(ratio for ratio, _ in results)
    assert worst >= 0.99, f"未改动的字样式保留率 {worst:.3f}"
    assert max(seconds for _, seconds in results) < 3.0


def test_opcodes_reconstruct_and_stay_fast_when_texts_differ_entirely():
    rng = random.Random(7)
    alphabet = [chr(code) for code in range(0x4E00, 0x4E00 + 3000)]
    a = "".join(rng.choice(alphabet) for _ in range(200_000))
    b = "".join(rng.choice(alphabet) for _ in range(200_000))
    started = time.perf_counter()
    ops = bounded_opcodes(a, b)
    assert time.perf_counter() - started < 2.0
    rebuilt, ia = [], 0
    for op, i1, i2, j1, j2 in ops:
        assert i1 == ia
        ia = i2
        rebuilt.append(a[i1:i2] if op == "equal" else b[j1:j2])
        if op == "equal":
            assert a[i1:i2] == b[j1:j2]
    assert ia == len(a) and "".join(rebuilt) == b
