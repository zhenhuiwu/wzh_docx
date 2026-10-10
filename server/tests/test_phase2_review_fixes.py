"""阶段2 审查（10-阶段2审查.md）修复的回归测试。"""
import json

from app.db import connect
from tests.test_document_access import _as, members  # noqa: F401
from tests.test_document_concurrency import TEXT, current, pair, save, team  # noqa: F401
from tests.test_documents import document_context  # noqa: F401


def _resolve(client, did, cid, headers, base, items, key):
    return client.post(f"/api/v1/documents/{did}/conflicts/{cid}/resolve", headers=headers,
                       json={"clientSaveKey": key, "baseDocumentRevision": base, "items": items})


def test_s1_theirs_deleted_keep_mine_restores_in_place(team):
    client, did, _users = team
    assert save(client, did, 1, TEXT.replace("第二段。\n\n", ""), _as(client, "editor")).status_code == 200
    b = _as(client, "editor2")
    r = save(client, did, 1, TEXT.replace("第二段。", "B 改的第二段。"), b)
    assert r.status_code == 409
    details = r.json()["error"]["details"]
    assert details["conflicts"][0]["reason"] == "DELETED"
    resp = _resolve(client, did, details["conflictId"], b, 2, [{"itemIndex": 0, "action": "keep_mine"}], "s1-keep-mine-1")
    assert resp.status_code == 200, resp.text
    assert resp.json()["data"]["markdownContent"].startswith("# 1 范围\n\n第一段。\n\nB 改的第二段。\n\n第三段。")
    # 界面验收 M2：全部处理后 summary 仍给出已处理条数，前端据此保留“冲突”入口
    assert client.get(f"/api/v1/documents/{did}/conflicts/summary").json()["data"] == {"mineOpen": 0, "mineHandled": 1, "allOpen": None}


def test_s1_reconflict_keeps_object_ids_for_placement(team):
    client, did, _users = team
    assert save(client, did, 1, TEXT.replace("第二段。", "A 的第二段。"), _as(client, "editor")).status_code == 200
    r = save(client, did, 1, TEXT.replace("第二段。", "B 的第二段。"), _as(client, "editor2"))
    cid = r.json()["error"]["details"]["conflictId"]
    # 对方随后把这段删掉：keep_mine 再次冲突，但条目仍记住原对象（用于定位）
    base = current(client, did)["revision"]
    assert save(client, did, base, TEXT.replace("A 的第二段。\n\n", "").replace("第二段。\n\n", ""), _as(client, "editor")).status_code == 200
    b = _as(client, "editor2")
    again = _resolve(client, did, cid, b, base + 1, [{"itemIndex": 0, "action": "keep_mine"}], "s1-reconflict-1")
    assert again.status_code == 409
    with connect() as connection:
        payload = json.loads(connection.execute("SELECT payload_json FROM document_conflicts WHERE id=%s", (cid,)).fetchone()["payload_json"])
    snap = payload["attempts"][-1]["items"]["0"]
    assert snap["objectIds"] and snap["goneObjectIds"] == snap["objectIds"]
    keep_both = _resolve(client, did, cid, b, base + 1, [{"itemIndex": 0, "action": "keep_both"}], "s1-keep-both-1")
    assert keep_both.status_code == 200, keep_both.text
    assert keep_both.json()["data"]["markdownContent"].startswith("# 1 范围\n\n第一段。\n\nB 的第二段。")


# ---------------------------------------------------------------- 复审 M3：恢复为新段落的位置

import pytest  # noqa: E402


def _deleted_conflict(client, did, b_text):
    b = _as(client, "editor2")
    r = save(client, did, 1, b_text, b)
    assert r.status_code == 409, r.text
    details = r.json()["error"]["details"]
    assert details["conflicts"][0]["reason"] == "DELETED"
    return b, details["conflictId"]


@pytest.mark.parametrize("action", ["restore_as_new", "keep_mine"])
def test_m3_restore_after_inserts_above(team, action):
    """审查员复现用例 test_r2_restore_after_inserts_above：对方删第二段后又有人在开头插入 3 段，恢复后仍在“第一段。”之后。"""
    client, did, _users = team
    a = _as(client, "editor")
    assert save(client, did, 1, TEXT.replace("第二段。\n\n", ""), a).status_code == 200
    cur = current(client, did)
    assert save(client, did, cur["revision"], "插一。\n\n插二。\n\n插三。\n\n" + cur["markdownContent"], a).status_code == 200
    b, cid = _deleted_conflict(client, did, TEXT.replace("第二段。", "B 改的第二段。"))
    res = _resolve(client, did, cid, b, current(client, did)["revision"], [{"itemIndex": 0, "action": action}], f"m3-{action}-1")
    assert res.status_code == 200, res.text
    data = res.json()["data"]
    assert data["markdownContent"].startswith("插一。\n\n插二。\n\n插三。\n\n# 1 范围\n\n第一段。\n\nB 改的第二段。\n\n第三段。")
    assert not any(item.get("placementFallback") for item in data["items"])


def test_m3_restore_follows_chain_of_deleted_predecessors(team):
    """前驱“第一段。”在冲突产生后也被删了：沿前驱链找到标题，放在标题之后。"""
    client, did, _users = team
    a = _as(client, "editor")
    assert save(client, did, 1, TEXT.replace("第二段。\n\n", ""), a).status_code == 200
    b, cid = _deleted_conflict(client, did, TEXT.replace("第二段。", "B 改的第二段。"))
    cur = current(client, did)
    a = _as(client, "editor")
    assert save(client, did, cur["revision"], "插一。\n\n" + cur["markdownContent"].replace("第一段。\n\n", ""), a).status_code == 200
    b = _as(client, "editor2")
    res = _resolve(client, did, cid, b, current(client, did)["revision"], [{"itemIndex": 0, "action": "restore_as_new"}], "m3-chain-1")
    assert res.status_code == 200, res.text
    assert res.json()["data"]["markdownContent"].startswith("插一。\n\n# 1 范围\n\nB 改的第二段。\n\n第三段。")


def test_m3_legacy_deleted_object_falls_back_to_heading_path_and_ordinal(team):
    client, did, _users = team
    a = _as(client, "editor")
    assert save(client, did, 1, TEXT.replace("第二段。\n\n", ""), a).status_code == 200
    cur = current(client, did)
    assert save(client, did, cur["revision"], "插一。\n\n插二。\n\n" + cur["markdownContent"], a).status_code == 200
    with connect() as connection:  # 模拟 114 之前删除的对象：没有前驱记录
        connection.execute("UPDATE document_objects SET deleted_after_ref=NULL WHERE document_id=%s", (did,))
        connection.commit()
    b, cid = _deleted_conflict(client, did, TEXT.replace("第二段。", "B 改的第二段。"))
    res = _resolve(client, did, cid, b, current(client, did)["revision"], [{"itemIndex": 0, "action": "keep_mine"}], "m3-legacy-1")
    assert res.status_code == 200, res.text
    data = res.json()["data"]
    assert "# 1 范围\n\n第一段。\n\nB 改的第二段。\n\n第三段。" in data["markdownContent"]
    assert not any(item.get("placementFallback") for item in data["items"])


def test_m3_falls_back_to_section_end_and_flags_it(team):
    """前驱没有记录（114 之前删除）且本节已没有足够的段落：放到所属章节末尾并标记 placementFallback。"""
    client, did, _users = team
    a = _as(client, "editor")
    assert save(client, did, 1, TEXT.replace("第二段。\n\n", ""), a).status_code == 200
    b, cid = _deleted_conflict(client, did, TEXT.replace("第二段。", "B 改的第二段。"))
    cur = current(client, did)
    a = _as(client, "editor")
    assert save(client, did, cur["revision"], "# 1 范围\n\n# 2 引用\n\n引用段。\n", a).status_code == 200
    with connect() as connection:
        connection.execute("UPDATE document_objects SET deleted_after_ref=NULL WHERE document_id=%s", (did,))
        connection.commit()
    b = _as(client, "editor2")
    res = _resolve(client, did, cid, b, current(client, did)["revision"], [{"itemIndex": 0, "action": "restore_as_new"}], "m3-end-1")
    assert res.status_code == 200, res.text
    data = res.json()["data"]
    assert "# 1 范围\n\nB 改的第二段。\n\n# 2 引用" in data["markdownContent"]
    assert data["items"][0]["placementFallback"] is True


def test_m3_deleted_after_ref_recorded(team):
    client, did, _users = team
    a = _as(client, "editor")
    assert save(client, did, 1, TEXT.replace("# 1 范围\n\n", "").replace("第二段。\n\n", ""), a).status_code == 200
    with connect() as connection:
        rows = {row["body"]: row["deleted_after_ref"] for row in connection.execute(
            "SELECT body, deleted_after_ref FROM document_objects WHERE document_id=%s AND deleted_at IS NOT NULL", (did,)).fetchall()}
        first = connection.execute("SELECT id FROM document_objects WHERE document_id=%s AND body='第一段。' AND deleted_at IS NULL",
                                   (did,)).fetchone()["id"]
    assert str(rows["# 1 范围"]) == "^"
    assert str(rows["第二段。"]) == str(first)


def test_preview_keeps_existing_noncharacter_apart_from_tabs(tmp_path):
    """复审：正文原有的 U+FDEF 不能被预览当成 tab 还原；行内 tab 仍保留。"""
    from docx import Document as WordDocument

    from app.docx_parser import parse_docx
    from app.docx_preview import render_document_html

    source = tmp_path / "mark.docx"
    doc = WordDocument()
    doc.add_paragraph("甲\ufdef乙\t丙\ufdee丁")
    doc.save(source)
    markdown = parse_docx(str(source)).markdown
    assert "\ufdef" in markdown and "\t" in markdown
    html_text = render_document_html(source, markdown)
    assert html_text.count("\t") == 1
    assert html_text.count("\ufdef") == 1 and html_text.count("\ufdee") == 1


def _bookmark_doc(path, build):
    from docx import Document as WordDocument
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    def mark(paragraph, kind, ident, name=None, parent=None):
        node = OxmlElement(f"w:bookmark{kind}")
        node.set(qn("w:id"), str(ident))
        if name:
            node.set(qn("w:name"), name)
        (parent if parent is not None else paragraph._p).append(node)

    doc = WordDocument()
    doc.add_heading("概述", 1)
    build(doc, mark)
    doc.save(path)


def _children(path, index):
    from docx import Document as WordDocument
    from docx.oxml.ns import qn

    out = []
    for child in WordDocument(path).paragraphs[index]._p:
        tag = child.tag.split("}")[1]
        if tag == "r":
            out.append("".join(node.text or "" for node in child.iter(qn("w:t"))))
        elif tag.startswith("bookmark"):
            out.append(f"{tag[8:]}:{child.get(qn('w:id'))}")
    return out


def _add_ref(doc, name):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    q = doc.add_paragraph("详见")
    for kind, text in (("begin", None), ("instr", f" REF {name} \\h "), ("separate", None), ("text", "表"), ("end", None)):
        run = q.add_run()
        if kind in ("begin", "separate", "end"):
            node = OxmlElement("w:fldChar")
            node.set(qn("w:fldCharType"), kind)
            run._r.append(node)
        elif kind == "instr":
            node = OxmlElement("w:instrText")
            node.text = text
            run._r.append(node)
        else:
            run.text = text


def test_bookmarks_keep_range_when_formatting_changes(tmp_path):
    """复审：改格式/段首加字时，部分覆盖与嵌套的书签保持原范围，不扩到整段、不丢。"""
    from app.docx_export import edit_original_docx
    from app.docx_parser import parse_docx

    source, out = tmp_path / "bm.docx", tmp_path / "out.docx"

    def build(doc, mark):
        p = doc.add_paragraph()
        mark(p, "Start", 1, "_Ref1")
        p.add_run("表 3-1")
        mark(p, "Start", 2, "_Ref2")
        p.add_run(" 系统")
        mark(p, "End", 2)
        mark(p, "End", 1)
        p.add_run("参数表")
        _add_ref(doc, "_Ref1")
        _add_ref(doc, "_Ref2")

    _bookmark_doc(source, build)
    markdown = parse_docx(str(source)).markdown
    for new, expected in (
        (markdown.replace("系统参数表", "系统**参数**表"),
         ["Start:1", "表 3-1", "Start:2", " 系统", "End:2", "End:1", "参", "数", "表"]),
        (markdown.replace("表 3-1 系统参数表", "附表 3-1 系统参数表"),
         ["附", "Start:1", "表 3-1", "Start:2", " 系统", "End:2", "End:1", "参数表"]),
        (markdown.replace("表 3-1 系统参数表", "**表** 3-1 系统参数表"),
         ["Start:1", "表", " 3-1", "Start:2", " 系统", "End:2", "End:1", "参数表"]),
    ):
        edit_original_docx(new, source, out)
        got = _children(out, 1)
        assert [item for item in got if item.startswith(("Start", "End"))] == [
            item for item in expected if item.startswith(("Start", "End"))]
        assert "".join(item for item in got if ":" not in item) == "".join(item for item in expected if ":" not in item)
        # 书签内文字不变
        start, end = got.index("Start:1"), got.index("End:1")
        inside = "".join(item for item in got[start:end] if ":" not in item).replace("**", "")
        assert inside.endswith("3-1 系统") and not inside.startswith("附"), got


def test_bookmark_spanning_paragraphs_keeps_its_end(tmp_path):
    from app.docx_export import edit_original_docx
    from app.docx_parser import parse_docx

    source, out = tmp_path / "span.docx", tmp_path / "out.docx"

    def build(doc, mark):
        first = doc.add_paragraph("第一段")
        mark(first, "Start", 5, "_Ref5")
        first.add_run("开始")
        second = doc.add_paragraph("第二段")
        mark(second, "End", 5)
        second.add_run("之后")
        _add_ref(doc, "_Ref5")

    _bookmark_doc(source, build)
    markdown = parse_docx(str(source)).markdown
    edit_original_docx(markdown.replace("第二段之后", "第二段**之后**"), source, out)
    got = _children(out, 2)
    assert got.index("End:5") == 1 and got[0] == "第二段", got


def test_referenced_bookmark_with_ambiguous_edge_stays_protected(tmp_path):
    import pytest as _pytest

    from app.docx_export import WordEditError, edit_original_docx
    from app.docx_parser import parse_docx

    source = tmp_path / "amb.docx"

    def build(doc, mark):
        p = doc.add_paragraph()
        p.add_run("前文")
        mark(p, "Start", 3, "_Ref3")
        p.add_run("表 3-1")
        mark(p, "End", 3)
        _add_ref(doc, "_Ref3")

    _bookmark_doc(source, build)
    markdown = parse_docx(str(source)).markdown
    with _pytest.raises((WordEditError, ValueError)):
        edit_original_docx(markdown.replace("前文表 3-1", "前言附录 3-1"), source, tmp_path / "out.docx")
    # 同样的改动落在未被引用的 _Toc 书签上可以改
    source2 = tmp_path / "toc.docx"

    def build2(doc, mark):
        p = doc.add_paragraph()
        p.add_run("前文")
        mark(p, "Start", 3, "_Toc3")
        p.add_run("表 3-1")
        mark(p, "End", 3)

    _bookmark_doc(source2, build2)
    edit_original_docx(parse_docx(str(source2)).markdown.replace("前文表 3-1", "前言附录 3-1"), source2, tmp_path / "out2.docx")
    assert "".join(item for item in _children(tmp_path / "out2.docx", 1) if ":" not in item) == "前言附录 3-1"


def test_bookmark_nested_in_other_element_stays_protected(tmp_path):
    import pytest as _pytest
    from docx.oxml import OxmlElement

    from app.docx_export import edit_original_docx
    from app.docx_parser import parse_docx

    source = tmp_path / "nested.docx"

    def build(doc, mark):
        p = doc.add_paragraph("正文")
        wrapper = OxmlElement("w:customXml")
        p._p.append(wrapper)
        mark(p, "Start", 9, "_Ref9", parent=wrapper)
        mark(p, "End", 9, parent=wrapper)
        p.add_run("其余")

    _bookmark_doc(source, build)
    markdown = parse_docx(str(source)).markdown
    with _pytest.raises(ValueError):
        edit_original_docx(markdown.replace("正文其余", "正文**其余**"), source, tmp_path / "out.docx")
