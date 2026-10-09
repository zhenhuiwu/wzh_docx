"""Checked native Visio drafts cannot overwrite a changing or inaccessible document."""

import base64
from copy import deepcopy
from io import BytesIO
import json
import sys
from uuid import UUID, uuid4
from zipfile import ZipFile

from docx import Document
from docx.shared import Inches
from fastapi.testclient import TestClient
import pytest
from pydantic import ValidationError

from app.config import clear_settings_cache
from app.db import connect, migrate
from app.document_diagrams import DiagramGraph, _diagram_path
from app.errors import AppError
from app.identity import create_user
from app.main import create_app


GRAPH = {"name": "订单处理流程", "nodes": [
    {"id": "start", "label": "开始", "type": "start"},
    {"id": "receive", "label": "接收订单", "type": "process"},
    {"id": "check", "label": "库存是否充足", "type": "decision"},
    {"id": "send", "label": "安排发货", "type": "process"},
    {"id": "wait", "label": "通知补货", "type": "process"},
    {"id": "end", "label": "结束", "type": "end"},
], "edges": [
    {"source": "start", "target": "receive", "label": ""},
    {"source": "receive", "target": "check", "label": ""},
    {"source": "check", "target": "send", "label": "是"},
    {"source": "check", "target": "wait", "label": "否"},
    {"source": "send", "target": "end", "label": ""},
    {"source": "wait", "target": "end", "label": ""},
]}


def test_graph_rejects_malformed_references_and_executable_extras():
    assert DiagramGraph.model_validate(GRAPH).model_dump() == GRAPH
    for change in (
        lambda graph: graph["nodes"].append(dict(graph["nodes"][0])),
        lambda graph: graph["edges"][0].update(target="missing"),
        lambda graph: graph["edges"][0].update(target="start"),
        lambda graph: graph["edges"].append(dict(graph["edges"][0])),
        lambda graph: graph["nodes"][0].update(label=" "),
        lambda graph: graph["nodes"][0].update(id="../outside"),
        lambda graph: graph["nodes"][0].update(label="a" * 61),
        lambda graph: graph.update(script="execute code"),
        lambda graph: graph.update(nodes=graph["nodes"] * 5),
    ):
        value = deepcopy(GRAPH)
        change(value)
        with pytest.raises(ValidationError):
            DiagramGraph.model_validate(value)
    with pytest.raises(AppError):
        _diagram_path(uuid4(), "../outside.png")


@pytest.mark.skipif(sys.platform != "win32", reason="Native Visio rendering requires Windows and installed Microsoft Visio")
def test_native_visio_api_adoption_and_conflicts(tmp_path, monkeypatch):
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path / 'app.db'}")
    monkeypatch.setenv("FILE_DATA_DIR", str(tmp_path / "files"))
    clear_settings_cache()
    try:
        assert "105_document_diagrams.sql" in migrate()
        owner = create_user(account="visio-admin", name="图表测试", password="test-password-1", roles=["system_admin"])
        create_user(account="visio-other", name="其他作者", password="test-password-1", roles=["project_creator"])
        generated_inputs = []
        def graph_draft(**kwargs):
            generated_inputs.append(kwargs)
            return deepcopy(GRAPH)
        monkeypatch.setattr("app.document_diagrams.generate_graph", graph_draft)
        with TestClient(create_app()) as client:
            csrf = client.get("/api/v1/auth/csrf").json()["data"]["csrfToken"]
            login = client.post("/api/v1/auth/login", json={"account": "visio-admin", "password": "test-password-1"}, headers={"X-CSRF-Token": csrf})
            assert login.status_code == 200
            headers = {"X-CSRF-Token": login.json()["data"]["csrfToken"]}
            project = client.post("/api/v1/projects", json={"name": "中性图表验收", "ownerId": str(owner), "status": "active"},
                                  headers={**headers, "Idempotency-Key": "visio-project"})
            assert project.status_code == 201, project.text
            pid = project.json()["data"]["id"]
            original = Document()
            original.add_heading("处理流程", 1)
            original.add_paragraph("接收订单，检查库存。库存充足则发货，否则通知补货。")
            original.add_picture(BytesIO(base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII=")), width=Inches(0.2))
            original.add_heading("说明", 1)
            original.add_paragraph("保持现有正文及原始图片。")
            source = BytesIO()
            original.save(source)
            upload = client.post("/api/v1/files", files={"file": ("中性说明.docx", source.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                                 data={"purpose": "attachment", "projectId": pid}, headers={**headers, "Idempotency-Key": "visio-source"})
            assert upload.status_code == 201, upload.text
            response = client.post(f"/api/v1/projects/{pid}/documents", json={"name": "订单流程说明", "sourceType": "upload", "wordFileId": upload.json()["data"]["id"]}, headers=headers)
            assert response.status_code == 200, response.text
            doc = response.json()["data"]
            did, revision = doc["id"], doc["revision"]
            path = f"/api/v1/documents/{did}/diagrams"
            assert client.get(path).json()["data"]["items"] == []
            body = {"expectedRevision": revision, "headingIndex": 0, "kind": "flowchart", "instruction": "生成订单处理流程"}
            assert client.post(path, json={**body, "expectedRevision": revision + 1}, headers=headers).status_code == 409
            assert not generated_inputs
            created = client.post(path, json=body, headers=headers)
            assert created.status_code == 201, created.text
            diagram = created.json()["data"]
            assert diagram["sourceRevision"] == revision and diagram["graph"] == GRAPH
            assert generated_inputs[0]["heading"] == "处理流程" and "说明" not in generated_inputs[0]["content"]
            assert client.get(path).json()["data"]["items"][0]["id"] == diagram["id"]
            assert client.get(f"/api/v1/documents/{did}").json()["data"]["revision"] == revision
            png = client.get(diagram["previewUrl"])
            assert png.status_code == 200 and png.content.startswith(b"\x89PNG")
            vsdx = client.get(diagram["downloadUrl"])
            assert vsdx.status_code == 200 and "vsdx" in vsdx.headers["content-disposition"]
            with ZipFile(BytesIO(vsdx.content)) as package:
                assert "visio/pages/page1.xml" in package.namelist()
            meta = json.loads((_diagram_path(UUID(diagram["id"]), "diagram.png").parent / "render-meta.json").read_text(encoding="utf-8"))
            assert meta["nodes"] == 6 and meta["connectors"] == 6 and meta["verifiedEditable"]
            assert client.get(diagram["downloadUrl"].replace("vsdx", "svg")).status_code == 422
            insert_path = path + f"/{diagram['id']}/insert"
            assert client.post(insert_path, json={"expectedRevision": revision + 1, "headingIndex": 0}, headers=headers).status_code == 409
            inserted = client.post(insert_path, json={"expectedRevision": revision, "headingIndex": 0}, headers=headers)
            assert inserted.status_code == 200, inserted.text
            updated = inserted.json()["data"]
            assert updated["revision"] == revision + 1
            images = updated["formatTemplate"]["images"]
            assert len(images) == 2 and images[0]["ref"].startswith("image_")
            assert images[1]["ref"] == f"docx-image:diagram-{diagram['id']}"
            assert updated["markdownContent"].index(images[1]["ref"]) < updated["markdownContent"].index("# 说明")
            assert client.post(insert_path, json={"expectedRevision": revision + 1, "headingIndex": 0}, headers=headers).status_code == 409
            exported = client.post(f"/api/v1/documents/{did}/export", json={"expectedRevision": revision + 1}, headers=headers)
            assert exported.status_code == 200, exported.text
            downloaded = client.get(f"/api/v1/files/{exported.json()['data']['fileId']}/content")
            assert downloaded.status_code == 200
            output = Document(BytesIO(downloaded.content))
            assert len(output.inline_shapes) == 2
            assert [paragraph.text for paragraph in output.paragraphs if paragraph.style.name == "Heading 1"] == ["处理流程", "说明"]
            with ZipFile(source) as old, ZipFile(BytesIO(downloaded.content)) as new:
                assert old.read("word/styles.xml") == new.read("word/styles.xml")
                assert old.read("word/media/image1.png") == new.read("word/media/image1.png")
            layout = client.get(f"/api/v1/documents/{did}/word-preview?scope=edited&expectedRevision={revision + 1}")
            assert layout.status_code == 200 and layout.content.startswith(b"%PDF"), layout.text[:200]
            saved = client.patch(f"/api/v1/documents/{did}/content", json={"expectedRevision": revision + 1, "markdownContent": updated["markdownContent"]}, headers=headers)
            assert saved.status_code == 200 and saved.json()["data"]["revision"] == revision + 2, saved.text
            # A chapter rewrite must preserve both source pictures and inserted diagrams.
            section = updated["markdownContent"].split("# 说明")[0].removeprefix("# 处理流程").strip()
            adopted = client.post(f"/api/v1/documents/{did}/sections/insert", json={
                "expectedRevision": revision + 2, "headingIndex": 0, "mode": "replace",
                "content": section.replace("接收订单", "核对订单"), "instruction": "仅修改接收订单为核对订单",
            }, headers=headers)
            assert adopted.status_code == 200 and adopted.json()["data"]["revision"] == revision + 3, adopted.text
            assert all(image["ref"] in adopted.json()["data"]["markdownContent"] for image in images)
            final_export = client.post(f"/api/v1/documents/{did}/export", json={"expectedRevision": revision + 3}, headers=headers)
            assert final_export.status_code == 200, final_export.text
            final_bytes = client.get(f"/api/v1/files/{final_export.json()['data']['fileId']}/content").content
            assert len(Document(BytesIO(final_bytes)).inline_shapes) == 2
            with ZipFile(source) as old, ZipFile(BytesIO(final_bytes)) as new:
                assert old.read("word/media/image1.png") == new.read("word/media/image1.png")
            # The graph draft and its native files are discarded if the source changes during rendering.
            def concurrent_update(_graph, directory):
                directory.mkdir(parents=True)
                with connect() as connection:
                    connection.execute("UPDATE documents SET revision=revision+1 WHERE id=%s", (did,))
                    connection.commit()
            monkeypatch.setattr("app.document_diagrams.render_visio", concurrent_update)
            conflict = client.post(path, json={**body, "expectedRevision": revision + 3}, headers=headers)
            assert conflict.status_code == 409
            assert len(client.get(path).json()["data"]["items"]) == 1
            assert len(list((tmp_path / "files" / "diagrams").iterdir())) == 1
            with connect() as connection:
                connection.execute("UPDATE documents SET markdown_content=%s WHERE id=%s", ("x" * 16001, did))
                connection.commit()
            assert client.post(path, json={**body, "expectedRevision": revision + 4, "headingIndex": None}, headers=headers).status_code == 422
            assert len(generated_inputs) == 2
        with TestClient(create_app()) as outsider:
            csrf = outsider.get("/api/v1/auth/csrf").json()["data"]["csrfToken"]
            other_login = outsider.post("/api/v1/auth/login", json={"account": "visio-other", "password": "test-password-1"}, headers={"X-CSRF-Token": csrf})
            assert other_login.status_code == 200
            for inaccessible in (path, diagram["downloadUrl"], diagram["previewUrl"]):
                assert outsider.get(inaccessible).status_code == 404
    finally:
        clear_settings_cache()
