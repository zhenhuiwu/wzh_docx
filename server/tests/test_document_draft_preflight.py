"""Finished chapter drafts must offer only adoption modes the native Word can keep."""

from __future__ import annotations

import asyncio
import json
from types import SimpleNamespace
from unittest.mock import MagicMock
from uuid import uuid4

import pytest
from docx import Document
from docx.oxml import OxmlElement

import app.documents as module
from app.docx_export import edit_original_docx
from app.docx_parser import parse_docx
from app.errors import AppError


@pytest.fixture
def route_context(tmp_path, monkeypatch):
    source = tmp_path / "original.docx"
    word = Document()
    word.add_heading("目标", level=1)
    paragraph = word.add_paragraph("旧正文，日期域必须保留。")
    for tag, value in (("w:fldChar", "begin"), ("w:instrText", " DATE "), ("w:fldChar", "end")):
        run, node = OxmlElement("w:r"), OxmlElement(tag)
        if tag == "w:fldChar":
            node.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fldCharType", value)
        else:
            node.text = value
        run.append(node)
        paragraph._p.append(run)
    word.add_heading("其他章节", level=1)
    word.add_paragraph("其他章节保持不变。")
    word.save(source)
    row = [None] * 16
    row[0], row[1], row[2], row[7], row[12], row[15] = uuid4(), uuid4(), "中性文档", parse_docx(str(source)).markdown, 7, ""
    settings = SimpleNamespace(ai_base_url="http://localhost/v1", ai_model="offline-test", ai_api_key="", ai_context_tokens=32000)
    request = SimpleNamespace(state=SimpleNamespace(request_id="preflight-test"))
    monkeypatch.setattr(module, "connect", lambda: MagicMock())
    monkeypatch.setattr(module, "section_memory_targets", lambda *args: {})
    monkeypatch.setattr(module, "document_principal", lambda *args: SimpleNamespace(roles=("system_admin",), user_id=uuid4()))
    monkeypatch.setattr(module, "_document_row", lambda *args: row)
    monkeypatch.setattr(module, "get_settings", lambda: settings)
    monkeypatch.setattr(module, "draft_section_text", lambda **kwargs: "清晰的新正文。")

    async def generate(**kwargs):
        yield "清晰的"
        yield "清晰的新正文。"

    monkeypatch.setattr(module, "stream_section_text", generate)
    checked = []

    def check_word(connection, current, content):
        checked.append(content)
        try:
            edit_original_docx(content, source, tmp_path / "checked.docx")
        except ValueError as exc:
            raise AppError(422, "AI_FORMAT_INVALID", str(exc)) from None

    monkeypatch.setattr(module, "_check_word", check_word)
    return request, row, source, tmp_path, checked


def stream_events(response):
    async def collect():
        return [json.loads(chunk) async for chunk in response.body_iterator]

    return asyncio.run(collect())


@pytest.mark.parametrize("streaming", [False, True])
def test_protected_word_chapter_returns_append_only_without_blocking_generation(route_context, streaming):
    request, row, source, tmp_path, checked = route_context
    original_bytes = source.read_bytes()
    response = module.draft_section(request, str(row[0]), module.SectionDraftBody(expectedRevision=7, headingIndex=0), stream=streaming)
    if streaming:
        events = stream_events(response)
        assert events[-1]["type"] == "done", events
        result = events[-1]
    else:
        result = response["data"]
    assert result["content"] == "清晰的新正文。" and result["allowedModes"] == ["append"]
    assert len(checked) == 2 and row[12] == 7 and source.read_bytes() == original_bytes
    adopted = module._section_update(row, 0, result["content"], "append")
    edit_original_docx(adopted, source, tmp_path / "adopted.docx")
    final_word = Document(tmp_path / "adopted.docx")
    assert final_word.element.body.xpath(".//w:fldChar")
    assert "旧正文，日期域必须保留。" in adopted and "其他章节保持不变。" in adopted
    with pytest.raises(AppError):
        module._check_word(None, row, module._section_update(row, 0, result["content"], "replace"))


@pytest.mark.parametrize("streaming", [False, True])
def test_draft_has_no_success_event_when_neither_word_adoption_mode_is_available(route_context, monkeypatch, streaming):
    request, row, source, _, _ = route_context
    before = source.read_bytes()

    def blocked(*args):
        raise AppError(422, "AI_FORMAT_INVALID", "原稿无法安全修改。")

    monkeypatch.setattr(module, "_check_word", blocked)
    body = module.SectionDraftBody(expectedRevision=7, headingIndex=0)
    if streaming:
        events = stream_events(module.draft_section(request, str(row[0]), body, stream=True))
        assert events[-1]["type"] == "error" and events[-1]["message"] == "原稿无法安全修改。"
        assert not any(event["type"] == "done" for event in events)
    else:
        with pytest.raises(AppError) as failure:
            module.draft_section(request, str(row[0]), body)
        assert failure.value.code == "AI_FORMAT_INVALID"
    assert source.read_bytes() == before and row[12] == 7


def test_plain_word_draft_can_offer_both_adoption_modes(route_context):
    _, row, source, _, _ = route_context
    word = Document()
    word.add_heading("目标", level=1)
    word.add_paragraph("旧正文。")
    word.save(source)
    row[7] = parse_docx(str(source)).markdown
    assert module._section_adoption_modes(None, row, 0, "新正文。") == ["replace", "append"]


def test_selection_draft_does_not_claim_preflight_without_a_real_selection_range(route_context):
    request, row, _, _, checked = route_context
    with pytest.raises(AppError) as rejected:
        module.draft_section(request, str(row[0]), module.SectionDraftBody(expectedRevision=7, headingIndex=0,
                             selectedText="旧正文", scope="selection"))
    assert rejected.value.code == "SELECTION_REQUIRED" and checked == []


def test_stream_progress_keeps_draft_and_checks_word_before_done(route_context, monkeypatch):
    request, row, _, _, checked = route_context

    async def generate(**kwargs):
        assert kwargs["include_progress"] is True
        yield {"type": "progress", "stage": "planning", "message": "正在确认本次修改范围"}
        yield "清晰的新正文。"

    monkeypatch.setattr(module, "stream_section_text", generate)
    events = stream_events(module.draft_section(request, str(row[0]), module.SectionDraftBody(expectedRevision=7, headingIndex=0), stream=True))
    assert [event["type"] for event in events] == ["start", "progress", "draft", "progress", "done"]
    assert [(event["stage"], event["message"]) for event in events if event["type"] == "progress"] == [
        ("planning", "正在确认本次修改范围"), ("validation", "正在检查正文规范与 Word 排版")]
    assert all(event["revision"] == 7 for event in events)
    assert events[-1]["content"] == "清晰的新正文。" and len(checked) == 2
